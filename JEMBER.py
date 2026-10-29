import hmac
import json, os, glob, datetime, re, base64, random, calendar as pycal, urllib.request, urllib.error, urllib.parse
import pandas as pd
import streamlit as st
import streamlit.components.v1 as components

st.set_page_config(page_title='JEMBER BONDOWOSO ANALYSIS — Dashboard Dealer', page_icon='🐺', layout='wide')
st.markdown("""<style>
[data-testid="stSidebar"] { background-color: #3B18FF !important; min-width: 245px !important; max-width: 265px !important; }
[data-testid="stSidebar"] > div:first-child { padding: 1.2rem 1.1rem 0 1.1rem !important; }
[data-testid="stSidebar"] *:not([data-testid="stIconMaterial"]):not(.material-icons) { color: #FFFFFF !important; }
[data-testid="stSidebar"] hr { margin: 12px 0px !important; border: none !important; border-top: 1px solid rgba(255,255,255,0.22) !important; }
[data-testid="stSidebar"] h3 { font-size: 16.5px !important; font-weight: 800 !important; margin: 2px 0 6px 0 !important; padding-bottom: 0px !important; letter-spacing: 0.3px; }
[data-testid="stSidebar"] div[role="radiogroup"] { gap: 2px !important; }
[data-testid="stSidebar"] div[role="radiogroup"] label { padding: 3px 4px !important; margin: 0px !important; border: none !important; font-size: 13.5px !important; font-weight: 700 !important; min-height: 28px !important; align-items: center !important; }
[data-testid="stSidebar"] div[role="radiogroup"] label p { font-size: 13.5px !important; font-weight: 700 !important; margin: 0px !important; line-height: 1.2 !important; }
[data-testid="stSidebar"] div[role="radiogroup"] label:hover { background-color: rgba(255,255,255,0.1) !important; border-radius: 5px !important; }
[data-testid="stSidebar"] div[data-testid="stDateInput"] label p { font-size: 12.5px !important; font-weight: 700 !important; margin-bottom: 4px !important; color: #FFFFFF !important; }
[data-testid="stSidebar"] div[data-baseweb="input"] { background-color: #FFFFFF !important; border: 1.5px solid #0E1117 !important; border-radius: 7px !important; min-height: 36px !important; }
[data-testid="stSidebar"] div[data-baseweb="input"] input { color: #0E1117 !important; -webkit-text-fill-color: #0E1117 !important; font-weight: 800 !important; font-size: 13.5px !important; background-color: #FFFFFF !important; padding: 6px 10px !important; }
[data-testid="stSidebar"] div[data-baseweb="calendar"] * { color: #0E1117 !important; }
[data-testid="stSidebar"] div[data-baseweb="textarea"], [data-testid="stSidebar"] div[data-baseweb="textarea"] textarea { background-color: #FFFFFF !important; color: #0E1117 !important; -webkit-text-fill-color: #0E1117 !important; font-weight: 700 !important; font-size: 12.5px !important; }
.azura-logo-box { display: flex; flex-direction: column; align-items: flex-start; justify-content: center; margin: -8px 0 10px 0; padding-bottom: 12px; border-bottom: 1px solid rgba(255,255,255,0.22); }
.azura-logo-img { max-width: 195px; max-height: 90px; width: auto; height: auto; object-fit: contain; border-radius: 6px; box-shadow: 0 2px 8px rgba(0,0,0,0.35); background-color: transparent; }
.azura-fallback-badge { width: 72px; height: 72px; border-radius: 12px; background: radial-gradient(circle, #0051FF 0%, #050B14 85%); border: 2px solid #94A3B8; box-shadow: 0 0 10px rgba(0,0,0,0.45); display: flex; flex-direction: column; align-items: center; justify-content: center; }
.target-ideal-caption { font-size: 12.5px !important; font-weight: 700 !important; color: rgba(255,255,255,0.85) !important; margin-top: 6px; padding-bottom: 10px; border-bottom: 1px solid rgba(255,255,255,0.22); }
.main-title { font-size: 44px !important; font-weight: 800 !important; color: #4318FF !important; margin-bottom: 2px; line-height: 1.1; }
.sub-title { font-size: 15px; font-weight: 600; color: #475569 !important; margin-bottom: 15px; }
.section-title { font-size: 22px !important; font-weight: 700 !important; color: #4318FF !important; margin: 20px 0 10px 0; }
.cap-title-banner { background: linear-gradient(90deg, #4318FF 0%, #1E293B 100%); color: #FFFFFF !important; font-family: 'Source Sans Pro', sans-serif; font-size: 15px; font-weight: 800; padding: 9px 14px; border-bottom: 2px solid #3B18FF; letter-spacing: 0.4px; text-align: left; }
.compact-table-wrap { width: 100%; overflow-x: auto; margin-bottom: 4px; border-radius: 8px; border: 1px solid #2D3748; background-color: #0E1117; }
table.compact-tbl { width: 100%; border-collapse: collapse; font-family: 'Source Sans Pro', sans-serif; font-size: 13px; background-color: #0E1117; }
table.compact-tbl th { background-color: #1A202C; color: #CBD5E1; font-weight: 700; padding: 7px 6px; border-bottom: 1px solid #2D3748; border-right: 1px solid #1E293B; text-align: right; white-space: nowrap; }
table.compact-tbl th:first-child, table.compact-tbl td:first-child { text-align: left; padding-left: 10px; }
table.compact-tbl td { padding: 6px 6px; border-bottom: 1px solid #1E293B; border-right: 1px solid #1E293B; text-align: right; white-space: nowrap; font-weight: 600; }
table.compact-tbl tr:hover td, table.ultra-compact-tbl tr:hover td { background-color: rgba(255,255,255,0.04); }
.ultra-compact-wrap { width: 100%; overflow-x: hidden; margin-bottom: 4px; border-radius: 8px; border: 1px solid #2D3748; background-color: #0E1117; }
table.ultra-compact-tbl { width: 100%; table-layout: fixed; border-collapse: collapse; font-family: 'Source Sans Pro', sans-serif; font-size: 13.5px; background-color: #0E1117; }
table.ultra-compact-tbl th { background-color: #1A202C; color: #CBD5E1; font-weight: 800; padding: 6px 3px; border-bottom: 1px solid #2D3748; border-right: 1px solid #1E293B; text-align: center; white-space: normal; line-height: 1.15; font-size: 11.5px; word-break: break-word; }
table.ultra-compact-tbl th:first-child, table.ultra-compact-tbl td:first-child { width: 17.5%; text-align: left; padding-left: 8px; font-size: 12px; font-weight: 700; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }
table.ultra-compact-tbl td { padding: 5px 5px; border-bottom: 1px solid #1E293B; border-right: 1px solid #1E293B; text-align: right; white-space: nowrap; font-weight: 700; font-size: 13.5px; overflow: hidden; }
table.narrative-tbl { width: 100%; border-collapse: collapse; font-family: 'Source Sans Pro', sans-serif; font-size: 13.5px; background-color: #0E1117; }
table.narrative-tbl th { background-color: #4318FF; color: #FFFFFF; font-weight: 700; padding: 10px 14px; border: 1px solid #2D3748; text-align: left; font-size: 13.5px; }
table.narrative-tbl td { padding: 10px 14px; border: 1px solid #2D3748; text-align: left; vertical-align: top; color: #E2E8F0; line-height: 1.55; font-weight: 400; white-space: normal; }
table.narrative-tbl td.item-col { width: 15%; color: #E2E8F0; text-align: left; background-color: #161B26; white-space: nowrap; font-weight: 600; }
table.narrative-tbl td.cond-col { width: 42.5%; background-color: #0E1117; }
table.narrative-tbl td.plan-col { width: 42.5%; background-color: #111622; }
p.nar-p { margin: 0 0 6px 0; padding: 0; color: #E2E8F0; font-weight: 400; line-height: 1.55; text-align: left; }
p.nar-p:last-of-type { margin-bottom: 4px; }
.focus-alert-box { margin-top: 14px; padding: 11px 16px; background: linear-gradient(90deg, rgba(255,0,127,0.14) 0%, rgba(14,17,23,0.95) 100%); border-left: 4px solid #FF007F; border-top: 1px solid #2D3748; border-right: 1px solid #2D3748; border-bottom: 1px solid #2D3748; border-radius: 6px; font-size: 13.5px; line-height: 1.65; color: #E2E8F0; }
.unmapped-sales-box { margin: 4px 0 12px 0; padding: 10px 14px; background: linear-gradient(90deg, rgba(255,234,0,0.12) 0%, rgba(14,17,23,0.96) 100%); border-left: 4px solid #FFEA00; border-top: 1px solid #2D3748; border-right: 1px solid #2D3748; border-bottom: 1px solid #2D3748; border-radius: 6px; font-size: 13px; line-height: 1.6; color: #E2E8F0; }
.unmapped-badge { display: inline-block; background-color: #1E293B; color: #FFEA00; border: 1px solid #475569; border-radius: 4px; padding: 2px 8px; margin: 2px 4px 2px 0; font-size: 12px; font-weight: 700; }
</style>""", unsafe_allow_html=True)
st.markdown("""<style>
[data-testid="stSidebar"] div[role="radiogroup"] label:has(input:checked) { background-color: rgba(255,255,255,0.20) !important; border-radius: 6px !important; box-shadow: inset 3px 0 0 #FFEA00 !important; padding-left: 9px !important; }
div[data-testid="stTabs"] button[data-baseweb="tab"], .stTabs button[data-baseweb="tab"] { padding: 10px 16px !important; }
div[data-testid="stTabs"] button[data-baseweb="tab"] p, .stTabs button[data-baseweb="tab"] p { font-size: 15px !important; font-weight: 700 !important; }
div[data-testid="stTabs"] button[aria-selected="true"] p, .stTabs button[aria-selected="true"] p { color: #9F8CFF !important; }
div[data-testid="stTabs"] [data-baseweb="tab-highlight"], .stTabs [data-baseweb="tab-highlight"] { background-color: #4318FF !important; height: 3px !important; }
@media (max-width: 768px) { .main-title { font-size: 28px !important; } .section-title { font-size: 18px !important; } .block-container { padding: 1rem 0.8rem 3rem 0.8rem !important; } }
</style>""", unsafe_allow_html=True)

# ====================================================================================================
# ===================== UI v2 — grafik interaktif, kartu progres, insight, leaderboard ==============
# (tanpa dependensi baru: grafik dibuat sebagai HTML/SVG ringan & dirender lewat components.html)
# ====================================================================================================
import html as _html

BSG_GRN, BSG_YEL, BSG_RED, BSG_CYN, BSG_PUR = '#00FF66', '#FFEA00', '#FF4D8D', '#00E5FF', '#4318FF'
_bsg_fragment = getattr(st, 'fragment', None) or getattr(st, 'experimental_fragment', None) or (lambda f: f)

def _esc(x): return _html.escape(str(x), quote=True)

def bsg_num(x, default=0.0):
    try:
        if x is None: return default
        if isinstance(x, str):
            s = x.replace('%', '').replace(',', '').strip()
            return float(s) if s not in ('', 'nan', 'NaN', 'None') else default
        v = float(x)
        return default if v != v else v
    except Exception: return default

def bsg_stat_color(v, good, warn): return BSG_GRN if v >= good else (BSG_YEL if v >= warn else BSG_RED)

def bsg_ach_bar(v, ideal):
    """Skala bar kartu: 0..1,5x ideal (marker ideal di 66,7%)."""
    try: return (max(0.0, min(float(v) / float(ideal), 1.5)) / 1.5 * 100.0, 66.67)
    except Exception: return (None, None)

def bsg_safe(fn, *a, label='Visual', **kw):
    try: return fn(*a, **kw)
    except Exception as e:
        try: st.caption(f'⚠️ {label} tidak dapat ditampilkan ({type(e).__name__}: {str(e)[:90]}).')
        except Exception: pass
        return None

def bsg_toast(msg, icon='✅'):
    try: st.toast(msg, icon=icon)
    except Exception: pass

def bsg_toast_once(key, msg, icon='✅'):
    flag = f'_toast_{key}'
    if not st.session_state.get(flag):
        st.session_state[flag] = True
        bsg_toast(msg, icon)

def bsg_segmented(label, options, key, default=None):
    default = default if default is not None else options[0]
    try:
        if hasattr(st, 'segmented_control'):
            v = st.segmented_control(label, options, default=default, key=key, label_visibility='collapsed')
            return v if v in options else default
    except Exception: pass
    return st.radio(label, options, index=options.index(default), horizontal=True, key=key, label_visibility='collapsed')

# ---------------------------------------------------------------- kartu dengan progress bar ----
def bsg_prog_bar(prog, color, ref=None):
    try: p = max(0.0, min(100.0, float(prog)))
    except Exception: return ''
    mk = ''
    if ref is not None:
        try: mk = f'<div title="Acuan ideal" style="position:absolute;left:{max(0.0, min(100.0, float(ref))):.1f}%;top:-3px;bottom:-3px;width:2px;background:#FFFFFF;opacity:.9;border-radius:1px;"></div>'
        except Exception: mk = ''
    return (f'<div style="position:relative;height:6px;border-radius:4px;background:#1E293B;margin-top:8px;">'
            f'<div style="width:{p:.1f}%;height:100%;border-radius:4px;background:{color};"></div>{mk}</div>')

def bsg_card(icon, label, value, sub, color, prog=None, ref=None):
    bar = bsg_prog_bar(prog, color, ref) if prog is not None else ''
    return (f'<div style="background:linear-gradient(160deg,#151B2B 0%,#0E1117 100%);border:1px solid #232B3E;border-top:3px solid {color};border-radius:12px;padding:12px 14px;min-height:98px;box-sizing:border-box;">'
            f'<div style="font-size:11.5px;font-weight:700;letter-spacing:0.6px;color:#94A3B8;text-transform:uppercase;white-space:nowrap;">{icon} {label}</div>'
            f'<div style="font-size:27px;font-weight:800;color:{color};line-height:1.2;margin:4px 0 2px 0;">{value}</div>'
            f'<div style="font-size:12px;font-weight:500;color:#CBD5E1;line-height:1.45;">{sub}</div>{bar}</div>')

# ---------------------------------------------------------------- kerangka iframe grafik ----
_BSG_FRAME = """<!DOCTYPE html><html><head><meta charset="utf-8"><style>
*{box-sizing:border-box}
html,body{margin:0;padding:0;background:transparent;font-family:'Source Sans Pro','Segoe UI',Arial,sans-serif;color:#E2E8F0}
.card{background:linear-gradient(160deg,#151B2B 0%,#0E1117 100%);border:1px solid #232B3E;border-radius:12px;padding:12px 14px 10px 14px;width:100%}
.ttl{font-size:14px;font-weight:800;color:#E2E8F0;line-height:1.25}
.sub{font-size:11.5px;color:#94A3B8;margin:2px 0 8px 0;line-height:1.4}
.lg{display:inline-flex;align-items:center;gap:5px;font-size:11px;color:#94A3B8;margin:0 12px 6px 0}
.dot{width:9px;height:9px;border-radius:50%;display:inline-block}
.mk{display:inline-block;width:2px;height:11px;background:#fff;opacity:.9}
.ds{display:inline-block;width:16px;height:0;border-top:2px dashed #94A3B8}
.chip{display:inline-block;font-size:11.5px;font-weight:700;padding:2px 8px;border-radius:999px;background:#1B2336;border:1px solid #2D3748;margin:0 6px 6px 0;color:#CBD5E1}
#tip{position:fixed;left:0;top:0;pointer-events:none;background:#0B1020;border:1px solid #475569;border-radius:8px;padding:6px 9px;font-size:12px;line-height:1.5;opacity:0;transition:opacity .08s;z-index:99;max-width:280px;color:#E2E8F0}
@keyframes grow{from{transform:scaleX(0)}to{transform:scaleX(1)}}
__CSS__
</style></head><body>__INNER__<div id="tip"></div>
<script>(function(){var t=document.getElementById('tip');if(!t)return;
document.querySelectorAll('[data-tip]').forEach(function(e){
e.addEventListener('mousemove',function(ev){t.innerHTML=e.getAttribute('data-tip');t.style.opacity=1;var x=ev.clientX+14,y=ev.clientY+14;
if(x+t.offsetWidth>window.innerWidth-6)x=Math.max(4,ev.clientX-t.offsetWidth-14);if(y+t.offsetHeight>window.innerHeight-4)y=Math.max(4,ev.clientY-t.offsetHeight-10);
t.style.left=x+'px';t.style.top=y+'px';});
e.addEventListener('mouseleave',function(){t.style.opacity=0;});});})();</script>
<script>(function(){function setFrameH(h){var fe=window.frameElement;if(!fe)return;var old=fe.getBoundingClientRect().height;if(Math.abs(old-h)<=2)return;var els=[],el=fe.parentElement,i=0;while(el&&i<4){if(Math.abs(el.getBoundingClientRect().height-old)<=3)els.push(el);el=el.parentElement;i++;}els.forEach(function(e){e.style.height=h+'px';e.style.minHeight=h+'px';});fe.style.height=h+'px';fe.setAttribute('height',h);}
function fit(){try{var c=document.querySelector('.card');if(!c)return;setFrameH(Math.ceil(c.getBoundingClientRect().height)+4);}catch(e){}}
fit();window.addEventListener('resize',function(){setTimeout(fit,60);});setTimeout(fit,200);setTimeout(fit,800);})();</script>
__JS__</body></html>"""

def _bsg_frame(inner, css='', js=''):
    return _BSG_FRAME.replace('__CSS__', css).replace('__INNER__', inner).replace('__JS__', js)

def bsg_show(res):
    if not res: return
    html_str, h = res
    components.html(html_str, height=int(h), scrolling=False)

def _hdr(title, subtitle=''):
    return f'<div class="ttl">{_esc(title)}</div>' + (f'<div class="sub">{_esc(subtitle)}</div>' if subtitle else '<div style="height:8px"></div>')

# ---------------------------------------------------------------- 1. ranking bar horizontal ----
_RANK_CSS = """.rk{display:flex;align-items:center;gap:8px;height:26px}
.rk-n{flex:0 0 36%;max-width:215px;font-size:12px;font-weight:700;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;color:#E2E8F0}
.rk-t{position:relative;flex:1;height:14px;border-radius:7px;background:#1B2336}
.rk-f{height:100%;border-radius:7px;transform-origin:left center;animation:grow .7s ease-out both}
.rk-m{position:absolute;top:-3px;bottom:-3px;width:2px;background:#fff;opacity:.9;border-radius:1px}
.rk-v{flex:0 0 58px;text-align:right;font-size:12px;font-weight:800}
.rk-x{flex:0 0 24px;text-align:right;font-size:12px;font-weight:800;color:#FFEA00}
.rk:hover .rk-t{background:#242E45}
.rk:hover .rk-n{color:#fff}"""

def bsg_chart_rank(title, subtitle, rows, ideal=None, ideal_label='Ideal', unit='%', max_rows=None, fixed_color=None, ascending=False, parts=False, scale_max=None):
    """rows: list dict(name, val, label=None, tip=None). Warna otomatis terhadap ideal (hijau ≥ ideal, kuning ≥ 85%, merah di bawahnya).
    parts=True → kembalikan (inner_html, css, tinggi) untuk digabung ke panel; False → (dokumen_iframe, tinggi)."""
    rows = [r for r in rows if r is not None]
    if not rows: return None
    rows = sorted(rows, key=lambda r: r['val'], reverse=not ascending)
    if max_rows: rows = rows[:max_rows]
    mx = float(scale_max) if scale_max else max([r['val'] for r in rows] + ([ideal * 1.15] if ideal else []) + [1.0])
    def _c(v):
        if fixed_color: return fixed_color
        if ideal is None: return BSG_CYN
        return BSG_GRN if v >= ideal else (BSG_YEL if v >= 0.85 * ideal else BSG_RED)
    body = ''; has_x = any(r.get('extra') is not None for r in rows)
    for i, r in enumerate(rows):
        v = max(0.0, r['val']); c = _c(v)
        lab = r.get('label') or f'{v:.1f}{unit}'
        tip = r.get('tip') or f"<b>{_esc(r['name'])}</b><br>{lab}"
        mk = f'<div class="rk-m" style="left:{min(100.0, ideal / mx * 100.0):.2f}%"></div>' if ideal else ''
        body += (f'<div class="rk" data-tip="{_esc(tip)}"><div class="rk-n" title="{_esc(r["name"])}">{_esc(r["name"])}</div>'
                 f'<div class="rk-t"><div class="rk-f" style="width:{min(100.0, v / mx * 100.0):.2f}%;background:{c};animation-delay:{min(i, 12) * 40}ms"></div>{mk}</div>'
                 f'<div class="rk-v" style="color:{c}">{_esc(lab)}</div>' + (f'<div class="rk-x">{_esc(r["extra"]) if r.get("extra") is not None else ""}</div>' if has_x else '') + '</div>')
    leg = ''
    if ideal and not fixed_color:
        leg = f'<span class="lg"><span class="mk"></span>{_esc(ideal_label)} {ideal:.0f}{unit}</span>'   # hanya penanda garis ideal; warna hijau/kuning/merah sudah terlihat dari batang
    inner = f'<div class="card">{_hdr(title, subtitle)}{leg}{body}</div>'
    h = 59 + (22 if leg else 0) + (14 if subtitle else 0) + 26 * len(rows)
    if parts: return inner, _RANK_CSS, h
    return _bsg_frame(inner, css=_RANK_CSS), h

# ---------------------------------------------------------------- 2. funnel ----
def bsg_chart_funnel(title, subtitle, stages):
    """stages: list (label, value, color). Lebar bar relatif terhadap tahap pertama; tampil konversi antar tahap."""
    st_ = [(l, max(0.0, float(v)), c) for l, v, c in stages]
    if not st_ or st_[0][1] <= 0: return None
    base = st_[0][1]; body = ''
    for i, (l, v, c) in enumerate(st_):
        w = max(v / base * 100.0, 12.0) if v > 0 else 8.0
        if i > 0:
            prev = st_[i - 1][1]
            conv = (v / prev * 100.0) if prev > 0 else 0.0
            body += f'<div class="fn-c">↓ {conv:.1f}% dari {_esc(st_[i - 1][0])}</div>'
        tip = f'<b>{_esc(l)}</b><br>{int(v):,} ({v / base * 100.0:.1f}% dari {_esc(st_[0][0])})'
        body += (f'<div class="fn-r" data-tip="{_esc(tip)}"><div class="fn-l" title="{_esc(l)}">{_esc(l)}</div><div class="fn-t">'
                 f'<div class="fn-b" style="width:{w:.1f}%;background:linear-gradient(90deg,{c},{c}CC);animation-delay:{i * 90}ms">{int(v):,}</div></div></div>')
    css = (".fn-r{display:flex;align-items:center;gap:8px;height:28px}"
           ".fn-l{flex:0 0 38%;font-size:12px;font-weight:700;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;color:#E2E8F0}"
           ".fn-t{flex:1;display:flex;justify-content:center}"
           ".fn-b{height:26px;border-radius:7px;display:flex;align-items:center;justify-content:center;color:#06101F;font-size:12.5px;font-weight:900;"
           "transform-origin:center;animation:grow .6s ease-out both;white-space:nowrap;min-width:46px}"
           ".fn-r:hover .fn-b{filter:brightness(1.18)}"
           ".fn-c{margin-left:calc(38% + 8px);text-align:center;font-size:10.5px;color:#94A3B8;height:16px;line-height:16px}")
    inner = f'<div class="card">{_hdr(title, subtitle)}{body}</div>'
    n = len(st_)
    return _bsg_frame(inner, css=css), 60 + (14 if subtitle else 0) + n * 28 + (n - 1) * 16

# ---------------------------------------------------------------- 3. heatmap ----
def _bsg_heat_color(s):
    """skor 0..100+ → merah → kuning → hijau."""
    stops = [(0.0, (255, 77, 141)), (50.0, (255, 176, 32)), (100.0, (0, 230, 118))]
    s = max(0.0, min(100.0, float(s)))
    for (a, ca), (b, cb) in zip(stops, stops[1:]):
        if s <= b:
            t = (s - a) / (b - a) if b > a else 0.0
            return '#%02X%02X%02X' % tuple(int(round(ca[k] + (cb[k] - ca[k]) * t)) for k in range(3))
    return '#00E676'

def bsg_chart_heat(title, subtitle, row_names, col_names, cells, max_rows=24):
    """cells[(row, col)] = (skor | None, label, tip)."""
    row_names = list(row_names)[:max_rows]
    if not row_names: return None
    css = (".ht{width:100%;border-collapse:separate;border-spacing:3px}.ht th{font-size:11.5px;font-weight:800;color:#CBD5E1;padding:4px 2px;text-align:center}"
           ".ht td.n{font-size:12px;font-weight:700;color:#E2E8F0;white-space:nowrap;max-width:190px;overflow:hidden;text-overflow:ellipsis;padding-right:6px}"
           ".ht td.c{height:26px;text-align:center;font-size:12.5px;font-weight:800;border-radius:6px;color:#06101F}.ht td.c:hover{filter:brightness(1.15);outline:1px solid #fff}"
           ".ht td.na{background:#1B2336;color:#64748B}")
    th = '<tr><th></th>' + ''.join(f'<th>{_esc(c)}</th>' for c in col_names) + '</tr>'
    rows_html = ''
    for rn in row_names:
        tds = f'<td class="n" title="{_esc(rn)}">{_esc(rn)}</td>'
        for cn in col_names:
            sc, lab, tip = cells.get((rn, cn), (None, '–', ''))
            if sc is None: tds += f'<td class="c na" data-tip="{_esc(tip or "Tidak ada data")}">–</td>'
            else: tds += f'<td class="c" style="background:{_bsg_heat_color(sc)}" data-tip="{_esc(tip)}">{_esc(lab)}</td>'
        rows_html += f'<tr>{tds}</tr>'
    leg = ('<div style="display:flex;align-items:center;gap:8px;margin-top:6px;font-size:10.5px;color:#94A3B8;">Skor vs ideal:'
           f'<span style="width:120px;height:8px;border-radius:4px;background:linear-gradient(90deg,{_bsg_heat_color(0)},{_bsg_heat_color(50)},{_bsg_heat_color(100)})"></span>0% → 100%+</div>')
    inner = f'<div class="card">{_hdr(title, subtitle)}<table class="ht">{th}{rows_html}</table>{leg}</div>'
    return _bsg_frame(inner, css=css), 67 + (14 if subtitle else 0) + 30 * (len(row_names) + 1)

# ---------------------------------------------------------------- 4. donut ----
BSG_PAL = ['#4318FF', '#00E5FF', '#FFB020', '#FF4D8D', '#00E676', '#B388FF', '#FF8A65', '#4DD0E1', '#FFD54F', '#90A4AE', '#F06292', '#AED581']

def bsg_chart_donut(title, subtitle, items, center_label='Total', parts=False, keep_zero=False):
    """items: list (label, value, color). keep_zero=True → item bernilai 0 tetap tampil di legenda (0 · 0%)."""
    all_its = [(l, float(v or 0), c) for l, v, c in items]
    its = [t for t in all_its if t[1] > 0]
    total = sum(v for _, v, _ in its)
    if total <= 0 and not keep_zero: return None
    if not keep_zero: all_its = its
    R, SW = 44.0, 16.0; C = 2 * 3.141592653589793 * R; off = 0.0; segs = ''
    for l, v, c in its:
        ln = v / total * C; gap = 0.8 if len(its) > 1 else 0.0
        tip = f'<b>{_esc(l)}</b><br>{int(v):,} ({v / total * 100.0:.1f}%)'
        segs += (f'<circle class="sg" cx="60" cy="60" r="{R}" fill="none" stroke="{c}" stroke-width="{SW}" stroke-dasharray="{max(ln - gap, 0.1):.3f} {C - max(ln - gap, 0.1):.3f}" '
                 f'stroke-dashoffset="{-off:.3f}" data-tip="{_esc(tip)}"/>')
        off += ln
    leg = ''.join(f'<div class="dl"><span class="dot" style="background:{c}"></span><span class="dn" title="{_esc(l)}">{_esc(l)}</span><b>{int(v):,}</b><i>{(v / total * 100.0 if total > 0 else 0.0):.0f}%</i></div>' for l, v, c in all_its)
    css = (".dw{display:flex;align-items:center;gap:14px}.dc{flex:0 0 132px}.sg{transition:stroke-width .12s}.sg:hover{stroke-width:20}"
           ".dlg{flex:1;min-width:0}.dl{display:flex;align-items:center;gap:7px;font-size:12px;min-height:22px}.dn{flex:1;overflow:hidden;text-overflow:ellipsis;white-space:nowrap;color:#CBD5E1}"
           ".dl b{font-weight:800}.dl i{font-style:normal;color:#94A3B8;width:34px;text-align:right}")
    ff = 'font-family="Source Sans Pro,Segoe UI,Arial,sans-serif"'
    svg = (f'<svg class="dc" viewBox="0 0 120 120" width="132" height="132"><circle cx="60" cy="60" r="{R}" fill="none" stroke="#1B2336" stroke-width="{SW}"/>'
           f'<g transform="rotate(-90 60 60)">{segs}</g><text {ff} x="60" y="61" text-anchor="middle" font-size="17" font-weight="800" fill="#E2E8F0">{int(total):,}</text>'
           f'<text {ff} x="60" y="76" text-anchor="middle" font-size="9" fill="#94A3B8">{_esc(center_label)}</text></svg>')
    inner = f'<div class="card">{_hdr(title, subtitle)}<div class="dw">{svg}<div class="dlg">{leg}</div></div></div>'
    h = 60 + (14 if subtitle else 0) + max(132, 22 * len(all_its))
    if parts: return inner, css, h + 22 * len(all_its) + 6   # +legenda: pada layar sempit/stack3 legenda berada di bawah donut
    return _bsg_frame(inner, css=css), h

# ---------------------------------------------------------------- 5. kurva target ideal vs realisasi (JS, responsif) ----
_CURVE_JS = """<script>(function(){
var CFG=__CFG__, wrap=document.getElementById('cv'), H=CFG.h;
function nice(x){var p=Math.pow(10,Math.floor(Math.log10(x||1)));var f=x/p;return (f<=1?1:f<=2?2:f<=5?5:10)*p;}
function fmt(v){return (Math.round(v*10)/10).toLocaleString('id-ID');}
function draw(){
 var W=wrap.clientWidth||600, m={l:46,r:16,t:12,b:26}, iw=W-m.l-m.r, ih=H-m.t-m.b, D=CFG.maxd;
 var vs=CFG.ideal.slice(); CFG.real.forEach(function(p){vs.push(p[1]);}); CFG.proj.forEach(function(p){vs.push(p[1]);}); vs.push(CFG.target); (CFG.ext||[]).forEach(function(e){vs.push(e.v);});
 var mx=Math.max.apply(null,vs)*1.06; if(!(mx>0)) mx=1; var step=nice(mx/4), top=Math.ceil(mx/step)*step;
 function X(d){return m.l+(d-1)/Math.max(1,D-1)*iw;} function Y(v){return m.t+ih-v/top*ih;}
 wrap.setAttribute('data-w',W); var s='<svg viewBox="0 0 '+W+' '+H+'" width="'+W+'" height="'+H+'" style="display:block;width:100%;height:auto">';
 for(var v=0; v<=top+1e-9; v+=step){var y=Y(v); s+='<line x1="'+m.l+'" x2="'+(W-m.r)+'" y1="'+y+'" y2="'+y+'" stroke="#1E293B"/><text x="'+(m.l-7)+'" y="'+(y+4)+'" text-anchor="end" font-size="10.5" fill="#64748B">'+fmt(v)+'</text>';}
 var xs=D>20?3:(D>10?2:1);
 for(var d=1; d<=D; d++){ if(d===1||d%xs===0||d===D){ s+='<text x="'+X(d)+'" y="'+(H-8)+'" text-anchor="middle" font-size="10.5" fill="#64748B">'+d+'</text>'; } }
 if(CFG.target>0){ s+='<line x1="'+m.l+'" x2="'+(W-m.r)+'" y1="'+Y(CFG.target)+'" y2="'+Y(CFG.target)+'" stroke="#475569" stroke-dasharray="2 4"/><text x="'+(W-m.r)+'" y="'+(Y(CFG.target)-5)+'" text-anchor="end" font-size="10.5" fill="#94A3B8">Target '+fmt(CFG.target)+'</text>'; }
 var p=''; CFG.ideal.forEach(function(v,i){p+=(i?'L':'M')+X(i+1)+','+Y(v);});
 s+='<path d="'+p+'" fill="none" stroke="#94A3B8" stroke-width="2" stroke-dasharray="6 4"/>';
 if(CFG.proj.length>1){ p=''; CFG.proj.forEach(function(q,i){p+=(i?'L':'M')+X(q[0])+','+Y(q[1]);}); s+='<path d="'+p+'" fill="none" stroke="#00E5FF" stroke-width="2.2" stroke-dasharray="2 4" stroke-linecap="round"/>'; }
 (CFG.ext||[]).forEach(function(e){ p=''; e.p.forEach(function(q,i){p+=(i?'L':'M')+X(q[0])+','+Y(q[1]);}); s+='<path d="'+p+'" fill="none" stroke="'+e.c+'" stroke-width="2.2" stroke-dasharray="'+e.d+'" stroke-linecap="round"/>'; });
 if(CFG.real.length){
  p=''; CFG.real.forEach(function(q,i){p+=(i?'L':'M')+X(q[0])+','+Y(q[1]);});
  var last=CFG.real[CFG.real.length-1], first=CFG.real[0];
  s+='<path d="'+p+'L'+X(last[0])+','+Y(0)+'L'+X(first[0])+','+Y(0)+'Z" fill="rgba(0,255,102,0.10)"/>';
  s+='<path d="'+p+'" fill="none" stroke="#00FF66" stroke-width="2.8" stroke-linejoin="round"/>';
  s+='<circle cx="'+X(last[0])+'" cy="'+Y(last[1])+'" r="5" fill="#00FF66" stroke="#0E1117" stroke-width="2"/>';
 }
 s+='<line x1="'+X(CFG.today)+'" x2="'+X(CFG.today)+'" y1="'+m.t+'" y2="'+(m.t+ih)+'" stroke="#FFEA00" stroke-dasharray="3 3" stroke-opacity=".7"/><text x="'+X(CFG.today)+'" y="'+(m.t+9)+'" text-anchor="middle" font-size="10" fill="#FFEA00" font-weight="700" dx="'+(CFG.today>D-3?-16:16)+'">Tgl '+CFG.today+'</text>';
 s+='<line id="hl" y1="'+m.t+'" y2="'+(m.t+ih)+'" stroke="#E2E8F0" stroke-opacity=".35" style="display:none"/><circle id="h1" r="4.5" fill="#94A3B8" style="display:none"/><circle id="h2" r="5" fill="#00FF66" stroke="#0E1117" stroke-width="2" style="display:none"/>';
 s+='<rect id="ov" x="'+m.l+'" y="'+m.t+'" width="'+iw+'" height="'+ih+'" fill="transparent"/></svg>';
 wrap.innerHTML=s;
 var ov=document.getElementById('ov'), hl=document.getElementById('hl'), h1=document.getElementById('h1'), h2=document.getElementById('h2');
 var rmap={}; CFG.real.forEach(function(q){rmap[q[0]]=q[1];});
 ov.addEventListener('mousemove',function(ev){
  var tip=document.getElementById('tip'), r=ov.getBoundingClientRect(), sc=(r.width/iw)||1, d=Math.round((ev.clientX-r.left)/sc/iw*(D-1))+1; d=Math.max(1,Math.min(D,d));
  var idv=CFG.ideal[d-1], rv=rmap[d], x=X(d);
  hl.setAttribute('x1',x); hl.setAttribute('x2',x); hl.style.display='';
  h1.setAttribute('cx',x); h1.setAttribute('cy',Y(idv)); h1.style.display='';
  var t='<b>Tgl '+d+'</b><br>Target ideal: <b>'+fmt(idv)+'</b>';
  if(rv!==undefined){ var g=rv-idv; h2.setAttribute('cx',x); h2.setAttribute('cy',Y(rv)); h2.style.display=''; t+='<br>Realisasi STU: <b style="color:#00FF66">'+fmt(rv)+'</b><br>Gap: <b style="color:'+(g>=0?'#00FF66':'#FF4D8D')+'">'+(g>=0?'+':'')+fmt(g)+'</b>'; } else { h2.style.display='none'; }
  if(d>CFG.today){ var lt=CFG.real.length?CFG.real[CFG.real.length-1]:[CFG.today,0], fr=(d-CFG.today)/Math.max(1,D-CFG.today);
   if(CFG.proj.length>1){ t+='<br><span style="color:#00E5FF">Forecast closing: <b>'+fmt(lt[1]+(CFG.proj[1][1]-lt[1])*fr)+'</b></span>'; }
   (CFG.ext||[]).forEach(function(e){ t+='<br><span style="color:'+e.c+'">'+e.n+': <b>'+fmt(lt[1]+(e.v-lt[1])*fr)+'</b></span>'; }); }
  tip.innerHTML=t; tip.style.opacity=1; var tx=ev.clientX+14, ty=ev.clientY+14; if(tx+tip.offsetWidth>window.innerWidth-6) tx=ev.clientX-tip.offsetWidth-14; tip.style.left=tx+'px'; tip.style.top=ty+'px';
 });
 ov.addEventListener('mouseleave',function(){var tip=document.getElementById('tip');hl.style.display='none';h1.style.display='none';h2.style.display='none';if(tip)tip.style.opacity=0;});
}
draw(); var rt; window.addEventListener('resize',function(){clearTimeout(rt); rt=setTimeout(draw,80);});
try{ if(window.ResizeObserver){ new ResizeObserver(function(){ if(String(wrap.clientWidth)!==wrap.getAttribute('data-w')){clearTimeout(rt); rt=setTimeout(draw,60);} }).observe(wrap); } }catch(e){}
})();</script>"""

def bsg_forc_lines(f_nota=None, f7=None, f3=None):
    """Garis forecast tambahan untuk kurva STU: (label, nilai akhir bulan, warna, pola garis)."""
    return [('Forc Nota', f_nota, '#FFB020', '6 4'), ('Forc 7', f7, '#B388FF', '1 5'), ('Forc 3', f3, '#FF4D8D', '8 3 2 3')]

def bsg_chart_curve(title, subtitle, df_stu, target, tgl_akhir, max_day, forc=None, h=280, extra=None, parts=False):
    """Kurva kumulatif STU: target ideal harian (dari tabel_target_harian) vs realisasi + proyeksi (forecast closing)."""
    target = bsg_num(target)
    if target <= 0: return None
    days = list(range(1, int(max_day) + 1))
    ideal = [round(target * tabel_target_harian.get(d, 0.0) / 100.0, 2) for d in days]
    cnt = {}
    if df_stu is not None and not df_stu.empty and 'hari' in df_stu.columns:
        for hh, cc in df_stu['hari'].dropna().astype(int).value_counts().items(): cnt[int(hh)] = int(cc)
    real, run = [], 0
    for d in range(1, int(tgl_akhir) + 1):
        run += cnt.get(d, 0); real.append([d, run])
    proj = [[int(tgl_akhir), run], [int(max_day), float(forc)]] if (forc is not None and bsg_num(forc) > 0 and int(tgl_akhir) < int(max_day)) else []
    ext = []
    if int(tgl_akhir) < int(max_day):
        for lb_, v_, c_, d_ in (extra or []):
            if v_ is not None and bsg_num(v_) > 0: ext.append({'n': lb_, 'v': float(bsg_num(v_)), 'c': c_, 'd': d_, 'p': [[int(tgl_akhir), run], [int(max_day), float(bsg_num(v_))]]})
    gap = run - ideal[int(tgl_akhir) - 1]
    gcol = BSG_GRN if gap >= 0 else BSG_RED
    chips = (f'<span class="chip">Realisasi <b style="color:{BSG_GRN}">{run:,}</b></span><span class="chip">Ideal tgl {int(tgl_akhir)} <b>{ideal[int(tgl_akhir) - 1]:,.0f}</b></span>'
             f'<span class="chip">Gap <b style="color:{gcol}">{gap:+,.0f}</b></span><span class="chip">Target <b>{int(target):,}</b></span>'
             + (f'<span class="chip">Forecast closing <b style="color:{BSG_CYN}">{int(bsg_num(forc)):,}</b></span>' if proj else ''))
    leg = (f'<span class="lg"><span class="ds"></span>Target ideal</span><span class="lg"><span class="dot" style="background:{BSG_GRN}"></span>Realisasi STU</span>'
           + (f'<span class="lg"><span class="dot" style="background:{BSG_CYN}"></span>Forecast {int(bsg_num(forc)):,}</span>' if proj else '')
           + ''.join(f'<span class="lg"><span class="ds" style="border-top-color:{e["c"]}"></span><span style="color:{e["c"]};font-weight:700">{_esc(e["n"])} {int(e["v"]):,}</span></span>' for e in ext))
    cfg = json.dumps({'maxd': int(max_day), 'ideal': ideal, 'real': real, 'proj': proj, 'ext': ext, 'today': int(tgl_akhir), 'target': target, 'h': int(h)})
    inner = f'<div class="card">{_hdr(title, subtitle)}<div>{chips}</div><div>{leg}</div><div id="cv" style="width:100%;height:{int(h)}px"></div></div>'
    tot_h = 104 + (14 if subtitle else 0) + int(h) + 6 + (22 if ext else 0)
    if parts: return inner + _CURVE_JS.replace('__CFG__', cfg), '', tot_h   # varian untuk panel (tabel + grafik)
    return _bsg_frame(inner, js=_CURVE_JS.replace('__CFG__', cfg)), tot_h

# ---------------------------------------------------------------- 6. podium / leaderboard & insight otomatis ----
def bsg_html_podium(rows, ideal):
    """rows: list dict(name, stu, target, pct) terurut dari terbaik."""
    if not rows: return ''
    medals = ['🥇', '🥈', '🥉']
    def _row(i, r, warn=False):
        col = bsg_stat_color(r['pct'], ideal, 0.85 * ideal) if ideal else BSG_CYN
        tag = '⚠️' if warn else (medals[i] if i < 3 else f'#{i + 1}')
        bar = bsg_prog_bar(min(r['pct'], 100.0), col, ideal if ideal else None)
        return (f'<div style="display:flex;align-items:center;gap:10px;background:linear-gradient(160deg,#151B2B 0%,#0E1117 100%);border:1px solid #232B3E;border-left:4px solid {col};border-radius:10px;padding:7px 12px;">'
                f'<div style="font-size:20px;width:26px;text-align:center;">{tag}</div>'
                f'<div style="flex:1;min-width:0;"><div style="font-size:12.5px;font-weight:800;color:#E2E8F0;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;" title="{_esc(r["name"])}">{_esc(r["name"])}</div>'
                f'<div style="font-size:11px;color:#94A3B8;">STU {int(r["stu"]):,} / Target {int(r["target"]):,}</div>{bar}</div>'
                f'<div style="font-size:20px;font-weight:900;color:{col};min-width:62px;text-align:right;">{r["pct"]:.1f}%</div></div>')
    top = rows[:3]; top_names = {r['name'] for r in top}
    bottom = [r for r in rows[::-1][:3] if r['name'] not in top_names]
    out = '<div style="font-size:13px;font-weight:800;color:#E2E8F0;margin:2px 0 6px 0;">🏅 Leaderboard Pencapaian</div>'
    out += '<div style="display:grid;gap:7px;">' + ''.join(_row(i, r) for i, r in enumerate(top)) + '</div>'
    if bottom:
        out += '<div style="font-size:13px;font-weight:800;color:#E2E8F0;margin:12px 0 6px 0;">🚨 Perlu Perhatian</div>'
        out += '<div style="display:grid;gap:7px;">' + ''.join(_row(i, r, True) for i, r in enumerate(bottom)) + '</div>'
    return out

def bsg_build_insights(ctx, rows):
    """rows: list dict(name, stu, target, ideal_u, pct) terurut dari terbaik. Mengembalikan list (warna, ikon, html)."""
    out = []
    r_p, tgl, ideal = ctx['r_tot_p'], ctx['tanggal_akhir'], float(ctx['ideal_pct'])
    stu, ideal_u, targ = bsg_num(r_p.get('STU')), bsg_num(r_p.get('IDEAL')), bsg_num(r_p.get('TARGET'))
    if targ > 0:
        gap = stu - ideal_u; col = BSG_GRN if gap >= 0 else (BSG_YEL if stu >= 0.85 * ideal_u else BSG_RED)
        out.append((col, '🎯', f"<b>Pace tgl {tgl}:</b> STU {int(stu):,} vs ideal {int(ideal_u):,} → <b style='color:{col}'>{gap:+,.0f} unit</b> ({stu / targ * 100:.1f}% dari target {int(targ):,})."))
        sisa, ha = int(ctx.get('sisa_hari', 0) or 0), max(1, int(ctx.get('hari_aktif', 1) or 1)); butuh = max(0.0, targ - stu)
        if butuh > 0 and sisa > 0: out.append((BSG_CYN, '⏱️', f"Untuk menutup target butuh <b>{butuh / sisa:.1f} STU/hari kerja</b> selama {sisa} hari kerja tersisa (rata-rata saat ini {stu / ha:.1f}/hari)."))
        elif butuh > 0: out.append((BSG_RED, '⏱️', f"Hari kerja habis, target masih kurang <b>{int(butuh):,} unit</b>."))
        else: out.append((BSG_GRN, '🏆', f"Target bulanan <b>sudah terlampaui</b> (+{int(stu - targ):,} unit)."))
    if rows:
        below = [r for r in rows if r['pct'] < ideal]
        out.append((BSG_YEL if below else BSG_GRN, '📉' if below else '✅', f"<b>{len(below)} dari {len(rows)}</b> {_esc(ctx.get('kol_label', 'entitas'))} berada di bawah pace ideal ({ideal:.0f}%)."))
        lag = rows[::-1][:3]
        out.append((BSG_RED, '🚨', '<b>Paling tertinggal:</b> ' + '; '.join(f"{_esc(r['name'])} ({r['pct']:.1f}%, {r['stu'] - r['ideal_u']:+,.0f} vs ideal)" for r in lag)))
        b = rows[0]
        out.append((BSG_GRN, '🌟', f"<b>Terbaik:</b> {_esc(b['name'])} — {b['pct']:.1f}% ({int(b['stu']):,}/{int(b['target']):,})."))
    for lbl, key, ideal_q, flt in ctx.get('quality_checks', []):
        try:
            cand = [(n, v) for n, v, ok in key if ok]
            if cand:
                n, v = min(cand, key=lambda t: t[1])
                out.append((bsg_stat_color(v, ideal_q, 0.7 * ideal_q), '🔎', f"<b>{lbl} terendah:</b> {_esc(n)} ({v:.1f}% vs ideal {ideal_q:.0f}%)."))
        except Exception: pass
    return out

def bsg_html_insights(items):
    if not items: return ''
    cells = ''.join(f'<div style="border-left:4px solid {c};background:linear-gradient(90deg,rgba(255,255,255,0.04) 0%,rgba(14,17,23,0.95) 100%);border-top:1px solid #2D3748;border-right:1px solid #2D3748;border-bottom:1px solid #2D3748;border-radius:6px;padding:8px 12px;font-size:13px;line-height:1.55;color:#E2E8F0;">{ic} {tx}</div>' for c, ic, tx in items)
    return '<div style="font-size:13px;font-weight:800;color:#E2E8F0;margin:2px 0 6px 0;">💡 Insight Otomatis</div><div style="display:grid;grid-template-columns:1fr;gap:7px;">' + cells + '</div>'

def bsg_celebrate(key, achieved, msg='🎉 Target tercapai!'):
    """Efek balon + toast sekali saja per kunci (scope + bulan + tanggal)."""
    flag = f'_celeb_{key}'
    if achieved and not st.session_state.get(flag):
        st.session_state[flag] = True
        try: st.balloons()
        except Exception: pass
        bsg_toast(msg, '🏆')

# ---------------------------------------------------------------- drill-down (fragment: tidak me-rerun seluruh halaman) ----
@_bsg_fragment
def bsg_drilldown(ctx, names):
    kol, ideal = ctx['kol_nm'], float(ctx['ideal_pct'])
    df_p, df_pr, df_s = ctx['df_p_disp'], ctx['df_pr_disp'], ctx['df_s_disp']
    sk = re.sub(r'[^A-Za-z0-9]', '_', str(ctx['scope_key']))
    P = {str(r[kol]): r for _, r in df_p.iterrows()}
    PR = {str(r[kol]): r for _, r in df_pr.iterrows()} if (df_pr is not None and not df_pr.empty) else {}
    SP = {str(r[kol]): r for _, r in df_s.iterrows()} if (df_s is not None and not df_s.empty) else {}
    names = [n for n in names if n in P]
    if not names: st.caption('Tidak ada baris untuk ditampilkan pada filter saat ini.'); return
    recs = []
    for n in names:
        r, pr, sp = P[n], PR.get(n), SP.get(n)
        recs.append({kol: n, 'TARGET': int(bsg_num(r.get('TARGET'))), 'STU': int(bsg_num(r.get('STU'))), 'IDEAL': int(bsg_num(r.get('IDEAL'))),
                     'GAP vs IDEAL': int(bsg_num(r.get('STU')) - bsg_num(r.get('IDEAL'))), 'Pencapaian': round(bsg_num(r.get('_pct')), 2),
                     'DB Q (%)': round(bsg_num(pr.get('DB Q')) if pr is not None else 0.0, 2), 'STU% (%)': round(bsg_num(pr.get('STU%')) if pr is not None else 0.0, 2),
                     'ACC% (%)': round(bsg_num(sp.get('_acc_val')) if sp is not None else 0.0, 2)})
    view = pd.DataFrame(recs)
    cfg = {}
    try:
        cfg = {'Pencapaian': st.column_config.ProgressColumn('Pencapaian', format='%.1f%%', min_value=0.0, max_value=float(max(100.0, view['Pencapaian'].max()))),
               'DB Q (%)': st.column_config.NumberColumn('DB Q', format='%.1f%%'), 'STU% (%)': st.column_config.NumberColumn('STU%', format='%.1f%%'),
               'ACC% (%)': st.column_config.NumberColumn('ACC%', format='%.1f%%')}
    except Exception: cfg = {}
    sel_name = None
    try:
        ev = st.dataframe(view, hide_index=True, use_container_width=True, column_config=cfg, on_select='rerun', selection_mode='single-row', key=f'bsg_dd_{sk}', height=min(max(120, calc_table_height(view)), 420))
        rows_sel = list(ev.selection.rows) if hasattr(ev, 'selection') else []
        if rows_sel and rows_sel[0] < len(view): sel_name = str(view.iloc[rows_sel[0]][kol])
        else: st.caption('👆 Klik satu baris untuk melihat kartu detail & kurva harian-nya.')
    except Exception:
        st.dataframe(view, hide_index=True, use_container_width=True, column_config=cfg)
        sel_name = st.selectbox('Pilih untuk detail:', ['—'] + names, key=f'bsg_dd_sel_{sk}')
        sel_name = None if sel_name == '—' else sel_name
    if not sel_name: return
    r, pr, sp = P[sel_name], PR.get(sel_name), SP.get(sel_name)
    stu, targ, idl, pct = bsg_num(r.get('STU')), bsg_num(r.get('TARGET')), bsg_num(r.get('IDEAL')), bsg_num(r.get('_pct'))
    dbq, stup = (bsg_num(pr.get('DB Q')) if pr is not None else 0.0), (bsg_num(pr.get('STU%')) if pr is not None else 0.0)
    acc = bsg_num(sp.get('_acc_val')) if sp is not None else 0.0
    st.markdown(f'<div class="section-title" style="font-size:18px !important;">🔎 Detail — {_esc(sel_name)}</div>', unsafe_allow_html=True)
    c1, c2 = st.columns([2, 3])
    with c1:
        cards = [bsg_card('🎯', 'Target vs Result', f'{int(targ):,} / {int(stu):,}', f'<b>{pct:.1f}%</b> tercapai · Ideal {int(idl):,}', bsg_stat_color(stu, idl, 0.85 * idl) if idl > 0 else BSG_CYN, min(pct, 100.0), ideal),
                 bsg_card('✨', 'DB Quality', f'{dbq:.1f}%', 'Ideal 30%', bsg_stat_color(dbq, 30, 20), *bsg_ach_bar(dbq, 30)),
                 bsg_card('📈', 'STU Ratio', f'{stup:.1f}%', 'Ideal 10%', bsg_stat_color(stup, 10, 7), *bsg_ach_bar(stup, 10)),
                 bsg_card('✅', 'Approval', f'{acc:.1f}%', 'Ideal 50%', bsg_stat_color(acc, 50, 40), *bsg_ach_bar(acc, 50))]
        st.markdown('<div style="display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:10px;">' + ''.join(cards) + '</div>', unsafe_allow_html=True)
    with c2:
        f_c = ctx.get('f_c'); col = 'Nama Dealer' if ctx['is_total'] else 'Koor'
        sub = f_c[f_c[col] == sel_name] if (f_c is not None and not f_c.empty and col in f_c.columns) else None
        bsg_show(bsg_chart_curve(f'Kurva STU — {sel_name}', 'Kumulatif harian vs target ideal', sub, targ, ctx['tanggal_akhir'], ctx['max_day'], forc=bsg_num(r.get('FORC')) or None, extra=bsg_forc_lines(r.get('F. NOTA'), r.get('FORC 7'), r.get('FORC 3')), h=230))

# ---------------------------------------------------------------- halaman "Ringkasan Visual" tab STU ----
def bsg_render_ringkasan(ctx):
    kol, ideal, tgl = ctx['kol_nm'], float(ctx['ideal_pct']), ctx['tanggal_akhir']
    df_p, df_pr, df_s = ctx['df_p_disp'], ctx['df_pr_disp'], ctx['df_s_disp']
    r_p, r_pr, r_s = ctx['r_tot_p'], ctx['r_tot_pr'], ctx['r_tot_s']
    sk = re.sub(r'[^A-Za-z0-9]', '_', str(ctx['scope_key']))
    if df_p is None or df_p.empty:
        st.info('ℹ️ Belum ada data penjualan untuk divisualisasikan pada periode ini.'); return
    P = {str(r[kol]): r for _, r in df_p.iterrows()}
    PR = {str(r[kol]): r for _, r in df_pr.iterrows()} if (df_pr is not None and not df_pr.empty) else {}
    SP = {str(r[kol]): r for _, r in df_s.iterrows()} if (df_s is not None and not df_s.empty) else {}
    names_all = list(P.keys())

    # ---- filter interaktif ----
    c_f1, c_f2 = st.columns([2, 3])
    with c_f1:
        st.caption('Tampilkan baris:')
        mode = bsg_segmented('Tampilkan', ['Semua', 'Di bawah ideal', 'Di atas ideal'], f'bsg_mode_{sk}')
    with c_f2:
        pick = st.multiselect('Filter baris (kosongkan = semua)', names_all, key=f'bsg_pick_{sk}')
    keep = []
    for n in names_all:
        if pick and n not in pick: continue
        pc = bsg_num(P[n].get('_pct'))
        if mode == 'Di bawah ideal' and not pc < ideal: continue
        if mode == 'Di atas ideal' and not pc >= ideal: continue
        keep.append(n)

    rows_all = []   # untuk leaderboard & insight (hanya yang punya target), terurut dari terbaik
    for n in names_all:
        r = P[n]; targ = bsg_num(r.get('TARGET'))
        if targ > 0: rows_all.append({'name': n, 'stu': bsg_num(r.get('STU')), 'target': targ, 'ideal_u': bsg_num(r.get('IDEAL')), 'pct': bsg_num(r.get('_pct'))})
    rows_all.sort(key=lambda x: x['pct'], reverse=True)

    # ---- celebrate bila target total tercapai ----
    try: bsg_celebrate(f"{sk}_{ctx['period_key']}", bsg_num(r_p.get('TARGET')) > 0 and bsg_num(r_p.get('STU')) >= bsg_num(r_p.get('TARGET')), '🎉 Target STU bulan ini tercapai!')
    except Exception: pass

    # ---- insight + leaderboard ----
    qchecks = []
    try:
        qchecks = [('DB Quality', [(n, bsg_num(PR[n].get('DB Q')), bsg_num(PR[n].get('DB')) >= 10) for n in names_all if n in PR], 30.0, None),
                   ('STU Ratio', [(n, bsg_num(PR[n].get('_stu_pct')), bsg_num(PR[n].get('DB')) >= 10) for n in names_all if n in PR], 10.0, None),
                   ('Approval SPK', [(n, bsg_num(SP[n].get('_acc_val')), bsg_num(SP[n].get('_spk_val')) >= 3) for n in names_all if n in SP], 50.0, None)]
    except Exception: qchecks = []
    cctx = dict(ctx); cctx['quality_checks'] = qchecks
    cctx['kol_label'] = 'dealer' if ctx['is_total'] else 'koordinator/counter'
    c_i1, c_i2 = st.columns([3, 2])
    with c_i1: st.markdown(bsg_html_insights(bsg_build_insights(cctx, rows_all)), unsafe_allow_html=True)
    with c_i2: st.markdown(bsg_html_podium(rows_all, ideal), unsafe_allow_html=True)

    # ---- kurva target vs realisasi ----
    st.markdown('<div style="height:6px"></div>', unsafe_allow_html=True)
    bsg_show(bsg_safe(bsg_chart_curve, f'Kurva Pencapaian STU — {ctx["dlr_label"]}', 'Arahkan kursor ke grafik untuk membaca nilai per tanggal. Garis putus-putus abu = target ideal harian.',
                      ctx.get('f_c'), bsg_num(r_p.get('TARGET')), tgl, ctx['max_day'], forc=bsg_num(r_p.get('FORC')) or None, extra=bsg_forc_lines(ctx.get('forc_nota'), ctx.get('forc7'), ctx.get('forc3')), label='Kurva pencapaian'))

    # ---- ranking + funnel ----
    c_r, c_f = st.columns([3, 2])
    with c_r:
        rank_rows = []
        for n in keep:
            r = P[n]; targ = bsg_num(r.get('TARGET'))
            if targ <= 0: continue
            pct, stu, idu = bsg_num(r.get('_pct')), bsg_num(r.get('STU')), bsg_num(r.get('IDEAL'))
            rank_rows.append({'name': n, 'val': pct, 'tip': f"<b>{_esc(n)}</b><br>STU <b>{int(stu):,}</b> / Target {int(targ):,}<br>Ideal tgl {tgl}: {int(idu):,} → <b>{stu - idu:+,.0f}</b><br>Pencapaian <b>{pct:.1f}%</b>"})
        res = bsg_safe(bsg_chart_rank, 'Ranking Pencapaian Target', f'% STU terhadap target bulanan · garis putih = ideal tgl {tgl} ({ideal:.0f}%)', rank_rows, ideal=ideal, label='Ranking')
        if res: bsg_show(res)
        else: st.caption('Belum ada baris bertarget pada filter ini.')
    with c_f:
        regs = 0; f_c = ctx.get('f_c')
        if f_c is not None and not f_c.empty and 'Status_Registrasi_Rating' in f_c.columns: regs = int((f_c['Status_Registrasi_Rating'] == 'SUDAH REGISTRASI').sum())
        stages = [('Database (DB)', bsg_num(r_pr.get('DB')), '#4318FF'), ('Hot Prospect', bsg_num(r_pr.get('HOT')), '#6D4DFF'), ('STU', bsg_num(r_p.get('STU')), '#00B8D9')]
        if regs > 0: stages.append(('Sudah Registrasi', float(regs), '#00E5FF'))
        bsg_show(bsg_safe(bsg_chart_funnel, 'Funnel Konversi Prospek', 'DB → Hot → STU (→ registrasi rating)', stages, label='Funnel'))
        s_c = ctx.get('s_c')
        if s_c is not None and not s_c.empty and 'Status_SPK_Clean' in s_c.columns:
            vc = s_c['Status_SPK_Clean'].value_counts()
            bsg_show(bsg_safe(bsg_chart_donut, 'Status SPK Leasing', 'Disetujui vs ditolak vs terbuka',
                              [('Disetujui', int(vc.get('DISETUJUI', 0)), '#00E676'), ('Ditolak', int(vc.get('DITOLAK', 0)), '#FF4D8D'), ('Terbuka', int(vc.get('TERBUKA', 0)), '#FFB020')], 'SPK', label='Donut SPK'))

    # ---- heatmap kualitas ----
    cols = ['Target', 'DB Quality', 'STU Ratio', 'Approval']
    cells, hr = {}, []
    for n in keep:
        r = P[n]; pr, sp = PR.get(n), SP.get(n); hr.append(n)
        targ = bsg_num(r.get('TARGET')); pc = bsg_num(r.get('_pct'))
        if targ > 0: cells[(n, 'Target')] = (pc / ideal * 100.0 if ideal > 0 else 0.0, f'{pc:.0f}%', f'<b>{_esc(n)}</b><br>Pencapaian {pc:.1f}% (ideal {ideal:.0f}%)')
        if pr is not None and bsg_num(pr.get('DB')) > 0:
            v = bsg_num(pr.get('DB Q')); cells[(n, 'DB Quality')] = (v / 30.0 * 100.0, f'{v:.0f}%', f'<b>{_esc(n)}</b><br>DB Quality {v:.1f}% (ideal 30%)')
            v = bsg_num(pr.get('_stu_pct')); cells[(n, 'STU Ratio')] = (v / 10.0 * 100.0, f'{v:.1f}%', f'<b>{_esc(n)}</b><br>STU Ratio {v:.1f}% (ideal 10%)')
        if sp is not None and bsg_num(sp.get('_spk_val')) > 0:
            v = bsg_num(sp.get('_acc_val')); cells[(n, 'Approval')] = (v / 50.0 * 100.0, f'{v:.0f}%', f'<b>{_esc(n)}</b><br>Approval {v:.1f}% (ideal 50%)')
    bsg_show(bsg_safe(bsg_chart_heat, 'Peta Panas Kualitas — capaian vs acuan ideal', 'Merah = jauh di bawah ideal · hijau = memenuhi/melampaui ideal (Target vs pace, DB Q 30%, STU% 10%, ACC% 50%)', hr, cols, cells, label='Heatmap'))

    # ---- drill-down ----
    st.markdown('<div class="section-title" style="font-size:18px !important;">🔎 Drill-down Interaktif</div>', unsafe_allow_html=True)
    bsg_safe(bsg_drilldown, ctx, keep, label='Drill-down')

# ---------------------------------------------------------------- grafik tambahan per tab ----
def bsg_render_leasing_charts(df_s, r_tot_s, kol_nm):
    """Tab STU › SPK & Leasing: ranking ACC% per baris + porsi SPK per leasing."""
    rows = []
    if df_s is not None and not df_s.empty:
        for _, r in df_s.iterrows():
            if bsg_num(r.get('_spk_val')) <= 0: continue
            acc = bsg_num(r.get('_acc_val'))
            rows.append({'name': str(r[kol_nm]), 'val': acc, 'tip': f"<b>{_esc(r[kol_nm])}</b><br>ACC% <b>{acc:.1f}%</b><br>SPK {int(bsg_num(r.get('_spk_val'))):,} (LM {int(bsg_num(r.get('_spk_lm_val'))):,}) · Terbuka {int(bsg_num(r.get('TERBUKA'))):,}"})
    c1, c2 = st.columns([3, 2])
    with c1:
        res = bsg_chart_rank('Approval Rate (ACC%) per baris', 'Garis putih = acuan ideal ACC% 50%', rows, ideal=50.0)
        if res: bsg_show(res)
        else: st.caption('Belum ada SPK pada periode ini.')
    with c2:
        items = [(k, bsg_num(r_tot_s.get(k)), c) for k, c in [('BAF', '#4318FF'), ('ADIRA', '#00E5FF'), ('SOF', '#FFB020'), ('IMFI', '#FF4D8D')]]
        bsg_show(bsg_chart_donut('Porsi SPK per Leasing', 'Total periode berjalan', items, 'SPK'))

def bsg_render_type_charts(df_td_i, ideal):
    """Tab STU › Per Type: komposisi STU per type + capaian terhadap target type."""
    if df_td_i is None or df_td_i.empty or 'TYPE' not in df_td_i.columns: return
    pal = {'AT PREMIUM': '#4318FF', 'AT CLASSY': '#00E5FF', 'AT LPM': '#FFB020', 'MOPED': '#FF4D8D', 'OTHERS': '#94A3B8'}
    items = [(str(r['TYPE']), bsg_num(r.get('STU')), pal.get(str(r['TYPE']), '#64748B')) for _, r in df_td_i.iterrows()]
    rows = []
    for _, r in df_td_i.iterrows():
        tg = bsg_num(r.get('TARGET'))
        if tg <= 0: continue
        ach = bsg_num(r.get('ACHIEV'))
        rows.append({'name': str(r['TYPE']), 'val': ach, 'tip': f"<b>{_esc(r['TYPE'])}</b><br>STU <b>{int(bsg_num(r.get('STU'))):,}</b> / Target {int(tg):,}<br>Capaian <b>{ach:.1f}%</b> · LM {int(bsg_num(r.get('LM'))):,}"})
    c1, c2 = st.columns([2, 3])
    with c1: bsg_show(bsg_chart_donut('Komposisi STU per Type', 'Porsi unit terjual periode berjalan', items, 'STU'))
    with c2:
        res = bsg_chart_rank('Capaian per Type vs Target Type', f'Garis putih = ideal tgl {int(bsg_num(ideal))} ({ideal:.0f}%)', rows, ideal=float(ideal))
        if res: bsg_show(res)

def bsg_render_salesman_charts(df_smi, ideal, n=10):
    """Tab STU › Salesman: top-N dan N salesman perlu dorongan (hanya yang bertarget)."""
    if df_smi is None or df_smi.empty or '_pct_val' not in df_smi.columns: return
    rows = []
    for _, r in df_smi.iterrows():
        tg = bsg_num(r.get('TARGET'))
        if tg <= 0: continue
        pc = bsg_num(r.get('_pct_val'))
        rows.append({'name': str(r['NAMA SALESMAN']), 'val': pc, 'tip': f"<b>{_esc(r['NAMA SALESMAN'])}</b><br>STU <b>{int(bsg_num(r.get('STU'))):,}</b> / Target {int(tg):,} ({pc:.1f}%)<br>DB {int(bsg_num(r.get('DB'))):,} · DB Q {_esc(r.get('DB Q', '-'))} · STU% {_esc(r.get('STU%', '-'))}"})
    if not rows: return
    c1, c2 = st.columns(2)
    with c1: bsg_show(bsg_chart_rank(f'Top {min(n, len(rows))} Salesman', f'% STU terhadap target · ideal {ideal:.0f}%', rows, ideal=float(ideal), max_rows=n))
    with c2: bsg_show(bsg_chart_rank(f'{min(n, len(rows))} Salesman Perlu Dorongan', 'Capaian terendah (hanya yang punya target)', rows, ideal=float(ideal), max_rows=n, ascending=True))

def bsg_render_reg_donut(sudah, belum):
    c1, c2 = st.columns([2, 3])
    with c1: bsg_show(bsg_chart_donut('Status Registrasi Rating (MYDR)', 'Konsumen faktur periode berjalan', [('Sudah Registrasi', sudah, '#00E676'), ('Belum Registrasi', belum, '#FF4D8D')], 'Konsumen'))

def bsg_render_spk_online_charts(df_sit, df_bwi, df_lsg_full):
    """Dashboard Leasing (tab SPK & Leasing): ranking ACC%, growth SPK vs bulan lalu, dan porsi leasing."""
    d = pd.concat([x for x in (df_sit, df_bwi) if x is not None and not x.empty], ignore_index=True) if (df_sit is not None or df_bwi is not None) else pd.DataFrame()
    if d.empty: return
    acc_rows, grw_rows = [], []
    for _, r in d.iterrows():
        nm = str(r['NAMA DEALER']); spk, lm = bsg_num(r.get('_spk_val')), bsg_num(r.get('_spk_lm_val'))
        if spk > 0: acc_rows.append({'name': nm, 'val': bsg_num(r.get('_acc_val')), 'tip': f"<b>{_esc(nm)}</b><br>ACC% <b>{bsg_num(r.get('_acc_val')):.1f}%</b><br>SPK {int(spk):,} · Terbuka {int(bsg_num(r.get('TERBUKA'))):,}"})
        if lm > 0: grw_rows.append({'name': nm, 'val': spk / lm * 100.0, 'tip': f"<b>{_esc(nm)}</b><br>SPK {int(spk):,} vs LM {int(lm):,}<br>Growth <b>{spk / lm * 100.0:.1f}%</b>"})
    st.markdown('<div class="section-title">📊 Visual Ringkas Leasing</div>', unsafe_allow_html=True)
    c1, c2, c3 = st.columns(3)
    with c1:
        res = bsg_chart_rank('Approval Rate (ACC%) per Dealer', 'Putih = ideal 50%', acc_rows, ideal=50.0)
        if res: bsg_show(res)
    with c2:
        res = bsg_chart_rank('Growth SPK vs Bulan Lalu', 'Putih = sama dengan bulan lalu (100%)', grw_rows, ideal=100.0)
        if res: bsg_show(res)
    with c3:
        tot = df_lsg_full.iloc[-1] if (df_lsg_full is not None and not df_lsg_full.empty) else {}
        items = [(k, bsg_num(tot.get(k)), c) for k, c in [('BAF', '#4318FF'), ('ADIRA', '#00E5FF'), ('SOF', '#FFB020'), ('IMFI', '#FF4D8D')]]
        bsg_show(bsg_chart_donut('Porsi SPK per Leasing', 'Total seluruh dealer', items, 'SPK'))

def bsg_render_master_kpi_cards(tot_d, ideal):
    """Kartu ringkas Master KPI dengan progress bar; warna status mengacu pada pace ideal tanggal berjalan."""
    g = lambda k: bsg_num(tot_d.get(k))
    pc = lambda r, t: (r / t * 100.0) if t > 0 else 0.0
    idl = float(ideal) if ideal else 100.0
    p_stu, p_max, p_mp, p_acc = pc(g('R. STU'), g('T. STU')), pc(g('R. MAXI'), g('T. MAXI')), pc(g('R. MP'), g('T. MP')), pc(g('R. ACC%'), g('T. ACC%'))
    pace = lambda p: bsg_stat_color(p, idl, 0.85 * idl)
    k1, k2, k3, k4 = st.columns(4)
    k1.markdown(bsg_card('🎯', 'KPI STU Dealer', f"{int(g('R. STU')):,} / {int(g('T. STU')):,}", f'<b>{p_stu:.1f}%</b> Tercapai · ideal {idl:.0f}%', pace(p_stu), min(p_stu, 100.0), idl), unsafe_allow_html=True)
    k2.markdown(bsg_card('🏍️', 'KPI MAXI & CLASSY', f"MAXI: {int(g('R. MAXI'))}/{int(g('T. MAXI'))}", f"<b>CLASSY: {int(g('R. CLASSY'))} / {int(g('T. CLASSY'))} Unit</b> · MAXI {p_max:.0f}%", pace(p_max), min(p_max, 100.0), idl), unsafe_allow_html=True)
    k3.markdown(bsg_card('👥', 'KPI MP & INCOMING', f"MP: {int(g('R. MP'))}/{int(g('T. MP'))}", f"<b>INC: {int(g('R. INC'))} / {int(g('T. INC'))} SPK</b> · MP {p_mp:.0f}%", pace(p_mp), min(p_mp, 100.0), idl), unsafe_allow_html=True)
    k4.markdown(bsg_card('✅', 'KPI ACC% & MYDR/GREV', f"ACC: {g('R. ACC%'):.1f}%", f"<b>Target ACC {g('T. ACC%'):.1f}% · MYDR: {int(g('R. MYDR'))}/{int(g('T. MYDR'))} | G.REV: {int(g('R. G REV'))}/{int(g('T. G REV'))}</b>", bsg_stat_color(p_acc, 100, 85), min(p_acc, 100.0), None), unsafe_allow_html=True)

# ====================================================================================================
# ============ PANEL GABUNGAN: tabel (kiri/atas) + grafik pendamping + tombol Copy/Download HD ========
# ====================================================================================================
_BSG_BASE_CSS = _BSG_FRAME.split('<style>')[1].split('__CSS__')[0]
_BSG_TIPJS = '<script>' + _BSG_FRAME.split('<div id="tip"></div>')[1].split('<script>')[1].split('</script>')[0] + '</script>'

_PN_CSS = """
.pn-grid{display:grid;gap:12px;align-items:start}
.pn-charts{display:grid;gap:12px;min-width:0;grid-template-columns:minmax(0,1fr)}
.pn-table{min-width:0;overflow:hidden;border-radius:8px;border:1px solid #2D3748;background:#0E1117}
.pn-table div{margin:0;border:0;border-radius:0;overflow:visible}
.pn-table table{border-collapse:collapse;table-layout:auto;width:max-content;background:#0E1117;font-family:'Source Sans Pro','Segoe UI',Arial,sans-serif}
.pn-table th{background:#1A202C;color:#CBD5E1;font-weight:800;padding:5px 4px;border-bottom:1px solid #2D3748;border-right:1px solid #1E293B;text-align:right;white-space:normal;line-height:1.15;font-size:10.5px;width:auto}
.pn-table td{padding:4px 4px;border-bottom:1px solid #1E293B;border-right:1px solid #1E293B;text-align:right;white-space:nowrap;font-weight:700;font-size:12px;overflow:visible;width:auto}
.pn-table th:first-child,.pn-table td:first-child{text-align:left;white-space:normal;min-width:110px;max-width:135px;line-height:1.2;padding-left:7px;font-size:11.5px;width:auto;overflow:visible;text-overflow:clip}
.pn-table tr:hover td{background:rgba(255,255,255,.05)}
.L-side .pn-grid{grid-template-columns:minmax(0,1.9fr) minmax(0,1fr)}
.L-side2 .pn-grid{grid-template-columns:minmax(0,1.7fr) minmax(0,1fr);align-items:stretch}
.L-side2 .pn-table td{font-size:13.5px;padding:5px 3px}
.L-side2 .pn-table th{font-size:11.5px;padding:6px 3px}
.L-side2 .pn-table th:first-child,.L-side2 .pn-table td:first-child{font-size:13px;max-width:150px;padding-left:7px}
.L-side2 .pn-table{align-self:stretch}
.L-stack .pn-grid{grid-template-columns:minmax(0,1fr)}
.L-stack .pn-charts{grid-template-columns:minmax(0,3fr) minmax(0,2fr)}
.L-eq .pn-charts{grid-template-columns:minmax(0,1fr) minmax(0,1fr)}
.L-stack3 .pn-grid{grid-template-columns:minmax(0,1fr)}
.L-stack3 .pn-charts{grid-template-columns:minmax(0,1.7fr) minmax(0,1fr) minmax(0,1fr)}
.L-stack3 .dw{flex-direction:column;align-items:center;gap:6px}
.L-stack3 .dlg{width:100%}
.L-stack3.L-leg .dw{flex-direction:row;align-items:center;gap:14px}
.L-stack3.L-leg .dlg{width:auto}
.pn-charts,.pn-bottom,.pn-sub{align-items:stretch}
.pn-bottom{display:grid;gap:12px;margin-top:12px;min-width:0}
.pn-left{min-width:0;display:flex;flex-direction:column;gap:0}
.pn-left .pn-bottom{margin-top:12px;flex:1}
.pn-sub{display:grid;gap:12px;min-width:0;height:100%;grid-auto-rows:1fr}
.pn-charts>.card,.pn-bottom>.card,.pn-sub>.card{display:flex;flex-direction:column;min-width:0}
.pn-charts .dw,.pn-bottom .dw,.pn-sub .dw{flex:1}
.L-leg .card>.rk{height:auto;flex:1 1 26px;max-height:56px}
.narrow .pn-bottom,.narrow .pn-sub{grid-template-columns:minmax(0,1fr) !important}
.L-charts .pn-grid{grid-template-columns:minmax(0,1fr)}
.L-charts .pn-charts{grid-template-columns:repeat(auto-fit,minmax(320px,1fr))}
.narrow .pn-grid{grid-template-columns:minmax(0,1fr) !important}
.narrow .pn-charts{grid-template-columns:minmax(0,1fr) !important}
.cap-only{display:none}.cap .cap-only{display:block}
.cap *{animation:none !important;transition:none !important}
.pn-banner{background:linear-gradient(90deg,#4318FF 0%,#1E293B 100%);color:#fff;font-size:16px;font-weight:800;padding:10px 14px;border-radius:8px;margin-bottom:12px;letter-spacing:.3px}
.tb{display:flex;align-items:center;gap:8px;margin-top:8px;flex-wrap:wrap}
.btn{background:#1E293B;color:#E2E8F0;border:1px solid #334155;border-radius:6px;padding:6px 12px;font-size:12.5px;font-weight:600;cursor:pointer;font-family:inherit}
.btn:hover{background:#4318FF;border-color:#4318FF;color:#fff}
#st{font-size:12px;font-weight:600;margin-left:4px}
#hold{position:fixed;left:-99999px;top:0;width:__CAPW__px;background:#0E1117;padding:14px;z-index:-1}
"""

_PN_DOC = """<!DOCTYPE html><html><head><meta charset="utf-8">
<script src="https://cdnjs.cloudflare.com/ajax/libs/html2canvas/1.4.1/html2canvas.min.js"></script>
<style>__BASE__
__PNCSS__
__CHARTCSS__</style></head><body>
<div id="fit"><div id="pn" class="pn __LAYOUT__">__CAPHDR__<div class="cap-only pn-banner">🐺 __TITLE__</div><div class="pn-grid"__GRIDSTYLE__>__TABLE__<div class="pn-charts"__CHSTYLE__>__CHARTS__</div></div>__BOTTOM__
__FOOTER__</div>
<div class="tb"><button class="btn" onclick="pnCap('copy')">📸 Copy (Tabel + Grafik · HD)</button><button class="btn" onclick="pnCap('jpg')">📥 Download JPG</button><span id="st"></span></div></div>
<div id="hold"></div><div id="tip"></div>__TIPJS__
<script>(function(){
var FN="__FNAME__";
function fitTables(root){root.querySelectorAll('.pn-table').forEach(function(w){var t=w.querySelector('table');if(!t)return;t.style.zoom=1;t.style.width='max-content';
 var nat=t.getBoundingClientRect().width,av=w.clientWidth;if(nat>av+1){t.style.zoom=(av/nat).toFixed(4);}else{t.style.width=av+'px';}});}
function setFrameH(h){var fe=window.frameElement;if(!fe)return;var old=fe.getBoundingClientRect().height;if(Math.abs(old-h)<=2)return;var els=[],el=fe.parentElement,i=0;while(el&&i<4){if(Math.abs(el.getBoundingClientRect().height-old)<=3)els.push(el);el=el.parentElement;i++;}els.forEach(function(e){e.style.height=h+'px';e.style.minHeight=h+'px';});fe.style.height=h+'px';fe.setAttribute('height',h);}
function fitHeight(){try{var c=document.getElementById('fit');setFrameH(Math.ceil(c.getBoundingClientRect().height)+8);}catch(e){}}
function stretch(root){try{if(!root.classList.contains('L-side2')||root.classList.contains('narrow'))return;var w=root.querySelector('.pn-table'),c=root.querySelector('.pn-charts'),t=w&&w.querySelector('table');if(!t||!c)return;
 var cells=t.querySelectorAll('td,th'),rows=t.querySelectorAll('tr');cells.forEach(function(x){x.style.paddingTop='';x.style.paddingBottom='';});
 var z=parseFloat(t.style.zoom)||1,th=t.getBoundingClientRect().height,ch=c.getBoundingClientRect().height,wb=w.offsetHeight-w.clientHeight;
 var ex=(ch-wb-th)/Math.max(1,rows.length)/z;if(!(ex>0.5))return;ex=Math.min(ex,28);
 cells.forEach(function(x){var cs=getComputedStyle(x);x.style.paddingTop=(parseFloat(cs.paddingTop)+ex/2)+'px';x.style.paddingBottom=(parseFloat(cs.paddingBottom)+ex/2)+'px';});}catch(e){}}
function layout(){var pn=document.getElementById('pn');var W=document.getElementById('fit').clientWidth;pn.classList.toggle('narrow',W<820);fitTables(pn);stretch(pn);fitHeight();}
var rt;window.addEventListener('resize',function(){clearTimeout(rt);rt=setTimeout(layout,80);});
layout();setTimeout(layout,250);setTimeout(layout,900);setTimeout(layout,2000);
try{if(window.ResizeObserver){var ro=new ResizeObserver(function(){clearTimeout(rt);rt=setTimeout(function(){stretch(document.getElementById('pn'));fitHeight();},60);});ro.observe(document.getElementById('fit'));}}catch(e){}
window.pnCap=async function(mode){
 var st=document.getElementById('st');st.style.color='#FFEA00';st.textContent='Memproses gambar HD...';
 try{
  if(typeof html2canvas==='undefined')throw new Error('html2canvas belum termuat (cek koneksi internet)');
  var hold=document.getElementById('hold');hold.innerHTML='';
  var c=document.getElementById('pn').cloneNode(true);c.removeAttribute('id');c.classList.remove('narrow');c.classList.add('cap');
  hold.appendChild(c);fitTables(c);stretch(c);await new Promise(function(r){setTimeout(r,120);});
  var canvas=await html2canvas(hold,{backgroundColor:'#0E1117',scale:2,useCORS:true,logging:false});hold.innerHTML='';
  function dl(){var a=document.createElement('a');a.href=canvas.toDataURL('image/jpeg',0.95);a.download=FN+'.jpg';a.click();}
  if(mode==='copy'){canvas.toBlob(async function(blob){try{await navigator.clipboard.write([new ClipboardItem({'image/png':blob})]);
    st.style.color='#00FF66';st.textContent='✅ Tabel + grafik disalin (HD). Siap Ctrl+V di WA/Word/Excel';}catch(e){dl();st.style.color='#00E5FF';st.textContent='📥 Browser membatasi clipboard, gambar diunduh sebagai JPG.';}
    setTimeout(function(){st.textContent='';},5000);},'image/png');}
  else{dl();st.style.color='#00FF66';st.textContent='✅ '+FN+'.jpg berhasil diunduh!';setTimeout(function(){st.textContent='';},4500);}
 }catch(e){st.style.color='#FF4D8D';st.textContent='Gagal membuat gambar: '+e.message;}
};
window.pnCapData=async function(){var hold=document.getElementById('hold');hold.innerHTML='';var c=document.getElementById('pn').cloneNode(true);c.removeAttribute('id');c.classList.remove('narrow');c.classList.add('cap');
  hold.appendChild(c);fitTables(c);stretch(c);await new Promise(function(r){setTimeout(r,150);});var canvas=await html2canvas(hold,{backgroundColor:'#0E1117',scale:2,useCORS:true,logging:false});hold.innerHTML='';return canvas.toDataURL('image/jpeg',0.95);};
})();</script></body></html>"""

def _pn_bottom_html(bottom, cols=None):
    """Baris grafik di BAWAH tabel+grafik utama. bottom: list item = chart (inner, css, tinggi) atau dict(group=[chart...], cols=n) untuk blok pie n kolom. Return (html, css, tinggi)."""
    htm, csl, hs = [], [], []
    for b in (bottom or []):
        if not b: continue
        if isinstance(b, dict):
            g = [c for c in b.get('group', []) if c]
            if not g: continue
            n = int(b.get('cols', 2)); rws = -(-len(g) // n)
            htm.append(f'<div class="pn-sub" style="grid-template-columns:repeat({n},minmax(0,1fr))">' + ''.join(c[0] for c in g) + '</div>')
            csl += [c[1] for c in g]; hs.append(rws * max(c[2] for c in g) + (rws - 1) * 12)
        else: htm.append(b[0]); csl.append(b[1]); hs.append(b[2])
    if not htm: return '', '', 0
    tpl = cols or f'repeat({len(htm)},minmax(0,1fr))'
    return f'<div class="pn-bottom" style="grid-template-columns:{tpl}">' + ''.join(htm) + '</div>', '\n'.join(dict.fromkeys(csl)), max(hs)

def bsg_panel(table_html, charts, layout, title, file_label, cap_header='', n_rows=0, cap_width=1480, footer_html='', bottom=None, bottom_cols=None, chart_cols=None, grid_cols=None, bottom_left=False):
    """charts: list (inner_html, css, tinggi). layout: 'side' (tabel kiri, grafik kanan) | 'stack' (tabel atas, grafik bawah) | 'charts' (hanya grafik)."""
    charts = [c for c in (charts or []) if c]
    bt_html, bt_css, bt_h = _pn_bottom_html(bottom, bottom_cols)
    css_all = '\n'.join(dict.fromkeys([c[1] for c in charts] + ([bt_css] if bt_css else [])))
    ch_html = ''.join(c[0] for c in charts)
    tbl = f'<div class="pn-table">{table_html}</div>' if table_html else ''
    th = 44 + 31 * max(1, n_rows) if table_html else 0
    if bottom_left and bt_html and tbl:
        tbl = f'<div class="pn-left">{tbl}{bt_html}</div>'; th += bt_h + 12; bt_html, bt_h = '', 0
    if layout.startswith('side'): body_h = max(th, sum(c[2] + 12 for c in charts))
    elif layout.startswith('stack'): body_h = th + 12 + (max([c[2] for c in charts]) if charts else 0)
    else: body_h = max([c[2] for c in charts]) if charts else 0
    if bt_h: body_h += bt_h + 12
    doc = (_PN_DOC.replace('__BASE__', _BSG_BASE_CSS).replace('__PNCSS__', _PN_CSS.replace('__CAPW__', str(int(cap_width)))).replace('__CHARTCSS__', css_all)
           .replace('__LAYOUT__', {'side': 'L-side', 'side2': 'L-side L-side2', 'stack': 'L-stack', 'stack_eq': 'L-stack L-eq', 'stack3': 'L-stack3', 'stack3s': 'L-stack3 L-leg', 'charts': 'L-charts'}.get(layout, 'L-stack')).replace('__CAPHDR__', f'<div class="cap-only" style="margin-bottom:12px">{cap_header}</div>' if cap_header else '')
           .replace('__TITLE__', _esc(title)).replace('__TABLE__', tbl).replace('__CHARTS__', ch_html).replace('__BOTTOM__', bt_html)
           .replace('__GRIDSTYLE__', (' style="' + (f'grid-template-columns:{grid_cols};' if grid_cols else '') + ('align-items:stretch' if bottom_left else '') + '"' if (grid_cols or bottom_left) else '')).replace('__CHSTYLE__', (f' style="grid-template-columns:{chart_cols}"' if chart_cols else '')).replace('__TIPJS__', _BSG_TIPJS)
           .replace('__FNAME__', re.sub(r'[^a-zA-Z0-9_-]', '_', file_label)))
    doc = doc.replace('__FOOTER__', footer_html or '')
    return doc, int(body_h + 70 + (110 if footer_html else 0))

_PANEL_REG = {}   # nama_file -> dokumen HTML panel yang dirender pada run ini (dipakai tombol Download SEMUA gambar PPT)

def bsg_panel_table(df_items, df_full, columns, name_col, metric_col, key_suffix, title_text, chart_fn=None, chart_args=(), layout='side', wrap_first=False, highlight_names=None, cap_header='', footer_html='', bottom_fn=None, bottom_args=(), bottom_cols=None, chart_cols=None, grid_cols=None, bottom_left=False):
    """Tabel ringkas + grafik pendamping dalam SATU panel; tombol Copy/Download menangkap tabel dan semua grafik (HD)."""
    if df_full is None or df_full.empty: return
    charts = []
    if chart_fn is not None:
        try: charts = [c for c in (chart_fn(*chart_args) or []) if c]
        except Exception as e: st.caption(f'⚠️ Grafik pendamping tidak dapat ditampilkan ({type(e).__name__}: {str(e)[:80]}).')
    bottom = []
    if bottom_fn is not None:
        try: bottom = [c for c in (bottom_fn(*bottom_args) or []) if c]
        except Exception as e: st.caption(f'⚠️ Grafik tambahan tidak dapat ditampilkan ({type(e).__name__}: {str(e)[:80]}).')
    dom_id = 'PN_' + re.sub(r'[^a-zA-Z0-9]', '_', key_suffix)
    tbl = build_compact_html_string(df_items, df_full, columns, name_col, metric_col, dom_id, ultra_compact=True, title_text=None, highlight_names=highlight_names, wrap_first=wrap_first)
    doc, h = bsg_panel(tbl, charts, layout if charts else 'stack', title_text, key_suffix, cap_header, n_rows=len(df_full), footer_html=footer_html, bottom=bottom, bottom_cols=bottom_cols, chart_cols=chart_cols, grid_cols=grid_cols, bottom_left=bottom_left)
    _PANEL_REG[re.sub(r'[^a-zA-Z0-9_-]', '_', key_suffix)] = doc
    components.html(doc, height=h, scrolling=False)

def bsg_panel_charts(title, charts, file_label):
    charts = [c for c in charts if c]
    if not charts: return
    doc, h = bsg_panel('', charts, 'charts', title, file_label)
    components.html(doc, height=h, scrolling=False)

# ---------------------------------------------------------------- grafik pendamping tiap tabel ----
def _pn_items(rows_iter):
    items = []
    for n, v in rows_iter:
        if v > 0: items.append((n, v, BSG_PAL[len(items) % len(BSG_PAL)]))
    return items

def bsg_pn_sales_charts(df_p, kol_nm, ideal, who):
    items = _pn_items((str(r[kol_nm]), bsg_num(r.get('STU'))) for _, r in df_p.iterrows())
    rows = []
    for _, r in df_p.iterrows():
        n, stu, targ, pct, idu = str(r[kol_nm]), bsg_num(r.get('STU')), bsg_num(r.get('TARGET')), bsg_num(r.get('_pct')), bsg_num(r.get('IDEAL'))
        if targ > 0: rows.append({'name': n, 'val': pct, 'label': f'{pct:.1f}%', 'tip': f"<b>{_esc(n)}</b><br>STU <b>{int(stu):,}</b> / Target {int(targ):,}<br>Ideal {int(idu):,} → <b>{stu - idu:+,.0f}</b><br>Pencapaian <b>{pct:.1f}%</b>"})
    return [bsg_chart_rank('Pencapaian vs Target', f'% STU terhadap target bulanan · putih = ideal {ideal:.0f}%', rows, ideal=float(ideal), parts=True),
            bsg_chart_donut(f'Kontribusi Penjualan per {who}', 'Porsi unit STU terhadap total', items, 'STU', parts=True)]

def bsg_pn_curve_charts(df_stu, target, tgl, max_day, forc, label, fx):
    """Pendamping Tabel Penjualan: kurva pencapaian STU (+ garis Forc Nota / Forc 7 / Forc 3) di sisi kanan tabel."""
    return [bsg_chart_curve(f'Kurva Pencapaian STU — {label}', 'Garis putus-putus abu = target ideal harian',
                            df_stu, bsg_num(target), tgl, max_day, forc=bsg_num(forc) or None, h=190, extra=bsg_forc_lines(*fx), parts=True)]

def bsg_pn_process_charts(df_pr, kol_nm):
    dbq, stup = [], []
    for _, r in df_pr.iterrows():
        if bsg_num(r.get('DB')) <= 0: continue
        n = str(r[kol_nm]); db, hot, stu = int(bsg_num(r.get('DB'))), int(bsg_num(r.get('HOT'))), int(bsg_num(r.get('STU')))
        dbq.append({'name': n, 'val': bsg_num(r.get('DB Q')), 'tip': f"<b>{_esc(n)}</b><br>DB Q <b>{bsg_num(r.get('DB Q')):.2f}%</b><br>HOT {hot:,} / DB {db:,}"})
        stup.append({'name': n, 'val': bsg_num(r.get('_stu_pct')), 'tip': f"<b>{_esc(n)}</b><br>STU% <b>{bsg_num(r.get('_stu_pct')):.2f}%</b><br>STU {stu:,} / DB {db:,}"})
    return [bsg_chart_rank('Database Quality (DB Q)', 'HOT ÷ DB · putih = ideal 30%', dbq, ideal=30.0, parts=True),
            bsg_chart_rank('STU Ratio (STU%)', 'STU ÷ DB · putih = ideal 10%', stup, ideal=10.0, parts=True)]

def bsg_pn_process_bottom(df_pr, df_s, kol_nm, hari):
    """Di bawah tabel Proses Kerja: AVG Database, AVG HOT Prospect (per MP per hari) dan Approval Ratio dari tiap koordinator / counter."""
    avgdb, avghot, acc = [], [], []
    for _, r in df_pr.iterrows():
        mp, db, hot = int(bsg_num(r.get('JML MP'))), int(bsg_num(r.get('DB'))), int(bsg_num(r.get('HOT')))
        if mp <= 0 or db <= 0: continue
        n = str(r[kol_nm]); a_db = bsg_num(r.get('AVG DB')); a_ht = hot / mp / max(1, hari)
        avgdb.append({'name': n, 'val': a_db, 'label': f'{a_db:.2f}', 'tip': f"<b>{_esc(n)}</b><br>AVG DB <b>{a_db:.2f}</b> /MP/hari<br>DB {db:,} · {mp} MP"})
        avghot.append({'name': n, 'val': a_ht, 'label': f'{a_ht:.2f}', 'tip': f"<b>{_esc(n)}</b><br>AVG HOT <b>{a_ht:.2f}</b> /MP/hari<br>HOT {hot:,} · {mp} MP"})
    if df_s is not None and hasattr(df_s, 'iterrows') and not df_s.empty:
        for _, r in df_s.iterrows():
            if bsg_num(r.get('_spk_val')) <= 0: continue
            n, a_ = str(r[kol_nm]), bsg_num(r.get('_acc_val'))
            acc.append({'name': n, 'val': a_, 'tip': f"<b>{_esc(n)}</b><br>ACC% <b>{a_:.1f}%</b><br>SPK {int(bsg_num(r.get('_spk_val'))):,} · Terbuka {int(bsg_num(r.get('TERBUKA'))):,}"})
    return [bsg_chart_rank('AVG Database per MP / hari', 'DB ÷ jumlah MP ÷ hari kerja', avgdb, unit='', parts=True),
            bsg_chart_rank('AVG HOT Prospect per MP / hari', 'HOT ÷ jumlah MP ÷ hari kerja', avghot, unit='', parts=True),
            bsg_chart_rank('Approval Ratio (ACC%)', 'Putih = acuan ideal ACC% 50%', acc, ideal=50.0, parts=True)]

def bsg_pn_activity_charts(df_pt, ideal):
    _pal = {str(r['PROSPECT TYPE']): BSG_PAL[i_ % len(BSG_PAL)] for i_, (_, r) in enumerate(df_pt.iterrows())}   # warna per aktifitas dibuat tetap agar donut STU & HOT searah
    items = [(str(r['PROSPECT TYPE']), bsg_num(r.get('STU')), _pal[str(r['PROSPECT TYPE'])]) for _, r in df_pt.iterrows() if bsg_num(r.get('STU')) > 0]
    items_hot = [(str(r['PROSPECT TYPE']), bsg_num(r.get('HOT')), _pal[str(r['PROSPECT TYPE'])]) for _, r in df_pt.iterrows() if bsg_num(r.get('HOT')) > 0]
    rows = []
    for _, r in df_pt.iterrows():
        n, tg, stu = str(r['PROSPECT TYPE']), bsg_num(r.get('TARGET')), bsg_num(r.get('STU'))
        if tg > 0:
            ach = bsg_num(r.get('ACHIEV')); rows.append({'name': n, 'val': ach, 'label': f'{ach:.1f}%', 'tip': f"<b>{_esc(n)}</b><br>STU <b>{int(stu):,}</b> / Target {int(tg):,}<br>Capaian <b>{ach:.1f}%</b>"})
    return [bsg_chart_rank('Pencapaian Aktifitas vs Target', f'% STU terhadap target aktifitas · putih = ideal {ideal:.0f}%', rows, ideal=float(ideal), parts=True),
            bsg_chart_donut('Kontribusi Penjualan per Aktifitas', 'Porsi unit STU terhadap total', items, 'STU', parts=True),
            bsg_chart_donut('Kontribusi HOT per Aktifitas', 'Porsi HOT Prospect terhadap total', items_hot, 'HOT', parts=True)]

def bsg_pn_sub_charts(df_sub, sel_pt):
    items = _pn_items((str(r['SUB PROSPECT TYPE']), bsg_num(r.get('STU'))) for _, r in df_sub.iterrows())
    return [bsg_chart_donut(f'Kontribusi Penjualan per Sub Prospect — {sel_pt}', 'Porsi unit STU terhadap total tipe ini', items, 'STU', parts=True)]

def bsg_pn_leasing_charts(df_s, r_tot_s, kol_nm, f_c=None):
    rows = []
    for _, r in df_s.iterrows():
        if bsg_num(r.get('_spk_val')) <= 0: continue
        n, acc = str(r[kol_nm]), bsg_num(r.get('_acc_val'))
        rows.append({'name': n, 'val': acc, 'tip': f"<b>{_esc(n)}</b><br>ACC% <b>{acc:.1f}%</b><br>SPK {int(bsg_num(r.get('_spk_val'))):,} (LM {int(bsg_num(r.get('_spk_lm_val'))):,}) · Terbuka {int(bsg_num(r.get('TERBUKA'))):,}"})
    items = [(k, bsg_num(r_tot_s.get(k)), c) for k, c in [('BAF', '#4318FF'), ('ADIRA', '#00E5FF'), ('SOF', '#FFB020'), ('IMFI', '#FF4D8D')]]
    out = [bsg_chart_rank('Approval Rate (ACC%) per baris', 'Putih = acuan ideal ACC% 50%', rows, ideal=50.0, parts=True),
           bsg_chart_donut('Porsi SPK per Leasing', f'BAF, ADIRA, SOF, IMFI · dari total {int(bsg_num(r_tot_s.get("SPK"))):,} SPK', items, 'SPK 4 leasing', parts=True)]
    if f_c is not None and not f_c.empty and 'Leasing_Faktur_Clean' in f_c.columns:
        vc = f_c['Leasing_Faktur_Clean'].value_counts()
        baf, adr, sof = int(vc.get('BAF', 0)), int(vc.get('ADIRA', 0)), int(vc.get('SOF', 0))
        kredit = int(sum(v for k, v in vc.items() if k not in ('CASH', '', None)))
        lain = max(0, kredit - baf - adr - sof)
        its = [('BAF', baf, '#4318FF'), ('ADIRA', adr, '#00E5FF'), ('SOF', sof, '#FFB020')] + ([('Leasing lain', lain, '#94A3B8')] if lain > 0 else [])
        out.append(bsg_chart_donut('Kontribusi STU Kredit per Leasing', f'Dari {kredit:,} faktur kredit (cash {int(vc.get("CASH", 0)):,} tidak dihitung)', its, 'Faktur kredit', parts=True))
    return out

def bsg_pn_leasing_bottom(r_tot_s, s_c):
    """Baris bawah panel SPK & Leasing: Approval Rate per leasing (kiri) + 4 pie Reject Reason (BAF, ADIRA, SOF, IMFI) di sisi kanan."""
    LS = [('BAF', '#4318FF'), ('ADIRA', '#00E5FF'), ('SOF', '#FFB020'), ('IMFI', '#FF4D8D')]
    ok = (s_c is not None and hasattr(s_c, 'empty') and not s_c.empty and all(c in s_c.columns for c in ('Leasing_SPK_Clean', 'Status_SPK_Clean', 'Alasan_Reject_Clean')))
    n_open = {l: (int(((s_c['Leasing_SPK_Clean'] == l) & (s_c['Status_SPK_Clean'] == 'TERBUKA')).sum()) if ok else 0) for l, _ in LS}
    rows = [{'name': l, 'val': bsg_num(r_tot_s.get(f'%{l}')), 'extra': str(n_open[l]), 'tip': f"<b>{l}</b><br>ACC% <b>{bsg_num(r_tot_s.get(f'%{l}')):.1f}%</b><br>SPK {int(bsg_num(r_tot_s.get(l))):,} · <span style='color:#FFEA00'>SPK terbuka {n_open[l]:,}</span>"} for l, _ in LS]
    out = [bsg_chart_rank('Approval Rate per Leasing', 'Putih = acuan ideal ACC% 50% · angka kuning = SPK terbuka', rows, ideal=50.0, parts=True)]
    rej = {l: (s_c[(s_c['Leasing_SPK_Clean'] == l) & (s_c['Status_SPK_Clean'] == 'DITOLAK')] if ok else None) for l, _ in LS}
    tot_all = {}   # urutan & warna alasan dibuat tetap di keempat pie agar alasan yang sama berwarna sama
    for d_ in rej.values():
        if d_ is not None and not d_.empty:
            for k_, v_ in d_['Alasan_Reject_Clean'].value_counts().items(): tot_all[str(k_)] = tot_all.get(str(k_), 0) + int(v_)
    pal = {k_: BSG_PAL[i_ % len(BSG_PAL)] for i_, (k_, _) in enumerate(sorted(tot_all.items(), key=lambda t: -t[1]))}
    pies = []
    for l, _ in LS:
        d_ = rej[l]; n_spk = int((s_c['Leasing_SPK_Clean'] == l).sum()) if ok else 0
        if d_ is None or d_.empty:
            pies.append(bsg_chart_donut(f'Reject Reason — {l}', f'Tidak ada SPK ditolak · dari {n_spk:,} SPK {l}', [('Tidak ada reject', 0, '#64748B')], 'Reject', parts=True, keep_zero=True)); continue
        vc = d_['Alasan_Reject_Clean'].value_counts()
        pies.append(bsg_chart_donut(f'Reject Reason — {l}', f'{int(vc.sum()):,} SPK ditolak · dari {n_spk:,} SPK {l}', [(str(k_), int(v_), pal.get(str(k_), '#64748B')) for k_, v_ in vc.items()], 'Reject', parts=True))
    out.append({'group': pies, 'cols': 2})
    return out

def bsg_pn_type_bottom(df_td_i, f_cg):
    """Di bawah tabel Penjualan per Type: 3 pie kontribusi leasing KREDIT (BAF, ADIRA, SOF) untuk AT PREMIUM, AT CLASSY, AT LPM. CASH tidak dihitung."""
    LS = [('BAF', '#4318FF'), ('ADIRA', '#00E5FF'), ('SOF', '#FFB020')]
    ok = (f_cg is not None and hasattr(f_cg, 'empty') and not f_cg.empty and 'Type_Grp' in f_cg.columns and 'Leasing_Faktur_Clean' in f_cg.columns)
    pies = []
    for t_ in ['AT PREMIUM', 'AT CLASSY', 'AT LPM']:
        d_ = f_cg[f_cg['Type_Grp'] == t_] if ok else None
        vc = d_['Leasing_Faktur_Clean'].value_counts() if (d_ is not None and not d_.empty) else pd.Series(dtype='int64')
        kr = {str(k_): int(v_) for k_, v_ in vc.items() if str(k_).strip() and str(k_).upper() not in ('CASH', 'NAN', 'NONE', '-')}
        its = [(k_, kr.get(k_, 0), c_) for k_, c_ in LS]
        lain = sum(v_ for k_, v_ in kr.items() if k_ not in dict(LS))
        if lain > 0: its.append(('Lainnya', lain, '#94A3B8'))
        n_cash = int(vc.get('CASH', 0))
        pies.append(bsg_chart_donut(f'Kontribusi Leasing — {t_}', f'{sum(kr.values()):,} STU kredit (cash {n_cash:,} tidak dihitung)', its, 'STU kredit', parts=True, keep_zero=True))
    return pies

def bsg_pn_type_charts(df_td_i, ideal):
    pal = {'AT PREMIUM': '#4318FF', 'AT CLASSY': '#00E5FF', 'AT LPM': '#FFB020', 'MOPED': '#FF4D8D', 'OTHERS': '#94A3B8'}
    items = [(str(r['TYPE']), bsg_num(r.get('STU')), pal.get(str(r['TYPE']), '#64748B')) for _, r in df_td_i.iterrows()]
    rows = []
    for _, r in df_td_i.iterrows():
        tg = bsg_num(r.get('TARGET'))
        if tg <= 0: continue
        ach = bsg_num(r.get('ACHIEV'))
        rows.append({'name': str(r['TYPE']), 'val': ach, 'tip': f"<b>{_esc(r['TYPE'])}</b><br>STU <b>{int(bsg_num(r.get('STU'))):,}</b> / Target {int(tg):,}<br>Capaian <b>{ach:.1f}%</b> · LM {int(bsg_num(r.get('LM'))):,}"})
    return [bsg_chart_rank('Capaian per Type vs Target Type', f'Putih = ideal {ideal:.0f}%', rows, ideal=float(ideal), parts=True),
            bsg_chart_donut('Komposisi STU per Type', 'Porsi unit terjual periode berjalan', items, 'STU', parts=True)]

def bsg_pn_salesman_charts(df_smi, ideal, n=10):
    """Dua grafik Top-N & N Salesman Perlu Dorongan (pendamping tabel detail salesman)."""
    rows = []
    for _, r in df_smi.iterrows():
        tg = bsg_num(r.get('TARGET'))
        if tg <= 0: continue
        pc = bsg_num(r.get('_pct_val'))
        rows.append({'name': str(r['NAMA SALESMAN']), 'val': pc, 'tip': f"<b>{_esc(r['NAMA SALESMAN'])}</b><br>STU <b>{int(bsg_num(r.get('STU'))):,}</b> / Target {int(tg):,} ({pc:.1f}%)<br>DB {int(bsg_num(r.get('DB'))):,} · DB Q {_esc(r.get('DB Q', '-'))} · STU% {_esc(r.get('STU%', '-'))}"})
    if not rows: return []
    k = min(n, len(rows))
    return [bsg_chart_rank(f'Top {k} Salesman', f'% STU terhadap target · ideal {ideal:.0f}%', rows, ideal=float(ideal), max_rows=n, parts=True),
            bsg_chart_rank(f'{k} Salesman Perlu Dorongan', 'Capaian terendah (hanya yang punya target)', rows, ideal=float(ideal), max_rows=n, ascending=True, parts=True)]

def bsg_render_salesman_panel(df_smi, ideal, title, n=10):
    """Top-N & N salesman perlu dorongan + tombol Copy/Download HD."""
    if df_smi is None or df_smi.empty or '_pct_val' not in df_smi.columns: return
    rows = []
    for _, r in df_smi.iterrows():
        tg = bsg_num(r.get('TARGET'))
        if tg <= 0: continue
        pc = bsg_num(r.get('_pct_val'))
        rows.append({'name': str(r['NAMA SALESMAN']), 'val': pc, 'tip': f"<b>{_esc(r['NAMA SALESMAN'])}</b><br>STU <b>{int(bsg_num(r.get('STU'))):,}</b> / Target {int(tg):,} ({pc:.1f}%)<br>DB {int(bsg_num(r.get('DB'))):,} · DB Q {_esc(r.get('DB Q', '-'))} · STU% {_esc(r.get('STU%', '-'))}"})
    if not rows: return
    k = min(n, len(rows))
    bsg_panel_charts(f'Top & Low Performance Salesman — {title}', [
        bsg_chart_rank(f'Top {k} Salesman', f'% STU terhadap target · ideal {ideal:.0f}%', rows, ideal=float(ideal), max_rows=n, parts=True),
        bsg_chart_rank(f'{k} Salesman Perlu Dorongan', 'Capaian terendah (hanya yang punya target)', rows, ideal=float(ideal), max_rows=n, ascending=True, parts=True)], 'Top_Low_Salesman_' + str(title))

def bsg_render_reg_panel(sudah, belum, f_c, title):
    """Donut sudah vs belum registrasi rating + ranking koordinator dengan % registrasi tertinggi."""
    donut = bsg_chart_donut('Status Registrasi Rating (MYDR)', 'Persentase konsumen faktur yang sudah vs belum registrasi', [('Sudah Registrasi', sudah, '#00E676'), ('Belum Registrasi', belum, '#FF4D8D')], 'Konsumen', parts=True, keep_zero=True)
    rows = []
    if f_c is not None and not f_c.empty and 'Koor' in f_c.columns and 'Status_Registrasi_Rating' in f_c.columns:
        for nm, g in f_c.groupby('Koor'):
            n = clean_str(nm)
            if not n or n in ('NAN', 'NONE', 'LAIN-LAIN'): continue
            tot, sdh = len(g), int((g['Status_Registrasi_Rating'] == 'SUDAH REGISTRASI').sum())
            pc = sdh / tot * 100.0 if tot else 0.0
            rows.append({'name': n, 'val': pc, 'label': f'{pc:.1f}%', 'tip': f"<b>{_esc(n)}</b><br>Sudah registrasi <b>{sdh:,}</b> dari {tot:,} konsumen<br><b>{pc:.1f}%</b>"})
    rank = bsg_chart_rank('Koordinator dengan Registrasi Rating Tertinggi', 'Persentase konsumen yang sudah registrasi (10 teratas)', rows, ideal=None, fixed_color='#00E676', max_rows=10, scale_max=100.0, parts=True)
    bsg_panel_charts(f'Registrasi Rating MYDR — {title}', [donut, rank], 'Registrasi_MYDR_' + str(title))


# ====================================================================================================
# ============ ARSIP DATA: EXPORT (semua data & pengaturan → 1 file .zip) & IMPORT (buka dashboard lama) ===
# ====================================================================================================
_ARSIP_JSON = ['dealer_targets_saved.json', 'koor_targets_saved.json', 'kpi_targets_saved.json', 'libur_nasional.json', 'kpi_dashboard_manual.json']

def _wib_now(): return datetime.datetime.now(datetime.timezone.utc) + datetime.timedelta(hours=7)

def bsg_build_export_zip(info):
    """Kemas seluruh file sumber (saved_raw_files) + pengaturan target/libur/KPI manual + manifest ke 1 file zip.
    API key (ai_config.json) SENGAJA tidak ikut diekspor."""
    import io, zipfile
    buf, n_files, now = io.BytesIO(), 0, _wib_now()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED, compresslevel=6) as z:
        for root, _d, files in os.walk(UPLOAD_DIR):
            for fn in files:
                fp = os.path.join(root, fn)
                z.write(fp, 'saved_raw_files/' + os.path.relpath(fp, UPLOAD_DIR).replace(os.sep, '/')); n_files += 1
        for nm in _ARSIP_JSON:
            fp = _ABS(nm)
            if os.path.exists(fp): z.write(fp, nm); n_files += 1
        man = {'app': 'JEMBER BONDOWOSO ANALYSIS', 'format': 1, 'exported_at': now.strftime('%d %b %Y, %H:%M WIB'), 'jumlah_file': n_files}
        man.update(info or {})
        z.writestr('manifest.json', json.dumps(man, ensure_ascii=False, indent=2))
    name = f"JEMBER_BONDOWOSO_Arsip_{re.sub(r'[^A-Za-z0-9]+', '_', str(man.get('periode', 'Data')))}_Tgl{man.get('tanggal_akhir', '')}_{now.strftime('%Y%m%d_%H%M')}.zip"
    return buf.getvalue(), name

def bsg_import_archive(upfile, max_bytes=1_500_000_000):
    """Ekstrak arsip ke folder sementara (aman dari zip-slip) lalu arahkan path metadata ke folder itu. Return (folder, manifest)."""
    import zipfile, tempfile, shutil
    with zipfile.ZipFile(upfile) as z:
        if 'manifest.json' not in z.namelist(): raise ValueError('bukan file arsip JEMBER BONDOWOSO (manifest.json tidak ditemukan)')
        man = json.loads(z.read('manifest.json').decode('utf-8'))
        if sum(m.file_size for m in z.infolist()) > max_bytes: raise ValueError('ukuran arsip terlalu besar')
        dest = tempfile.mkdtemp(prefix='jember_bondowoso_arsip_'); root = os.path.normpath(dest) + os.sep
        for m in z.infolist():
            if m.is_dir(): continue
            target = os.path.normpath(os.path.join(dest, m.filename))
            if not target.startswith(root): continue
            os.makedirs(os.path.dirname(target), exist_ok=True)
            with z.open(m) as src, open(target, 'wb') as out: shutil.copyfileobj(src, out)
    mp = os.path.join(dest, 'saved_raw_files', 'metadata.json')
    if os.path.exists(mp):
        try:
            with open(mp, 'r') as fp: meta = json.load(fp)
            for k, v in meta.items():
                if not isinstance(v, dict): continue
                if v.get('saved_path'): v['saved_path'] = os.path.join(dest, 'saved_raw_files', os.path.basename(str(v['saved_path']).replace('\\', '/')))
                if isinstance(v.get('saved_paths'), list):
                    sub = 'database_files' if k == 'database_multi' else ('rating_files' if k == 'rating_multi' else '')
                    v['saved_paths'] = [os.path.join(dest, 'saved_raw_files', sub, os.path.basename(str(p).replace('\\', '/'))) for p in v['saved_paths']]
            with open(mp, 'w') as fp: json.dump(meta, fp)
        except Exception: pass
    return dest, man

def bsg_arsip_reset():
    import shutil
    old = st.session_state.get('arsip_dir')
    if old and os.path.isdir(old): shutil.rmtree(old, ignore_errors=True)
    for k in ('arsip_dir', 'arsip_meta', 'arsip_sig', 'arsip_name'): st.session_state.pop(k, None)

def bsg_arsip_activate(upfile):
    import shutil
    dest, man = bsg_import_archive(upfile)
    old = st.session_state.get('arsip_dir')
    if old and old != dest and os.path.isdir(old): shutil.rmtree(old, ignore_errors=True)
    st.session_state.update({'arsip_dir': dest, 'arsip_meta': man, 'arsip_sig': f'{upfile.name}_{upfile.size}', 'arsip_name': upfile.name, 'arsip_show_import': False})
    bsg_toast('Arsip dibuka: ' + str(man.get('periode', '')), '📦')


# ====================================================================================================
# ============ CREATE PPT: dashboard KPI → PowerPoint (masalah global → dealer → tim terlemah) =========
# Butuh library `python-pptx` (tambahkan ke requirements.txt). Semua fungsi di sini murni (tanpa Streamlit).
# ====================================================================================================
_PP = dict(NAVY='0B2A5B', TEAL='1B5E75', INK='0F172A', MUTED='64748B', LINE='E2E8F0', SOFT='F1F5F9', WHITE='FFFFFF',
           OK='16A34A', WARN='D97706', BAD='DC2626', INFO='0891B2', AMBER='F59E0B', ICE='CADCFC', PEN='FFFBEB')
_PP_GRADE = {'A': '16A34A', 'B': '0891B2', 'C': 'D97706', 'D': 'DC2626'}
_PP_TINT = {'ok': ('DCFCE7', '166534'), 'warn': ('FEF3C7', '92400E'), 'bad': ('FEE2E2', '991B1B'), 'na': ('F1F5F9', '64748B')}
_PP_FONT = 'Calibri'

def _pp_status(r):
    if r.get('weight', 0) == 0: return 'na'
    if r.get('fixed') or r['point'] >= r['weight']: return 'ok'
    return 'warn' if r['point'] > 0 else 'bad'

def _pp_fmt(code, v):
    if code in ('APV', 'MYM'): return f'{v * 100:.1f}%'
    if code == 'GRT': return f'{v:.1f}'
    if code == 'PRD': return f'{v:.2f}'
    return f'{v:,.0f}'

# ---------------------------------------------------------------------------------------------------
# Pembuat PPT berbasis TEMPLATE (TEMPLATE_PPT_JEMBER_BONDOWOSO.pptx): slide, layout, judul & tabel Condition/Problem
# dipakai dari template; gambar screenshot lama diganti tabel/grafik native hasil hitung dashboard.
# ---------------------------------------------------------------------------------------------------
_TPL_NAME = 'TEMPLATE_PPT_JEMBER_BONDOWOSO.pptx'
_TP = dict(BG='0E1117', CARD='111827', LINE='2D3748', HDR='4318FF', TOT='3B18FF', R1='0E1117', R2='161B26', TXT='E2E8F0', MUT='94A3B8',
           GRN='00FF66', YEL='FFEA00', RED='FF4D8D', CYN='00E5FF', WHT='FFFFFF', DARK='1F2937')
_TP_KN = 'KOOR / COUNTER / DEALER'

def _tp_find_template():
    """Cari file template di folder app.py (repo GitHub), folder kerja, atau ./templates. Return path / None."""
    cands = []
    try: cands.append(os.path.join(os.path.dirname(os.path.abspath(__file__)), _TPL_NAME))
    except NameError: pass
    cands += [os.path.join(os.getcwd(), _TPL_NAME), os.path.join(os.getcwd(), 'templates', _TPL_NAME)]
    return next((p_ for p_ in cands if os.path.isfile(p_)), None)

def _tp_rgb(h):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(h)

def _tp_col(v, good, warn): return _TP['GRN'] if v >= good else (_TP['YEL'] if v >= warn else _TP['RED'])

# ---------- primitif ----------
def _tp_text(sl, x, y, w, h, paras, size=11, color=None, bold=False, align='l', anchor='t', italic=False):
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    al = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
    if isinstance(paras, str): paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = al
        for t_, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run(); r.text = t_; f = r.font; f.name = 'Calibri'; f.size = Pt(o.get('s', size)); f.bold = o.get('b', bold)
            f.italic = o.get('i', italic); f.color.rgb = _tp_rgb(o.get('c', color or _TP['TXT']))
    return tb

def _tp_rect(sl, x, y, w, h, fill=None, line=None, lw=0.75, rad=None):
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    s = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rad else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if rad: s.adjustments[0] = rad
    if fill: s.fill.solid(); s.fill.fore_color.rgb = _tp_rgb(fill)
    else: s.fill.background()
    if line: s.line.color.rgb = _tp_rgb(line); s.line.width = Pt(lw)
    else: s.line.fill.background()
    s.shadow.inherit = False
    return s

def _tp_ensure_bg(sl):
    """Pastikan ada latar gelap penuh (slide yang tadinya hanya berisi gambar tidak punya)."""
    from pptx.util import Emu
    W, H = 12192000, 6858000
    for sh in sl.shapes:
        if sh.shape_type == 1 and sh.left is not None and sh.width is not None and sh.width >= W * 0.97 and sh.height >= H * 0.97: return
    r = _tp_rect(sl, 0, 0, 13.333, 7.5, fill='000000')
    sl.shapes._spTree.remove(r._element); sl.shapes._spTree.insert(2, r._element)

def _tp_clear_pictures(sl):
    for sh in list(sl.shapes):
        if sh.shape_type == 13: sh._element.getparent().remove(sh._element)

def _tp_head(sl, title, sub=None, w=7.3):
    _tp_text(sl, 0.25, 0.08, w, 0.5, title, size=13, color=_TP['CYN'], bold=True)
    if sub: _tp_text(sl, 0.25, 0.4, w, 0.2, sub, size=8, color=_TP['MUT'])

def _tp_table(sl, x, y, col_w, header, rows, row_h=0.26, hdr_h=0.3, size=9, total_last=True, aligns=None):
    """rows: list of list; sel = str atau dict(t=, c=warna teks, f=fill, b=bold, a=align). Baris terakhir = TOTAL bila total_last."""
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    nr, nc = len(rows) + 1, len(col_w)
    gs = sl.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(col_w)), Inches(hdr_h + row_h * len(rows)))
    tbl = gs.table; tbl.first_row = False; tbl.horz_banding = False
    for ch in gs._element.graphic.graphicData.tbl.tblPr:
        if ch.tag.endswith('tableStyleId'): ch.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
    for j, w_ in enumerate(col_w): tbl.columns[j].width = Inches(w_)
    for i in range(nr): tbl.rows[i].height = Inches(hdr_h if i == 0 else row_h)
    al_map = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}
    aligns = aligns or (['l'] + ['c'] * (nc - 1))
    def _put(cell, spec, fill, color, bold, al):
        spec = spec if isinstance(spec, dict) else dict(t=spec)
        cell.margin_left = cell.margin_right = Inches(0.05); cell.margin_top = cell.margin_bottom = Inches(0.02); cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid(); cell.fill.fore_color.rgb = _tp_rgb(spec.get('f', fill))
        tf = cell.text_frame; tf.word_wrap = True; p = tf.paragraphs[0]; p.alignment = al_map[spec.get('a', al)]
        r = p.add_run(); r.text = str(spec.get('t', '')); f = r.font; f.name = 'Calibri'; f.size = Pt(spec.get('s', size)); f.bold = spec.get('b', bold); f.color.rgb = _tp_rgb(spec.get('c', color))
    for j, h_ in enumerate(header): _put(tbl.cell(0, j), h_, _TP['HDR'], _TP['WHT'], True, 'c' if j else 'l')
    for i, row in enumerate(rows, 1):
        is_tot = total_last and i == len(rows)
        for j, spec in enumerate(row): _put(tbl.cell(i, j), spec, _TP['TOT'] if is_tot else (_TP['R1'] if i % 2 else _TP['R2']), _TP['WHT'] if is_tot else _TP['TXT'], is_tot or (j == 0), aligns[j])
    return tbl

def _tp_bar(sl, x, y, w, h, title, cats, vals, ideal=None, fmt='0.0"%"', colors=None, maxv=None):
    """Bar horizontal native. Warna otomatis terhadap ideal: >=ideal hijau, >=70% ideal kuning, selebihnya merah."""
    from pptx.util import Inches, Pt
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    _tp_text(sl, x, y, w, 0.22, title, size=10, color=_TP['CYN'], bold=True)
    if not cats: _tp_text(sl, x, y + 0.35, w, 0.3, 'Tidak ada data.', size=9, color=_TP['MUT']); return
    cd = CategoryChartData(); cd.categories = [str(c)[:26] for c in cats]; cd.add_series('v', [round(float(v), 2) for v in vals])
    gf = sl.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y + 0.22), Inches(w), Inches(max(0.5, h - 0.22)), cd); ch = gf.chart
    ch.has_legend = False; ch.has_title = False; ch.font.size = Pt(9); ch.font.name = 'Calibri'; ch.font.color.rgb = _tp_rgb(_TP['TXT'])
    pl = ch.plots[0]; pl.gap_width = 45; pl.vary_by_categories = False
    va, ca = ch.value_axis, ch.category_axis
    va.visible = False; va.has_major_gridlines = False; va.minimum_scale = 0
    top = maxv or max(max([float(v) for v in vals] + [0.0]) * 1.25, (ideal or 0) * 1.15, 1.0); va.maximum_scale = top
    ca.reverse_order = True; ca.has_major_gridlines = False; ca.format.line.fill.background(); ca.tick_labels.font.size = Pt(9); ca.tick_labels.font.color.rgb = _tp_rgb(_TP['TXT'])
    s0 = pl.series[0]; s0.format.fill.solid(); s0.format.fill.fore_color.rgb = _tp_rgb(_TP['CYN'])
    for i, v in enumerate(vals):
        c_ = (colors[i] if colors else (_tp_col(float(v), ideal, 0.7 * ideal) if ideal else _TP['CYN']))
        pt = s0.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = _tp_rgb(c_)
    dl = s0.data_labels; dl.show_value = True; dl.number_format = fmt; dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END; dl.font.size = Pt(9); dl.font.bold = True; dl.font.color.rgb = _tp_rgb(_TP['WHT'])

def _tp_card(sl, x, y, w, h, label, value, sub, color, prog=None):
    _tp_rect(sl, x, y, w, h, fill=_TP['CARD'], line=_TP['LINE'], lw=1, rad=0.08)
    _tp_text(sl, x + 0.1, y + 0.05, w - 0.2, 0.2, label.upper(), size=8, color=_TP['MUT'], bold=True)
    _tp_text(sl, x + 0.1, y + 0.24, w - 0.2, 0.4, value, size=18 if h >= 0.9 else 15, color=color, bold=True)
    _tp_text(sl, x + 0.1, y + h - (0.3 if prog is not None else 0.27), w - 0.2, 0.22, sub, size=7.5, color=_TP['MUT'])
    if prog is not None:
        _tp_rect(sl, x + 0.1, y + h - 0.09, w - 0.2, 0.05, fill=_TP['LINE'])
        _tp_rect(sl, x + 0.1, y + h - 0.09, max(0.02, (w - 0.2) * min(max(prog, 0.0), 100.0) / 100.0), 0.05, fill=color)

# ---------- gambar dashboard (hasil tombol Download JPG) ----------
def _tp_san(x): return re.sub(r'[^a-zA-Z0-9_-]', '_', str(x))

_TP_IMG = {2: [('Picture 3', 'Review_1_Penjualan_', 'dealer')], 3: [('Picture 2', 'Review_5_Salesman_', 'dealer')], 5: [('Picture 2', 'Review_2_Proses_', 'dealer')],
           6: [('Picture 2', 'Review_Aktifitas_PT_', 'dealer'), ('Picture 1', 'Review_Aktifitas_Sub_', 'dealer')], 8: [('Picture 2', 'Review_1_Penjualan_', 'koor')], 9: [('Picture 9', 'Review_5_Salesman_', 'koor')],
           10: [('Picture 2', 'Review_Aktifitas_PT_', 'koor'), ('Picture 3', 'Review_Aktifitas_Sub_', 'koor')], 12: [('Picture 2', 'Review_1_Penjualan_', 'counter')],
           13: [('Picture 2', 'Review_2_Proses_', 'counter'), ('Picture 5', 'Review_Aktifitas_PT_', 'counter')], 15: [('Picture 2', 'Review_3_Leasing_', 'dealer')],
           16: [('Picture 2', 'Review_4_Type_', 'dealer')], 17: [('Picture 6', 'Dashboard_KPI_', 'kpi')]}

_TP_KPI_ENT = ('kpi_koor', 'kpi_counter')   # gambar Dashboard KPI koordinator / counter TERLEMAH (nama file memakai nama entitas saja)

def _tp_img_base(D, key):
    if key in _TP_KPI_ENT:
        S_ = D['S_koor'] if key == 'kpi_koor' else D['S_counter']
        return _tp_san(S_['koor']) if S_ else None
    if key in ('dealer', 'kpi'): return _tp_san(D['scope'])
    S_ = D['S_koor'] if key == 'koor' else D['S_counter']
    return _tp_san(f"{S_['dealer_lbl']} — {S_['koor']}") if S_ else None

def _tp_find_img(D, prefix, key):
    """Cari gambar unggahan sesuai nama file bawaan tombol Download (mis. Review_1_Penjualan_SIP_AMBULU.jpg). Return (bytes|None, nama_diharapkan)."""
    base = _tp_img_base(D, key)
    if base is None: return None, None
    _kk = key in ('kpi',) + _TP_KPI_ENT
    exp = (prefix + base + ('_Tgl_' if _kk else '')) + ('.jpg' if not _kk else '<tgl>.jpg')
    if key == 'dealer': pat = re.compile('^' + re.escape(prefix + base) + r'(?:_(?!__)|$)')
    elif _kk: pat = re.compile('^' + re.escape(prefix + base + '_Tgl_'))
    else: pat = re.compile('^' + re.escape(prefix + base) + r'(?:_|$)')
    for nm_, b_ in (D.get('images') or {}).items():
        stem = re.sub(r'\s*\(\d+\)$', '', os.path.splitext(os.path.basename(nm_))[0])
        if pat.search(stem): return b_, exp
    return None, exp

def _tp_swap_pic(sl, name, img):
    """Ganti ISI gambar template bernama `name` dengan gambar unggahan. Ukuran, posisi, urutan tumpukan & crop kotak template TIDAK diubah.
    Return pesan peringatan bila rasio gambar jauh berbeda dari kotak (gambar akan tampak meregang), selain itu None."""
    import io
    from pptx.oxml.ns import qn
    from PIL import Image
    old = next((s_ for s_ in sl.shapes if s_.name == name and s_.shape_type == 13), None)
    if old is None: raise ValueError(f'gambar template "{name}" tidak ditemukan')
    iw, ih = Image.open(io.BytesIO(img)).size
    _, rid = sl.part.get_or_add_image_part(io.BytesIO(img))
    blip = old._element.xpath('.//a:blip')[0]; old_rid = blip.get(qn('r:embed')); blip.set(qn('r:embed'), rid)
    if old_rid and old_rid != rid:
        try: sl.part.drop_rel(old_rid)
        except Exception: pass
    box = old.width / old.height; im = iw / ih
    old.crop_left = old.crop_right = old.crop_top = old.crop_bottom = 0.0
    if abs(im / box - 1) > 0.04:   # rasio beda: pasang gambar UTUH di tengah kotak (ruang sisa transparan) agar tidak meregang; ukuran & posisi kotak template tetap
        if im < box: pad = (box / im - 1.0) / 2.0; old.crop_left = old.crop_right = -pad
        else: pad = (im / box - 1.0) / 2.0; old.crop_top = old.crop_bottom = -pad
    return (f'rasio gambar {im:.2f} vs kotak {box:.2f}' if abs(im / box - 1) > 0.12 else None)

# ---------- pengisi tabel template (CONDITION / PROBLEM / PERBAIKAN / TARGET / TIME LINE) ----------
def _tp_fit(texts_by_row, col_w, avail_h, sizes=(11, 10, 9, 8, 7)):
    import math
    for sz in sizes:
        cw, tot = sz * 0.0072, 0.0
        for row in texts_by_row:
            tot += max(math.ceil(max(1, len(str(t))) * cw / max(0.3, w_ - 0.15)) + str(t).count('\n') for t, w_ in zip(row, col_w)) * sz * 1.22 / 72 + 0.1
        if tot <= avail_h: return sz
    return sizes[-1]

def _tp_fill_analysis(sl, rows, bottom=7.35, top=None, name=None):
    """Isi HANYA kolom CONDITION pada tabel template; PROBLEM / PERBAIKAN / TARGET PERBAIKAN / TIME LINE dibiarkan kosong (diisi manual)."""
    from pptx.util import Inches, Pt
    shp = next((s_ for s_ in sl.shapes if s_.has_table and name and s_.name == name), None) or next(s_ for s_ in sl.shapes if s_.has_table); tbl = shp.table
    nb = len(tbl.rows) - 1
    if top is not None: shp.top = Inches(top)
    rows = list(rows)[:nb] + [('', 'Tidak ada data pada periode ini.')] * max(0, nb - len(rows))
    col0 = tbl.columns[0].width / 914400; top_in = shp.top / 914400 + tbl.rows[0].height / 914400
    sz = _tp_fit([[(r_[0] + ' ' if r_[0] else '') + r_[1]] for r_ in rows], [col0], bottom - top_in, sizes=(11, 10, 9, 8))
    for i, r_ in enumerate(rows, 1):
        lb, cd = r_[0], r_[1]
        cell = tbl.cell(i, 0); tf = cell.text_frame; tf.clear(); tf.word_wrap = True
        cell.margin_left = cell.margin_right = Inches(0.06); cell.margin_top = cell.margin_bottom = Inches(0.03)
        p = tf.paragraphs[0]
        if lb:
            r0 = p.add_run(); r0.text = lb + ' '; r0.font.bold = True; r0.font.size = Pt(sz); r0.font.name = 'Calibri'; r0.font.color.rgb = _tp_rgb('1B1B6B')
        r = p.add_run(); r.text = cd; r.font.size = Pt(sz); r.font.name = 'Calibri'; r.font.color.rgb = _tp_rgb(_TP['DARK'])
        for j in range(1, len(tbl.columns)): tbl.cell(i, j).text_frame.clear()
    shp.height = sum(r_.height for r_ in tbl.rows)

# ---------- pengumpul data per scope (dealer / koordinator / counter) ----------
def _ppt_scope(dlist, koor=None, label='', total_view=False):
    """Hitung ulang seluruh tabel tab STU untuk dealer(dlist) [+ filter koordinator/counter]. Memakai rumus & data yang sama dengan dashboard."""
    import math
    g = globals(); saved_td, saved_tk = load_json(TARGET_DEALER_FILE), load_json(TARGET_KOOR_FILE)
    cdn = 'Nama Dealer' if (g.get('df_nota') is not None and 'Nama Dealer' in g['df_nota'].columns) else 'Dealer'
    def _sel(df, dcol, kcol='Koor'):
        if df is None or not hasattr(df, 'empty') or df.empty or dcol not in df.columns: return pd.DataFrame()
        d_ = df[df[dcol].isin(dlist)]
        if koor and kcol in d_.columns: d_ = d_[d_[kcol] == koor]
        return d_
    f_c, f_l = _sel(g.get('df_fak_curr_g'), 'Nama Dealer'), _sel(g.get('df_fak_lm_g'), 'Nama Dealer')
    n_c, n_t = _sel(g.get('df_nota_curr_g'), cdn), _sel(g.get('df_nota_tod_g'), cdn)
    db_c = _sel(g.get('df_db_curr_g'), 'Nama_Dealer_Clean')
    s_c, s_l = _sel(g.get('df_spk_curr_g'), 'Nama_Dealer_SPK', 'Koor_SPK'), _sel(g.get('df_spk_lm_g'), 'Nama_Dealer_SPK', 'Koor_SPK')
    def _ek(df, e, col='Koor'): return df[df[col] == e] if (df is not None and not df.empty and col in df.columns) else pd.DataFrame()
    tgt = get_koor_or_counter_target(koor, saved_tk) if koor else sum(safe_int(saved_td.get(d_, 0)) for d_ in dlist)
    r_p = calc_sales_row(label, _TP_KN, tgt, f_c, f_l, n_c, n_t); r_pr = calc_proses_row(label, _TP_KN, db_c, len(f_c), n_c); r_s = calc_leasing_row(label, s_c, s_l, _TP_KN)
    # --- entitas (koordinator / counter) ---
    ents = set()
    for ds, dc, kc in [(g.get('df_faktur'), 'Nama Dealer', 'Koor'), (g.get('df_nota'), cdn, 'Koor'), (g.get('df_database'), 'Nama_Dealer_Clean', 'Koor'), (g.get('df_spk'), 'Nama_Dealer_SPK', 'Koor_SPK')]:
        d_ = _sel(ds, dc, kc)
        if not d_.empty and kc in d_.columns: ents.update(clean_str(v) for v in d_[kc].dropna().unique() if clean_str(v) and clean_str(v) not in ['NAN', 'NONE', 'LAIN-LAIN'])
    ent_p, ent_pr, ent_s = [], [], []
    for e in sorted(ents):
        rp_ = calc_sales_row(e, _TP_KN, get_koor_or_counter_target(e, saved_tk), _ek(f_c, e), _ek(f_l, e), _ek(n_c, e), _ek(n_t, e))
        if rp_['TARGET'] == 0 and rp_['STU'] == 0 and rp_['NOTA'] == 0 and rp_['STU LM'] == 0: continue
        ent_p.append(rp_); ent_pr.append(calc_proses_row(e, _TP_KN, _ek(db_c, e), len(_ek(f_c, e)), _ek(n_c, e))); ent_s.append(calc_leasing_row(e, _ek(s_c, e, 'Koor_SPK'), _ek(s_l, e, 'Koor_SPK'), _TP_KN))
    dlr_p, dlr_pr, dlr_s = [], [], []
    if total_view:
        for d_ in dlist:
            a_, b_ = _ek(f_c, d_, 'Nama Dealer'), _ek(f_l, d_, 'Nama Dealer'); c_, e_ = _ek(n_c, d_, cdn), _ek(n_t, d_, cdn); dbd = _ek(db_c, d_, 'Nama_Dealer_Clean')
            dlr_p.append(calc_sales_row(d_, _TP_KN, safe_int(saved_td.get(d_, 0)), a_, b_, c_, e_)); dlr_pr.append(calc_proses_row(d_, _TP_KN, dbd, len(a_), c_))
            dlr_s.append(calc_leasing_row(d_, _ek(s_c, d_, 'Nama_Dealer_SPK'), _ek(s_l, d_, 'Nama_Dealer_SPK'), _TP_KN))
    srt = lambda L, k1, k2: sorted(L, key=lambda r_: (r_[k1], r_[k2]))
    ent_p, ent_pr, ent_s = srt(ent_p, '_pct', 'STU'), srt(ent_pr, '_stu_pct', 'STU'), srt(ent_s, '_acc_val', '_spk_val')
    dlr_p, dlr_pr, dlr_s = srt(dlr_p, '_pct', 'STU'), srt(dlr_pr, '_stu_pct', 'STU'), srt(dlr_s, '_acc_val', '_spk_val')
    # --- salesman ---
    sm_set = set()
    for ds in [f_c, n_c, db_c, s_c]:
        if not ds.empty and 'Salesman_Clean' in ds.columns: sm_set.update(clean_str(v) for v in ds['Salesman_Clean'].dropna().unique() if clean_str(v) and clean_str(v) not in ['NAN', 'NONE', 'TOTAL'])
    dbg = g.get('df_database'); act_sm = set(clean_str(v) for v in dbg['Salesman_Clean'].dropna().unique() if clean_str(v)) if (dbg is not None and not dbg.empty and 'Salesman_Clean' in dbg.columns) else set()
    if act_sm: sm_set = {m for m in sm_set if m in act_sm}
    def _cnt(df, col, val): return int((df[col] == val).sum()) if (df is not None and not df.empty and col in df.columns) else 0
    sm_rows = []
    for sm in sorted(sm_set):
        tg_ = get_auto_sales_target(sm); stu = _cnt(f_c, 'Salesman_Clean', sm); sdb = _ek(db_c, sm, 'Salesman_Clean'); sspk = _ek(s_c, sm, 'Salesman_Clean')
        db_v, hot_v = len(sdb), _cnt(sdb, 'Data_Type_Clean', 'DEAL PROSPECT'); acc_v, rej_v = _cnt(sspk, 'Status_SPK_Clean', 'DISETUJUI'), _cnt(sspk, 'Status_SPK_Clean', 'DITOLAK')
        sm_rows.append(dict(name=sm, target=tg_, stu=stu, pct=(stu / tg_ * 100.0 if tg_ > 0 else 0.0), avg_db=db_v / hari_aktif, db=db_v, hot=hot_v, dbq=(hot_v / db_v * 100.0 if db_v else 0.0),
                            spk=len(sspk), acc=(acc_v / (acc_v + rej_v) * 100.0 if (acc_v + rej_v) else 0.0), stup=(stu / db_v * 100.0 if db_v else 0.0), lm=_cnt(_ek(f_l, sm, 'Salesman_Clean'), 'Salesman_Clean', sm)))
    sm_rows.sort(key=lambda r_: (0 if r_['target'] > 0 else 1, r_['pct'], r_['stu'], r_['db']))
    mp_tgt_dlr, mp_tgt_sm = safe_int(r_p['TARGET']), int(sum(r_['target'] for r_ in sm_rows))
    mp_raw = (mp_tgt_dlr * 1.25 - mp_tgt_sm) / 2.0; mp_need = int(-(-mp_raw // 1)) if mp_raw > 0 else 0
    ideal_ = float(target_persen_harian)
    mp = dict(need=mp_need, tgt_dlr=mp_tgt_dlr, tgt_sm=mp_tgt_sm, raw=max(mp_raw, 0.0), grn=[r_['name'] for r_ in sm_rows if r_['target'] > 0 and r_['pct'] >= ideal_], red=[r_['name'] for r_ in sm_rows if r_['target'] > 0 and r_['pct'] < 0.85 * ideal_])
    # --- aktivitas (prospect type & sub prospect) ---
    def _act_vals(d_x):
        n_db = len(d_x); n_hot = int(d_x['Data_Type_Clean'].isin(['DEAL PROSPECT', 'SEAL PROSPECT']).sum()) if (n_db and 'Data_Type_Clean' in d_x.columns) else 0
        n_stu = int((d_x['STU'] == 'STU').sum()) if (n_db and 'STU' in d_x.columns) else 0
        return n_db, n_hot, (n_hot / n_db * 100.0 if n_db else 0.0), n_stu, (n_stu / n_db * 100.0 if n_db else 0.0)
    _cnt_v = bool(koor) and str(koor).upper().startswith('COUNTER')   # khusus COUNTER: Walk In 60% | Media Sosial 30% | Event 7% | Service 3%
    ACT = {'WALKIN': 60, 'MEDIASOSIAL': 30, 'EVENT': 7, 'SERVICE': 3} if _cnt_v else {'EVENT': 30, 'WALKIN': 20, 'MEDIASOSIAL': 25, 'EVENTNONNACS': 10, 'KUNJUNGANRUMAH': 10, 'SERVICE': 5}
    def _pct_of(nm):
        a_n = alnum_str(nm)
        if a_n in ACT: return ACT[a_n]
        if 'SOSMED' in a_n or 'MEDIASOSIAL' in a_n: return 30 if _cnt_v else 25
        if _cnt_v: return None
        return 10 if 'NONNACS' in a_n else None
    tgt_all = safe_int(r_p['TARGET'])
    def _act_rows(db_df, grp, with_target):
        out = []
        for gv in grp.unique():
            v = _act_vals(db_df[grp == gv]); r_ = dict(name=str(gv), db=v[0], hot=v[1], dbq=v[2], stu=v[3], stup=v[4], target=None, ach=None)
            if with_target:
                pc = _pct_of(gv)
                if pc is not None: r_['target'] = int(round(tgt_all * pc / 100.0)); r_['ach'] = (v[3] / r_['target'] * 100.0 if r_['target'] > 0 else 0.0)
            out.append(r_)
        out.sort(key=lambda r_: -r_['db'])
        v = _act_vals(db_df); tot = dict(name='TOTAL', db=v[0], hot=v[1], dbq=v[2], stu=v[3], stup=v[4], target=(tgt_all if with_target else None), ach=((v[3] / tgt_all * 100.0 if tgt_all > 0 else 0.0) if with_target else None))
        return out, tot
    pt_rows, pt_tot, sub_rows, sub_tot, sub_pt = [], None, [], None, ''
    if db_c is not None and not db_c.empty and db_c.shape[1] >= 8:
        pt_g = db_c.iloc[:, 6].apply(lambda v: clean_str(v) or '(KOSONG)'); st_g = db_c.iloc[:, 7].apply(lambda v: strip_sub_code(v) or '(KOSONG)')
        pt_rows, pt_tot = _act_rows(db_c, pt_g, True)
        if pt_rows:
            sub_pt = next((r_['name'] for r_ in pt_rows if alnum_str(r_['name']) == 'EVENT'), pt_rows[0]['name'])
            sub_rows, sub_tot = _act_rows(db_c[pt_g == sub_pt], st_g[pt_g == sub_pt], False)
    # --- per type ---
    TYPE_TGT = {'AT PREMIUM': 60, 'AT CLASSY': 25, 'AT LPM': 10, 'MOPED': 3, 'OTHERS': 2}
    def _grp(d_x):
        if d_x is None or d_x.empty or 'Type_Clean' not in d_x.columns: return d_x
        d_x = d_x.copy(); d_x['Type_Grp'] = d_x['Type_Clean'].where(d_x['Type_Clean'].isin(['AT PREMIUM', 'AT CLASSY', 'AT LPM', 'MOPED']), 'OTHERS'); return d_x
    typ_rows, typ_tot = [], None
    try:
        df_ti, df_tf, _t = build_type_table(_grp(f_c), _grp(f_l), _grp(db_c), _grp(s_c), [(k, 'Type_Grp', 'Type_Grp', 'Type_Grp') for k in TYPE_TGT], 'TYPE', tot_label='TOTAL', sort_by_growth=True)
        if df_tf is not None and not df_tf.empty:
            for r_ in df_tf.to_dict('records'):
                tg_ = tgt_all if r_['TYPE'] == 'TOTAL' else int(round(tgt_all * TYPE_TGT.get(r_['TYPE'], 0) / 100.0)); r_['TARGET'] = tg_; r_['ACHIEV'] = (safe_int(r_['STU']) / tg_ * 100.0 if tg_ > 0 else 0.0)
                (typ_rows if r_['TYPE'] != 'TOTAL' else []).append(r_)
                if r_['TYPE'] == 'TOTAL': typ_tot = r_
    except Exception: pass
    lsg_stu = {l_: _cnt(f_c, 'Leasing_Faktur_Clean', l_) for l_ in ['BAF', 'ADIRA', 'SOF', 'IMFI', 'CASH']}
    return dict(label=label, koor=koor, dlist=list(dlist), total_view=total_view, tgt=tgt, r_p=r_p, r_pr=r_pr, r_s=r_s, nm=nota_metrics(n_c), ent_p=ent_p, ent_pr=ent_pr, ent_s=ent_s, dlr_p=dlr_p, dlr_pr=dlr_pr, dlr_s=dlr_s,
                sm=sm_rows, mp=mp, pt=pt_rows, pt_tot=pt_tot, sub=sub_rows, sub_tot=sub_tot, sub_pt=sub_pt, typ=typ_rows, typ_tot=typ_tot, lsg_stu=lsg_stu)

# ---------- teks analisa otomatis (Condition / Problem / Perbaikan / Target / Time line) ----------
def _tx_ceil(x): import math; return int(math.ceil(x - 1e-9))

def _tx_weak(rows, key, name='name', minv=1, dbkey=None, low=True):
    c = [r_ for r_ in rows if (r_.get(dbkey, 1) if dbkey else 1) >= minv]
    if not c: return None
    return sorted(c, key=lambda r_: bsg_num(r_[key]))[0] if low else sorted(c, key=lambda r_: -bsg_num(r_[key]))[0]

def _tx_conv(db, hot, stu, avg=None, mp=None):
    """Return (problem_text, fix_text) dari DB Q (ideal 30%), STU% (ideal 10%) & produktivitas input DB."""
    if db <= 0: return 'Belum ada data Database pada periode ini.', 'Pastikan prospek harian diinput ke Database & file Database terbaru sudah diunggah.'
    dbq, stup = hot / db * 100.0, stu / db * 100.0; pr, fx = [], []
    if dbq < 30: pr.append(f'DB Q {dbq:.1f}% < 30% (kurang {max(0, _tx_ceil(db * 0.3) - hot)} HOT)'); fx.append('saring prospek saat input (kebutuhan, budget, rencana beli) agar HOT naik')
    if stup < 10: pr.append(f'STU% {stup:.1f}% < 10% (kurang {max(0, _tx_ceil(db * 0.1) - stu)} STU dari DB yang ada)'); fx.append('follow-up seluruh HOT maks 1x24 jam hingga SPK/STU')
    if avg is not None and avg < 4: pr.append(f'input {avg:.1f} DB/MP/hari < standar 4'); fx.append('disiplin input minimal 4 DB/MP/hari')
    if not pr: return 'DB Q dan STU% sudah memenuhi standar (30% / 10%).', 'Pertahankan kualitas database & disiplin follow-up HOT.'
    return '; '.join(pr) + '.', (fx[0][0].upper() + fx[0][1:] + ('; ' + '; '.join(fx[1:]) if len(fx) > 1 else '') + '.')

def _tx_rows_dbq(S, H, addf, ai):
    pr = S['r_pr']; db, hot, stu, mp = int(pr['DB']), int(pr['HOT']), int(pr['STU']), int(pr['JML MP'])
    avg = bsg_num(pr['AVG DB']); dbq, stup = bsg_num(pr['DB Q']), bsg_num(pr['STU%'])
    cond = f'DB {db:,} dengan HOT {hot:,} → DB Q {dbq:.1f}% (ideal 30%); STU {stu:,} → STU% {stup:.1f}% (ideal 10%); input {avg:.1f} DB/MP/hari oleh {mp} MP.'
    wk = _tx_weak(S['ent_pr'], 'DB Q', minv=1, dbkey='DB'); prb, fx = _tx_conv(db, hot, stu, avg if mp else None)
    if wk and len(S['ent_pr']) > 1: fx += f" Prioritas pendampingan: {wk['KOOR / COUNTER / DEALER']} (DB Q {bsg_num(wk['DB Q']):.1f}%)."
    try: add, goal = addf(int(S['r_p']['STU']), int(S['r_p']['TARGET']), max(1, mp))
    except Exception: add, goal = 0, 0
    tg = f'DB Q ≥30%, STU% ≥10%' + (f', tambahan +{add} STU (menjadi {goal} unit)' if add > 0 else '')
    return [('DB QUALITY:', cond, prb, ai.get('plan_dealer') or fx, tg, H['timeline'])]

def _tx_sub_row(lbl, pt, r_, H, addf=None):
    db, hot, stu = r_['db'], r_['hot'], r_['stu']; prb, fx = _tx_conv(db, hot, stu)
    miss = max(0, _tx_ceil(db * 0.1) - stu)
    return (lbl, f"{pt} › {r_['name']}: DB {db}, HOT {hot} (DB Q {r_['dbq']:.1f}%), STU {stu} (STU% {r_['stup']:.1f}%).", prb, fx, f'STU% ≥10%' + (f' (+{miss} STU)' if miss > 0 else '') + ', DB Q ≥30%', H['timeline'])

def _tx_rows_sub(S, H, ai, n=3, extra_media=False):
    rows_ = [r_ for r_ in S['sub'] if r_['db'] > 0]
    rows_.sort(key=lambda r_: (-(max(0, _tx_ceil(r_['db'] * 0.1) - r_['stu'])), -r_['db']))   # potensi STU yang belum terkonversi terbesar dulu
    out = [_tx_sub_row(f'SUB PROSPECT {i}:', S['sub_pt'], r_, H) for i, r_ in enumerate(rows_[:n], 1)]
    if ai.get('plan_md'): out = [(a, b, c, ai['plan_md'], e, f) for (a, b, c, d, e, f) in out]
    while len(out) < n: out.append((f'SUB PROSPECT {len(out) + 1}:', 'Tidak ada data sub prospect pada periode ini.', '-', '-', '-', '-'))
    if extra_media:
        mr = next((r_ for r_ in S['pt'] if 'SOSMED' in alnum_str(r_['name']) or 'MEDIASOSIAL' in alnum_str(r_['name'])), None)
        if mr:
            prb, fx = _tx_conv(mr['db'], mr['hot'], mr['stu']); tg_s = f" target {mr['target']} STU, realisasi {mr['stu']} ({mr['ach']:.1f}%)" if mr['target'] else f" STU {mr['stu']}"
            out.append(('MEDIA SOSIAL:', f"Media Sosial:{tg_s}; DB {mr['db']}, HOT {mr['hot']} (DB Q {mr['dbq']:.1f}%), STU% {mr['stup']:.1f}%.", prb, ai.get('plan_sosmed') or fx, 'DB Q ≥30%, STU% ≥10%, capai target kontribusi Media Sosial', H['timeline']))
        else: out.append(('MEDIA SOSIAL:', 'Tidak ada data Media Sosial pada periode ini.', '-', '-', '-', '-'))
    return out

def _tx_rows_salesman(S, H, addf, ai, n=3):
    out = []
    cand = [r_ for r_ in S['sm'] if r_['target'] > 0 and not is_counter_personnel(r_['name'])] or [r_ for r_ in S['sm'] if r_['target'] > 0]
    for i, r_ in enumerate(cand[:n], 1):
        prb, fx = _tx_conv(r_['db'], r_['hot'], r_['stu'], r_['avg_db'])
        try: add, goal = addf(r_['stu'], r_['target'], 1)
        except Exception: add, goal = 1, r_['stu'] + 1
        cond = f"{r_['name']}: STU {r_['stu']}/{r_['target']} ({r_['pct']:.1f}%), DB {r_['db']}, HOT {r_['hot']} (DB Q {r_['dbq']:.1f}%), SPK {r_['spk']}, ACC {r_['acc']:.0f}%, STU% {r_['stup']:.1f}%."
        out.append((f'SALESMAN TERBURUK {i}:', cond, prb, ai.get('plan_mp') or ('Coaching harian & ride-along; ' + fx[:1].lower() + fx[1:]), f'+{add} STU (menjadi {goal}), +{max(add, 1) * 2} SPK, DB Q ≥30%', H['timeline']))
    while len(out) < n: out.append((f'SALESMAN TERBURUK {len(out) + 1}:', 'Tidak ada data salesman.', '-', '-', '-', '-'))
    return out

def _tx_rows_counter(S, H, addf, ai):
    out = _tx_rows_dbq(S, H, addf, {}); out = [('DB QUALITY:',) + out[0][1:]]
    if ai.get('plan_counter'): out = [(out[0][0], out[0][1], out[0][2], ai['plan_counter'], out[0][4], out[0][5])]
    rows_ = [r_ for r_ in S['pt'] if r_['db'] > 0]; rows_.sort(key=lambda r_: (-(max(0, _tx_ceil(r_['db'] * 0.1) - r_['stu'])), -r_['db']))
    for i, r_ in enumerate(rows_[:3], 1):
        prb, fx = _tx_conv(r_['db'], r_['hot'], r_['stu']); ach_t = (f"; capaian {r_['stu']}/{r_['target']} ({r_['ach']:.0f}%) dari target aktivitas" if r_['target'] else '')
        out.append((f'AKTIFITAS {i}:', f"{r_['name']}: DB {r_['db']}, HOT {r_['hot']} (DB Q {r_['dbq']:.1f}%), STU {r_['stu']} (STU% {r_['stup']:.1f}%){ach_t}.", prb, fx, 'DB Q ≥30%, STU% ≥10%', H['timeline']))
    while len(out) < 4: out.append((f'AKTIFITAS {len(out)}:', 'Tidak ada data aktivitas pada periode ini.', '-', '-', '-', '-'))
    return out

def _tx_rows_leasing(S, H, ai):
    out = []; rs = S['r_s']; tot_stu = max(1, sum(S['lsg_stu'].get(l_, 0) for l_ in ['BAF', 'ADIRA', 'SOF', 'IMFI']))
    accs = {l_: bsg_num(rs.get(f'%{l_}')) for l_ in ['BAF', 'ADIRA', 'SOF'] if rs.get(l_, 0) > 0}; best = max(accs, key=accs.get) if accs else None
    for l_ in ['BAF', 'ADIRA', 'SOF']:
        n_spk, acc, stu_l = int(rs.get(l_, 0)), bsg_num(rs.get(f'%{l_}')), S['lsg_stu'].get(l_, 0)
        cond = f'SPK {n_spk}, ACC% {acc:.1f}% (ideal 50%), STU kredit cair {stu_l} unit ({stu_l / tot_stu * 100:.0f}% dari STU kredit); total SPK dealer {int(rs["SPK"])} (LM {int(rs["SPK LM"])}), SPK terbuka {int(rs["TERBUKA"])}.'
        if n_spk == 0: prb, fx, tg = f'Belum ada SPK ke {l_} pada periode ini.', f'Ajukan SPK ke {l_} untuk konsumen yang sesuai kriteria; bandingkan ACC% antar leasing.', f'≥{max(3, int(rs["SPK"] * 0.2))} SPK ke {l_}, ACC% ≥50%'
        elif acc < 50: prb, fx, tg = f'ACC% {acc:.1f}% < 50% — banyak SPK ditolak/terbuka.', f'Pra-screening konsumen sesuai kriteria {l_}, lengkapi dokumen, selesaikan SPK terbuka 1x24 jam' + (f'; alihkan ke {best} (ACC% {accs[best]:.0f}%)' if best and best != l_ else '') + '.', 'ACC% ≥50%, tutup SPK terbuka'
        else: prb, fx, tg = f'ACC% {acc:.1f}% sudah ≥50%.', f'Pertahankan kualitas pengajuan ke {l_}; tambah volume SPK.', 'Pertahankan ACC% ≥50%, tambah SPK'
        out.append((f'LEASING {l_}:', cond, prb, ai.get('plan_leasing') or fx, tg, H['timeline']))
    return out

def _tx_rows_type(S, H, addf, ai):
    out = []
    for t_ in ['AT PREMIUM', 'AT CLASSY', 'AT LPM']:
        r_ = next((x for x in S['typ'] if x['TYPE'] == t_), None)
        if not r_: out.append((f'{t_}:', 'Tidak ada data pada periode ini.', '-', '-', '-', '-')); continue
        tg_, stu, lm, ach = int(r_['TARGET']), int(r_['STU']), int(r_['STU LM']), float(r_['ACHIEV']); ideal = float(target_persen_harian)
        prb, fx = _tx_conv(int(r_['DB']), int(r_['HOT']), stu)
        pr_all = []
        if ach < ideal: pr_all.append(f'capaian {ach:.1f}% di bawah ideal {ideal:.0f}% (gap {stu - round(tg_ * ideal / 100.0):+d} unit)')
        if stu < lm: pr_all.append(f'turun vs LM ({stu} vs {lm})')
        pr_all.append(prb)
        try: add, goal = addf(stu, tg_, 1)
        except Exception: add, goal = 0, stu
        cond = f"STU {stu}/{tg_} ({ach:.1f}%), LM {lm} (growth {bsg_num(r_['GROWTH']):.0f}%), DB {int(r_['DB'])}, HOT {int(r_['HOT'])} (DB Q {bsg_num(r_['DB Q']):.1f}%), SPK {int(r_['SPK'])}, ACC {bsg_num(r_['ACC%']):.0f}%."
        out.append((f'{t_}:', cond, ' '.join(pr_all), ai.get('plan_type') or fx, (f'+{add} STU (menjadi {goal} unit), ' if add > 0 else '') + 'DB Q ≥30%, STU% ≥10%', H['timeline']))
    return out

# ---------- slide builders ----------
def _tp_f1(v): return f'{bsg_num(v):.1f}%'

def _tp_cards_block(sl, S, y1=0.7):
    p, pr, s, nm = S['r_p'], S['r_pr'], S['r_s'], S['nm']; tg = safe_int(p['TARGET'])
    dbq, stup, acc, grw = bsg_num(pr['DB Q']), bsg_num(pr['STU%']), bsg_num(s['ACC%']), bsg_num(s['GROWTH']); ideal = float(target_persen_harian)
    forc_pct = (p['FORC'] / tg * 100.0) if tg > 0 else 0.0
    c1 = [('Target vs Result', f"{tg:,} / {int(p['STU']):,}", f"{p['_pct']:.1f}% tercapai · Ideal {int(p['IDEAL']):,}", _TP['GRN'] if p['STU'] >= p['IDEAL'] else (_TP['YEL'] if p['_pct'] >= 0.85 * ideal else _TP['RED']), min(p['_pct'], 100.0)),
          ('DB Quality', pr['DB Q'], f"HOT {int(pr['HOT']):,} / {int(pr['DB']):,} DB · ideal 30%", _tp_col(dbq, 30, 20), min(dbq / 30.0 * 100, 100)),
          ('STU Ratio', pr['STU%'], f"STU {int(pr['STU']):,} / {int(pr['DB']):,} DB · ideal 10%", _tp_col(stup, 10, 7), min(stup / 10.0 * 100, 100)),
          ('SPK Growth', s['GROWTH'], f"SPK {int(s['SPK']):,} vs LM {int(s['SPK LM']):,}", _tp_col(grw, 100, 80), min(grw, 100)),
          ('Approval', s['ACC%'], f"SPK terbuka {int(s['TERBUKA']):,} · ideal 50%", _tp_col(acc, 50, 40), min(acc / 50.0 * 100, 100)),
          ('Forecast Closing', f"{int(p['FORC']):,} unit", f'{forc_pct:.1f}% dari target', _tp_col(forc_pct, 100, 85), min(forc_pct, 100))]
    def _fc(v):
        if tg <= 0: return ('target belum diatur', _TP['CYN'])
        pc = v / tg * 100.0; return (f'{pc:.1f}% dari target', _tp_col(pc, 100, 85))
    c2 = [('AVG NOTA', f"{nm['avg']:.2f} /hari", f"{nm['n']:,} nota ÷ {hari_aktif} hari kerja", _TP['CYN'], None),
          ('AVG 7', f"{nm['avg7']:.2f} /hari", f"{nm['n7']:,} nota ÷ {nm['d7']} hari kerja", _TP['GRN'] if nm['avg7'] >= nm['avg'] else _TP['RED'], None),
          ('AVG 3', f"{nm['avg3']:.2f} /hari", f"{nm['n3']:,} nota ÷ {nm['d3']} hari kerja", _TP['GRN'] if nm['avg3'] >= nm['avg'] else _TP['RED'], None),
          ('Forc Nota', f"{nm['f_nota']:,}", _fc(nm['f_nota'])[0], _fc(nm['f_nota'])[1], None), ('Forc 7', f"{nm['f7']:,}", _fc(nm['f7'])[0], _fc(nm['f7'])[1], None), ('Forc 3', f"{nm['f3']:,}", _fc(nm['f3'])[0], _fc(nm['f3'])[1], None)]
    for i, (lb, v, sb, c, pg) in enumerate(c1): _tp_card(sl, 0.25 + i * 2.15, y1, 2.05, 0.95, lb, str(v), sb, c, pg)
    for i, (lb, v, sb, c, pg) in enumerate(c2): _tp_card(sl, 0.25 + i * 2.15, y1 + 1.02, 2.05, 0.78, lb, str(v), sb, c, pg)
    return y1 + 1.85

def _tp_focus_block(sl, S, y, lines, h=None):
    h = h or (0.18 + 0.2 * len(lines)); _tp_rect(sl, 0.25, y, 12.8, h, fill='1B1030', line=_TP['LINE']); _tp_rect(sl, 0.25, y, 0.06, h, fill='FF007F')
    _tp_text(sl, 0.42, y + 0.07, 12.5, h - 0.1, [[(a + ' ', {'b': True, 'c': _TP['WHT']}), (b, {'c': _TP['TXT']})] for a, b in lines], size=9)
    return y + h + 0.08

def _tp_row_style(pct, ideal): return _tp_col(pct, ideal, 0.85 * ideal)

def _tp_sales_table(sl, S, y, rows=None, size=9, maxn=8):
    p = S['r_p']; ideal = float(target_persen_harian); base = rows if rows is not None else (S['dlr_p'] if S['total_view'] else S['ent_p'])
    hdr = ['KOOR / COUNTER / DEALER', 'TARGET', 'IDEAL', 'STU', 'NOTA', 'TODAY', col_name_pct, 'GAP TG', 'STU LM', '% LM', 'GAP LM', 'FORC', 'FORC 7']
    body = []
    for r_ in (base[:maxn - 1] + [p]):
        body.append([r_[_TP_KN], r_['TARGET'], r_['IDEAL'], r_['STU'], r_['NOTA'], r_['TODAY'], dict(t=f"{r_['_pct']:.2f}%", c=_tp_row_style(r_['_pct'], ideal), b=True), dict(t=f"{r_['GAP TG']:+d}", c=_TP['GRN'] if r_['GAP TG'] >= 0 else _TP['RED']),
                     r_['STU LM'], r_['% LM'], dict(t=f"{r_['GAP LM']:+d}", c=_TP['GRN'] if r_['GAP LM'] >= 0 else _TP['RED']), r_['FORC'], r_['FORC 7']])
    body[-1][0] = S['label'] if str(S['label']).startswith('TOTAL') else f"TOTAL ({S['label']})"
    return _tp_table(sl, 0.25, y, [2.8] + [0.833] * 12, hdr, body, size=size)

def _tp_focus_lines(S):
    lines = []; ideal = int(round(float(target_persen_harian)))
    def _ln(r_): return f"{r_[_TP_KN]} — Target {int(r_['TARGET'])} / STU {int(r_['STU'])} ({r_['_pct']:.2f}% | Ideal {int(r_['IDEAL'])})"
    if not S['koor']:
        ko = sorted([r_ for r_ in S['ent_p'] if not str(r_[_TP_KN]).upper().startswith('COUNTER') and r_['TARGET'] > 0], key=lambda r_: (r_['_pct'], r_['STU']))
        cn = sorted([r_ for r_ in S['ent_p'] if str(r_[_TP_KN]).upper().startswith('COUNTER') and r_['TARGET'] > 0], key=lambda r_: (r_['_pct'], r_['STU']))
        lines.append(('Focus Koordinator :', _ln(ko[0]) if ko else 'Tidak ada data Koordinator'))
        lines.append(('Focus Counter :', _ln(cn[0]) if cn else 'Tidak ada data Counter'))
    sm = [r_ for r_ in S['sm'] if r_['target'] > 0 and not is_counter_personnel(r_['name']) and r_['db'] >= 2]
    lines.append(('Focus Salesman :', (f"{sm[0]['name']} — Target {sm[0]['target']} / STU {sm[0]['stu']} ({sm[0]['pct']:.2f}% | Ideal {round(sm[0]['target'] * persen_desimal)}) | DB {sm[0]['db']} (DB Q {sm[0]['dbq']:.2f}%)") if sm else 'Tidak ada salesman aktif (non-counter dengan DB ≥ 2)'))
    if not S['koor']: lines.append(('Kebutuhan Salesman Baru :', (f"{S['mp']['need']} orang" if S['mp']['need'] > 0 else 'Tidak diperlukan') + f" — (Target {S['mp']['tgt_dlr']:,} × 125% − Target total MP {S['mp']['tgt_sm']:,}) ÷ 2 = {S['mp']['raw']:.2f}"))
    return lines

def _tp_detail_slide(sl, S, title, D):
    _tp_clear_pictures(sl); _tp_ensure_bg(sl)
    _tp_head(sl, f"DETAIL DEALER — {title} (1 s.d {D['tgl']} {D['periode']})", None, w=7.3 if S['koor'] else 12.8)
    y = _tp_cards_block(sl, S, y1=0.68); y = _tp_focus_block(sl, S, y, _tp_focus_lines(S))
    _tp_text(sl, 0.25, y, 12.8, 0.22, f"1. TABEL PENJUALAN — {title}", size=10, color=_TP['CYN'], bold=True)
    _tp_sales_table(sl, S, y + 0.26)
    base = S['dlr_p'] if S['total_view'] else S['ent_p']; n_rows = min(8, len(base) + 1); yc = y + 0.26 + 0.3 + 0.26 * n_rows + 0.2; ideal = float(target_persen_harian)
    if S['koor']: pts_ = [(r_['name'], r_['pct']) for r_ in S['sm'] if r_['target'] > 0][:6]; ttl = 'Pencapaian Salesman vs Target (%)'
    else: pts_ = [(r_[_TP_KN], r_['_pct']) for r_ in base if r_['TARGET'] > 0][:6]; ttl = ('Pencapaian Dealer vs Target (%)' if S['total_view'] else 'Pencapaian Koordinator / Counter vs Target (%)')
    if pts_ and yc < 6.6: _tp_bar(sl, 0.25, yc, 8.5, min(7.3 - yc, 0.3 + 0.27 * len(pts_) + 0.1), f'{ttl} — ideal {ideal:.0f}%', [a for a, _ in pts_], [b for _, b in pts_], ideal, fmt='0"%"')

def _tp_salesman_table(sl, S, y, maxn, size=9):
    ideal = float(target_persen_harian); tot_ = S['sm']; rows = []
    for r_ in tot_[:maxn]: rows.append([r_['name'], r_['target'], r_['stu'], dict(t=f"{r_['pct']:.2f}%", c=_tp_row_style(r_['pct'], ideal) if r_['target'] > 0 else _TP['MUT'], b=True), f"{r_['avg_db']:.2f}", r_['db'], r_['hot'], dict(t=f"{r_['dbq']:.2f}%", c=_tp_col(r_['dbq'], 30, 20)), r_['spk'], f"{r_['acc']:.2f}%", dict(t=f"{r_['stup']:.2f}%", c=_tp_col(r_['stup'], 10, 7))])
    tt, ts, tdb, thot, tspk = sum(r_['target'] for r_ in tot_), sum(r_['stu'] for r_ in tot_), sum(r_['db'] for r_ in tot_), sum(r_['hot'] for r_ in tot_), sum(r_['spk'] for r_ in tot_)
    rows.append([f"TOTAL ({len(tot_)} MP)", tt, ts, f"{(ts / tt * 100 if tt else 0):.2f}%", f"{(tdb / max(1, len(tot_)) / hari_aktif):.2f}", tdb, thot, f"{(thot / tdb * 100 if tdb else 0):.2f}%", tspk, '', f"{(ts / tdb * 100 if tdb else 0):.2f}%"])
    return _tp_table(sl, 0.25, y, [2.9] + [0.99] * 10, ['NAMA SALESMAN', 'TARGET', 'STU', col_name_pct, 'AVG DB', 'DB', 'HOT', 'DB Q', 'SPK', 'ACC%', 'STU%'], rows, row_h=0.26, size=size)

def _tp_manpower_slide(sl, S, title, D):
    _tp_clear_pictures(sl); _tp_ensure_bg(sl)
    _tp_head(sl, f"5. TABEL DETAIL SALESMAN — {title} (1 s.d {D['tgl']} {D['periode']})", None, w=7.3)
    if not S['sm']: _tp_text(sl, 0.25, 1.0, 12.8, 0.4, 'Belum ada data salesman pada periode ini.', size=14, color=_TP['MUT']); return
    tbl = _tp_salesman_table(sl, S, 0.68, 16, size=9); y = 0.68 + 0.3 + 0.26 * (min(16, len(S['sm'])) + 1) + 0.15
    mp = S['mp']; red = ', '.join(mp['red']) or '—'; grn = ', '.join(mp['grn']) or '—'
    ln = [('1.', (f"Dengan kondisi man power saat ini dibutuhkan {mp['need']} orang salesman baru untuk mencapai target {title} bulan ini." if mp['need'] > 0 else f"Dengan kondisi man power saat ini tidak diperlukan penambahan salesman baru untuk mencapai target {title} bulan ini.")),
          ('2.', f"Maksimalkan salesman grade hijau (≥ ideal {float(target_persen_harian):.0f}%): {grn}."), ('3.', f"Pendampingan & coaching harian untuk grade merah (< 85% ideal): {red}.")]
    _tp_focus_block(sl, S, min(y, 6.0), ln, h=0.2 + 0.24 * 3)

def _tp_process_slide(sl, S, title, D):
    _tp_clear_pictures(sl); _tp_ensure_bg(sl)
    _tp_head(sl, f"2. TABEL PROSES KERJA — {title} (1 s.d {D['tgl']} {D['periode']})", None, w=7.0)
    base = S['dlr_pr'] if S['total_view'] else S['ent_pr']; pr = S['r_pr']; show = base[:6]
    body = [[r_[_TP_KN], r_['JML MP'], r_['AVG DB'], r_['DB'], r_['HOT'], dict(t=r_['DB Q'], c=_tp_col(bsg_num(r_['DB Q']), 30, 20), b=True), r_['STU'], dict(t=r_['STU%'], c=_tp_col(bsg_num(r_['STU%']), 10, 7), b=True), r_['AVG NOTA'], r_['AVG 7'], r_['AVG 3']] for r_ in show]
    body.append([f"TOTAL ({title})", pr['JML MP'], pr['AVG DB'], pr['DB'], pr['HOT'], pr['DB Q'], pr['STU'], pr['STU%'], pr['AVG NOTA'], pr['AVG 7'], pr['AVG 3']])
    _tp_table(sl, 0.25, 0.62, [2.9] + [0.99] * 10, ['KOOR / COUNTER / DEALER', 'JML MP', 'AVG DB', 'DB', 'HOT', 'DB Q', 'STU', 'STU%', 'AVG NOTA', 'AVG 7', 'AVG 3'], body, row_h=0.25, size=9)
    yb = 0.62 + 0.3 + 0.25 * len(body) + 0.12
    cats = [r_[_TP_KN] for r_ in show]; _tp_bar(sl, 0.25, yb, 6.3, 4.5 - yb, 'Database Quality (DB Q) — ideal 30%', cats, [bsg_num(r_['DB Q']) for r_ in show], 30.0)
    _tp_bar(sl, 6.75, yb, 6.3, 4.5 - yb, 'STU Ratio (STU%) — ideal 10%', cats, [bsg_num(r_['STU%']) for r_ in show], 10.0)

def _tp_activity_blocks(sl, S, y0, with_sub=True, max_pt=8):
    ideal_pct = float(target_persen_harian)
    if not S['pt']: _tp_text(sl, 0.25, y0, 12.8, 0.4, 'Belum ada data Database / Prospect Type pada periode ini.', size=13, color=_TP['MUT']); return
    _tp_text(sl, 0.25, y0, 8.0, 0.22, 'TABEL REVIEW AKTIFITAS (PROSPECT TYPE)', size=10, color=_TP['CYN'], bold=True)
    pts = S['pt'][:max_pt]; body = [[r_['name'], ('-' if r_['target'] is None else r_['target']), r_['stu'], dict(t=('-' if r_['ach'] is None else f"{r_['ach']:.2f}%"), c=(_TP['MUT'] if r_['ach'] is None else _tp_row_style(r_['ach'], ideal_pct)), b=True), r_['db'], r_['hot'], dict(t=f"{r_['dbq']:.2f}%", c=_tp_col(r_['dbq'], 30, 20)), dict(t=f"{r_['stup']:.2f}%", c=_tp_col(r_['stup'], 10, 7))] for r_ in pts]
    t_ = S['pt_tot']; body.append([f"TOTAL", t_['target'], t_['stu'], f"{t_['ach']:.2f}%", t_['db'], t_['hot'], f"{t_['dbq']:.2f}%", f"{t_['stup']:.2f}%"])
    _tp_table(sl, 0.25, y0 + 0.25, [2.4, 0.8, 0.7, 0.9, 0.7, 0.7, 0.9, 0.9], ['PROSPECT TYPE', 'TARGET', 'STU', 'ACHIEV', 'DB', 'HOT', 'DB Q', 'STU%'], body, row_h=0.25, size=9)
    ach_rows = [r_ for r_ in pts if r_['ach'] is not None]
    _tp_bar(sl, 8.4, y0, 4.65, 0.3 + 0.28 * max(2, len(ach_rows)) + 0.3, 'Pencapaian Aktifitas vs Target (%)', [r_['name'] for r_ in ach_rows], [r_['ach'] for r_ in ach_rows], ideal_pct)
    if not with_sub: return
    y1 = y0 + 0.25 + 0.3 + 0.25 * len(body) + 0.25
    _tp_text(sl, 0.25, y1, 8.0, 0.22, f"TABEL SUB PROSPECT TYPE — {S['sub_pt']}", size=10, color=_TP['CYN'], bold=True)
    subs = S['sub'][:8]; b2 = [[r_['name'], r_['db'], r_['hot'], dict(t=f"{r_['dbq']:.2f}%", c=_tp_col(r_['dbq'], 30, 20)), r_['stu'], dict(t=f"{r_['stup']:.2f}%", c=_tp_col(r_['stup'], 10, 7), b=True)] for r_ in subs]
    t2 = S['sub_tot']; b2.append([f"TOTAL {S['sub_pt']}", t2['db'], t2['hot'], f"{t2['dbq']:.2f}%", t2['stu'], f"{t2['stup']:.2f}%"])
    _tp_table(sl, 0.25, y1 + 0.25, [4.0, 0.8, 0.8, 0.9, 0.8, 0.9], ['SUB PROSPECT TYPE', 'DB', 'HOT', 'DB Q', 'STU', 'STU%'], b2, row_h=0.25, size=9)
    _tp_bar(sl, 8.4, y1, 4.65, 0.3 + 0.28 * max(2, len(subs)) + 0.3, 'STU% per Sub Prospect — ideal 10%', [r_['name'] for r_ in subs], [r_['stup'] for r_ in subs], 10.0)

def _tp_activity_slide(sl, S, title, D):
    _tp_clear_pictures(sl); _tp_ensure_bg(sl)
    _tp_head(sl, f"3. REVIEW AKTIFITAS — {title} (1 s.d {D['tgl']} {D['periode']})", None, w=7.3)
    _tp_activity_blocks(sl, S, 0.62, with_sub=True)

def _tp_counter_process_slide(sl, S, title, D):
    _tp_clear_pictures(sl); _tp_ensure_bg(sl)
    _tp_head(sl, f"2. TABEL PROSES KERJA & AKTIFITAS — {title} (1 s.d {D['tgl']} {D['periode']})", None, w=7.3)
    pr = S['r_pr']; rows = [[S['koor'] or S['label'], pr['JML MP'], pr['AVG DB'], pr['DB'], pr['HOT'], dict(t=pr['DB Q'], c=_tp_col(bsg_num(pr['DB Q']), 30, 20), b=True), pr['STU'], dict(t=pr['STU%'], c=_tp_col(bsg_num(pr['STU%']), 10, 7), b=True), pr['AVG NOTA'], pr['AVG 7'], pr['AVG 3']]]
    _tp_table(sl, 0.25, 0.62, [2.9] + [0.99] * 10, ['KOOR / COUNTER / DEALER', 'JML MP', 'AVG DB', 'DB', 'HOT', 'DB Q', 'STU', 'STU%', 'AVG NOTA', 'AVG 7', 'AVG 3'], rows, row_h=0.26, size=9, total_last=False)
    dbq_, stup_ = bsg_num(pr['DB Q']), bsg_num(pr['STU%'])
    _tp_card(sl, 0.25, 1.35, 3.1, 0.95, 'Database Quality (DB Q)', pr['DB Q'], f"HOT {int(pr['HOT']):,} / {int(pr['DB']):,} DB · ideal 30%", _tp_col(dbq_, 30, 20), min(dbq_ / 30.0 * 100, 100))
    _tp_card(sl, 3.5, 1.35, 3.1, 0.95, 'STU Ratio (STU%)', pr['STU%'], f"STU {int(pr['STU']):,} / {int(pr['DB']):,} DB · ideal 10%", _tp_col(stup_, 10, 7), min(stup_ / 10.0 * 100, 100))
    _tp_activity_blocks(sl, S, 2.6, with_sub=False, max_pt=6)

def _tp_leasing_slide(sl, S, title, D):
    _tp_clear_pictures(sl); _tp_ensure_bg(sl)
    _tp_head(sl, f"3. TABEL SPK & LEASING — {title} (1 s.d {D['tgl']} {D['periode']})", None, w=7.3)
    base = S['dlr_s'] if S['total_view'] else S['ent_s']; r_s = S['r_s']; show = base[:6]
    def _row(r_, nm_): return [nm_, r_['SPK'], r_['SPK LM'], r_['GROWTH'], dict(t=r_['ACC%'], c=_tp_col(bsg_num(r_['ACC%']), 50, 40), b=True), r_['TERBUKA'], r_['BAF'], r_['%BAF'], r_['ADIRA'], r_['%ADIRA'], r_['SOF'], r_['%SOF'], r_['IMFI'], r_['%IMFI']]
    body = [_row(r_, r_[_TP_KN]) for r_ in show] + [_row(r_s, f"TOTAL ({title})")]
    _tp_table(sl, 0.25, 0.62, [2.6] + [0.784] * 13, ['KOOR / COUNTER / DEALER', 'SPK', 'SPK LM', 'GROWTH', 'ACC%', 'TERBUKA', 'BAF', '%BAF', 'ADIRA', '%ADIRA', 'SOF', '%SOF', 'IMFI', '%IMFI'], body, row_h=0.25, size=9)
    yb = 0.62 + 0.3 + 0.25 * len(body) + 0.15; hb = max(0.9, 3.85 - yb)
    ls = [l_ for l_ in ['BAF', 'ADIRA', 'SOF', 'IMFI'] if r_s.get(l_, 0) > 0]
    _tp_bar(sl, 0.25, yb, 6.3, hb, 'Approval Rate (ACC%) per Leasing — ideal 50%', ls, [bsg_num(r_s[f'%{l_}']) for l_ in ls], 50.0)
    st_ = [l_ for l_ in ['CASH', 'BAF', 'ADIRA', 'SOF', 'IMFI'] if S['lsg_stu'].get(l_, 0) > 0]; tot_ = max(1, sum(S['lsg_stu'][l_] for l_ in st_))
    _tp_bar(sl, 6.75, yb, 6.3, hb, 'Kontribusi STU per Pembayaran (%)', st_, [S['lsg_stu'][l_] / tot_ * 100.0 for l_ in st_], None, colors=[_TP['CYN']] * len(st_))

def _tp_type_slide(sl, S, title, D):
    _tp_clear_pictures(sl); _tp_ensure_bg(sl)
    _tp_head(sl, f"4. TABEL PENJUALAN PER TYPE — {title} (1 s.d {D['tgl']} {D['periode']})", None, w=7.3)
    ideal = float(target_persen_harian)
    if not S['typ']: _tp_text(sl, 0.25, 0.8, 12.8, 0.4, 'Belum ada data penjualan per type pada periode ini.', size=13, color=_TP['MUT']); return
    def _row(r_, nm_): return [nm_, r_['TARGET'], r_['STU'], dict(t=f"{r_['ACHIEV']:.2f}%", c=_tp_row_style(r_['ACHIEV'], ideal), b=True), r_['STU LM'], r_['GROWTH'], r_['DB'], r_['HOT'], dict(t=r_['DB Q'], c=_tp_col(bsg_num(r_['DB Q']), 30, 20)), r_['STU CASH'], r_['STU BAF'], r_['STU ADIRA'], r_['STU SOF']]
    body = [_row(r_, r_['TYPE']) for r_ in S['typ']] + ([_row(S['typ_tot'], f"TOTAL ({title})")] if S['typ_tot'] else [])
    _tp_table(sl, 0.25, 0.62, [2.5] + [0.859] * 12, ['TYPE', 'TARGET', 'STU', 'ACHIEV', 'LM', 'GROWTH', 'DB', 'HOT', 'DB Q', 'STU CASH', 'STU BAF', 'STU ADIRA', 'STU SOF'], body, row_h=0.26, size=9)
    yb = 0.62 + 0.3 + 0.26 * len(body) + 0.15; hb = max(0.9, 3.85 - yb)
    _tp_bar(sl, 0.25, yb, 6.3, hb, f'Capaian per Type vs Target (ideal {ideal:.0f}%)', [r_['TYPE'] for r_ in S['typ']], [r_['ACHIEV'] for r_ in S['typ']], ideal, fmt='0"%"')
    tot_ = max(1, sum(r_['STU'] for r_ in S['typ'])); _tp_bar(sl, 6.75, yb, 6.3, hb, 'Komposisi STU per Type (%)', [r_['TYPE'] for r_ in S['typ']], [r_['STU'] / tot_ * 100.0 for r_ in S['typ']], None, colors=[_TP['CYN']] * len(S['typ']))

def _tp_kpi_slide(sl, P, D):
    _tp_clear_pictures(sl); _tp_ensure_bg(sl)
    if not P: _tp_text(sl, 0.5, 3.0, 12.3, 0.6, 'Data Dashboard KPI belum tersedia (periksa file Master KPI / Master Employee).', size=16, color=_TP['MUT'], align='c'); return
    tot, grade = P['total'], P['grade']; gcol = _PP_GRADE.get(grade, '0891B2')
    _tp_text(sl, 0.25, 0.1, 6.0, 0.25, 'KEY PERFORMANCE INDICATOR REPORT', size=9, color=_TP['MUT'], bold=True); _tp_text(sl, 0.25, 0.32, 6.0, 0.45, P['scope'], size=22, color=_TP['WHT'], bold=True)
    _tp_text(sl, 0.25, 0.8, 6.0, 0.22, f"Periode 1 s.d {P['date_lbl']}", size=9, color=_TP['MUT'])
    cards = [('GRADE', grade, gcol, 'Skala A–D'), ('TOTAL POINT', str(tot['point']), _TP['WHT'], 'dari 100 poin'), ('RESULT', str(tot['sales']), _TP['CYN'], f"dari {P['max_sales']} · {P['concl_res']}"), ('PROCESS', str(tot['process']), 'A78BFA', f"dari {P['max_process']} · {P['concl_prc']}")]
    for i, (lb, v, c, sb) in enumerate(cards): _tp_card(sl, 6.6 + i * 1.62, 0.1, 1.55, 0.95, lb, v, sb, c)
    hdr = ['CODE', 'PERFORMANCE CATEGORY', 'WEIGHT', 'MIN ACV', 'TARGET', 'RESULT', 'ACHIEVEMENT', 'POINT']
    body = []
    for r_ in P['rows']:
        stt = _pp_status(r_); c_ = {'ok': _TP['GRN'], 'warn': _TP['YEL'], 'bad': _TP['RED'], 'na': _TP['MUT']}[stt]
        body.append([r_['code'], r_['name'], r_['weight'], ('-' if r_.get('no_min') else f"{r_.get('min_acv', 0) * 100:.0f}%"), _pp_fmt(r_['code'], r_['target']), _pp_fmt(r_['code'], r_['result']), dict(t=f"{r_['achv'] * 100:.1f}%", c=c_, b=True), dict(t=f"{r_['point']}", c=c_, b=True)])
    body.append(['', 'TOTAL', sum(r_['weight'] for r_ in P['rows']), '', '', '', '', tot['point']])
    _tp_table(sl, 0.25, 1.2, [0.8, 4.2, 1.2, 1.2, 1.5, 1.5, 1.6, 0.8], hdr, body, row_h=0.36, size=10, aligns=['c', 'l', 'c', 'c', 'c', 'c', 'c', 'c'])
    _tp_text(sl, 0.25, 7.05, 12.8, 0.3, P.get('scheme_note', ''), size=7.5, color=_TP['MUT'], italic=True)

def _tp_kpi_entity_note(sl, S, kind_lbl):
    """Cadangan bila gambar Dashboard KPI koordinator / counter belum diunggah: slide dikosongkan dari gambar lama & diberi petunjuk."""
    _tp_clear_pictures(sl); _tp_ensure_bg(sl)
    nm = S['koor'] if S else '-'
    _tp_text(sl, 0.5, 2.9, 12.3, 0.5, f'Dashboard KPI {kind_lbl} terlemah: {nm}', size=20, color=_TP['WHT'], bold=True, align='c')
    _tp_text(sl, 0.5, 3.6, 12.3, 0.9, f'Gambar belum diunggah. Klik “Download SEMUA gambar untuk PPT” lalu unggah ZIP-nya (file Dashboard_KPI_{"Counter" if kind_lbl == "counter" else "Koordinator"}_{_tp_san(nm)}_Tgl_…jpg).', size=13, color=_TP['MUT'], align='c')

def _tp_recruit_slide(sl, S, D):
    mp = S['mp']; txt = (f"Tambah {mp['need']} salesman baru" if mp['need'] > 0 else 'Tidak perlu tambah salesman')
    for sh in sl.shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith('PERBAIKAN RECRUITMENT'):
            ps = sh.text_frame.paragraphs
            if len(ps) > 1:
                rs = ps[1].runs
                if rs:
                    rs[0].text = '1. ' + txt
                    for r_ in rs[1:]: r_._r.getparent().remove(r_._r)
                else: ps[1].add_run().text = '1. ' + txt

def _tp_cover(sl, D):
    for sh in sl.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() in (['SIP PANJI'] + list(daftar_dealer)):
            r = sh.text_frame.paragraphs[0].runs[0]; r.text = D['scope']
            from pptx.util import Pt
            n_ = len(D['scope']); r.font.size = Pt(66 if n_ <= 9 else (48 if n_ <= 11 else 40)); sh.text_frame.word_wrap = False

_TP_KIND = {('Review_1_Penjualan_', 'dealer'): 'detail_d', ('Review_5_Salesman_', 'dealer'): 'salesman_d', ('Review_2_Proses_', 'dealer'): 'process_d', ('Review_Aktifitas_PT_', 'dealer'): 'activity_d',
            ('Review_Aktifitas_Sub_', 'dealer'): 'activity_d', ('Review_1_Penjualan_', 'koor'): 'detail_k', ('Review_5_Salesman_', 'koor'): 'salesman_k', ('Review_Aktifitas_PT_', 'koor'): 'activity_k',
            ('Review_Aktifitas_Sub_', 'koor'): 'activity_k', ('Review_1_Penjualan_', 'counter'): 'detail_c', ('Review_2_Proses_', 'counter'): 'process_c', ('Review_Aktifitas_PT_', 'counter'): 'process_c',
            ('Review_3_Leasing_', 'dealer'): 'leasing', ('Review_4_Type_', 'dealer'): 'type', ('Dashboard_KPI_', 'kpi'): 'kpi',
            ('Dashboard_KPI_Koordinator_', 'kpi_koor'): 'kpi_k', ('Dashboard_KPI_Counter_', 'kpi_counter'): 'kpi_c'}
_TP_LEGACY_TBL = {5: 'dbq', 7: 'sub_d', 9: 'sales_k', 11: 'sub_k', 14: 'counter', 15: 'leasing', 16: 'type'}

def _tp_scan(sl):
    """Kenali isi slide dari NAMA SHAPE: gambar `BSG|<awalan_file>|<dealer/koor/counter/kpi>` dan tabel `BSGT|<jenis>`. Template lama tanpa penanda dibaca berdasarkan nomor slide."""
    pics, tbls, tagged = {}, {}, False
    for i, s_ in enumerate(sl):
        for sh in s_.shapes:
            nm = sh.name or ''
            if nm.startswith('BSG|') and sh.shape_type == 13 and nm.count('|') >= 2: _, pf, ky = nm.split('|', 2); pics.setdefault(i, []).append((nm, pf, ky)); tagged = True
            elif nm.startswith('BSGT|') and sh.has_table: tbls.setdefault(i, []).append((nm, nm.split('|', 1)[1])); tagged = True
    # Perbaikan bug: gambar slide 2 (DETAIL DEALER) pada template sering belum diberi nama penanda `BSG|...` sehingga tidak pernah diganti
    # (tetap menampilkan JPG dealer lama). Bila slide 2 tidak punya gambar bertanda, gambar terbesarnya dianggap Review_1_Penjualan_|dealer.
    if tagged and 1 not in pics and len(sl) >= 2:
        _cand = [sh for sh in sl[1].shapes if sh.shape_type == 13 and not (sh.name or '').startswith('BSG|')]
        if _cand:
            _big = max(_cand, key=lambda sh: (sh.width or 0) * (sh.height or 0))
            pics[1] = [(_big.name, 'Review_1_Penjualan_', 'dealer')]
    if not tagged:
        for n, specs in _TP_IMG.items():
            if n <= len(sl): pics[n - 1] = list(specs)
        for n, k in _TP_LEGACY_TBL.items():
            if n <= len(sl): tbls[n - 1] = [(None, k)]
    return pics, tbls

def bsg_ppt_tpl_build(D, tpl_src):
    """D = payload (lihat bsg_ppt_for_dealer). tpl_src = path / bytes template. Return (bytes_pptx, daftar_pesan)."""
    import io
    from pptx import Presentation
    prs = Presentation(io.BytesIO(tpl_src) if isinstance(tpl_src, (bytes, bytearray)) else tpl_src)
    sl = list(prs.slides); warn, info, native = [], [], set()
    Sd, Sk, Sc, H, addf, ai = D['S_dealer'], D['S_koor'], D['S_counter'], D['H'], D['addf'], D.get('ai') or {}
    title_d = D['scope']; title_k = (f"{Sk['dealer_lbl']} — {Sk['koor']}" if Sk else ''); title_c = (f"{Sc['dealer_lbl']} — {Sc['koor']}" if Sc else '')
    pics, tbls = _tp_scan(sl)
    def _no_data(s_, txt): _tp_clear_pictures(s_); _tp_ensure_bg(s_); _tp_text(s_, 0.5, 3.2, 12.3, 0.8, txt, size=18, color=_TP['MUT'], align='c')
    def ok(needs): return needs is None or (needs == 'k' and Sk) or (needs == 'c' and Sc)
    builders = {'detail_d': (None, lambda s_: _tp_detail_slide(s_, Sd, title_d, D)), 'salesman_d': (None, lambda s_: _tp_manpower_slide(s_, Sd, title_d, D)), 'process_d': (None, lambda s_: _tp_process_slide(s_, Sd, title_d, D)),
                'activity_d': (None, lambda s_: _tp_activity_slide(s_, Sd, title_d, D)), 'detail_k': ('k', lambda s_: _tp_detail_slide(s_, Sk, title_k, D)), 'salesman_k': ('k', lambda s_: _tp_salesman_slide_k(s_, Sk, title_k, D)),
                'activity_k': ('k', lambda s_: _tp_activity_slide(s_, Sk, title_k, D)), 'detail_c': ('c', lambda s_: _tp_detail_slide(s_, Sc, title_c, D)), 'process_c': ('c', lambda s_: _tp_counter_process_slide(s_, Sc, title_c, D)),
                'leasing': (None, lambda s_: _tp_leasing_slide(s_, Sd, title_d, D)), 'type': (None, lambda s_: _tp_type_slide(s_, Sd, title_d, D)), 'kpi': (None, lambda s_: _tp_kpi_slide(s_, D.get('kpi'), D)),
                'kpi_k': ('k', lambda s_: _tp_kpi_entity_note(s_, Sk, 'koordinator')), 'kpi_c': ('c', lambda s_: _tp_kpi_entity_note(s_, Sc, 'counter'))}
    for i, s_ in enumerate(sl):   # cover & recruitment dikenali dari teks
        for fn in (lambda x: _tp_cover(x, D), lambda x: _tp_recruit_slide(x, Sd, D)):
            try: fn(s_)
            except Exception as e_: warn.append(f'Slide {i + 1}: {type(e_).__name__}: {e_}')
    for i, specs in sorted(pics.items()):   # gambar dashboard (atau tabel/grafik otomatis bila gambar belum diunggah)
        kind = _TP_KIND.get((specs[0][1], specs[0][2])); s_ = sl[i]
        if kind is None: continue
        needs, fn = builders[kind]
        try:
            if not ok(needs): _no_data(s_, 'Tidak ada data ' + ('koordinator' if needs == 'k' else 'counter') + ' pada dealer / periode ini.'); continue
            found = [(nm_, _tp_find_img(D, pf, ky)) for nm_, pf, ky in specs]
            if D.get('images'):
                if all(f_[1][0] is not None for f_ in found):
                    for nm_, (img_, _e) in found:
                        w_ = _tp_swap_pic(s_, nm_, img_)
                        if w_: info.append(f'ℹ️ Slide {i + 1}: {w_} — gambar dipasang utuh di tengah kotak (tidak meregang); sesuaikan ukuran kotak gambar di template bila ingin memenuhi slide.')
                    _tp_ensure_bg(s_); continue
                info.append(f"ℹ️ Slide {i + 1}: memakai tabel otomatis — gambar belum diunggah: " + ', '.join(e_ for _, (b_, e_) in found if b_ is None and e_))
            fn(s_); native.add(i)
        except Exception as e_: warn.append(f'Slide {i + 1}: {type(e_).__name__}: {e_}')
    rows_by = {'dbq': (None, lambda: _tx_rows_dbq(Sd, H, addf, ai)), 'sub_d': (None, lambda: _tx_rows_sub(Sd, H, ai, 3, extra_media=True)), 'sales_k': ('k', lambda: _tx_rows_salesman(Sk, H, addf, ai)),
               'sub_k': ('k', lambda: _tx_rows_sub(Sk, H, ai, 3)), 'counter': ('c', lambda: _tx_rows_counter(Sc, H, addf, ai)), 'leasing': (None, lambda: _tx_rows_leasing(Sd, H, ai)), 'type': (None, lambda: _tx_rows_type(Sd, H, addf, ai))}
    for i, lst in sorted(tbls.items()):   # tabel Condition: hanya kolom CONDITION yang diisi
        for nm_, kind in lst:
            if kind not in rows_by: continue
            needs, rf = rows_by[kind]
            try: _tp_fill_analysis(sl[i], rf() if ok(needs) else [], top=(3.95 if (kind in ('leasing', 'type') and i in native) else None), name=nm_)
            except Exception as e_: warn.append(f'Slide {i + 1} (kolom Condition): {type(e_).__name__}: {e_}')
    if Sk or Sc: info.insert(0, 'ℹ️ Slide koordinator/counter memakai ' + (f"koordinator terlemah: {Sk['koor']}" if Sk else '') + (' · ' if (Sk and Sc) else '') + (f"counter terlemah: {Sc['koor']}" if Sc else '') + '. Untuk gambar dashboard slide tersebut, pilih nama itu di dropdown Koordinator / Counter lalu Download.')
    buf = io.BytesIO(); prs.save(buf)
    return buf.getvalue(), info + warn

def _tp_salesman_slide_k(sl, S, title, D):
    _tp_clear_pictures(sl); _tp_ensure_bg(sl)
    _tp_head(sl, f"5. TABEL DETAIL SALESMAN — {title}", None, w=7.3)
    if not S['sm']: _tp_text(sl, 0.25, 1.0, 12.8, 0.4, 'Belum ada data salesman pada koordinator ini.', size=14, color=_TP['MUT']); return
    _tp_salesman_table(sl, S, 0.62, 9, size=9)

# ====================================================================================================
# ============ Pengumpul data dari DASHBOARD KPI (memakai fungsi & data aplikasi) =====================
# ====================================================================================================
_PP_SHORT = {'STU': 'STU', 'CLS': 'Classy', 'MAX': 'Maxi', 'PRM': 'Promosi', 'SAL': 'Salesman', 'PRD': 'Produktivitas', 'WLK': 'Walk In', 'MYM': 'Member Point', 'DBS': 'Database', 'INC': 'Incoming', 'MYD': 'MYDR', 'GRV': 'G-Review'}

def bsg_team_stats(ent):
    """Angka analisa per tim (koordinator/counter): sama rumusnya dengan Tabel Penjualan, Proses Kerja & SPK di tab STU."""
    g = globals()
    def _sub(name, col):
        d_ = g.get(name)
        return d_[d_[col] == ent] if (d_ is not None and hasattr(d_, 'empty') and not d_.empty and col in d_.columns) else pd.DataFrame()
    f_c_, f_l_, n_c_, n_t_ = _sub('df_fak_curr_g', 'Koor'), _sub('df_fak_lm_g', 'Koor'), _sub('df_nota_curr_g', 'Koor'), _sub('df_nota_tod_g', 'Koor')
    db_c_, s_c_, s_l_ = _sub('df_db_curr_g', 'Koor'), _sub('df_spk_curr_g', 'Koor_SPK'), _sub('df_spk_lm_g', 'Koor_SPK')
    targ = get_koor_or_counter_target(ent, load_json(TARGET_KOOR_FILE))
    rp = calc_sales_row(ent, 'N', targ, f_c_, f_l_, n_c_, n_t_); rpr = calc_proses_row(ent, 'N', db_c_, len(f_c_), n_c_); rs = calc_leasing_row(ent, s_c_, s_l_, 'N')
    return dict(stu=int(rp['STU']), target=int(targ), ideal=int(rp['IDEAL']), pct=(rp['STU'] / targ * 100.0 if targ > 0 else 0.0), gap_ideal=int(rp['STU'] - rp['IDEAL']),
                db=int(rpr['DB']), hot=int(rpr['HOT']), mp=int(rpr['JML MP']), avg_db=bsg_num(rpr['AVG DB']), dbq=bsg_num(rpr['DB Q']), stup=bsg_num(rpr['_stu_pct']),
                spk=int(rs['SPK']), acc=bsg_num(rs['_acc_val']), open_spk=int(rs['TERBUKA']))

def _pp_team_text(x, tgl, tgl_sasaran, pct_sasaran, timeline):
    rows_e = x['rows']
    lost = sorted([(c, r['weight'] - r['point']) for c, r in rows_e.items() if not r.get('fixed') and r['point'] < r['weight']], key=lambda t: -t[1])
    probs = [f"{_PP_SHORT.get(c, c)} {_pp_fmt(c, rows_e[c]['result'])}/{_pp_fmt(c, rows_e[c]['target'])} ({rows_e[c]['achv'] * 100:.0f}%)" for c, _ in lost[:3]]
    try: st_ = bsg_team_stats(x['name'])
    except Exception: st_ = None
    cond = f"Poin KPI {x['point']} (Grade {x['grade']})."
    if st_ and st_['target'] > 0: cond += f" STU {st_['stu']}/{st_['target']} unit ({st_['pct']:.0f}%, gap ideal {st_['gap_ideal']:+d})."
    if st_ and st_['db'] > 0:
        cond += f" Input {st_['avg_db']:.1f} DB/MP/hari; DB Q {st_['dbq']:.1f}% (ideal 30%), STU% {st_['stup']:.1f}% (ideal 10%)" + (f", ACC% {st_['acc']:.0f}% (ideal 50%)" if st_['spk'] > 0 else '') + '.'
    if probs: cond += ' KPI tertinggal: ' + '; '.join(probs) + '.'
    target = '—'
    if st_ and st_['target'] > 0:
        try:
            add, goal = calc_rational_7d_addition(st_['stu'], st_['target'], tgl, tgl_sasaran, persen_desimal, pct_sasaran, active_mp=max(1, st_['mp']))
            if add > 0: target = f"+{add} STU (menjadi {goal} unit) · +{max(add + 1, round(add / 0.5))} SPK · DB Q ≥30%, STU% ≥10%, ACC% ≥50%"
            else: target = f"Pertahankan; capai target {st_['target']} unit · DB Q ≥30%, STU% ≥10%"
        except Exception: target = f"Capai target {st_['target']} unit · DB Q ≥30%, STU% ≥10%, ACC% ≥50%"
    return dict(cond=cond, target=target, timeline=timeline)

def bsg_collect_ppt_payload(sel_dash, scheme_dash, rows_sel, s_sel, p_sel, t_sel, dealer_tbl, inputs_dash, date_lbl, scheme_note, n_teams=5, with_teams=True):
    """Kumpulkan seluruh isi PPT dari DASHBOARD KPI: global (KPI + dealer) → dealer terlemah → koordinator & counter terlemah."""
    is_total = (sel_dash == 'TOTAL JEMBER BONDOWOSO'); tgl = int(tanggal_akhir)
    tgl_sasaran = 30 if tgl > 21 else min(30, tgl + 7)
    sisa_hari = max(1, sum(1 for d_n in range(tgl + 1, min(tgl_sasaran, max_day_in_month) + 1) if is_workday_day(d_n)))
    pct_sasaran = tabel_target_harian.get(tgl_sasaran, 100.0) / 100.0
    timeline = f'S.d Tgl {tgl_sasaran} {periode_bulan_tahun} ({sisa_hari} hari kerja)'
    def _row(r): return dict(code=r['code'], name=r['name'], weight=r['weight'], min_acv=r.get('min_acv', 0.0), target=r['target'], result=r['result'], achv=r['achv'], point=r['point'], fixed=bool(r.get('fixed')), no_min=bool(r.get('no_min')))
    w_ = KPI_WEIGHT_MDS if scheme_dash == 'MDS' else KPI_WEIGHT_DEALER
    concl = lambda v, a, b: 'VERY GOOD' if v >= a else ('GOOD' if v >= b else 'NOT GOOD')
    rk = ['STU', 'MAX', 'CLS', 'OPR', 'SAL', 'APV', 'INC', 'MYM', 'MYD', 'GRT', 'GRV']
    def _probs(rows_d):
        out = []
        for c in KPI_CODES:
            r = rows_d[c]
            if r['weight'] > 0 and r['point'] < r['weight']:
                out.append(dict(name=r['name'], lost=r['weight'] - r['point'], detail=f"{r['achv'] * 100:.0f}% (min {r['min_acv'] * 100:.0f}%)"))
        return sorted(out, key=lambda t: -t['lost'])
    dealers = [dict(name=x['name'], code=x['code'], point=x['point'], grade=x['grade'], sales=x['sales'], process=x['process'], achv=dict(x['achv']), note=x['note']) for x in dealer_tbl]
    weak = sorted(dealers, key=lambda x: (x['point'], x['process']))[:3]
    for x in weak:
        try: x['problems'] = _probs(kpi_score_entity(inputs_dash[x['name']], scheme_dash)[0])
        except Exception: x['problems'] = []
    rank_pos = next((i for i, x in enumerate(dealers, 1) if x['name'] == sel_dash), None)
    teams = {}
    koor_focus, cnt_focus = get_focus_entities() if with_teams else ([], [])
    for kind, ents in (('koor', koor_focus), ('counter', cnt_focus)):
        items = []
        for e in ents:
            dl = kpi_entity_dealer(e)
            if not is_total and dl != sel_dash: continue
            try:
                inp = get_entity_kpi_inputs(kind, e)
                if safe_float(inp['STU'][0]) <= 0 and safe_float(inp['DBS'][0]) <= 0: continue
                rows_e, s_e, p_e, t_e = kpi_score_entity_sim(inp, kind, detect_counter_type(e) if kind == 'counter' else 'ADMIN')
            except Exception: continue
            items.append(dict(name=e, dealer=dl, point=t_e, grade=kpi_grade(t_e), sales=s_e, process=p_e, rows=rows_e, kind=kind))
        items.sort(key=lambda x: (x['point'], x['process'], x['name']))
        items = items[:n_teams]
        for x in items: x.update(_pp_team_text(x, tgl, tgl_sasaran, pct_sasaran, timeline))
        teams[kind] = [{k: v for k, v in x.items() if k not in ('rows', 'kind')} for x in items]
    return dict(scope=sel_dash, scheme=scheme_dash, date_lbl=date_lbl, periode_lbl=periode_bulan_tahun, tanggal=tgl, ideal_pct=float(target_persen_harian), grade=kpi_grade(t_sel),
                total=dict(point=t_sel, sales=s_sel, process=p_sel), max_sales=sum(w_[c] for c in KPI_SALES_CODES), max_process=sum(w_[c] for c in KPI_PROCESS_CODES),
                concl_res=concl(s_sel, 38, 20), concl_prc=concl(p_sel, 57, 45), rows=[_row(rows_sel[c]) for c in KPI_CODES], dealers=dealers, dealers_weak=weak, rank_pos=rank_pos,
                rank_codes=rk, rank_min={c: KPI_MIN_ACV[c] for c in rk}, scheme_note=scheme_note, teams=teams)

DEFAULT_SHARED_API_KEY = ""   # JANGAN isi key di kode. Pakai st.secrets["GEMINI_API_KEY"] / env GEMINI_API_KEY / kolom API Key di menu Analisa.
IDEAL_DBQ_PCT, IDEAL_STU_PCT, IDEAL_ACC_PCT, MAX_NARRATIVE_WORDS = 30.0, 10.0, 50.0, 35
SUBTOTAL_ROW_LABELS = {'JEMBER', 'BONDOWOSO', 'SUB JEMBER', 'SUB BONDOWOSO'}
tabel_target_harian = {1:0.0, 2:1.0, 3:2.0, 4:5.0, 5:8.0, 6:12.0, 7:18.0, 8:21.0, 9:24.0, 10:29.0, 11:34.0, 12:36.0, 13:40.0, 14:44.0, 15:48.0, 16:51.0, 17:54.0, 18:57.0, 19:60.0, 20:63.0, 21:67.0, 22:70.0, 23:73.0, 24:77.0, 25:79.0, 26:82.0, 27:84.0, 28:88.0, 29:92.0, 30:100.0, 31:100.0}
FIXED_COUNTER_TARGETS = {}   # isi {'NAMA COUNTER': target} jika ada target tetap counter
FOKUS_KATEGORI_TYPE = ['AT PREMIUM', 'AT CLASSY', 'AT LPM']
daftar_dealer = ['SENTRAL YAMAHA JEMBER', 'SIP AMBULU', 'SIP BALUNG', 'SIP BONDOWOSO', 'SIP GAJAH MADA', 'SIP KEBONSARI', 'SIP NANGKAAN', 'SIP RAMBIPUJI', 'SIP SUMBERSARI', 'SIP TRUNOJOYO']
ARCHIVE_DIR = st.session_state.get('arsip_dir') if (st.session_state.get('arsip_dir') and os.path.isdir(st.session_state.get('arsip_dir'))) else None   # MODE ARSIP: baca data dari folder arsip hasil IMPORT
ARCHIVE_MODE, ARCHIVE_META = bool(ARCHIVE_DIR), (st.session_state.get('arsip_meta') or {})
_ABS = (lambda f_: os.path.join(ARCHIVE_DIR, f_)) if ARCHIVE_DIR else (lambda f_: f_)
TARGET_DEALER_FILE, TARGET_KOOR_FILE, KPI_SAVED_FILE, AI_CONFIG_FILE, UPLOAD_DIR = _ABS('dealer_targets_saved.json'), _ABS('koor_targets_saved.json'), _ABS('kpi_targets_saved.json'), 'ai_config.json', _ABS('saved_raw_files')
DB_Multi_DIR, RATING_Multi_DIR, RAW_META_FILE = os.path.join(UPLOAD_DIR, 'database_files'), os.path.join(UPLOAD_DIR, 'rating_files'), os.path.join(UPLOAD_DIR, 'metadata.json')
os.makedirs(DB_Multi_DIR, exist_ok=True)
os.makedirs(RATING_Multi_DIR, exist_ok=True)

DEFAULT_KPI_DEALER = {   # target default dealer fokus (silakan sesuaikan di menu Target)
    'SENTRAL YAMAHA JEMBER': {'T. STU': 0, 'T. MAXI': 0, 'T. CLASSY': 0, 'T. MP': 0, 'T. INC': 0, 'T. ACC%': 50.0, 'T. MYDR': 0, 'T. G REV': 0},
    'SIP AMBULU': {'T. STU': 0, 'T. MAXI': 0, 'T. CLASSY': 0, 'T. MP': 0, 'T. INC': 0, 'T. ACC%': 50.0, 'T. MYDR': 0, 'T. G REV': 0},
    'SIP BALUNG': {'T. STU': 0, 'T. MAXI': 0, 'T. CLASSY': 0, 'T. MP': 0, 'T. INC': 0, 'T. ACC%': 50.0, 'T. MYDR': 0, 'T. G REV': 0},
    'SIP BONDOWOSO': {'T. STU': 0, 'T. MAXI': 0, 'T. CLASSY': 0, 'T. MP': 0, 'T. INC': 0, 'T. ACC%': 50.0, 'T. MYDR': 0, 'T. G REV': 0},
    'SIP GAJAH MADA': {'T. STU': 0, 'T. MAXI': 0, 'T. CLASSY': 0, 'T. MP': 0, 'T. INC': 0, 'T. ACC%': 50.0, 'T. MYDR': 0, 'T. G REV': 0},
    'SIP KEBONSARI': {'T. STU': 0, 'T. MAXI': 0, 'T. CLASSY': 0, 'T. MP': 0, 'T. INC': 0, 'T. ACC%': 50.0, 'T. MYDR': 0, 'T. G REV': 0},
    'SIP NANGKAAN': {'T. STU': 0, 'T. MAXI': 0, 'T. CLASSY': 0, 'T. MP': 0, 'T. INC': 0, 'T. ACC%': 50.0, 'T. MYDR': 0, 'T. G REV': 0},
    'SIP RAMBIPUJI': {'T. STU': 0, 'T. MAXI': 0, 'T. CLASSY': 0, 'T. MP': 0, 'T. INC': 0, 'T. ACC%': 50.0, 'T. MYDR': 0, 'T. G REV': 0},
    'SIP SUMBERSARI': {'T. STU': 0, 'T. MAXI': 0, 'T. CLASSY': 0, 'T. MP': 0, 'T. INC': 0, 'T. ACC%': 50.0, 'T. MYDR': 0, 'T. G REV': 0},
    'SIP TRUNOJOYO': {'T. STU': 0, 'T. MAXI': 0, 'T. CLASSY': 0, 'T. MP': 0, 'T. INC': 0, 'T. ACC%': 50.0, 'T. MYDR': 0, 'T. G REV': 0},
}

# ---- Pembaca Excel yang tahan banting: file berekstensi .xls/.xlsx tetapi isinya sebenarnya HTML / CSV / TSV / rusak ----
_orig_read_excel = pd.read_excel
def _robust_read_excel(io, *args, **kwargs):
    try:
        return _orig_read_excel(io, *args, **kwargs)
    except ValueError as e_:
        if 'cannot be determined' not in str(e_) or not isinstance(io, (str, os.PathLike)): raise
        fp_ = os.fspath(io)
        try:
            if os.path.getsize(fp_) == 0: raise ValueError(f'File kosong (0 byte): {os.path.basename(fp_)} — hapus file ini lalu upload ulang.')
            with open(fp_, 'rb') as fh_: head_ = fh_.read(4096)
        except OSError: raise e_
        hdr_ = kwargs.get('header', 0); low_ = head_.lower()
        if head_[:2] == b'PK':   # xlsx
            try: return _orig_read_excel(io, *args, **{**kwargs, 'engine': 'openpyxl'})
            except Exception: pass
        if head_[:4] == b'\xd0\xcf\x11\xe0':   # xls lama
            try: return _orig_read_excel(io, *args, **{**kwargs, 'engine': 'xlrd'})
            except Exception: pass
        if b'<html' in low_ or b'<table' in low_ or b'<!doctype' in low_ or b'<?xml' in low_:   # export web berekstensi .xls
            try:
                tabs_ = pd.read_html(fp_, header=hdr_)
                if tabs_: return max(tabs_, key=lambda t_: t_.size)
            except Exception: pass
        for enc_ in (() if (b'\x00' in head_ and head_[:2] not in (b'\xff\xfe', b'\xfe\xff')) else ('utf-8-sig', 'utf-16', 'cp1252', 'latin-1')):   # teks (CSV/TSV) berekstensi .xls
            try: return pd.read_csv(fp_, sep=None, engine='python', header=hdr_, encoding=enc_)
            except Exception: continue
        raise ValueError(f'File "{os.path.basename(fp_)}" bukan Excel/CSV yang valid (mungkin rusak/terpotong saat diunduh). Hapus file ini dari folder saved_raw_files lalu upload ulang yang asli.') from e_
pd.read_excel = _robust_read_excel

def clean_str(val): return '' if pd.isna(val) else ' '.join(str(val).split()).upper()
def nospace_str(val): return '' if pd.isna(val) else ''.join(str(val).split()).upper()
def alnum_str(val): return '' if pd.isna(val) else ''.join(c for c in str(val).upper() if c.isalnum())
FIXED_COUNTER_TARGETS_NS = {alnum_str(k): v for k, v in FIXED_COUNTER_TARGETS.items()}

def normalize_dealer_name(val):
    s = clean_str(val)
    if not s or s in ['NAN', 'NONE'] or s in daftar_dealer: return s
    s_ns = alnum_str(s)
    for k, v in [('NANGKAAN', 'SIP NANGKAAN'), ('BONDOWOSO', 'SIP BONDOWOSO'), ('AMBULU', 'SIP AMBULU'), ('BALUNG', 'SIP BALUNG'),
                 ('GAJAHMADA', 'SIP GAJAH MADA'), ('KEBONSARI', 'SIP KEBONSARI'), ('RAMBIPUJI', 'SIP RAMBIPUJI'),
                 ('SUMBERSARI', 'SIP SUMBERSARI'), ('TRUNOJOYO', 'SIP TRUNOJOYO'), ('SENTRAL', 'SENTRAL YAMAHA JEMBER')]:
        if k in s_ns: return v
    return s

def safe_int(val, default=0):
    try: return default if pd.isna(val) else int(float(str(val).replace(',', '').replace('%', '').strip()))
    except Exception: return default

def safe_float(val, default=0.0):
    try: return default if pd.isna(val) else float(str(val).replace(',', '').replace('%', '').strip())
    except Exception: return default

def strip_formatting_tags(text):
    return ' '.join(re.sub(r'<[^>]+>', '', str(text).replace('<br>', ' ').replace('<br/>', ' ').replace('</p>', ' ')).split()) if text else ''

def limit_words(text, max_words=MAX_NARRATIVE_WORDS):
    clean_txt = strip_formatting_tags(text)
    if not clean_txt: return ''
    words = clean_txt.split()
    if len(words) <= max_words: return clean_txt
    truncated = ' '.join(words[:max_words]).rstrip(',;:-')
    return truncated if truncated.endswith('.') else truncated + '.'

def format_paragraphs_html(text):
    clean_txt = limit_words(text, MAX_NARRATIVE_WORDS)
    return f'<p class="nar-p">{clean_txt}</p>' if clean_txt else ''

def calc_rational_7d_addition(curr_stu, monthly_targ, curr_day, next_day, pct_curr_dec, pct_next_dec, active_mp=1):
    sisa_hr, hr_jalan, delta_pct = max(1, next_day - curr_day), max(1, curr_day), max(0.12, pct_next_dec - pct_curr_dec)
    normal_period_quota = monthly_targ * delta_pct if monthly_targ > 0 else max(2.0, active_mp * 1.5)
    stretch_from_rate = ((curr_stu / hr_jalan) * sisa_hr * 1.8) + max(1.0, active_mp * 0.6)
    ideal_now = monthly_targ * pct_curr_dec
    if ideal_now > 0 and curr_stu < (ideal_now * 0.5):
        add_stu = min(round((normal_period_quota * 0.65) + (stretch_from_rate * 0.35)), max(curr_stu + 4, round(normal_period_quota * 1.15), round(active_mp * 1.5)))
    else: add_stu = round(normal_period_quota + (max(0.0, ideal_now - curr_stu) * 0.30))
    add_stu = max(max(2, round(active_mp * 0.8)) if monthly_targ >= 10 else 1, add_stu)
    return int(add_stu), int(curr_stu + add_stu)

def calc_table_height(df, row_h=35, head_h=38): return 150 if (df is None or df.empty) else (len(df) * row_h) + head_h + 3
def get_compact_col_config(columns, name_col, editable_col=None):
    return {c: (st.column_config.TextColumn(c, width="medium") if c == name_col else (st.column_config.NumberColumn(c, width="small", format="%d") if c == editable_col else st.column_config.Column(c, width="small"))) for c in columns}

def get_neon_gradient_hex(ratio):
    r = max(0.0, min(1.0, float(ratio)))
    return '#{:02X}{:02X}{:02X}'.format(255, int((r/0.5)*234), int(127*(1.0-r/0.5))) if r <= 0.5 else '#{:02X}{:02X}{:02X}'.format(int(255*(1.0-(r-0.5)/0.5)), int(234+((r-0.5)/0.5)*21), int(((r-0.5)/0.5)*102))

def build_compact_html_string(df_items, df_full, columns, name_col, metric_col, table_dom_id, header_labels=None, ultra_compact=False, title_text=None, highlight_names=None, wrap_first=False):
    if df_full is None or df_full.empty: return ''
    hdrs = header_labels if (header_labels is not None and len(header_labels) == len(columns)) else columns
    min_v = float(df_items[metric_col].min()) if (df_items is not None and not df_items.empty and metric_col in df_items.columns) else 0.0
    max_v = float(df_items[metric_col].max()) if (df_items is not None and not df_items.empty and metric_col in df_items.columns) else 0.0
    n_items = len(df_items) if df_items is not None else 0
    wrap_cls, tbl_cls = ('ultra-compact-wrap', 'ultra-compact-tbl') if ultra_compact else ('compact-table-wrap', 'compact-tbl')
    banner_html = f'<div class="cap-title-banner">🐺 {title_text}</div>' if title_text else ''
    html = [f'<div id="{table_dom_id}" class="{wrap_cls}">{banner_html}<table class="{tbl_cls}"><thead><tr>'] + [(f'<th style="white-space:normal;font-size:11px;padding:5px 4px;text-align:left;line-height:1.2;">{h}</th>' if wrap_first == 'all' else f'<th>{h}</th>') for h in hdrs] + ['</tr></thead><tbody>']
    for _, row in df_full.iterrows():
        nm = str(row[name_col])
        style = 'background-color: #4318FF; color: #FFFFFF; font-weight: 800;' if (nm.startswith('TOTAL') or nm.startswith('CREDIT SHARE')) else ('background-color: #1E293B; color: #FFFFFF; font-weight: 800;' if (nm in SUBTOTAL_ROW_LABELS or nm.startswith('SUB ') or (highlight_names is not None and nm in highlight_names)) else '')
        if not style and '_row_ratio' in df_full.columns and pd.notna(row['_row_ratio']):
            style = f'color: {get_neon_gradient_hex(float(row["_row_ratio"]))};'
        if not style and df_items is not None and not df_items.empty and metric_col in df_items.columns:
            m_idx = df_items.index[df_items[name_col] == nm].tolist()
            if m_idx:
                v = float(df_items.loc[m_idx[0], metric_col])
                style = f'color: {get_neon_gradient_hex(((v - min_v) / (max_v - min_v)) if max_v > min_v else (m_idx[0] / max(1, n_items - 1)))};'
        if wrap_first == 'all' and nm.startswith('TOTAL') and '_summary' in df_full.columns and pd.notna(row['_summary']):   # baris total: 1 sel gabungan selebar tabel agar terbaca
            html.append(f'<tr style="{style}"><td colspan="{len(columns)}" style="white-space:normal;text-align:center;font-size:12.5px;padding:9px 8px;line-height:1.4;letter-spacing:0.3px;">{row["_summary"]}</td></tr>')
            continue
        html.append(f'<tr style="{style}">' + ''.join((f'<td style="white-space:normal;word-break:break-word;font-size:11.5px;padding:4px 5px;text-align:left;line-height:1.25;{"width:30px;" if ci_c == 0 else ""}">{row[c]}</td>' if wrap_first == 'all' else (f'<td style="white-space:normal;word-break:break-word;min-width:150px;max-width:260px;line-height:1.25;">{row[c]}</td>' if (wrap_first and ci_c == 0) else f'<td>{row[c]}</td>')) for ci_c, c in enumerate(columns)) + '</tr>')
    return ''.join(html + ['</tbody></table></div>'])

def render_table_image_buttons(table_dom_id, file_label, standalone_html=None):
    safe_fname, hidden_payload = re.sub(r'[^a-zA-Z0-9_-]', '_', file_label), json.dumps(standalone_html or '')
    components.html(f"""<html><head><script src="https://cdnjs.cloudflare.com/ajax/libs/html2canvas/1.4.1/html2canvas.min.js"></script><style>
        body {{ margin: 0; padding: 2px 0; background: transparent; font-family: 'Source Sans Pro', sans-serif; display: flex; align-items: center; gap: 8px; }}
        .act-btn {{ background-color: #1E293B; color: #E2E8F0; border: 1px solid #334155; border-radius: 6px; padding: 5px 11px; font-size: 12px; font-weight: 600; cursor: pointer; display: inline-flex; align-items: center; gap: 5px; transition: all 0.15s; }}
        .act-btn:hover {{ background-color: #4318FF; color: #FFFFFF; border-color: #4318FF; }}
        #status_msg {{ font-size: 12px; color: #00FF66; font-weight: 600; margin-left: 6px; }}
        #offscreen_holder {{ position: fixed; left: -9999px; top: 0; background: #0E1117; padding: 10px; z-index: -1; }}
        .compact-table-wrap, .ultra-compact-wrap {{ border-radius: 8px; border: 1px solid #2D3748; background-color: #0E1117; overflow: hidden; }}
        .cap-title-banner {{ background: linear-gradient(90deg, #4318FF 0%, #1E293B 100%); color: #FFFFFF !important; font-family: 'Source Sans Pro', sans-serif; font-size: 15px; font-weight: 800; padding: 9px 14px; border-bottom: 2px solid #3B18FF; letter-spacing: 0.4px; text-align: left; }}
        table.compact-tbl, table.ultra-compact-tbl {{ border-collapse: collapse; font-family: 'Source Sans Pro', sans-serif; font-size: 13.5px; background-color: #0E1117; width: 100%; }}
        table.compact-tbl th, table.ultra-compact-tbl th {{ background-color: #1A202C; color: #CBD5E1; font-weight: 800; padding: 6px 6px; border: 1px solid #2D3748; text-align: right; white-space: nowrap; }}
        table.compact-tbl th:first-child, table.compact-tbl td:first-child, table.ultra-compact-tbl th:first-child, table.ultra-compact-tbl td:first-child {{ text-align: left; padding-left: 8px; }}
        table.compact-tbl td, table.ultra-compact-tbl td {{ padding: 5px 6px; border: 1px solid #1E293B; text-align: right; white-space: nowrap; font-weight: 700; }}
    </style></head><body>
    <button class="act-btn" onclick="captureTbl('copy')">📸 Copy Tabel (Gambar)</button>
    <button class="act-btn" onclick="captureTbl('jpg')">📥 Download JPG</button>
    <span id="status_msg"></span><div id="offscreen_holder"></div>
    <script>
    const fallbackHtml = {hidden_payload};
    async function captureTbl(mode) {{
        const msg = document.getElementById('status_msg'); msg.style.color = '#FFEA00'; msg.innerText = 'Memproses gambar...';
        try {{
            const holder = document.getElementById('offscreen_holder');
            let targetEl = fallbackHtml ? (holder.innerHTML = fallbackHtml, holder.firstElementChild) : window.parent.document.getElementById('{table_dom_id}');
            if (!targetEl) {{ msg.style.color = '#FF007F'; msg.innerText = 'Tabel tidak ditemukan.'; return; }}
            const canvas = await html2canvas(targetEl, {{ backgroundColor: '#0E1117', scale: 2, useCORS: true, logging: false }});
            if (mode === 'copy') {{
                canvas.toBlob(async (blob) => {{
                    try {{
                        await navigator.clipboard.write([new ClipboardItem({{ 'image/png': blob }})]);
                        msg.style.color = '#00FF66'; msg.innerText = '✅ Gambar berhasil disalin ke Clipboard! (Siap Ctrl+V di WA/Excel)';
                    }} catch (err) {{
                        const a = document.createElement('a'); a.href = canvas.toDataURL('image/jpeg', 0.95); a.download = '{safe_fname}.jpg'; a.click();
                        msg.style.color = '#00E5FF'; msg.innerText = '📥 Browser membatasi clipboard, gambar otomatis diunduh sebagai JPG.';
                    }}
                    setTimeout(() => msg.innerText = '', 4500);
                }}, 'image/png');
            }} else {{
                const a = document.createElement('a'); a.href = canvas.toDataURL('image/jpeg', 0.95); a.download = '{safe_fname}.jpg'; a.click();
                msg.style.color = '#00FF66'; msg.innerText = '✅ File {safe_fname}.jpg berhasil diunduh!';
                setTimeout(() => msg.innerText = '', 4000);
            }}
            holder.innerHTML = '';
        }} catch (e) {{ msg.style.color = '#FF007F'; msg.innerText = 'Gagal membuat gambar: ' + e.message; }}
    }}
    </script></body></html>""", height=36)

def render_compact_html_table(df_items, df_full, columns, name_col, metric_col, key_suffix='tbl', header_labels=None, ultra_compact=False, title_text=None, custom_capture_html=None, highlight_names=None, wrap_first=False):
    if df_full is None or df_full.empty: return
    dom_id = f"Capture_{re.sub(r'[^a-zA-Z0-9]', '_', key_suffix)}"
    html_display_str = build_compact_html_string(df_items, df_full, columns, name_col, metric_col, dom_id, header_labels=header_labels, ultra_compact=ultra_compact, title_text=None, highlight_names=highlight_names, wrap_first=wrap_first)
    cap_title = title_text if title_text else key_suffix.replace('_', ' ')
    html_capture_str = custom_capture_html if custom_capture_html is not None else build_compact_html_string(df_items, df_full, columns, name_col, metric_col, f"{dom_id}_cap", header_labels=header_labels, ultra_compact=ultra_compact, title_text=cap_title, highlight_names=highlight_names, wrap_first=wrap_first)
    if ultra_compact:
        safe_fname, comp_height, hidden_uc_payload = re.sub(r'[^a-zA-Z0-9_-]', '_', key_suffix), (len(df_full) * 29) + 100, json.dumps(html_capture_str)
        components.html(f"""<html><head><script src="https://cdnjs.cloudflare.com/ajax/libs/html2canvas/1.4.1/html2canvas.min.js"></script><style>
            body {{ margin: 0; padding: 0; background: transparent; font-family: 'Source Sans Pro', sans-serif; color: #E2E8F0; }}
            .ultra-compact-wrap {{ width: 100%; overflow-x: hidden; margin-bottom: 6px; border-radius: 8px; border: 1px solid #2D3748; background-color: #0E1117; }}
            .cap-title-banner {{ background: linear-gradient(90deg, #4318FF 0%, #1E293B 100%); color: #FFFFFF !important; font-family: 'Source Sans Pro', sans-serif; font-size: 15px; font-weight: 800; padding: 9px 14px; border-bottom: 2px solid #3B18FF; letter-spacing: 0.4px; text-align: left; }}
            table.ultra-compact-tbl {{ width: 100%; table-layout: fixed; border-collapse: collapse; font-family: 'Source Sans Pro', sans-serif; font-size: 13.5px; background-color: #0E1117; }}
            table.ultra-compact-tbl th {{ background-color: #1A202C; color: #CBD5E1; font-weight: 800; padding: 6px 3px; border-bottom: 1px solid #2D3748; border-right: 1px solid #1E293B; text-align: center; white-space: normal; line-height: 1.15; font-size: 11.5px; word-break: break-word; }}
            table.ultra-compact-tbl th:first-child, table.ultra-compact-tbl td:first-child {{ width: 17.5%; text-align: left; padding-left: 8px; font-size: 12px; font-weight: 700; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }}
            table.ultra-compact-tbl td {{ padding: 5px 5px; border-bottom: 1px solid #1E293B; border-right: 1px solid #1E293B; text-align: right; white-space: nowrap; font-weight: 700; font-size: 13.5px; overflow: hidden; }}
            table.ultra-compact-tbl tr:hover td {{ background-color: rgba(255,255,255,0.04); }}
            .toolbar {{ display: flex; align-items: center; gap: 8px; padding-top: 2px; }}
            .act-btn {{ background-color: #1E293B; color: #E2E8F0; border: 1px solid #334155; border-radius: 6px; padding: 5px 11px; font-size: 12px; font-weight: 600; cursor: pointer; display: inline-flex; align-items: center; gap: 5px; transition: all 0.15s; }}
            .act-btn:hover {{ background-color: #4318FF; color: #FFFFFF; border-color: #4318FF; }}
            #status_msg {{ font-size: 12px; color: #00FF66; font-weight: 600; margin-left: 6px; }}
            #offscreen_uc {{ position: fixed; left: -9999px; top: 0; width: 1350px; background: #0E1117; padding: 10px; z-index: -1; }}
        </style></head><body>{html_display_str}
        <div class="toolbar"><button class="act-btn" onclick="captureUC('copy')">📸 Copy Tabel (Gambar)</button><button class="act-btn" onclick="captureUC('jpg')">📥 Download JPG</button><span id="status_msg"></span></div>
        <div id="offscreen_uc"></div>
        <script>
        const ucCaptureHtml = {hidden_uc_payload};
        async function captureUC(mode) {{
            const msg = document.getElementById('status_msg'); msg.style.color = '#FFEA00'; msg.innerText = 'Memproses gambar...';
            try {{
                const holder = document.getElementById('offscreen_uc'); holder.innerHTML = ucCaptureHtml;
                const canvas = await html2canvas(holder.firstElementChild, {{ backgroundColor: '#0E1117', scale: 2, useCORS: true, logging: false }});
                holder.innerHTML = '';
                if (mode === 'copy') {{
                    canvas.toBlob(async (blob) => {{
                        try {{
                            await navigator.clipboard.write([new ClipboardItem({{ 'image/png': blob }})]);
                            msg.style.color = '#00FF66'; msg.innerText = '✅ Gambar tabel + Judul disalin ke Clipboard! (Siap Ctrl+V di WA/Excel)';
                        }} catch (err) {{
                            const a = document.createElement('a'); a.href = canvas.toDataURL('image/jpeg', 0.95); a.download = '{safe_fname}.jpg'; a.click();
                            msg.style.color = '#00E5FF'; msg.innerText = '📥 Gambar otomatis diunduh sebagai JPG.';
                        }}
                        setTimeout(() => msg.innerText = '', 4500);
                    }}, 'image/png');
                }} else {{
                    const a = document.createElement('a'); a.href = canvas.toDataURL('image/jpeg', 0.95); a.download = '{safe_fname}.jpg'; a.click();
                    msg.style.color = '#00FF66'; msg.innerText = '✅ File {safe_fname}.jpg berhasil diunduh!';
                    setTimeout(() => msg.innerText = '', 4000);
                }}
            }} catch (e) {{ msg.style.color = '#FF007F'; msg.innerText = 'Gagal membuat gambar: ' + e.message; }}
        }}
        </script></body></html>""", height=comp_height, scrolling=False)
    else:
        st.markdown(html_display_str, unsafe_allow_html=True)
        render_table_image_buttons(dom_id, key_suffix, standalone_html=html_capture_str)

def load_json(f):
    if os.path.exists(f):
        try:
            with open(f, 'r') as fp: return json.load(fp)
        except Exception: return {}
    return {}

def save_json(data, f):
    try:
        with open(f, 'w') as fp: json.dump(data, fp)
    except Exception: pass

def resolve_gemini_api_key():
    if st.session_state.get('shared_gemini_api_key'): return st.session_state['shared_gemini_api_key'].strip()
    cfg_k = load_json(AI_CONFIG_FILE).get('gemini_api_key', '').strip()
    if cfg_k: return cfg_k
    try:
        for k_sec in ['GEMINI_API_KEY', 'gemini_api_key', 'GOOGLE_API_KEY']:
            if k_sec in st.secrets and str(st.secrets[k_sec]).strip(): return str(st.secrets[k_sec]).strip()
    except Exception: pass
    return os.environ.get('GEMINI_API_KEY', DEFAULT_SHARED_API_KEY).strip()

def save_raw_file(up_file, prefix):
    ext = os.path.splitext(up_file.name)[1].lower()
    for old in glob.glob(os.path.join(UPLOAD_DIR, f"{prefix}_*")):
        try: os.remove(old)
        except Exception: pass
    fpath = os.path.join(UPLOAD_DIR, f"{prefix}_latest{ext}")
    with open(fpath, 'wb') as f: f.write(up_file.getbuffer())
    meta = load_json(RAW_META_FILE)
    meta[prefix] = {'original_name': up_file.name, 'saved_path': fpath, 'updated_at': datetime.datetime.now().strftime('%d %b %Y, %H:%M WIB')}
    save_json(meta, RAW_META_FILE)

def save_database_files(up_files):
    meta = load_json(RAW_META_FILE)
    now_str = datetime.datetime.now().strftime('%d %b %Y, %H:%M WIB')
    file_times = meta.get('database_multi', {}).get('file_times', {})
    updated_dealers = []
    for uf in up_files:
        new_dlr = normalize_dealer_name(os.path.splitext(uf.name)[0])
        for old_p in glob.glob(os.path.join(DB_Multi_DIR, "*.*")):
            old_bname = os.path.basename(old_p)
            old_dlr = normalize_dealer_name(os.path.splitext(old_bname)[0])
            if (new_dlr in daftar_dealer and old_dlr == new_dlr) or (old_bname.upper() == uf.name.upper()):
                try: os.remove(old_p)
                except Exception: pass
        p = os.path.join(DB_Multi_DIR, uf.name)
        with open(p, 'wb') as f: f.write(uf.getbuffer())
        file_times[uf.name] = now_str
        updated_dealers.append(f"{uf.name} ({new_dlr})")
    all_saved_paths = sorted(glob.glob(os.path.join(DB_Multi_DIR, "*.*")))
    all_names = [os.path.basename(p) for p in all_saved_paths]
    file_times = {k: file_times.get(k, now_str) for k in all_names}
    meta['database_multi'] = {'file_names': all_names, 'saved_paths': all_saved_paths, 'file_times': file_times, 'count': len(all_names), 'updated_at': now_str}
    save_json(meta, RAW_META_FILE)
    return updated_dealers, len(all_names)

def get_saved_file_path(prefix):
    meta = load_json(RAW_META_FILE)
    if prefix in meta and os.path.exists(meta[prefix].get('saved_path', '')): return meta[prefix]['saved_path']
    m = glob.glob(os.path.join(UPLOAD_DIR, f"{prefix}_latest.*"))
    return m[0] if m else None

def get_saved_database_paths():
    return sorted(glob.glob(os.path.join(DB_Multi_DIR, "*.*")))

def save_rating_files(up_files):
    meta = load_json(RAW_META_FILE)
    now_str = datetime.datetime.now().strftime('%d %b %Y, %H:%M WIB')
    file_times = meta.get('rating_multi', {}).get('file_times', {})
    updated_dealers = []
    for uf in up_files:
        new_dlr = normalize_dealer_name(os.path.splitext(uf.name)[0])
        for old_p in glob.glob(os.path.join(RATING_Multi_DIR, "*.*")):
            old_bname = os.path.basename(old_p)
            old_dlr = normalize_dealer_name(os.path.splitext(old_bname)[0])
            if (new_dlr in daftar_dealer and old_dlr == new_dlr) or (old_bname.upper() == uf.name.upper()):
                try: os.remove(old_p)
                except Exception: pass
        p = os.path.join(RATING_Multi_DIR, uf.name)
        with open(p, 'wb') as f: f.write(uf.getbuffer())
        file_times[uf.name] = now_str
        updated_dealers.append(f"{uf.name} ({new_dlr})")
    all_saved_paths = sorted(glob.glob(os.path.join(RATING_Multi_DIR, "*.*")))
    all_names = [os.path.basename(p) for p in all_saved_paths]
    file_times = {k: file_times.get(k, now_str) for k in all_names}
    meta['rating_multi'] = {'file_names': all_names, 'saved_paths': all_saved_paths, 'file_times': file_times, 'count': len(all_names), 'updated_at': now_str}
    save_json(meta, RAW_META_FILE)
    return updated_dealers, len(all_names)

def get_saved_rating_paths():
    return sorted(glob.glob(os.path.join(RATING_Multi_DIR, "*.*")))

def get_azura_logo_path():
    saved_logo = get_saved_file_path('azura_logo')
    if saved_logo and os.path.exists(saved_logo): return saved_logo
    for cand in ['AZURA.jpg', 'AZURA.png', 'AZURA.jpeg', 'azura.jpg', 'azura.png', 'logo_azura.jpg', 'logo_azura.png']:
        if os.path.exists(cand): return cand
    for m in (glob.glob('*AZURA*.*') + glob.glob('*azura*.*')):
        if m.lower().endswith(('.jpg', '.jpeg', '.png', '.webp')): return m
    return None

def render_sidebar_azura_logo():
    logo_p = get_azura_logo_path()
    if logo_p and os.path.exists(logo_p):
        try:
            with open(logo_p, 'rb') as img_f: b64_str = base64.b64encode(img_f.read()).decode('utf-8')
            mime = 'image/png' if logo_p.lower().endswith('.png') else 'image/jpeg'
            st.sidebar.markdown(f'<div class="azura-logo-box"><img src="data:{mime};base64,{b64_str}" class="azura-logo-img" alt="AZURA"></div>', unsafe_allow_html=True)
            return
        except Exception: pass
    st.sidebar.markdown('<div class="azura-logo-box"><div class="azura-fallback-badge"><div style="font-size:22px; line-height:1;">🐺</div><div style="font-size:10px; font-weight:900; letter-spacing:1px; color:#00E5FF; margin-top:2px;">AZURA</div></div></div>', unsafe_allow_html=True)

def call_gemini_ai_table(api_key, raw_context_str):
    clean_key = api_key.strip()
    prompt = f"""Anda adalah Senior Sales Manager & Data Analyst Dealer Motor Yamaha.
Berikut adalah DATA MENTAH LENGKAP dari Tabel 1 s.d 5, Data Alasan Reject SPK, serta TARGET RASIONAL 7 HARI KE DEPAN:
{raw_context_str}

TUGAS ANDA:
Lakukan analisa silang yang tajam (temukan akar masalah/bottleneck tersembunyi dari data di atas) dan susun Action Plan taktis terukur dalam format JSON murni dengan 16 key berikut.

ATURAN WAJIB (SANGAT KETAT):
1. MAKSIMAL 35 KATA UNTUK SETIAP KEY! Buat kalimat yang super padat, to-the-point, langsung ke akar masalah dan angka sasaran (1 paragraf pendek saja per key, maksimal 35 kata).
2. GUNAKAN ANGKA TARGET RASIONAL yang sudah disediakan di dalam konteks data (JANGAN membuat target tidak masuk akal seperti memaksa STU dari 2 langsung ke 25 dalam 7 hari; gunakan angka sasaran rasional & menantang yang telah dihitung di data konteks).
3. JANGAN gunakan tag HTML apa pun (teks polos murni).
4. Gunakan kata 'dealer' (bukan cabang) dan istilah 'zona kritis'.
5. Pada setiap key plan (kecuali plan_leasing), WAJIB sertakan frasa: TANPA MENURUNKAN KUALITAS DATABASE.
6. Gunakan acuan standar ideal: DB Quality (HOT/DB) = 30%, STU% (STU/DB) = 10%, Approval Leasing ACC% = 50%.
7. Untuk Leasing: bedah masalah bottleneck konversi SPK ke Faktur Kredit (BAF/ADIRA/SOF), risiko SPK Terbuka, serta sebutkan alasan penolakan (Reject) tertinggi beserta solusinya.

Format JSON:
{{
  "cond_dealer": "Maks 35 kata: diagnosa tajam akar masalah STU Dealer vs Ideal/LM & bottleneck konversi DB->HOT (ideal 30%) atau HOT->STU (ideal 10%).",
  "plan_dealer": "Maks 35 kata: target rasional tambahan STU (sesuai konteks), STU/hari, & DB/MP/hari TANPA MENURUNKAN KUALITAS DATABASE.",
  "cond_type": "Maks 35 kata: bedah segmen type terburuk (AT PREMIUM/CLASSY/LPM), penyebab anjloknya dilihat dari DB Q, ACC%, atau STU%.",
  "plan_type": "Maks 35 kata: target rasional tambahan STU, SPK, & DB harian fokus segmen terburuk TANPA MENURUNKAN KUALITAS DATABASE.",
  "cond_leasing": "Maks 35 kata: bedah gap SPK vs Faktur Kredit (BAF/ADIRA/SOF), bahaya SPK Terbuka, & alasan Reject tertinggi.",
  "plan_leasing": "Maks 35 kata: target eksekusi SPK Terbuka 1x24 jam, alokasi aplikasi ke leasing terkuat, & solusi mengatasi alasan Reject utama.",
  "cond_koor": "Maks 35 kata: diagnosa kelemahan utama Koordinator terburuk (apakah kurang DB lapangan, rendah DB Q <30%, atau macet di ACC%/STU%).",
  "plan_koor": "Maks 35 kata: target rasional tambahan STU, SPK, & DB/MP/hari tim koordinator tersebut TANPA MENURUNKAN KUALITAS DATABASE.",
  "cond_counter": "Maks 35 kata: diagnosa kelemahan Counter terburuk dari sisi konversi walk-in/leads digital (DB Q vs 30%, STU% vs 10%, ACC% vs 50%).",
  "plan_counter": "Maks 35 kata: target rasional tambahan STU, follow-up SPK terbuka, & input DB harian counter TANPA MENURUNKAN KUALITAS DATABASE.",
  "cond_mp": "Maks 35 kata: diagnosa pola masalah Salesman lapangan di zona kritis (target besar tapi STU/DB Q/ACC% rendah).",
  "plan_mp": "Maks 35 kata: kuota rasional +1 s.d +2 STU, SPK, & 4-5 DB/hari per salesman kritis TANPA MENURUNKAN KUALITAS DATABASE.",
  "cond_md": "Maks 35 kata: analisa aktifitas MD KHUSUS prospect type Event; sebutkan sub prospect Event yang TERLEMAH (DB, HOT/DB Q, STU/STU%) dan penyebabnya.",
  "plan_md": "Maks 35 kata: target rasional STU & follow-up HOT untuk sub Event terlemah TANPA MENURUNKAN KUALITAS DATABASE.",
  "cond_sosmed": "Maks 35 kata: analisa prospect dari Media Sosial (DB, HOT/DB Q, STU/STU%) vs standar 30% dan 10%.",
  "plan_sosmed": "Maks 35 kata: target rasional STU & follow-up HOT Media Sosial TANPA MENURUNKAN KUALITAS DATABASE."
}}"""
    if clean_key.startswith('AIza'):
        auth_variants = [
            ({'Content-Type': 'application/json', 'x-goog-api-key': clean_key}, f"?key={urllib.parse.quote(clean_key)}"),
            ({'Content-Type': 'application/json', 'x-goog-api-key': clean_key}, "")
        ]
    else:
        auth_variants = [
            ({'Content-Type': 'application/json', 'Authorization': f'Bearer {clean_key}'}, ""),
            ({'Content-Type': 'application/json', 'x-goog-api-key': clean_key}, ""),
            ({'Content-Type': 'application/json', 'x-goog-api-key': clean_key}, f"?key={urllib.parse.quote(clean_key)}")
        ]
    candidate_models = [('v1beta', 'gemini-2.5-flash'), ('v1beta', 'gemini-2.0-flash'), ('v1beta', 'gemini-1.5-flash'), ('v1', 'gemini-2.5-flash'), ('v1', 'gemini-2.0-flash'), ('v1', 'gemini-1.5-flash')]
    last_err = None
    for api_ver, model_name in candidate_models:
        for hdrs, q_str in auth_variants:
            url = f"https://generativelanguage.googleapis.com/{api_ver}/models/{model_name}:generateContent{q_str}"
            gen_cfg = {"temperature": 0.45, "responseMimeType": "application/json"} if api_ver == 'v1beta' else {"temperature": 0.45}
            try:
                req = urllib.request.Request(url, data=json.dumps({"contents": [{"parts": [{"text": prompt}]}], "generationConfig": gen_cfg}).encode('utf-8'), headers=hdrs)
                with urllib.request.urlopen(req, timeout=25) as resp:
                    txt = json.loads(resp.read().decode('utf-8'))['candidates'][0]['content']['parts'][0]['text'].strip()
                    if txt.startswith("```"): txt = txt.strip("`").replace("json", "", 1).strip()
                    s_i, e_i = txt.find('{'), txt.rfind('}')
                    if s_i != -1 and e_i != -1: txt = txt[s_i:e_i+1]
                    return {k: limit_words(v, MAX_NARRATIVE_WORDS) for k, v in json.loads(txt).items()}
            except urllib.error.HTTPError as he:
                err_body = ""
                try: err_body = he.read().decode('utf-8')
                except Exception: pass
                last_err = f"HTTP {he.code} ({api_ver}/{model_name}): {err_body[:160] or he.reason}"
                continue
            except Exception as e:
                last_err = f"{model_name}: {str(e)}"
                continue
    raise RuntimeError(last_err or "Gagal menghubungi Gemini API.")

render_sidebar_azura_logo()
st.sidebar.markdown('### ⚔️ JEMBER BONDOWOSO ANALYSIS')
MAIN_TABS, UPLOAD_TAB, MASTER_TAB = ['🏢 STU', '📊 DASHBOARD KPI'], '📁 UPLOAD FILE', '🎯 MASTER KPI'
UPLOAD_MENU = [UPLOAD_TAB, MASTER_TAB]
if st.session_state.get('active_menu_tab') not in MAIN_TABS + UPLOAD_MENU:
    st.session_state['active_menu_tab'] = '🏢 STU'
if st.session_state.get('nav_main_radio') not in MAIN_TABS + [None]:
    del st.session_state['nav_main_radio']

def on_main_nav_change():
    if st.session_state.get('nav_main_radio') is not None: st.session_state['active_menu_tab'], st.session_state['nav_upload_radio'] = st.session_state['nav_main_radio'], None
def on_upload_nav_change():
    if st.session_state.get('nav_upload_radio') is not None: st.session_state['active_menu_tab'], st.session_state['nav_main_radio'] = st.session_state['nav_upload_radio'], None

st.sidebar.radio('Menu Jember Bondowoso Analysis:', MAIN_TABS, index=MAIN_TABS.index(st.session_state['active_menu_tab']) if st.session_state['active_menu_tab'] in MAIN_TABS else None, key='nav_main_radio', on_change=on_main_nav_change, label_visibility='collapsed')
st.sidebar.markdown('<hr>', unsafe_allow_html=True)
st.sidebar.radio('Menu Upload File:', UPLOAD_MENU, index=UPLOAD_MENU.index(st.session_state['active_menu_tab']) if st.session_state['active_menu_tab'] in UPLOAD_MENU else None, key='nav_upload_radio', on_change=on_upload_nav_change, label_visibility='collapsed')
st.sidebar.markdown('<hr>', unsafe_allow_html=True)
# ---------------------------------------------------------------- KODE AKSES ----
# Kode SPV dapat diganti lewat st.secrets['KODE_AKSES_SPV'] atau environment variable BSG_KODE_AKSES_SPV (default: SPV).
def _resolve_kode_spv():
    try: v_ = st.secrets.get('KODE_AKSES_SPV', None)
    except Exception: v_ = None
    return str(v_ or os.environ.get('BSG_KODE_AKSES_SPV') or 'SPV').strip()
st.sidebar.markdown('### 🔑 Kode Akses')
_kode_in = (st.sidebar.text_input('Kode Akses:', type='password', key='kode_akses_input', label_visibility='collapsed', placeholder='Masukkan kode akses') or '').strip()
IS_SPV = bool(_kode_in) and hmac.compare_digest(_kode_in.encode('utf-8'), _resolve_kode_spv().encode('utf-8'))
if IS_SPV: st.sidebar.success('✅ Akses SPV — semua menu terlihat')
elif _kode_in: st.sidebar.error('❌ Kode akses salah — tab Analisa & Action Plan disembunyikan')
else: st.sidebar.caption('Kode kosong = akses standar (tab 🧠 Analisa & Action Plan disembunyikan).')
st.sidebar.markdown('<hr>', unsafe_allow_html=True)
st.sidebar.markdown('### 🗄️ Arsip Data')
_arsip_do_export = st.sidebar.button('📤 EXPORT', use_container_width=True, key='btn_arsip_export', help='Simpan seluruh data & pengaturan ke 1 file arsip (.zip) agar dashboard bulan ini bisa dibuka lagi kapan saja.')
if st.sidebar.button('📥 IMPORT', use_container_width=True, key='btn_arsip_import', help='Buka kembali dashboard dari file arsip hasil EXPORT.'): st.session_state['arsip_show_import'] = not st.session_state.get('arsip_show_import', False)
_arsip_slot = st.sidebar.container()
with _arsip_slot:
    if ARCHIVE_MODE:
        st.markdown(f'<div style="font-size:12px;font-weight:700;color:#FFEA00;line-height:1.5;">📦 MODE ARSIP aktif<br>{ARCHIVE_META.get("periode", "-")}</div>', unsafe_allow_html=True)
        if st.button('↩️ Kembali ke data terkini', use_container_width=True, key='btn_arsip_back'):
            bsg_arsip_reset(); st.rerun()
    if st.session_state.get('arsip_show_import'):
        _arsip_up = st.file_uploader('Pilih file arsip (.zip) hasil EXPORT', type=['zip'], key='arsip_uploader')
        if _arsip_up is not None and st.session_state.get('arsip_sig') != f'{_arsip_up.name}_{_arsip_up.size}':
            try:
                bsg_arsip_activate(_arsip_up); st.rerun()
            except Exception as _e_arsip: st.error(f'Gagal membuka arsip: {_e_arsip}')
st.sidebar.markdown('<hr>', unsafe_allow_html=True)
if ARCHIVE_MODE:
    st.markdown(f'<div style="margin:0 0 12px 0;padding:10px 14px;border:1px solid #2D3748;border-left:5px solid #FFEA00;border-radius:8px;background:linear-gradient(90deg,rgba(255,234,0,0.12) 0%,rgba(14,17,23,0.96) 100%);font-size:13.5px;line-height:1.6;color:#E2E8F0;">📦 <b>MODE ARSIP</b> — dashboard periode <b>{ARCHIVE_META.get("periode", "-")}</b> · diekspor {ARCHIVE_META.get("exported_at", "-")} · file <code>{st.session_state.get("arsip_name", "")}</code>. Data terkini Anda <b>tidak berubah</b>. Klik <b>↩️ Kembali ke data terkini</b> di sidebar untuk keluar dari mode ini.</div>', unsafe_allow_html=True)
menu_pilihan = st.session_state['active_menu_tab']

def smart_parse_datetime(s):
    if s is None or len(s) == 0 or pd.api.types.is_datetime64_any_dtype(s): return pd.to_datetime(s, errors='coerce')
    s_str = s.dropna().astype(str).str.strip()
    if s_str.empty: return pd.to_datetime(s, errors='coerce')
    return pd.to_datetime(s, errors='coerce', yearfirst=True) if s_str.str.match(r'^\d{4}[-/]\d{1,2}[-/]\d{1,2}').mean() > 0.5 else pd.to_datetime(s, errors='coerce', dayfirst=True)

def make_unique_columns(cols):
    seen, res = {}, []
    for i, c in enumerate(cols):
        cs = str(c).strip()
        if not cs or cs.upper() in ['NAN', 'NONE']: cs = f'UNNAMED_{i}'
        seen[cs] = seen.get(cs, -1) + 1
        res.append(f"{cs}_{seen[cs]}" if seen[cs] > 0 else cs)
    return res

def parse_date_cols(df):
    if df is None or df.empty: return df
    cu = [str(c).strip().upper() for c in df.columns]
    idx = next((i for i, x in enumerate(cu) if x in ['TANGGAL', 'TGL PENGAJUAN', 'TGL NOTA', 'TGL FAKTUR', 'TGL SPK', 'TANGGAL SPK', 'SPK DATE', 'DATE', 'TGL', 'PROSPECT DATE', 'TGL INPUT', 'TGL DATABASE', 'TANGGAL INPUT', 'TGL PROSPEK']), None)
    if idx is None: idx = next((i for i, c in enumerate(cu) if any(k in c for k in ['TGL', 'TANGGAL', 'DATE']) and 'LAHIR' not in c), 4 if df.shape[1] > 4 else None)
    if idx is not None:
        df['Tanggal_Parsed'] = smart_parse_datetime(df.iloc[:, idx])
        df['bulan'], df['hari'], df['tahun'] = df['Tanggal_Parsed'].dt.month, df['Tanggal_Parsed'].dt.day, df['Tanggal_Parsed'].dt.year
    return df

def normalize_leasing_val(val):
    s = clean_str(val)
    if not s or s in ['NAN', 'NONE', '-']: return ''
    if s in ['CASH', 'TUNAI', 'KSM MANDIRI', 'KSMBCA', 'MICROFINANCE INDO MOBILE FINANCE INDONESIA'] or any(k in s for k in ['MICROFINANCE', 'KSM MANDIRI', 'KSMBCA', 'KSM BCA', 'CASH', 'TUNAI']): return 'CASH'
    if s in ['ADIRA', 'ADIRA FINANCE'] or 'ADIRA' in s: return 'ADIRA'
    if s in ['BAF', 'BUSSAN AUTO FINANCE'] or any(k in s for k in ['BAF', 'BUSSAN']): return 'BAF'
    if s in ['SOF', 'SUMMIT OTO FINANCE'] or any(k in s for k in ['SOF', 'SUMMIT', 'OTO MULTIARTHA']): return 'SOF'
    if s in ['IMFI', 'INDOMOBIL FINANCE INDONESIA'] or any(k in s for k in ['IMFI', 'INDOMOBIL']): return 'IMFI'
    return 'OTHERS'

def assign_leasing_faktur(df):
    if df is None or df.empty:
        if df is not None: df['Leasing_Faktur_Clean'] = ''
        return df
    cu = [str(c).strip().upper() for c in df.columns]
    best_c = df.columns[cu.index('LEASING')] if 'LEASING' in cu else (df.columns[20] if df.shape[1] > 20 else None)
    if best_c is None:
        ign = {'NAMA DEALER', 'TANGGAL_PARSED', 'BULAN', 'HARI', 'TAHUN', 'KOOR', 'TYPE_CLEAN', 'LEASING_FAKTUR_CLEAN', 'SALESMAN_CLEAN'}
        best_sc = 0
        for i, cn in enumerate(df.columns):
            if str(cn).strip().upper() in ign: continue
            sc = int(df.iloc[:, i].dropna().astype(str).apply(normalize_leasing_val).isin(['BAF', 'ADIRA', 'SOF', 'IMFI', 'OTHERS', 'CASH']).sum())
            if sc > best_sc: best_sc, best_c = sc, cn
    if best_c is None: best_c = next((df.columns[cu.index(c)] for c in ['LEASING', 'NAMA LEASING', 'FINCOY', 'FINANCE COMPANY', 'NAMA FINCOY', 'KREDITUR', 'PEMBIAYAAN', 'VIA'] if c in cu), None)
    if best_c is None: best_c = next((df.columns[i] for i, c in enumerate(cu) if any(k in c for k in ['LEASING', 'FINCOY', 'FINANCE', 'PEMBIAYAAN', 'KREDIT'])), None)
    s_t = df[best_c].iloc[:, 0] if (best_c is not None and isinstance(df[best_c], pd.DataFrame)) else (df[best_c] if best_c is not None else None)
    df['Leasing_Faktur_Clean'] = s_t.apply(normalize_leasing_val) if s_t is not None else ''
    return df

def detect_type_col(df, map_d, map_ns):
    if df is None or df.empty: return None
    ign = {'NAMA_DEALER_CLEAN', 'DATA_TYPE_CLEAN', 'TANGGAL_PARSED', 'BULAN', 'HARI', 'TAHUN', 'KOOR', 'KOOR_SPK', 'NAMA_DEALER_SPK', 'LEASING_SPK_CLEAN', 'STATUS_SPK_CLEAN', 'ALASAN_REJECT_CLEAN', 'SALESMAN_CLEAN', 'TYPE_CLEAN', 'DATA TYPE', 'LEASING_FAKTUR_CLEAN'}
    cu = [str(c).strip().upper() for c in df.columns]
    if map_d:
        ku, kns, best_c, best_sc = set(map_d.keys()), set(map_ns.keys()), None, 0
        for i, cn in enumerate(df.columns):
            if cu[i] in ign or 'DATA TYPE' in cu[i]: continue
            s = df.iloc[:, i].dropna().astype(str)
            if s.empty: continue
            sc = max(len(set(s.str.split().str.join(' ').str.upper().unique()).intersection(ku)), len(set(s.str.split().str.join('').str.upper().unique()).intersection(kns)))
            if sc > best_sc: best_sc, best_c = sc, cn
        if best_sc > 0: return best_c
    for cand in ['TYPE', 'TIPE', 'NAMA TYPE', 'NAMA TIPE', 'TYPE MOTOR', 'TIPE MOTOR', 'MODEL', 'MODEL NAME', 'PRODUCT NAME', 'UNIT', 'NAMA UNIT', 'DESKRIPSI UNIT', 'DESKRIPSI', 'NAMA BARANG', 'TYPE UNIT', 'TIPE UNIT', 'INTERESTED MODEL', 'PRODUCT', 'MOTOR', 'SERIES']:
        if cand in cu: return df.columns[cu.index(cand)]
    return next((df.columns[i] for i, c in enumerate(cu) if c not in ign and not any(x in c for x in ['DATA TYPE', 'PAYMENT', 'BAYAR', 'TRANSAKSI']) and any(k in c for k in ['TYPE', 'TIPE', 'MODEL', 'PRODUK', 'PRODUCT', 'UNIT', 'MOTOR'])), None)

def assign_type_from_lookup(df, map_d, map_ns, unmapped=None):
    if df is None or df.empty:
        if df is not None: df['Type_Clean'] = 'LAIN-LAIN'
        return df
    tc = detect_type_col(df, map_d, map_ns)
    if tc is not None and map_d:
        def map_fn(val):
            k = clean_str(val)
            if not k or k in ['NAN', 'NONE', '', 'NA', '-', 'TOTAL', 'JUMLAH']: return 'LAIN-LAIN'
            if k in map_d: return map_d[k]
            k_ns = nospace_str(val)
            if k_ns in map_ns: return map_ns[k_ns]
            if unmapped is not None: unmapped.add(k)
            return 'LAIN-LAIN'
        s_t = df[tc].iloc[:, 0] if isinstance(df[tc], pd.DataFrame) else df[tc]
        df['Type_Clean'] = s_t.apply(map_fn)
    else: df['Type_Clean'] = 'LAIN-LAIN'
    return df

def get_file_mtime(fpath):
    try: return os.path.getmtime(fpath) if fpath and os.path.exists(fpath) else 0.0
    except Exception: return 0.0

@st.cache_data(show_spinner=False)
def read_single_db_file_cached(fpath, mtime):
    df_raw = pd.read_csv(fpath, header=None) if fpath.lower().endswith('.csv') else pd.read_excel(fpath, header=None)
    df_raw = smart_text_to_columns(df_raw)
    if df_raw is None or df_raw.empty: return pd.DataFrame()
    h_idx = next((r for r in range(min(25, len(df_raw))) if any(m in [clean_str(v) for v in df_raw.iloc[r].values if pd.notna(v)] for m in ['CUSTOMER NAME', 'POINT NAME', 'PROSPECT DATE', 'DATA TYPE', 'KTP NO.', 'POINT CODE'])), 0)
    df = df_raw.iloc[h_idx + 1:].copy()
    df.columns = make_unique_columns([str(c).strip() for c in df_raw.iloc[h_idx].values])
    df = df.dropna(how='all').reset_index(drop=True)
    cu = [c.upper() for c in df.columns]
    s_chk = df.iloc[:, cu.index('CUSTOMER NAME')] if 'CUSTOMER NAME' in cu else df.iloc[:, 0]
    nm_up = s_chk.astype(str).str.strip().str.upper()
    _keep_hdr = ~nm_up.isin(['CUSTOMER NAME', 'NO', 'NO.', 'TOTAL', 'JUMLAH'])   # baris judul/total dibuang
    df = df[_keep_hdr.values].reset_index(drop=True)
    nm_up = nm_up[_keep_hdr.values].reset_index(drop=True)
    df['Nama_Dealer_Clean'] = normalize_dealer_name(os.path.splitext(os.path.basename(fpath))[0]) or 'TIDAK DIKETAHUI'
    col_k = df.iloc[:, cu.index('POINT NAME')].astype(str).str.strip().str.upper() if 'POINT NAME' in cu else (df.iloc[:, 10].astype(str).str.strip().str.upper() if df.shape[1] > 10 else None)
    df['Data_Type_Clean'] = df.iloc[:, cu.index('DATA TYPE')].astype(str).str.strip().str.upper() if 'DATA TYPE' in cu else (df.iloc[:, 5].astype(str).str.strip().str.upper() if df.shape[1] > 5 else '')
    df = parse_date_cols(df)
    _nm_ok = ~nm_up.isin(['NAN', 'NONE', '']).values
    # baris tanpa Customer Name tetap dihitung bila tanggal prospek valid (sama seperti hitung manual dari file asli)
    df = df[_nm_ok | df['Tanggal_Parsed'].notna().values].reset_index(drop=True) if 'Tanggal_Parsed' in df.columns else df[_nm_ok].reset_index(drop=True)
    return df

@st.cache_data(show_spinner=False)
def db_file_rawcount(fpath, mtime):
    """Jumlah baris data mentah di file Database (di luar judul kolom) — untuk rekonsiliasi."""
    try:
        df_raw = pd.read_csv(fpath, header=None) if fpath.lower().endswith('.csv') else pd.read_excel(fpath, header=None)
        df_raw = smart_text_to_columns(df_raw)
        if df_raw is None or df_raw.empty: return 0
        h_idx = next((r for r in range(min(25, len(df_raw))) if any(m in [clean_str(v) for v in df_raw.iloc[r].values if pd.notna(v)] for m in ['CUSTOMER NAME', 'POINT NAME', 'PROSPECT DATE', 'DATA TYPE', 'KTP NO.', 'POINT CODE'])), 0)
        return int(len(df_raw.iloc[h_idx + 1:].dropna(how='all')))
    except Exception: return 0

def read_single_db_file(fpath):
    return read_single_db_file_cached(fpath, get_file_mtime(fpath)).copy()

def smart_text_to_columns(df):
    if df is not None and df.shape[1] == 1 and len(df) > 0:
        sample_str = str(df.iloc[0, 0])
        for delim in ['\t', ';', ',', '|']:
            if delim in sample_str:
                df_split = df.iloc[:, 0].str.split(delim, expand=True)
                if df_split.shape[1] > 1:
                    return df_split
    return df

@st.cache_data(show_spinner=False)
def read_single_rating_file_cached(fpath, mtime):
    df_raw = pd.read_csv(fpath, header=None) if fpath.lower().endswith('.csv') else pd.read_excel(fpath, header=None)
    df_raw = smart_text_to_columns(df_raw)
    if df_raw is None or df_raw.empty: return pd.DataFrame()
    rating_markers = ['KODE DEALER', 'DEALER CODE', 'POINT CODE', 'POINT NAME', 'CUSTOMER NAME', 'NAMA CUSTOMER', 'NAMA KONSUMEN', 'RATING', 'REVIEW', 'STAR', 'BINTANG', 'SCORE', 'NILAI', 'SALESMAN', 'NAMA SALES', 'NAMA PEGAWAI', 'PRAMUNIAGA', 'TANGGAL', 'DATE', 'PROSPECT DATE']
    h_idx = next((r for r in range(min(25, len(df_raw))) if any(m in [clean_str(v) for v in df_raw.iloc[r].values if pd.notna(v)] for m in rating_markers)), 0)
    df = df_raw.iloc[h_idx + 1:].copy()
    df.columns = make_unique_columns([str(c).strip() for c in df_raw.iloc[h_idx].values])
    df = df.dropna(how='all').reset_index(drop=True)
    if df.empty: return pd.DataFrame()
    cu = [c.upper() for c in df.columns]
    chk_col = next((df.columns[cu.index(c)] for c in ['CUSTOMER NAME', 'NAMA CUSTOMER', 'NAMA KONSUMEN', 'NAMA PELANGGAN', 'NAMA SALES', 'SALESMAN', 'KODE DEALER', 'DEALER CODE', 'POINT CODE'] if c in cu), df.columns[0])
    s_chk = df[chk_col].iloc[:, 0] if isinstance(df[chk_col], pd.DataFrame) else df[chk_col]
    df = df[s_chk.notna() & ~s_chk.astype(str).str.strip().str.upper().isin(['CUSTOMER NAME', 'NAMA CUSTOMER', 'NAMA KONSUMEN', 'KODE DEALER', 'DEALER CODE', 'POINT CODE', 'NO', 'NO.', 'TOTAL', 'JUMLAH', 'NAN', 'NONE', ''])].reset_index(drop=True)
    fallback_dlr = normalize_dealer_name(os.path.splitext(os.path.basename(fpath))[0])
    df['Nama_Dealer_Clean'] = fallback_dlr if fallback_dlr in daftar_dealer else (fallback_dlr or 'TIDAK DIKETAHUI')
    return parse_date_cols(df)

def read_single_rating_file(fpath):
    return read_single_rating_file_cached(fpath, get_file_mtime(fpath)).copy()

def assign_rating_dealer_and_koor(df, map_dlr_code):
    if df is None or df.empty:
        if df is not None:
            df['Kode_Dealer_Clean'] = ''
            df['Salesman_Clean'] = None
            df['Koor'] = 'LAIN-LAIN'
            df['Nama_Dealer_Clean'] = 'TIDAK DIKETAHUI'
        return df
    
    df.columns = make_unique_columns([str(c).strip() for c in df.columns])
    cu = [str(c).strip().upper() for c in df.columns]
    fallback_file_dlr = df['Nama_Dealer_Clean'].iloc[0] if ('Nama_Dealer_Clean' in df.columns and len(df) > 0) else 'TIDAK DIKETAHUI'
    
    # Deteksi cerdas kolom Salesman dengan mencocokkan isi kolom terhadap daftar nama Salesman di Data Lookup
    sm_col = next((df.columns[cu.index(c)] for c in ['SALESMAN/TEKNISI', 'NAMA SALESMAN', 'SALESMAN', 'NAMA SALES', 'NAMA PEGAWAI', 'PRAMUNIAGA', 'SALES NAME', 'PIC'] if c in cu), None)
    if sm_col is None and (mapping_salesman_koor or mapping_salesman_dealer):
        known_sm_set = set(mapping_salesman_koor.keys()) | set(mapping_salesman_dealer.keys())
        known_sm_ns = {alnum_str(k) for k in known_sm_set if alnum_str(k)}
        best_c, best_sc = None, 0
        for i, cn in enumerate(df.columns):
            if cu[i] in {'NAMA_DEALER_CLEAN', 'TANGGAL_PARSED', 'BULAN', 'HARI', 'TAHUN'}: continue
            s_vals = df.iloc[:, i].dropna().astype(str).apply(clean_str)
            if s_vals.empty: continue
            sc = int(s_vals.isin(known_sm_set).sum()) + int(s_vals.apply(alnum_str).isin(known_sm_ns).sum())
            if sc > best_sc:
                best_sc, best_c = sc, cn
        if best_sc > 0:
            sm_col = best_c
    if sm_col is None and df.shape[1] > 9: sm_col = df.columns[9]
    if sm_col is None:
        sm_col = next((df.columns[i] for i, c in enumerate(cu) if any(k in c for k in ['SALES', 'PEGAWAI', 'PRAMUNIAGA', 'TEKNISI'])), None)

    kd_col = next((df.columns[cu.index(c)] for c in ['KODE DEALER', 'DEALER CODE', 'POINT CODE', 'KODE CABANG', 'BRANCH CODE', 'KD DEALER', 'KODE_DEALER', 'DEALER ID', 'OUTLET CODE', 'KODE OUTLET'] if c in cu), None)
    if kd_col is None and df.shape[1] > 2: kd_col = df.columns[2]
    
    dn_col = next((df.columns[cu.index(c)] for c in ['NAMA DEALER', 'DEALER', 'POINT NAME', 'CABANG'] if c in cu), None)
    if dn_col is None and df.shape[1] > 3: dn_col = df.columns[3]

    map_sm_dlr_ns = {alnum_str(k): v for k, v in mapping_salesman_dealer.items() if alnum_str(k)} if mapping_salesman_dealer else {}
    map_sm_koor_ns = {alnum_str(k): v for k, v in mapping_salesman_koor.items() if alnum_str(k)} if mapping_salesman_koor else {}

    def resolve_koor_from_sm(sm_c):
        if not sm_c or sm_c in ['NAN', 'NONE', '', '-']: return 'LAIN-LAIN'
        if mapping_salesman_koor:
            if sm_c in mapping_salesman_koor: return clean_str(mapping_salesman_koor[sm_c])
            sm_ns = alnum_str(sm_c)
            if sm_ns and sm_ns in map_sm_koor_ns: return clean_str(map_sm_koor_ns[sm_ns])
            for k_s, v_k in mapping_salesman_koor.items():
                if len(k_s) >= 4 and (k_s in sm_c or sm_c in k_s): return clean_str(v_k)
        return 'LAIN-LAIN'

    def resolve_dlr_and_koor(kd_val, dn_val, sm_val, kr_resolved):
        sm_c = clean_str(sm_val)
        # 1. Prioritas Utama: Lookup Dealer langsung dari Nama Salesman
        if sm_c and mapping_salesman_dealer:
            if sm_c in mapping_salesman_dealer: return mapping_salesman_dealer[sm_c]
            sm_ns = alnum_str(sm_c)
            if sm_ns and sm_ns in map_sm_dlr_ns: return map_sm_dlr_ns[sm_ns]
            for sm_k, dlr_v in mapping_salesman_dealer.items():
                if len(sm_k) >= 4 and (sm_k in sm_c or sm_c in sm_k): return dlr_v

        # 2. Prioritas Kedua: Lookup Dealer dari Koordinator Salesman tersebut
        if kr_resolved and kr_resolved != 'LAIN-LAIN':
            if 'mapping_koor_dealer' in globals() and kr_resolved in mapping_koor_dealer:
                return mapping_koor_dealer[kr_resolved]
            norm_kr = normalize_dealer_name(kr_resolved)
            if norm_kr in daftar_dealer:
                return norm_kr

        # 3. Fallback: Kode Dealer atau Nama Dealer di baris file
        k_code = clean_str(kd_val)
        if k_code.endswith('.0'): k_code = k_code[:-2].strip()
        k_alnum = alnum_str(k_code)
        if k_code:
            if k_code in map_dlr_code: dl = map_dlr_code[k_code]
            elif k_alnum in map_dlr_code: dl = map_dlr_code[k_alnum]
            else: dl = normalize_dealer_name(k_code)
            if dl in daftar_dealer: return dl
            
        d_name = clean_str(dn_val)
        if d_name:
            norm_dn = normalize_dealer_name(d_name)
            if norm_dn in daftar_dealer: return norm_dn
                
        return fallback_file_dlr if fallback_file_dlr in daftar_dealer else (normalize_dealer_name(d_name) or 'TIDAK DIKETAHUI')

    dealers, sm_cleans, koors, kode_dealers = [], [], [], []
    for _, row in df.iterrows():
        kd_v = row[kd_col] if kd_col is not None and kd_col in df.columns else ''
        dn_v = row[dn_col] if dn_col is not None and dn_col in df.columns else ''
        sm_v = row[sm_col] if sm_col is not None and sm_col in df.columns else ''
        
        if isinstance(kd_v, pd.Series): kd_v = kd_v.iloc[0]
        if isinstance(dn_v, pd.Series): dn_v = dn_v.iloc[0]
        if isinstance(sm_v, pd.Series): sm_v = sm_v.iloc[0]
        
        sm_clean = clean_str(sm_v)
        sm_clean = sm_clean if (sm_clean and sm_clean not in ['NAN', 'NONE', '', '-']) else None
        kr_val = resolve_koor_from_sm(sm_clean)
        if not kr_val or kr_val in ['NAN', 'NONE', '-']: kr_val = 'LAIN-LAIN'
        dlr_res = resolve_dlr_and_koor(kd_v, dn_v, sm_clean, kr_val)
                
        dealers.append(dlr_res)
        sm_cleans.append(sm_clean)
        koors.append(kr_val)
        kode_dealers.append(clean_str(kd_v)[:-2].strip() if clean_str(kd_v).endswith('.0') else clean_str(kd_v))

    df['Nama_Dealer_Clean'] = dealers
    df['Salesman_Clean'] = sm_cleans
    df['Koor'] = koors
    df['Kode_Dealer_Clean'] = kode_dealers
    return df

def load_employee_targets(fpath):
    t_exact, t_ns = dict(FIXED_COUNTER_TARGETS), dict(FIXED_COUNTER_TARGETS_NS)
    if not fpath or not os.path.exists(fpath): return t_exact, t_ns
    try:
        df_e = pd.read_csv(fpath, header=None) if fpath.lower().endswith('.csv') else pd.read_excel(fpath, header=None)
        if df_e is not None and df_e.shape[1] >= 6:
            for i in range(len(df_e)):
                nm, pos = clean_str(df_e.iloc[i, 2]), clean_str(df_e.iloc[i, 5])
                if not nm or nm in ['NAMA', 'NAN', 'NONE'] or nm in FIXED_COUNTER_TARGETS or alnum_str(nm) in FIXED_COUNTER_TARGETS_NS: continue
                tv = 5 if 'PLATINUM' in pos else (3 if 'GOLD' in pos else (2 if 'TRAINING' in pos else None))
                if tv is not None: t_exact[nm], t_ns[alnum_str(nm)] = tv, tv
    except Exception as e: st.error(f'Error File Employee: {e}')
    return t_exact, t_ns

if menu_pilihan == '📁 UPLOAD FILE':
    if ARCHIVE_MODE:
        st.warning('📦 MODE ARSIP aktif — menu Upload dinonaktifkan agar data terkini tidak tertimpa. Klik ↩️ Kembali ke data terkini di sidebar untuk mengunggah file baru.')
        st.stop()
    st.markdown('<div class="main-title">UPLOAD FILE</div><div class="sub-title">Unggah file sumber data sesuai kelompoknya (tersimpan otomatis &amp; permanen)</div>', unsafe_allow_html=True)
    meta_info = load_json(RAW_META_FILE)

    def _status_line(kp, lbl):
        m_ = meta_info.get(kp)
        if m_ and os.path.exists(m_.get('saved_path', '')):
            st.markdown(f"<div style='font-size:12px;color:#94A3B8;line-height:1.5;'>📌 <b>Aktif:</b> <code style='font-size:11px;'>{m_['original_name']}</code> · {m_['updated_at']}</div>", unsafe_allow_html=True)
        else: st.markdown("<div style='font-size:12px;color:#FBBF24;'>⚠️ Belum ada file.</div>", unsafe_allow_html=True)

    # ---------- Grup 1: Data Penjualan & Master ----------
    with st.container(border=True):
        st.markdown("#### 📒 Data Penjualan & Master")
        st.caption('Diupload oleh admin MD')
        up_defs = [('Faktur', 'faktur', ['xls', 'xlsx'], 'File Faktur penjualan (.xls / .xlsx).'),
                   ('Nota', 'nota', ['xls', 'xlsx'], 'File Nota (.xls / .xlsx).'),
                   ('SPK', 'spk', ['xls', 'xlsx', 'csv'], 'File SPK leasing (.xls / .xlsx / .csv).'),
                   ('Master Employee', 'employee', ['xls', 'xlsx', 'csv'], 'Target otomatis salesman — Kolom C: Nama, Kolom F: Posisi.'),
                   ('Data Lookup', 'lookup', ['xls', 'xlsx'], 'Mapping type, salesman, koordinator & dealer (.xls / .xlsx).')]
        for col_u, (lbl, kp, exts, hlp) in zip(st.columns(len(up_defs)), up_defs):
            with col_u:
                uf = st.file_uploader(lbl, type=exts, key=f'{kp}_up', help=hlp)
                if uf: save_raw_file(uf, kp); st.success(f"✅ {lbl} diperbarui: **{uf.name}**"); bsg_toast_once(f'up_{kp}_{uf.name}_{uf.size}', f'{lbl} diperbarui: {uf.name}'); meta_info = load_json(RAW_META_FILE)
                _status_line(kp, lbl)

    # ---------- Grup 2: Data Operasional Dealer ----------
    with st.container(border=True):
        st.markdown("#### 🗂️ Data Operasional Dealer")
        st.caption('Diupload setiap sore oleh admin dealer')
        col_db, col_rt = st.columns(2)
        with col_db:
            db_ups = st.file_uploader('Database — 1 file (parsial) atau 6 file sekaligus (.xls / .xlsx / .csv)', type=['xls', 'xlsx', 'csv'], accept_multiple_files=True, key='db_up')
            if db_ups:
                upd_list, total_active = save_database_files(db_ups)
                st.success(f"✅ **Update Parsial Berhasil!** File diperbarui: `{', '.join(upd_list)}`. File dealer lainnya tetap aman (Total **{total_active} File Database** aktif).")
                bsg_toast_once('up_db_' + '|'.join(f'{u.name}{u.size}' for u in db_ups), f'Database diperbarui ({len(db_ups)} file)')
            saved_db_paths_now = get_saved_database_paths()
            if saved_db_paths_now:
                ftimes = load_json(RAW_META_FILE).get('database_multi', {}).get('file_times', {})
                badge_items = [f"`{os.path.basename(p)}` → **{normalize_dealer_name(os.path.splitext(os.path.basename(p))[0])}** ({ftimes.get(os.path.basename(p), 'Tersimpan')})" for p in saved_db_paths_now]
                st.info(f"📌 **Database Aktif ({len(saved_db_paths_now)} File):** " + " | ".join(badge_items))
                if st.button('🗑️ Reset Semua DB', key='btn_rst_all_db', help='Hapus seluruh file database jika ingin mulai dari 0'):
                    for old_p in saved_db_paths_now:
                        try: os.remove(old_p)
                        except Exception: pass
                    meta_info.pop('database_multi', None); save_json(meta_info, RAW_META_FILE); st.rerun()
            else: st.warning("⚠️ Belum ada file Database.")
        with col_rt:
            rt_ups = st.file_uploader('Rating Dealer — 1 file (parsial) atau 6 file sekaligus (.xls / .xlsx / .csv)', type=['xls', 'xlsx', 'csv'], accept_multiple_files=True, key='rating_up')
            if rt_ups:
                upd_rt_list, total_rt_active = save_rating_files(rt_ups)
                st.success(f"✅ **Update Parsial Rating Berhasil!** File diperbarui: `{', '.join(upd_rt_list)}`. File dealer lainnya tetap aman (Total **{total_rt_active} File Rating** aktif).")
                bsg_toast_once('up_rt_' + '|'.join(f'{u.name}{u.size}' for u in rt_ups), f'Rating diperbarui ({len(rt_ups)} file)')
            saved_rt_paths_now = get_saved_rating_paths()
            if saved_rt_paths_now:
                rt_ftimes = load_json(RAW_META_FILE).get('rating_multi', {}).get('file_times', {})
                rt_badge_items = [f"`{os.path.basename(p)}` → **{normalize_dealer_name(os.path.splitext(os.path.basename(p))[0])}** ({rt_ftimes.get(os.path.basename(p), 'Tersimpan')})" for p in saved_rt_paths_now]
                st.info(f"📌 **File Rating Aktif ({len(saved_rt_paths_now)} File):** " + " | ".join(rt_badge_items))
                if st.button('🗑️ Reset Semua Rating', key='btn_rst_all_rt', help='Hapus seluruh file rating jika ingin mulai dari 0'):
                    for old_p in saved_rt_paths_now:
                        try: os.remove(old_p)
                        except Exception: pass
                    meta_info.pop('rating_multi', None); save_json(meta_info, RAW_META_FILE); st.rerun()
            else: st.warning("⚠️ Belum ada file Rating.")

    # ---------- Logo ----------
    with st.expander('🐺 Logo AZURA (pojok kiri atas)'):
        logo_uf = st.file_uploader('Upload Logo AZURA (.jpg / .png)', type=['jpg', 'jpeg', 'png'], key='azura_logo_up')
        if logo_uf: save_raw_file(logo_uf, 'azura_logo'); st.success(f"✅ Logo AZURA berhasil diperbarui: **{logo_uf.name}**"); st.rerun()

path_faktur, path_nota, path_lookup = get_saved_file_path('faktur'), get_saved_file_path('nota'), get_saved_file_path('lookup')
paths_database, paths_rating, path_spk = get_saved_database_paths(), get_saved_rating_paths(), get_saved_file_path('spk')
path_employee = get_saved_file_path('employee') or (glob.glob('*EMPLOYEE*.xls*')[0] if glob.glob('*EMPLOYEE*.xls*') else None)
mapping_emp_target, mapping_emp_target_ns = load_employee_targets(path_employee)

def get_fixed_counter_target(name_val):
    s = clean_str(name_val)
    if s in FIXED_COUNTER_TARGETS: return FIXED_COUNTER_TARGETS[s]
    s_no_prefix = clean_str(s.replace('COUNTER', '').replace('-', ' '))
    if s_no_prefix in FIXED_COUNTER_TARGETS: return FIXED_COUNTER_TARGETS[s_no_prefix]
    return next((v_fix for k_fix, v_fix in FIXED_COUNTER_TARGETS.items() if k_fix in s), None)

def is_counter_personnel(name_val):
    s = clean_str(name_val)
    return False if not s else (True if (s.startswith('COUNTER') or get_fixed_counter_target(s) is not None) else clean_str(mapping_salesman_koor.get(s, '')).startswith('COUNTER'))

def get_auto_sales_target(sm):
    fix_c = get_fixed_counter_target(sm)
    return fix_c if fix_c is not None else mapping_emp_target.get(clean_str(sm), mapping_emp_target_ns.get(alnum_str(clean_str(sm)), 0))

def get_koor_or_counter_target(ent_name, saved_tk_dict):
    fix_c = get_fixed_counter_target(ent_name)
    return fix_c if fix_c is not None else safe_int(saved_tk_dict.get(ent_name, 0))

mapping_salesman_koor, mapping_salesman_dealer, mapping_dealer_spk, mapping_type_lookup, mapping_type_lookup_nospace, daftar_kategori_type = {}, {}, {}, {}, {}, []
if path_lookup:
    try:
        xl_lk = pd.ExcelFile(path_lookup)
        sh_lk = next((s for s in xl_lk.sheet_names if 'LOOKUP' in s.strip().upper()), xl_lk.sheet_names[0])
        df_lk = pd.read_excel(path_lookup, sheet_name=sh_lk, header=None)
        if df_lk.shape[1] >= 2:
            skip_h = {'TYPE', 'TIPE', 'NAMA TYPE', 'NAMA TIPE', 'UNIT', 'MODEL', 'KATEGORI', 'GROUP TYPE', 'JENIS', 'ROW LABELS', 'NAN', 'NONE', ''}
            for i in range(len(df_lk)):
                ra, rb = clean_str(df_lk.iloc[i, 0]), clean_str(df_lk.iloc[i, 1])
                if ra and rb and ra not in skip_h and rb not in skip_h:
                    mapping_type_lookup[ra] = mapping_type_lookup_nospace[nospace_str(ra)] = rb
                    if rb not in daftar_kategori_type: daftar_kategori_type.append(rb)
        sc, kc, dc, hr = None, None, None, None
        for r in range(min(50, len(df_lk))):
            for c, val in enumerate(df_lk.iloc[r]):
                vs = clean_str(val)
                if any(k in vs for k in ['NAMA SALES', 'SALESMAN', 'PRAMUNIAGA', 'PEGAWAI']): sc = c
                elif 'KOOR' in vs: kc = c
                elif any(k in vs for k in ['DEALER', 'NAMA DEALER', 'CABANG', 'POINT NAME', 'OUTLET']): dc = c
            if sc is not None: hr = r; break
        if sc is not None:
            for i in range(hr + 1, len(df_lk)):
                sm = clean_str(df_lk.iloc[i, sc]) if sc is not None and sc < df_lk.shape[1] else ''
                kr = clean_str(df_lk.iloc[i, kc]) if kc is not None and kc < df_lk.shape[1] else ''
                dlr_lk = clean_str(df_lk.iloc[i, dc]) if dc is not None and dc < df_lk.shape[1] else ''
                if sm and sm not in ['NAN', 'NONE']:
                    if kr and kr not in ['NAN', 'NONE']: mapping_salesman_koor[sm] = kr
                    if dlr_lk and dlr_lk not in ['NAN', 'NONE']:
                        norm_dlr = normalize_dealer_name(dlr_lk)
                        if norm_dlr in daftar_dealer:
                            mapping_salesman_dealer[sm] = norm_dlr
        if df_lk.shape[1] > 9:
            for i in range(len(df_lk)):
                ri, rj = clean_str(df_lk.iloc[i, 8]), clean_str(df_lk.iloc[i, 9])
                if ri.endswith('.0'): ri = ri[:-2].strip()
                if rj.endswith('.0'): rj = rj[:-2].strip()
                if ri and rj and ri not in ['NAN', 'NONE'] and rj not in ['NAN', 'NONE']:
                    ni, nj = normalize_dealer_name(ri), normalize_dealer_name(rj)
                    if nj in daftar_dealer:
                        mapping_dealer_spk[ri] = mapping_dealer_spk[alnum_str(ri)] = nj
                    elif ni in daftar_dealer:
                        mapping_dealer_spk[rj] = mapping_dealer_spk[alnum_str(rj)] = ni
                    else:
                        mapping_dealer_spk[ri] = nj
        for c_idx in range(df_lk.shape[1] - 1):
            for r_idx in range(len(df_lk)):
                v1, v2 = clean_str(df_lk.iloc[r_idx, c_idx]), clean_str(df_lk.iloc[r_idx, c_idx + 1])
                if v1.endswith('.0'): v1 = v1[:-2].strip()
                if v2.endswith('.0'): v2 = v2[:-2].strip()
                if not v1 or not v2 or v1 in ['NAN', 'NONE'] or v2 in ['NAN', 'NONE']: continue
                n1, n2 = normalize_dealer_name(v1), normalize_dealer_name(v2)
                if n2 in daftar_dealer and v1 not in daftar_dealer:
                    mapping_dealer_spk[v1] = mapping_dealer_spk[alnum_str(v1)] = n2
                elif n1 in daftar_dealer and v2 not in daftar_dealer:
                    mapping_dealer_spk[v2] = mapping_dealer_spk[alnum_str(v2)] = n1
    except Exception as e: st.error(f'Error Lookup: {e}')

unmapped_types_all, sales_tanpa_koor, sales_tanpa_koor_db, sales_tanpa_koor_spk, sales_tanpa_koor_rating = set(), set(), set(), set(), set()

def assign_koor_and_sales(df, col_idx_override=None, tracker=None, koor_col_name='Koor'):
    if df is None: return df
    if df.empty:
        df['Salesman_Clean'], df[koor_col_name] = None, 'LAIN-LAIN'
        return df
    cu = [str(c).strip().upper() for c in df.columns]
    ign_col_kw = {'CUSTOMER', 'KONSUMEN', 'PELANGGAN', 'RATING', 'REVIEW', 'STAR', 'BINTANG', 'SCORE', 'NILAI', 'KOMENTAR', 'TANGGAL', 'DATE', 'KODE', 'PHONE', 'TELP', 'ALAMAT'}
    sc_name = df.columns[col_idx_override] if (col_idx_override is not None and df.shape[1] > col_idx_override) else None
    if sc_name is None:
        valid_cands = [x for x in ['NAMA PEGAWAI', 'NAMA SALES', 'NAMA PRAMUNIAGA', 'SALESMAN', 'PRAMUNIAGA', 'SALES', 'SALES NAME', 'NAMA SALESMAN', 'SALES PERSON', 'SALESPERSON', 'PIC', 'EMPLOYEE NAME', 'SALESMAN NAME', 'CREATED BY', 'USER INPUT'] if x in cu]
        if valid_cands: sc_name = df.columns[cu.index(valid_cands[0])]
    if sc_name is None:
        # Search columns that don't match ignored keywords
        for i, c_str in enumerate(cu):
            if not any(kw in c_str for kw in ign_col_kw) and any(k in c_str for k in ['NAMA', 'SALES', 'PEGAWAI', 'PRAMUNIAGA', 'PIC']):
                sc_name = df.columns[i]
                break
    if sc_name is not None:
        s_s = df[sc_name].iloc[:, 0] if isinstance(df[sc_name], pd.DataFrame) else df[sc_name]
        s_cl = s_s.astype(str).str.split().str.join(' ').str.upper()
        df['Salesman_Clean'] = s_cl.where(~s_cl.isin(['NAN', 'NONE', '', 'NA', '-', 'TOTAL', 'JUMLAH']), None)
        def map_kr(s):
            if s and pd.notna(s):
                kr_val = clean_str(mapping_salesman_koor.get(s, '')) if mapping_salesman_koor else ''
                if not kr_val or kr_val in ['NAN', 'NONE', '-', 'LAIN-LAIN']:
                    if tracker is not None: tracker.add(s)
                    return 'LAIN-LAIN'
                return kr_val
            return 'LAIN-LAIN'
        df[koor_col_name] = df['Salesman_Clean'].apply(map_kr)
    else: df['Salesman_Clean'], df[koor_col_name] = None, 'LAIN-LAIN'
    return df

df_faktur, df_faktur_all = None, None
if path_faktur:
    try:
        df_faktur_raw = pd.read_excel(path_faktur, header=8)
        if 'Tanggal' not in df_faktur_raw.columns and 'Tgl Pengajuan' not in df_faktur_raw.columns: df_faktur_raw = pd.read_excel(path_faktur, header=0)
        df_faktur_raw = assign_leasing_faktur(parse_date_cols(df_faktur_raw))
        df_faktur_all = df_faktur_raw.copy()
        df_faktur = df_faktur_raw[df_faktur_raw['Nama Dealer'].isin(daftar_dealer)].reset_index(drop=True) if 'Nama Dealer' in df_faktur_raw.columns else df_faktur_raw.copy()
        df_faktur = assign_koor_and_sales(assign_type_from_lookup(df_faktur, mapping_type_lookup, mapping_type_lookup_nospace, unmapped_types_all), tracker=sales_tanpa_koor)
    except Exception as e: st.error(f'Error Faktur: {e}')

df_nota = None
if path_nota:
    try:
        df_nt = pd.read_excel(path_nota, header=5)
        df_nota = parse_date_cols(pd.read_excel(path_nota, header=0) if not any(k in df_nt.columns for k in ['Nama Dealer', 'Dealer', 'Tanggal']) else df_nt)
        cd = 'Nama Dealer' if 'Nama Dealer' in df_nota.columns else ('Dealer' if 'Dealer' in df_nota.columns else None)
        if cd: df_nota = df_nota[df_nota[cd].isin(daftar_dealer)].reset_index(drop=True)
        df_nota = assign_koor_and_sales(assign_type_from_lookup(df_nota, mapping_type_lookup, mapping_type_lookup_nospace, unmapped_types_all), tracker=sales_tanpa_koor)
    except Exception as e: st.error(f'Error Nota: {e}')

df_database = None
if paths_database:
    l_db = [assign_type_from_lookup(read_single_db_file(p), mapping_type_lookup, mapping_type_lookup_nospace, unmapped_types_all) for p in paths_database]
    l_db = [d for d in l_db if not d.empty]
    if l_db: df_database = assign_koor_and_sales(pd.concat(l_db, ignore_index=True), tracker=sales_tanpa_koor_db)
df_database_alldlr = df_database   # sebelum difilter (untuk rekonsiliasi)
if df_database is not None and not df_database.empty and 'Nama_Dealer_Clean' in df_database.columns and df_database['Nama_Dealer_Clean'].isin(daftar_dealer).any():
    df_database = df_database[df_database['Nama_Dealer_Clean'].isin(daftar_dealer)].reset_index(drop=True)   # KPI & Proses memakai basis data yang sama

def strip_sub_code(v):
    """Buang kode di depan nama sub prospect (mis. '03 - Pameran' / 'EV01: Pameran' -> 'PAMERAN')."""
    t = clean_str(v)
    t2 = re.sub(r'^[\[\(]?[A-Z]{0,5}[-_./]?\d+[A-Z0-9._/-]*[\]\)]?\s*[-–:.)|]\s*', '', t).strip()
    t2 = t2 or re.sub(r'^[\[\(]?[A-Z]{0,5}[-_./]?\d+[A-Z0-9._/-]*[\]\)]?\s+', '', t).strip()
    return t2 or t

def ktp_norm(v):
    if v is None or (not isinstance(v, str) and pd.isna(v)): return ''
    t = str(v).strip()
    if t.endswith('.0'): t = t[:-2]
    t = re.sub(r'\D', '', t)
    return t if len(t) >= 10 else ''

# Kolom STU otomatis di Database: KTP (kolom C Database) dicocokkan dengan KTP kolom I file Faktur -> keterangan 'STU'
if df_database is not None and not df_database.empty:
    df_database['STU'] = ''
    if df_database.shape[1] > 2 and df_faktur_all is not None and not df_faktur_all.empty and df_faktur_all.shape[1] > 8:
        _ktp_fak = set(k for k in df_faktur_all.iloc[:, 8].apply(ktp_norm) if k)
        _ktp_db = df_database.iloc[:, 2].apply(ktp_norm)
        df_database['STU'] = ['STU' if (k and k in _ktp_fak) else '' for k in _ktp_db]

df_spk = None
if path_spk:
    try:
        df_sr = pd.read_csv(path_spk, header=None) if path_spk.lower().endswith('.csv') else pd.read_excel(path_spk, sheet_name=next((s for s in pd.ExcelFile(path_spk).sheet_names if s.strip().upper() == 'SPK'), 0), header=None)
        if df_sr is not None and not df_sr.empty:
            hi = next((r for r in range(min(25, len(df_sr))) if len([v for v in df_sr.iloc[r].values if pd.notna(v)]) >= 4 and any(k in ' '.join(clean_str(v) for v in df_sr.iloc[r].values if pd.notna(v)) for k in ['NAMA DEALER', 'NO TRANSAKSI', 'TGL', 'LEASING', 'STATUS'])), 0)
            df_spk = df_sr.iloc[hi + 1:].copy()
            df_spk.columns = make_unique_columns([str(c).strip() for c in df_sr.iloc[hi].values])
            df_spk = df_spk.dropna(how='all').reset_index(drop=True)
            df_spk['Nama_Dealer_SPK'] = df_spk.iloc[:, 3].apply(lambda x: mapping_dealer_spk.get(clean_str(x), mapping_dealer_spk.get(normalize_dealer_name(x), normalize_dealer_name(x)))) if df_spk.shape[1] > 3 else 'TIDAK DIKETAHUI'
            df_spk = df_spk[df_spk['Nama_Dealer_SPK'].isin(daftar_dealer)].reset_index(drop=True)
            df_spk['Leasing_SPK_Clean'] = df_spk.iloc[:, 16].apply(normalize_leasing_val) if df_spk.shape[1] > 16 else ''
            df_spk['Alasan_Reject_Clean'] = df_spk.iloc[:, 18].apply(lambda x: clean_str(x) if clean_str(x) not in ['', 'NAN', 'NONE', '-'] else 'TANPA KETERANGAN') if df_spk.shape[1] > 18 else 'TANPA KETERANGAN'
            if df_spk.shape[1] > 17:
                df_spk['Status_SPK_Clean'] = df_spk.iloc[:, 17].apply(clean_str)
                df_spk = df_spk[df_spk['Status_SPK_Clean'].isin(['DISETUJUI', 'DITOLAK', 'TERBUKA'])].reset_index(drop=True)
            else: df_spk['Status_SPK_Clean'] = ''
            if df_spk.shape[1] > 5:
                df_spk['Tanggal_Parsed'] = smart_parse_datetime(df_spk.iloc[:, 5])
                df_spk['bulan'], df_spk['hari'], df_spk['tahun'] = df_spk['Tanggal_Parsed'].dt.month, df_spk['Tanggal_Parsed'].dt.day, df_spk['Tanggal_Parsed'].dt.year
            df_spk = assign_type_from_lookup(assign_koor_and_sales(df_spk, col_idx_override=25, tracker=sales_tanpa_koor_spk, koor_col_name='Koor_SPK'), mapping_type_lookup, mapping_type_lookup_nospace, unmapped_types_all)
    except Exception as e: st.error(f'Error SPK: {e}')

# Lengkapi mapping_salesman_dealer dan mapping_koor_dealer dari seluruh sumber data (Database, Faktur, Nota, SPK)
mapping_koor_dealer = {}
for _df_s, _dlr_c, _kr_c in [
    (df_database, 'Nama_Dealer_Clean', 'Koor'),
    (df_faktur, 'Nama Dealer', 'Koor'),
    (df_nota, 'Nama Dealer' if (df_nota is not None and 'Nama Dealer' in df_nota.columns) else 'Dealer', 'Koor'),
    (df_spk, 'Nama_Dealer_SPK', 'Koor_SPK')
]:
    if _df_s is not None and not _df_s.empty and _dlr_c in _df_s.columns:
        if 'Salesman_Clean' in _df_s.columns:
            _sub_sm = _df_s[_df_s['Salesman_Clean'].notna() & _df_s[_dlr_c].isin(daftar_dealer)][['Salesman_Clean', _dlr_c]].drop_duplicates()
            for _, _r in _sub_sm.iterrows():
                _sm_n, _dl_n = clean_str(_r['Salesman_Clean']), clean_str(_r[_dlr_c])
                if _sm_n and _dl_n in daftar_dealer and _sm_n not in mapping_salesman_dealer:
                    mapping_salesman_dealer[_sm_n] = _dl_n
        if _kr_c in _df_s.columns:
            _sub_kr = _df_s[_df_s[_kr_c].notna() & (_df_s[_kr_c] != 'LAIN-LAIN') & _df_s[_dlr_c].isin(daftar_dealer)][[_kr_c, _dlr_c]].drop_duplicates()
            for _, _r in _sub_kr.iterrows():
                _kr_n, _dl_n = clean_str(_r[_kr_c]), clean_str(_r[_dlr_c])
                if _kr_n and _dl_n in daftar_dealer and _kr_n not in mapping_koor_dealer:
                    mapping_koor_dealer[_kr_n] = _dl_n

for _sm_k, _kr_v in mapping_salesman_koor.items():
    _kr_c = clean_str(_kr_v)
    if _sm_k not in mapping_salesman_dealer:
        if _kr_c in mapping_koor_dealer:
            mapping_salesman_dealer[_sm_k] = mapping_koor_dealer[_kr_c]
        else:
            _norm_k = normalize_dealer_name(_kr_c)
            if _norm_k in daftar_dealer:
                mapping_salesman_dealer[_sm_k] = _norm_k

df_rating = None
if paths_rating:
    l_rt = [assign_type_from_lookup(assign_rating_dealer_and_koor(read_single_rating_file(p), mapping_dealer_spk), mapping_type_lookup, mapping_type_lookup_nospace, unmapped_types_all) for p in paths_rating]
    l_rt = [d for d in l_rt if not d.empty]
    if l_rt: df_rating = pd.concat(l_rt, ignore_index=True)

def detect_customer_col(df, is_rating=False):
    if df is None or df.empty: return None
    cu = [str(c).strip().upper() for c in df.columns]
    exact_cands = [
        'NAMA KONSUMEN', 'NAMA CUSTOMER', 'CUSTOMER NAME', 'NAMA PELANGGAN',
        'NAMA PEMILIK', 'NAMA STNK', 'NAMA BPKB', 'NAMA PEMBELI', 'KONSUMEN',
        'CUSTOMER', 'PEMBELI', 'PEMILIK', 'NAMA USER', 'USER NAME'
    ]
    for cand in exact_cands:
        if cand in cu:
            return df.columns[cu.index(cand)]
    ign_kw = {'DEALER', 'CABANG', 'OUTLET', 'SALES', 'PEGAWAI', 'PRAMUNIAGA', 'TEKNISI', 'TYPE', 'TIPE', 'BARANG', 'UNIT', 'FINCOY', 'LEASING', 'KOOR', 'KATEGORI', 'FILE', 'POINT', 'STATUS'}
    for i, c_str in enumerate(cu):
        if not any(kw in c_str for kw in ign_kw) and any(k in c_str for k in ['KONSUMEN', 'CUSTOMER', 'PELANGGAN', 'PEMBELI', 'PEMILIK', 'STNK', 'NAMA']):
            return df.columns[i]
    if is_rating and df.shape[1] > 4:
        return df.columns[4]
    return None

def detect_phone_col(df):
    if df is None or df.empty: return None
    cu = [str(c).strip().upper() for c in df.columns]
    for cand in ['NO HP', 'NO. HP', 'NO TELP', 'NO. TELP', 'TELEPON', 'TELP', 'HP', 'HANDPHONE', 'NO HANDPHONE', 'PHONE', 'MOBILE', 'MOBILE PHONE', 'PHONE NO', 'CONTACT', 'WA', 'NO WA']:
        if cand in cu:
            return df.columns[cu.index(cand)]
    for i, c_str in enumerate(cu):
        if any(k in c_str for k in ['HP', 'TELP', 'TELEPON', 'PHONE', 'MOBILE', 'WA', 'KONTAK']) and 'KODE' not in c_str:
            return df.columns[i]
    return None

def detect_raw_type_col(df):
    if df is None or df.empty: return None
    tc = detect_type_col(df, mapping_type_lookup, mapping_type_lookup_nospace)
    if tc is not None: return tc
    cu = [str(c).strip().upper() for c in df.columns]
    for cand in ['NAMA BARANG', 'DESKRIPSI UNIT', 'TIPE MOTOR', 'TYPE MOTOR', 'NAMA TYPE', 'NAMA TIPE', 'TYPE', 'TIPE', 'MODEL', 'UNIT']:
        if cand in cu: return df.columns[cu.index(cand)]
    return None

def assign_faktur_rating_status(df_f, df_rt):
    if df_f is None or df_f.empty:
        if df_f is not None:
            df_f['Nama_Konsumen_Clean'] = ''
            df_f['No_HP_Clean'] = '-'
            df_f['Tipe_Unit_Display'] = '-'
            df_f['Status_Registrasi_Rating'] = 'BELUM REGISTRASI'
        return df_f

    cust_col_f = detect_customer_col(df_f, is_rating=False)
    if cust_col_f is not None:
        s_cf = df_f[cust_col_f].iloc[:, 0] if isinstance(df_f[cust_col_f], pd.DataFrame) else df_f[cust_col_f]
        df_f['Nama_Konsumen_Clean'] = s_cf.apply(clean_str)
    else:
        df_f['Nama_Konsumen_Clean'] = ''

    hp_col_f = detect_phone_col(df_f)
    if hp_col_f is not None:
        s_hp = df_f[hp_col_f].iloc[:, 0] if isinstance(df_f[hp_col_f], pd.DataFrame) else df_f[hp_col_f]
        df_f['No_HP_Clean'] = s_hp.apply(lambda x: clean_str(x)[:-2] if clean_str(x).endswith('.0') else (clean_str(x) or '-'))
    else:
        df_f['No_HP_Clean'] = '-'

    raw_tc = detect_raw_type_col(df_f)
    if raw_tc is not None:
        s_rt = df_f[raw_tc].iloc[:, 0] if isinstance(df_f[raw_tc], pd.DataFrame) else df_f[raw_tc]
        df_f['Tipe_Unit_Display'] = s_rt.apply(lambda x: clean_str(x) or '-')
    else:
        df_f['Tipe_Unit_Display'] = df_f['Type_Clean'] if 'Type_Clean' in df_f.columns else '-'

    rt_names_exact = set()
    rt_names_alnum = set()
    if df_rt is not None and not df_rt.empty:
        cust_col_rt = detect_customer_col(df_rt, is_rating=True)
        cols_to_scan = [cust_col_rt] if cust_col_rt is not None else []
        # Tambahkan juga pengecekan pada seluruh kolom teks di file rating jika nama konsumen berada di kolom lain
        for cn in df_rt.columns:
            if cn not in cols_to_scan and str(cn).upper() not in {'NAMA_DEALER_CLEAN', 'SALESMAN_CLEAN', 'KOOR', 'KODE_DEALER_CLEAN', 'TANGGAL_PARSED', 'BULAN', 'HARI', 'TAHUN', 'TYPE_CLEAN'}:
                cols_to_scan.append(cn)
        if cust_col_rt is not None:
            # Utamakan kolom konsumen utama terlebih dahulu
            s_cr = df_rt[cust_col_rt].iloc[:, 0] if isinstance(df_rt[cust_col_rt], pd.DataFrame) else df_rt[cust_col_rt]
            for val in s_cr.dropna():
                c_v = clean_str(val)
                if c_v and c_v not in {'NAN', 'NONE', '-', 'CUSTOMER NAME', 'NAMA KONSUMEN', 'NAMA CUSTOMER'}:
                    rt_names_exact.add(c_v)
                    a_v = alnum_str(c_v)
                    if len(a_v) >= 3: rt_names_alnum.add(a_v)
        # Jika hasil kecocokan dengan kolom utama masih 0, scan kolom potensial lainnya di df_rt
        f_names_sample = {x for x in df_f['Nama_Konsumen_Clean'].unique() if x and len(x) >= 3}
        if f_names_sample and not f_names_sample.intersection(rt_names_exact):
            for cn in cols_to_scan:
                s_c = df_rt[cn].iloc[:, 0] if isinstance(df_rt[cn], pd.DataFrame) else df_rt[cn]
                cand_set = {clean_str(v) for v in s_c.dropna() if clean_str(v)}
                if f_names_sample.intersection(cand_set):
                    for c_v in cand_set:
                        rt_names_exact.add(c_v)
                        a_v = alnum_str(c_v)
                        if len(a_v) >= 3: rt_names_alnum.add(a_v)

    def check_reg_status(nm):
        c_nm = clean_str(nm)
        if not c_nm or c_nm in {'NAN', 'NONE', '-', ''}:
            return 'BELUM REGISTRASI'
        if c_nm in rt_names_exact:
            return 'SUDAH REGISTRASI'
        a_nm = alnum_str(c_nm)
        if a_nm and len(a_nm) >= 3 and a_nm in rt_names_alnum:
            return 'SUDAH REGISTRASI'
        return 'BELUM REGISTRASI'

    df_f['Status_Registrasi_Rating'] = df_f['Nama_Konsumen_Clean'].apply(check_reg_status)
    return df_f

if df_faktur is not None:
    df_faktur = assign_faktur_rating_status(df_faktur, df_rating)
if df_faktur_all is not None:
    df_faktur_all = assign_faktur_rating_status(df_faktur_all, df_rating)


def get_unmapped_salesman_details():
    details = {}
    def collect_from_df(df_src, koor_col, dlr_col, src_label):
        if df_src is None or df_src.empty or 'Salesman_Clean' not in df_src.columns: return
        sub = df_src[df_src['Salesman_Clean'].notna() & (df_src[koor_col] == 'LAIN-LAIN')] if koor_col in df_src.columns else df_src[df_src['Salesman_Clean'].notna()]
        for _, r in sub.iterrows():
            sm = clean_str(r.get('Salesman_Clean', ''))
            if not sm or sm in ['NAN', 'NONE', 'TOTAL', 'JUMLAH', '-']: continue
            dlr = clean_str(r.get(dlr_col, '')) if dlr_col in df_src.columns else ''
            if sm not in details: details[sm] = {'dealers': set(), 'sources': set()}
            if dlr and dlr not in ['NAN', 'NONE']: details[sm]['dealers'].add(dlr)
            details[sm]['sources'].add(src_label)
    collect_from_df(df_database, 'Koor', 'Nama_Dealer_Clean', 'Database')
    collect_from_df(df_faktur, 'Koor', 'Nama Dealer', 'Faktur')
    cd_nt = 'Nama Dealer' if (df_nota is not None and 'Nama Dealer' in df_nota.columns) else 'Dealer'
    collect_from_df(df_nota, 'Koor', cd_nt, 'Nota')
    collect_from_df(df_spk, 'Koor_SPK', 'Nama_Dealer_SPK', 'SPK')
    collect_from_df(df_rating, 'Koor', 'Nama_Dealer_Clean', 'Rating')
    return details

def render_unmapped_salesman_alert(key_prefix='sales'):
    unmapped_info = get_unmapped_salesman_details()
    if not unmapped_info:
        st.markdown('<div class="unmapped-sales-box" style="border-left-color:#00FF66; background:linear-gradient(90deg, rgba(0,255,102,0.10) 0%, rgba(14,17,23,0.96) 100%);">✅ <b>Status Data Lookup Koordinator:</b> Seluruh nama salesman pada Database, Faktur, Nota, dan SPK sudah memiliki Koordinator.</div>', unsafe_allow_html=True)
        return
    sorted_names = sorted(unmapped_info.keys())
    db_only_names = [nm for nm in sorted_names if 'Database' in unmapped_info[nm]['sources']]
    badges_html = []
    rows_detail = []
    for nm in sorted_names:
        dlrs = ', '.join(sorted(unmapped_info[nm]['dealers'])) or '-'
        srcs = ', '.join(sorted(unmapped_info[nm]['sources']))
        badges_html.append(f'<span class="unmapped-badge">{nm} <span style="color:#94A3B8;font-weight:400;">({dlrs} | {srcs})</span></span>')
        rows_detail.append({'NAMA SALESMAN (BELUM ADA KOORDINATOR)': nm, 'DEALER': dlrs, 'SUMBER DATA': srcs})
    st.markdown(f"""<div class="unmapped-sales-box">
        <div>⚠️ <b>Ditemukan {len(sorted_names)} Nama Salesman Belum Memiliki Koordinator di Data Lookup ({len(db_only_names)} dari Database):</b></div>
        <div style="margin-top:6px;">{''.join(badges_html)}</div>
    </div>""", unsafe_allow_html=True)
    with st.expander(f"📋 Klik untuk Salin / Lihat Tabel {len(sorted_names)} Nama Salesman Tanpa Koordinator (Untuk Update Data Lookup)"):
        c_u1, c_u2 = st.columns([2, 1])
        with c_u1:
            st.dataframe(pd.DataFrame(rows_detail), hide_index=True, use_container_width=True)
        with c_u2:
            st.caption("Salin daftar nama di bawah ini langsung ke kolom Salesman di Excel **Data Lookup**:")
            st.code('\n'.join(sorted_names), language='text')

today_wib = (datetime.datetime.now(datetime.timezone.utc) + datetime.timedelta(hours=7)).date()
current_month, current_year = today_wib.month, today_wib.year
for dc in [df_faktur_all, df_faktur, df_nota, df_spk, df_database]:
    if dc is not None and not dc.empty and 'bulan' in dc.columns and dc['bulan'].notna().any():
        current_year = int(dc['tahun'].dropna().max())
        current_month = int(dc[dc['tahun'] == current_year]['bulan'].dropna().max())
        break
last_month, last_month_year = (current_month - 1, current_year) if current_month > 1 else (12, current_year - 1)
NAMA_BULAN_ID = {1: 'Januari', 2: 'Februari', 3: 'Maret', 4: 'April', 5: 'Mei', 6: 'Juni', 7: 'Juli', 8: 'Agustus', 9: 'September', 10: 'Oktober', 11: 'November', 12: 'Desember'}
nama_bulan_curr = NAMA_BULAN_ID.get(current_month, 'Oktober')
periode_bulan_tahun = f'{nama_bulan_curr} {current_year}'

max_day_in_month = pycal.monthrange(current_year, current_month)[1]
default_day_today = min(today_wib.day, max_day_in_month)
if ARCHIVE_MODE and ARCHIVE_META.get('tanggal_akhir'): default_day_today = min(max(1, int(ARCHIVE_META['tanggal_akhir'])), max_day_in_month)
st.sidebar.markdown('### 📅 Periode Analisa')
tgl_kal_pilihan = st.sidebar.date_input('Pilih Tanggal Akhir Periode (Kalender):', value=datetime.date(current_year, current_month, default_day_today), min_value=datetime.date(current_year, current_month, 1), max_value=datetime.date(current_year, current_month, max_day_in_month), format='DD/MM/YYYY')
tanggal_akhir = int(tgl_kal_pilihan[-1].day) if isinstance(tgl_kal_pilihan, (tuple, list)) and len(tgl_kal_pilihan) > 0 else (int(tgl_kal_pilihan.day) if hasattr(tgl_kal_pilihan, 'day') else default_day_today)

target_persen_harian = tabel_target_harian.get(tanggal_akhir, 0.0)
persen_desimal = target_persen_harian / 100.0
col_name_pct = f'{target_persen_harian:.0f}%' if target_persen_harian.is_integer() else f'{target_persen_harian}%'
LIBUR_FILE = _ABS('libur_nasional.json')
DEFAULT_LIBUR_NASIONAL = ['2026-01-01', '2026-01-16', '2026-02-17', '2026-03-19', '2026-03-21', '2026-03-22', '2026-04-03', '2026-04-05', '2026-05-01', '2026-05-14', '2026-05-27', '2026-05-31', '2026-06-01', '2026-06-16', '2026-08-17', '2026-08-25', '2026-12-25']
def _parse_libur(lst):
    out = set()
    for x in lst:
        try: out.add(datetime.date.fromisoformat(str(x).strip()))
        except Exception: pass
    return out
_libur_saved = load_json(LIBUR_FILE)
LIBUR_SET = _parse_libur(_libur_saved['dates'] if isinstance(_libur_saved, dict) and 'dates' in _libur_saved else DEFAULT_LIBUR_NASIONAL)
def is_workday_date(dt): return dt.weekday() != 6 and dt not in LIBUR_SET   # Minggu & libur nasional bukan hari aktif
def is_workday_day(d_n): return is_workday_date(datetime.date(current_year, current_month, d_n))
WORKDAYS_UPTO = [d_n for d_n in range(1, tanggal_akhir + 1) if is_workday_day(d_n)]
SISA_HARI_KERJA = sum(1 for d_n in range(tanggal_akhir + 1, max_day_in_month + 1) if is_workday_day(d_n))
hari_aktif = max(1, len(WORKDAYS_UPTO))
st.sidebar.markdown(f'<div class="target-ideal-caption">🎯 Target Ideal Tgl {tanggal_akhir}: {col_name_pct}</div>', unsafe_allow_html=True)
if _arsip_do_export:
    try:
        _arsip_bytes, _arsip_name = bsg_build_export_zip({'bulan': int(current_month), 'tahun': int(current_year), 'periode': periode_bulan_tahun, 'tanggal_akhir': int(tanggal_akhir)})
        st.session_state['arsip_export'] = (_arsip_name, _arsip_bytes); bsg_toast('Arsip siap diunduh', '📦')
    except Exception as _e_exp:
        st.session_state['arsip_export'] = None
        with _arsip_slot: st.error(f'Gagal membuat arsip: {_e_exp}')
if st.session_state.get('arsip_export'):
    with _arsip_slot:
        st.download_button('⬇️ Unduh file arsip', data=st.session_state['arsip_export'][1], file_name=st.session_state['arsip_export'][0], mime='application/zip', use_container_width=True, key='dl_arsip')
        st.caption(f"{st.session_state['arsip_export'][0]} · {len(st.session_state['arsip_export'][1]) / 1048576:.1f} MB")

# ===================== ALARM HARIAN — hanya notifikasi popup tepat di jam alarm (tanpa checklist / rekap / alarm susulan) =====================
DEFAULT_REMINDERS = [('briefing', 'Sudah Kirim Laporan Briefing Pagi ?', '08:30'), ('rekrutmen', 'Sudah Kirim Laporan Rekrutmen ?', '11:00'),
                     ('visit', 'Sudah Kirim Laporan Visit Market ?', '13:00'), ('daily', 'Sudah Kirim Laporan Daily Report ?', '16:00')]
_now_wib = datetime.datetime.now(datetime.timezone.utc) + datetime.timedelta(hours=7)
_today_s = _now_wib.date().isoformat()
_rm_pending = [{'key': k_r, 'label': lbl_r, 'time': t_r, 'minutes': int(t_r[:2]) * 60 + int(t_r[3:5])} for k_r, lbl_r, t_r in DEFAULT_REMINDERS] if is_workday_date(_now_wib.date()) else []

# popup realtime di browser (hanya saat aplikasi terbuka): muncul sekali tepat di menit jam alarm; jika dibuka setelah jam alarm, tidak ada yang dimunculkan
_rm_js = """<script>
(function(){ try {
  const P = window.parent, D = P.document, items = __ITEMS__, today = '__TODAY__';
  if (P.__jember_bondowosoAlarmTimer) P.clearInterval(P.__jember_bondowosoAlarmTimer);
  P.__jember_bondowosoFired = P.__jember_bondowosoFired || {};
  function holder(){ let b = D.getElementById('jember_bondowoso-alarm-box'); if (!b) { b = D.createElement('div'); b.id = 'jember_bondowoso-alarm-box'; b.style.cssText = 'position:fixed;top:70px;right:18px;z-index:99999;display:flex;flex-direction:column;gap:10px;max-width:340px;'; D.body.appendChild(b); } return b; }
  function beep(){ try { const A = new (P.AudioContext || P.webkitAudioContext)(), o = A.createOscillator(), g = A.createGain(); o.connect(g); g.connect(A.destination); o.frequency.value = 880; g.gain.value = 0.08; o.start(); setTimeout(() => { o.stop(); A.close(); }, 450); } catch(e) {} }
  function show(it){
    const d = D.createElement('div'); d.className = 'jember_bondowoso-alarm-pop'; d.dataset.key = it.key;
    d.style.cssText = 'background:#1E1B4B;color:#fff;border:2px solid #FFEA00;border-left:6px solid #FF007F;border-radius:10px;padding:12px 14px;font-family:sans-serif;box-shadow:0 6px 24px rgba(0,0,0,.5);';
    d.innerHTML = '<div style=\"font-weight:800;font-size:15px;\">🔔 ' + it.label + '</div><div style=\"font-size:12px;margin-top:4px;color:#CBD5E1;\">Jadwal ' + it.time + ' WIB</div><div style=\"margin-top:8px;display:flex;gap:8px;\"><button class=\"x\" style=\"cursor:pointer;\">Tutup</button></div>';
    d.querySelector('.x').onclick = () => d.remove();
    holder().appendChild(d); beep();
    try { if (P.Notification && P.Notification.permission === 'granted') new P.Notification(it.label, { body: 'Jadwal ' + it.time + ' WIB' }); } catch(e) {}
  }
  function check(){
    const n = new Date(new Date().toLocaleString('en-US', { timeZone: 'Asia/Jakarta' })), mins = n.getHours() * 60 + n.getMinutes();
    items.forEach(it => { const k = today + it.key; if (mins === it.minutes && !P.__jember_bondowosoFired[k]) { P.__jember_bondowosoFired[k] = true; show(it); } });
  }
  check(); P.__jember_bondowosoAlarmTimer = P.setInterval(check, 10000);
} catch(e) {} })();
</script>"""
with st.sidebar: components.html(_rm_js.replace('__ITEMS__', json.dumps(_rm_pending)).replace('__TODAY__', _today_s), height=0)

with st.sidebar.expander('📆 Hari Libur Nasional', expanded=False):
    st.caption('Hari Minggu & tanggal libur di bawah ini tidak dihitung sebagai hari aktif (AVG, AVG 3/7, Forecast).')
    _lib_txt = st.text_area('Tanggal libur (YYYY-MM-DD, satu per baris)', value='\n'.join(sorted(d_l.isoformat() for d_l in LIBUR_SET)), height=200, key='libur_txt')
    if st.button('💾 Simpan Hari Libur', key='libur_save', use_container_width=True):
        save_json({'dates': sorted(d_l.isoformat() for d_l in _parse_libur(re.split(r'[\s,;]+', _lib_txt)))}, LIBUR_FILE); st.rerun()

def nota_metrics(df_n):
    """AVG NOTA, AVG 7, AVG 3 (rata-rata per HARI KERJA: Minggu & libur nasional tidak dihitung) + forecast s.d akhir bulan."""
    last7, last3 = set(WORKDAYS_UPTO[-7:]), set(WORKDAYS_UPTO[-3:])
    n_all = len(df_n) if df_n is not None else 0
    ok = df_n is not None and not df_n.empty and 'hari' in df_n.columns
    n7 = int(df_n['hari'].isin(last7).sum()) if ok else 0
    n3 = int(df_n['hari'].isin(last3).sum()) if ok else 0
    d7, d3 = max(1, len(last7)), max(1, len(last3))
    avg, avg7, avg3 = n_all / hari_aktif, n7 / d7, n3 / d3
    return {'n': n_all, 'n7': n7, 'n3': n3, 'd7': d7, 'd3': d3, 'avg': avg, 'avg7': avg7, 'avg3': avg3,
            'f_nota': int(round(n_all + avg * SISA_HARI_KERJA)), 'f7': int(round(n_all + avg7 * SISA_HARI_KERJA)), 'f3': int(round(n_all + avg3 * SISA_HARI_KERJA))}

def calc_sales_row(name_label, name_col, targ, df_f_c, df_f_l, df_n_c, df_n_t):
    ideal, stu, stulm, nota, today = round(targ * persen_desimal), len(df_f_c), len(df_f_l), len(df_n_c), len(df_n_t)
    forc, pct = (round(nota / persen_desimal) if persen_desimal > 0 else 0), ((stu / targ * 100.0) if targ > 0 else 0.0)
    _nm = nota_metrics(df_n_c)
    return {name_col: name_label, 'TARGET': targ, 'IDEAL': ideal, 'STU': stu, 'NOTA': nota, 'TODAY': today, col_name_pct: f'{pct:.2f}%', 'GAP TG': stu - targ, 'STU LM': stulm, '% LM': f'{(stu / stulm * 100.0):.2f}%' if stulm > 0 else '0.00%', 'GAP LM': stu - stulm, 'FORC': forc, 'F. NOTA': _nm['f_nota'], 'FORC 7': _nm['f7'], 'FORC 3': _nm['f3'], '_nota': nota, '_ideal': ideal, '_pct': pct}

def calc_proses_row(name_label, name_col, df_db_sub, stu_cnt, df_n_sub=None):
    db_c = len(df_db_sub)
    hot_c = len(df_db_sub[df_db_sub['Data_Type_Clean'] == 'DEAL PROSPECT']) if (not df_db_sub.empty and 'Data_Type_Clean' in df_db_sub.columns) else 0
    mp_c = int(df_db_sub['Salesman_Clean'].dropna().nunique()) if (not df_db_sub.empty and 'Salesman_Clean' in df_db_sub.columns) else 0
    avg_db, db_q, stu_p = ((db_c / mp_c / hari_aktif) if mp_c > 0 else 0.0), ((hot_c / db_c * 100.0) if db_c > 0 else 0.0), ((stu_cnt / db_c * 100.0) if db_c > 0 else 0.0)
    _nm = nota_metrics(df_n_sub)
    return {name_col: name_label, 'JML MP': mp_c, 'AVG DB': f'{avg_db:.2f}', 'DB': db_c, 'HOT': hot_c, 'DB Q': f'{db_q:.2f}%', 'STU': stu_cnt, 'STU%': f'{stu_p:.2f}%', 'AVG NOTA': f"{_nm['avg']:.2f}", 'AVG 7': f"{_nm['avg7']:.2f}", 'AVG 3': f"{_nm['avg3']:.2f}", '_stu_pct': stu_p}

def calc_leasing_row(label_nama, df_c, df_l, name_col):
    spk_c, spk_l = len(df_c), len(df_l)
    acc = len(df_c[df_c['Status_SPK_Clean'] == 'DISETUJUI']) if (not df_c.empty and 'Status_SPK_Clean' in df_c.columns) else 0
    rej = len(df_c[df_c['Status_SPK_Clean'] == 'DITOLAK']) if (not df_c.empty and 'Status_SPK_Clean' in df_c.columns) else 0
    opn = len(df_c[df_c['Status_SPK_Clean'] == 'TERBUKA']) if (not df_c.empty and 'Status_SPK_Clean' in df_c.columns) else 0
    grw, acc_p = ((spk_c / spk_l * 100.0) if spk_l > 0 else 0.0), ((acc / (acc + rej) * 100.0) if (acc + rej) > 0 else 0.0)
    r = {name_col: label_nama, 'SPK': spk_c, 'SPK LM': spk_l, 'GROWTH': f'{grw:.2f}%', 'ACC%': f'{acc_p:.2f}%', 'TERBUKA': opn, '_acc_val': acc_p, '_spk_val': spk_c, '_spk_lm_val': spk_l}
    for lsg in ['BAF', 'ADIRA', 'SOF', 'IMFI']:
        sub_l = df_c[df_c['Leasing_SPK_Clean'] == lsg] if (not df_c.empty and 'Leasing_SPK_Clean' in df_c.columns) else pd.DataFrame()
        la, lr = (len(sub_l[sub_l['Status_SPK_Clean'] == 'DISETUJUI']) if not sub_l.empty else 0), (len(sub_l[sub_l['Status_SPK_Clean'] == 'DITOLAK']) if not sub_l.empty else 0)
        r[lsg], r[f'%{lsg}'] = len(sub_l), f'{(la / (la + lr) * 100.0):.2f}%' if (la + lr) > 0 else '0.00%'
    return r

def build_type_table(df_fc, df_fl, df_dbc, df_spkc, filter_col_vals, col_headers_name, tot_label='TOTAL', sort_by_growth=True):
    lsg_list, rows, t_acc, t_rej = ['BAF', 'ADIRA', 'SOF'], [], 0, 0
    t_lf, t_la, t_lr = {l: 0 for l in lsg_list}, {l: 0 for l in lsg_list}, {l: 0 for l in lsg_list}
    for item_val, f_col_f, f_col_db, f_col_spk in filter_col_vals:
        sf_c = df_fc[df_fc[f_col_f] == item_val] if (not df_fc.empty and f_col_f in df_fc.columns) else pd.DataFrame()
        sf_l = df_fl[df_fl[f_col_f] == item_val] if (not df_fl.empty and f_col_f in df_fl.columns) else pd.DataFrame()
        sdb = df_dbc[df_dbc[f_col_db] == item_val] if (not df_dbc.empty and f_col_db in df_dbc.columns) else pd.DataFrame()
        sspk = df_spkc[df_spkc[f_col_spk] == item_val] if (not df_spkc.empty and f_col_spk in df_spkc.columns) else pd.DataFrame()
        stu, stulm, db_v, spk_v = len(sf_c), len(sf_l), len(sdb), len(sspk)
        hot_v = len(sdb[sdb['Data_Type_Clean'] == 'DEAL PROSPECT']) if (not sdb.empty and 'Data_Type_Clean' in sdb.columns) else 0
        acc_v = len(sspk[sspk['Status_SPK_Clean'] == 'DISETUJUI']) if (not sspk.empty and 'Status_SPK_Clean' in sspk.columns) else 0
        rej_v = len(sspk[sspk['Status_SPK_Clean'] == 'DITOLAK']) if (not sspk.empty and 'Status_SPK_Clean' in sspk.columns) else 0
        t_acc += acc_v; t_rej += rej_v
        grw, dbq, accp, stup = ((stu / stulm * 100.0) if stulm > 0 else 0.0), ((hot_v / db_v * 100.0) if db_v > 0 else 0.0), ((acc_v / (acc_v + rej_v) * 100.0) if (acc_v + rej_v) > 0 else 0.0), ((stu / db_v * 100.0) if db_v > 0 else 0.0)
        r = {col_headers_name: item_val, 'STU': stu, 'STU LM': stulm, 'GROWTH': f'{grw:.2f}%', 'DB': db_v, 'HOT': hot_v, 'DB Q': f'{dbq:.2f}%', 'SPK': spk_v, 'ACC%': f'{accp:.2f}%', 'STU%': f'{stup:.2f}%', '_stu_val': stu, '_growth_val': grw}
        for l in lsg_list:
            fcnt = len(sf_c[sf_c['Leasing_Faktur_Clean'] == l]) if (not sf_c.empty and 'Leasing_Faktur_Clean' in sf_c.columns) else 0
            t_lf[l] += fcnt
            sl = sspk[sspk['Leasing_SPK_Clean'] == l] if (not sspk.empty and 'Leasing_SPK_Clean' in sspk.columns) else pd.DataFrame()
            la, lr = (len(sl[sl['Status_SPK_Clean'] == 'DISETUJUI']) if not sl.empty else 0), (len(sl[sl['Status_SPK_Clean'] == 'DITOLAK']) if not sl.empty else 0)
            t_la[l] += la; t_lr[l] += lr
            r[l], r[f'%{l}'] = fcnt, f'{(la / (la + lr) * 100.0):.2f}%' if (la + lr) > 0 else '0.00%'
            r[f'STU {l}'] = fcnt
        r['STU CASH'] = len(sf_c[sf_c['Leasing_Faktur_Clean'] == 'CASH']) if (not sf_c.empty and 'Leasing_Faktur_Clean' in sf_c.columns) else 0
        rows.append(r)
    if not rows: return pd.DataFrame(), pd.DataFrame(), {}
    df_it = pd.DataFrame(rows).sort_values(by=(['_growth_val', '_stu_val'] if sort_by_growth else ['_stu_val', 'DB']), ascending=([True, True] if sort_by_growth else [False, False])).reset_index(drop=True)
    ts, tsl, tdb, thot, tspk = int(df_it['STU'].sum()), int(df_it['STU LM'].sum()), int(df_it['DB'].sum()), int(df_it['HOT'].sum()), int(df_it['SPK'].sum())
    tgrw, tdbq, taccp, tstup = ((ts / tsl * 100.0) if tsl > 0 else 0.0), ((thot / tdb * 100.0) if tdb > 0 else 0.0), ((t_acc / (t_acc + t_rej) * 100.0) if (t_acc + t_rej) > 0 else 0.0), ((ts / tdb * 100.0) if tdb > 0 else 0.0)
    tot_r = {col_headers_name: tot_label, 'STU': ts, 'STU LM': tsl, 'GROWTH': f'{tgrw:.2f}%', 'DB': tdb, 'HOT': thot, 'DB Q': f'{tdbq:.2f}%', 'SPK': tspk, 'ACC%': f'{taccp:.2f}%', 'STU%': f'{tstup:.2f}%', '_stu_val': ts, '_growth_val': tgrw}
    for l in lsg_list: tot_r[l], tot_r[f'%{l}'] = t_lf[l], f'{(t_la[l] / (t_la[l] + t_lr[l]) * 100.0):.2f}%' if (t_la[l] + t_lr[l]) > 0 else '0.00%'
    for l in lsg_list: tot_r[f'STU {l}'] = t_lf[l]
    tot_r['STU CASH'] = int(df_it['STU CASH'].sum())
    return df_it, pd.concat([df_it, pd.DataFrame([tot_r])], ignore_index=True), tot_r

df_fak_curr_g = df_faktur[(df_faktur['bulan'] == current_month) & (df_faktur['tahun'] == current_year) & (df_faktur['hari'] <= tanggal_akhir)] if (df_faktur is not None and 'bulan' in df_faktur.columns) else pd.DataFrame()
df_fak_lm_g = df_faktur[(df_faktur['bulan'] == last_month) & (df_faktur['tahun'] == last_month_year) & (df_faktur['hari'] <= tanggal_akhir)] if (df_faktur is not None and 'bulan' in df_faktur.columns) else pd.DataFrame()
df_nota_curr_g = df_nota[(df_nota['bulan'] == current_month) & (df_nota['tahun'] == current_year) & (df_nota['hari'] <= tanggal_akhir)] if (df_nota is not None and 'bulan' in df_nota.columns) else pd.DataFrame()
df_nota_tod_g = df_nota[(df_nota['bulan'] == current_month) & (df_nota['tahun'] == current_year) & (df_nota['hari'] == tanggal_akhir)] if (df_nota is not None and 'bulan' in df_nota.columns) else pd.DataFrame()
if df_database is not None and not df_database.empty and 'hari' in df_database.columns and df_database['hari'].notna().any():
    if 'bulan' in df_database.columns and 'tahun' in df_database.columns and ((df_database['bulan'] == current_month) & (df_database['tahun'] == current_year)).any():
        df_db_curr_g = df_database[(df_database['bulan'] == current_month) & (df_database['tahun'] == current_year) & (df_database['hari'] >= 1) & (df_database['hari'] <= tanggal_akhir)]
    else:
        df_db_curr_g = df_database[(df_database['hari'] >= 1) & (df_database['hari'] <= tanggal_akhir)]
else:
    df_db_curr_g = df_database.copy() if df_database is not None else pd.DataFrame()
if df_rating is not None and not df_rating.empty and 'hari' in df_rating.columns and df_rating['hari'].notna().any():
    if 'bulan' in df_rating.columns and 'tahun' in df_rating.columns and ((df_rating['bulan'] == current_month) & (df_rating['tahun'] == current_year)).any():
        df_rating_curr_g = df_rating[(df_rating['bulan'] == current_month) & (df_rating['tahun'] == current_year) & (df_rating['hari'] >= 1) & (df_rating['hari'] <= tanggal_akhir)]
    else:
        df_rating_curr_g = df_rating[(df_rating['hari'] >= 1) & (df_rating['hari'] <= tanggal_akhir)]
else:
    df_rating_curr_g = df_rating.copy() if df_rating is not None else pd.DataFrame()
df_spk_curr_g = df_spk[(df_spk['bulan'] == current_month) & (df_spk['tahun'] == current_year) & (df_spk['hari'] <= tanggal_akhir)] if (df_spk is not None and 'bulan' in df_spk.columns) else pd.DataFrame()
df_spk_lm_g = df_spk[(df_spk['bulan'] == last_month) & (df_spk['tahun'] == last_month_year) & (df_spk['hari'] <= tanggal_akhir)] if (df_spk is not None and 'bulan' in df_spk.columns) else pd.DataFrame()

import math
saved_kpi = load_json(KPI_SAVED_FILE)
saved_td = load_json(TARGET_DEALER_FILE)
saved_tk = load_json(TARGET_KOOR_FILE)
KPI_MANUAL_FILE = _ABS('kpi_dashboard_manual.json')
DEFAULT_T_DB = {'koor': 200, 'counter': 100}   # default T. DB (masih bisa diedit & disimpan di MASTER KPI)
DEFAULT_T_WALKIN = 25                           # default T. Walk In counter (masih bisa diedit & disimpan)

def detect_prospect_type_cols(df):
    """Kolom Prospect Type di file database: cari via nama header (huruf besar/kecil bebas), fallback kolom G (indeks 6)."""
    if df is None or df.empty: return []
    cols = [c for c in df.columns if alnum_str(c).startswith(('PROSPECTTYPE', 'TIPEPROSPECT', 'JENISPROSPECT'))]
    if cols: return cols
    return [df.columns[6]] if df.shape[1] > 6 else []

def count_walkin(df_db_sub):
    """Jumlah baris database dengan Prospect Type (kolom G) = Walk In."""
    cols = detect_prospect_type_cols(df_db_sub)
    if not cols: return 0
    is_walk = pd.Series(False, index=df_db_sub.index)
    for c in cols:
        ser = df_db_sub[c]
        ser = ser.iloc[:, 0] if isinstance(ser, pd.DataFrame) else ser
        is_walk = is_walk | (ser.apply(alnum_str) == 'WALKIN')
    return int(is_walk.sum())

def kpi_entity_dealer(ent):
    """Dealer (dari 10 dealer fokus) tempat koordinator/counter bertugas; None jika di luar 10 dealer fokus."""
    e = clean_str(ent)
    d = mapping_koor_dealer.get(e)
    if d in daftar_dealer: return d
    n = normalize_dealer_name(e)
    return n if n in daftar_dealer else None

_ACTIVE_DB_SM = None
def get_active_db_salesmen():
    """Nama salesman yang masih ada di file Database (acuan MP aktif). Salesman yang sudah tidak ada di Database (resign) tidak dihitung."""
    global _ACTIVE_DB_SM
    if _ACTIVE_DB_SM is None:
        _ACTIVE_DB_SM = set()
        if df_database is not None and not df_database.empty and 'Salesman_Clean' in df_database.columns:
            _ACTIVE_DB_SM = set(clean_str(x) for x in df_database['Salesman_Clean'].dropna().unique() if clean_str(x))
    return _ACTIVE_DB_SM

def is_counter_row_name(nm, *dfs):
    """Counter: lolos is_counter_personnel ATAU di data baris salesman tsb terpetakan ke Koordinator berawalan COUNTER."""
    if is_counter_personnel(nm): return True
    for d_x in dfs:
        for kc in ('Koor', 'Koor_SPK'):
            if d_x is not None and not d_x.empty and kc in d_x.columns and 'Salesman_Clean' in d_x.columns:
                kk = d_x.loc[d_x['Salesman_Clean'] == nm, kc].dropna()
                if not kk.empty and kk.astype(str).str.upper().str.startswith('COUNTER').any(): return True
    return False

def compute_auto_kpi_results(df_f_sub, df_db_sub, df_spk_sub):
    r_stu = len(df_f_sub) if df_f_sub is not None else 0
    if df_f_sub is not None and not df_f_sub.empty and 'Type_Clean' in df_f_sub.columns:
        r_maxi = int(df_f_sub['Type_Clean'].isin(['AT PREMIUM', 'MAXI', 'AT MAXI']).sum())
        r_classy = int((df_f_sub['Type_Clean'] == 'AT CLASSY').sum())
    else:
        r_maxi, r_classy = 0, 0

    mp_set = set()
    for ds in [df_db_sub, df_f_sub, df_spk_sub]:
        if ds is not None and not ds.empty and 'Salesman_Clean' in ds.columns:
            mp_set.update(clean_str(x) for x in ds['Salesman_Clean'].dropna().unique() if clean_str(x) and clean_str(x) not in ['NAN', 'NONE', 'TOTAL'])
    _act = get_active_db_salesmen()
    if _act: mp_set = {m for m in mp_set if m in _act}   # acuan file Database: yang tidak ada di Database (resign) diabaikan
    mp_set = {m for m in mp_set if not is_counter_row_name(m, df_db_sub, df_f_sub, df_spk_sub)}   # counter tidak dihitung sbg QTY Salesman
    r_mp = len(mp_set)

    r_inc = len(df_spk_sub) if df_spk_sub is not None else 0
    acc_c = len(df_spk_sub[df_spk_sub['Status_SPK_Clean'] == 'DISETUJUI']) if (df_spk_sub is not None and not df_spk_sub.empty and 'Status_SPK_Clean' in df_spk_sub.columns) else 0
    rej_c = len(df_spk_sub[df_spk_sub['Status_SPK_Clean'] == 'DITOLAK']) if (df_spk_sub is not None and not df_spk_sub.empty and 'Status_SPK_Clean' in df_spk_sub.columns) else 0
    r_acc_pct = round((acc_c / (acc_c + rej_c) * 100.0), 1) if (acc_c + rej_c) > 0 else 0.0
    return r_stu, r_maxi, r_classy, r_mp, r_inc, r_acc_pct, acc_c, rej_c

_DLR_ROW_CACHE = {}
def get_dealer_master_row(dealer):
    """T./R. MYDR & T./R. G REV dari baris dealer induk di tabel MASTER KPI Dealer (dipakai Koordinator & Counter)."""
    zero = {'T. MYDR': 0, 'R. MYDR': 0, 'T. G REV': 0, 'R. G REV': 0}
    if not dealer or dealer not in daftar_dealer: return zero
    ck = (dealer, json.dumps(saved_kpi.get('dealer', {}).get(dealer, {}), sort_keys=True), tanggal_akhir)
    if ck not in _DLR_ROW_CACHE:
        df_d, _ = build_kpi_section_df([dealer], 'NAMA DEALER', 'dealer', 'dealer')
        _DLR_ROW_CACHE[ck] = zero if (df_d is None or df_d.empty) else {k: safe_int(df_d.iloc[0][k]) for k in zero}
    return _DLR_ROW_CACHE[ck]

def build_kpi_section_df(entity_list, name_col, section_key, filter_mode):
    sec_saved = saved_kpi.get(section_key, {})
    rows = []
    tot_acc_c, tot_rej_c = 0, 0
    for ent in entity_list:
        ent_saved = sec_saved.get(ent, {})
        if filter_mode == 'dealer':
            f_sub = df_fak_curr_g[df_fak_curr_g['Nama Dealer'] == ent] if not df_fak_curr_g.empty else pd.DataFrame()
            db_sub = df_db_curr_g[df_db_curr_g['Nama_Dealer_Clean'] == ent] if not df_db_curr_g.empty else pd.DataFrame()
            spk_sub = df_spk_curr_g[df_spk_curr_g['Nama_Dealer_SPK'] == ent] if not df_spk_curr_g.empty else pd.DataFrame()
            rt_sub = df_rating_curr_g[df_rating_curr_g['Nama_Dealer_Clean'] == ent] if not df_rating_curr_g.empty else pd.DataFrame()
            def_vals = DEFAULT_KPI_DEALER.get(ent, {})
            def_t_stu = safe_int(saved_td.get(ent, def_vals.get('T. STU', 0)))
            def_t_maxi = def_vals.get('T. MAXI', 0)
            def_t_classy = def_vals.get('T. CLASSY', 0)
            def_t_mp = def_vals.get('T. MP', 0)
            def_t_inc = def_vals.get('T. INC', 0)
            def_t_acc = def_vals.get('T. ACC%', 50.0)
        else:
            f_sub = df_fak_curr_g[df_fak_curr_g['Koor'] == ent] if not df_fak_curr_g.empty else pd.DataFrame()
            db_sub = df_db_curr_g[df_db_curr_g['Koor'] == ent] if not df_db_curr_g.empty else pd.DataFrame()
            spk_sub = df_spk_curr_g[df_spk_curr_g['Koor_SPK'] == ent] if not df_spk_curr_g.empty else pd.DataFrame()
            rt_sub = df_rating_curr_g[df_rating_curr_g['Koor'] == ent] if not df_rating_curr_g.empty else pd.DataFrame()
            def_t_stu = get_koor_or_counter_target(ent, saved_tk) or (20 if filter_mode == 'counter' else 25)
            def_t_maxi = 8 if filter_mode == 'counter' else 10
            def_t_classy = 3 if filter_mode == 'counter' else 5
            def_t_mp = 1 if filter_mode == 'counter' else 6
            def_t_inc = 15 if filter_mode == 'counter' else 18
            def_t_acc = 50.0

        r_stu, r_maxi, r_classy, r_mp, r_inc, r_acc_pct, acc_c, rej_c = compute_auto_kpi_results(f_sub, db_sub, spk_sub)
        tot_acc_c += acc_c; tot_rej_c += rej_c

        t_stu = safe_int(ent_saved.get('T. STU', def_t_stu))
        t_maxi = safe_int(ent_saved.get('T. MAXI', def_t_maxi))
        t_classy = safe_int(ent_saved.get('T. CLASSY', def_t_classy))
        t_mp = safe_int(ent_saved.get('T. MP', def_t_mp))
        t_inc = safe_int(ent_saved.get('T. INC', def_t_inc))
        t_acc = safe_float(ent_saved.get('T. ACC%', def_t_acc))
        t_mydr = int(round(r_stu * 0.30))
        r_mydr = len(rt_sub) if (df_rating is not None and not df_rating.empty) else safe_int(ent_saved.get('R. MYDR', 0))
        if filter_mode == 'dealer':
            t_grev = safe_int(ent_saved.get('T. G REV', 0))
            r_grev = safe_int(ent_saved.get('R. G REV', 0))
            extra_cols = {}
        else:
            # MYDR & Google Review koordinator / counter: target & result mengikuti dealer induk (baris dealer di MASTER KPI)
            _dm = get_dealer_master_row(kpi_entity_dealer(ent))
            t_mydr, r_mydr = _dm['T. MYDR'], _dm['R. MYDR']
            t_grev, r_grev = _dm['T. G REV'], _dm['R. G REV']
            # KPI Database (koordinator & counter) dan Walk In (counter) — result dari file database
            extra_cols = {'T. DB': safe_int(ent_saved.get('T. DB', DEFAULT_T_DB[filter_mode])), 'R. DB': len(db_sub), '_dealer': kpi_entity_dealer(ent)}
            if filter_mode == 'counter':
                extra_cols['T. Walk In'] = safe_int(ent_saved.get('T. Walk In', DEFAULT_T_WALKIN))
                extra_cols['R. Walk In'] = count_walkin(db_sub)

        pct_stu = (r_stu / t_stu * 100.0) if t_stu > 0 else 0.0
        rows.append({
            name_col: ent,
            'T. STU': t_stu, 'R. STU': r_stu,
            'T. MAXI': t_maxi, 'R. MAXI': r_maxi,
            'T. CLASSY': t_classy, 'R. CLASSY': r_classy,
            'T. MP': t_mp, 'R. MP': r_mp,
            'T. INC': t_inc, 'R. INC': r_inc,
            'T. ACC%': round(t_acc, 1), 'R. ACC%': round(r_acc_pct, 1),
            'T. MYDR': t_mydr, 'R. MYDR': r_mydr,
            'T. G REV': t_grev, 'R. G REV': r_grev,
            **extra_cols,
            '_pct_stu': pct_stu
        })
    if not rows:
        return pd.DataFrame(), pd.DataFrame()
    df_items = pd.DataFrame(rows).reset_index(drop=True)
    tot_t_acc = round(float(df_items['T. ACC%'].mean()), 1) if not df_items.empty else 50.0
    tot_r_acc = round((tot_acc_c / (tot_acc_c + tot_rej_c) * 100.0), 1) if (tot_acc_c + tot_rej_c) > 0 else 0.0
    tot_t_stu = int(df_items['T. STU'].sum())
    tot_r_stu = int(df_items['R. STU'].sum())
    tot_row = {
        name_col: 'TOTAL',
        'T. STU': tot_t_stu, 'R. STU': tot_r_stu,
        'T. MAXI': int(df_items['T. MAXI'].sum()), 'R. MAXI': int(df_items['R. MAXI'].sum()),
        'T. CLASSY': int(df_items['T. CLASSY'].sum()), 'R. CLASSY': int(df_items['R. CLASSY'].sum()),
        'T. MP': int(df_items['T. MP'].sum()), 'R. MP': int(df_items['R. MP'].sum()),
        'T. INC': int(df_items['T. INC'].sum()), 'R. INC': int(df_items['R. INC'].sum()),
        'T. ACC%': tot_t_acc, 'R. ACC%': tot_r_acc,
        'T. MYDR': int(df_items['T. MYDR'].sum()), 'R. MYDR': int(df_items['R. MYDR'].sum()),
        'T. G REV': int(df_items['T. G REV'].sum()), 'R. G REV': int(df_items['R. G REV'].sum()),
        '_pct_stu': (tot_r_stu / tot_t_stu * 100.0) if tot_t_stu > 0 else 0.0
    }
    if filter_mode != 'dealer' and '_dealer' in df_items.columns:
        _ud = df_items.drop_duplicates('_dealer')
        for _c in ['T. MYDR', 'R. MYDR', 'T. G REV', 'R. G REV']: tot_row[_c] = int(_ud[_c].sum())
    for _tc, _rc in [('T. DB', 'R. DB'), ('T. Walk In', 'R. Walk In')]:
        if _tc in df_items.columns:
            tot_row[_tc], tot_row[_rc] = int(df_items[_tc].sum()), int(df_items[_rc].sum())
    return df_items, pd.concat([df_items, pd.DataFrame([tot_row])], ignore_index=True)

# ===================== DASHBOARD KPI: skema poin (sesuai Excel KPI Report) =====================
KPI_CODES = ['STU','MAX','CLS','OPR','SAL','APV','INC','MYM','MYD','GRT','GRV','OPS']
KPI_NAMES = {'STU':'STU Total','MAX':'STU MAXI','CLS':'STU AT Classy','OPR':'Operational Dealer','SAL':'QTY Salesman','APV':'Approval Ratio Total','INC':'Incoming Total','MYM':'My Yamaha Member Point','MYD':'My Yamaha Dealer Rating','GRT':'Google Rating','GRV':'Google Review','OPS':'Operational Supervisor'}
KPI_MIN_ACV = {'STU':0.95,'MAX':0.95,'CLS':0.95,'OPR':0.90,'SAL':0.75,'APV':0.75,'INC':0.75,'MYM':0.75,'MYD':0.75,'GRT':1.0,'GRV':0.75,'OPS':0.90}
KPI_WEIGHT_MDS = {'STU':20,'MAX':10,'CLS':10,'OPR':5,'SAL':10,'APV':5,'INC':9,'MYM':3,'MYD':8,'GRT':5,'GRV':5,'OPS':10}
KPI_WEIGHT_DEALER = {'STU':20,'MAX':10,'CLS':10,'OPR':15,'SAL':10,'APV':5,'INC':9,'MYM':3,'MYD':8,'GRT':5,'GRV':5,'OPS':0}
KPI_SALES_CODES, KPI_PROCESS_CODES = ['STU','MAX','CLS'], ['OPR','SAL','APV','INC','MYM','MYD','GRT','GRV','OPS']
KPI_DEALER_CODES = {}   # isi {'SIP XXX': 'KODE'} jika ada kode dealer
KPI_MANUAL_COLS = ['T. OPR','R. OPR','T. MYM','R. MYM','T. GRT','R. GRT']
DEFAULT_KPI_OPS = {'T. OPS': 44, 'R. OPS': 48}
DEFAULT_KPI_MANUAL = {d_: {'T. OPR':156,'R. OPR':0,'T. MYM':70.0,'R. MYM':0.0,'T. GRT':5.0,'R. GRT':0.0} for d_ in daftar_dealer}   # nilai awal; edit di menu KPI

def kpi_point_generic(achv, min_acv, weight, split=0.5):
    try:
        if achv >= 1: return int(weight)
        if min_acv == 0: return int(math.floor(achv * weight + 1e-9))
        if achv >= min_acv:
            base = math.ceil(split * weight - 1e-9)
            extra = math.floor((achv - min_acv) / (1 - min_acv) * math.floor((1 - split) * weight + 1e-9) + 1e-9)
            return int(base + extra)
        return 0
    except Exception:
        return 0

def kpi_point_grt(target, result, weight, scheme):
    if scheme == 'DEALER':
        if target >= 4.8 and result >= 4.8: return weight
        if target >= 4.8 and result < target and result > 0: return int(math.ceil(weight * 0.5 - 1e-9))
        if target < 4.8 and target > 0 and result >= target: return weight
        return 0
    if target >= 4.8 and result < target: return int(math.ceil(weight * 0.5 - 1e-9))
    if result >= 4.8 and target >= 4.8: return weight
    if result < 4.8 and result >= target: return weight
    return 0

def kpi_score_entity(tr, scheme='MDS'):
    weights = KPI_WEIGHT_MDS if scheme == 'MDS' else KPI_WEIGHT_DEALER
    rows = {}
    for c in KPI_CODES:
        t, r = tr.get(c, (0, 0))
        t, r = float(t or 0), float(r or 0)
        achv = (r / t) if t > 0 else 0.0
        w = weights[c]
        if c == 'GRT': pt = kpi_point_grt(t, r, w, scheme)
        else: pt = kpi_point_generic(achv, KPI_MIN_ACV[c], w, 0.75 if (scheme == 'DEALER' and c in KPI_SALES_CODES) else 0.5)
        rows[c] = dict(code=c, name=KPI_NAMES[c], weight=w, min_acv=KPI_MIN_ACV[c], target=t, result=r, achv=achv, point=pt)
    sales = sum(rows[c]['point'] for c in KPI_SALES_CODES)
    process = sum(rows[c]['point'] for c in KPI_PROCESS_CODES)
    return rows, sales, process, sales + process

def kpi_grade(p): return 'D' if p < 60 else ('C' if p < 70 else ('B' if p < 95 else 'A'))
def kpi_note(process): return '' if process >= 45 else ('CUT 30%' if process >= 30 else 'CUT 50%')

def get_dashboard_kpi_inputs():
    """Target & Result per KPI. STU/MAX/CLS/SAL/APV/INC/MYD/GRV dari MASTER KPI; OPR/MYM/GRT/OPS dari input manual."""
    df_it, df_fl = build_kpi_section_df(daftar_dealer, 'NAMA DEALER', 'dealer', 'dealer')
    man = load_json(KPI_MANUAL_FILE)
    man_d, man_ops = man.get('dealers', {}), {**DEFAULT_KPI_OPS, **man.get('ops', {})}
    ops_pair = (safe_float(man_ops.get('T. OPS')), safe_float(man_ops.get('R. OPS')))
    inputs = {}
    def _g(r_x, k_x): return safe_float(r_x[k_x]) if r_x is not None else 0.0
    for d in daftar_dealer:
        mv = {**DEFAULT_KPI_MANUAL.get(d, {}), **man_d.get(d, {})}
        r = df_it[df_it['NAMA DEALER'] == d].iloc[0] if (df_it is not None and not df_it.empty and d in df_it['NAMA DEALER'].values) else None
        inputs[d] = {'STU': (_g(r, 'T. STU'), _g(r, 'R. STU')), 'MAX': (_g(r, 'T. MAXI'), _g(r, 'R. MAXI')), 'CLS': (_g(r, 'T. CLASSY'), _g(r, 'R. CLASSY')),
                     'OPR': (safe_float(mv.get('T. OPR')), safe_float(mv.get('R. OPR'))), 'SAL': (_g(r, 'T. MP'), _g(r, 'R. MP')),
                     'APV': (_g(r, 'T. ACC%') / 100.0, _g(r, 'R. ACC%') / 100.0), 'INC': (_g(r, 'T. INC'), _g(r, 'R. INC')),
                     'MYM': (safe_float(mv.get('T. MYM')) / 100.0, safe_float(mv.get('R. MYM')) / 100.0), 'MYD': (_g(r, 'T. MYDR'), _g(r, 'R. MYDR')),
                     'GRT': (safe_float(mv.get('T. GRT')), safe_float(mv.get('R. GRT'))), 'GRV': (_g(r, 'T. G REV'), _g(r, 'R. G REV')), 'OPS': ops_pair}
    rt = df_fl[df_fl['NAMA DEALER'] == 'TOTAL'].iloc[0] if (df_fl is not None and not df_fl.empty and 'TOTAL' in df_fl['NAMA DEALER'].values) else None
    ds = list(daftar_dealer)
    stu_w = [max(inputs[d]['STU'][1], 0.0) for d in ds]
    def _mean(idx, pos): return sum(inputs[d][idx][pos] for d in ds) / len(ds)
    mym_r = (sum(inputs[d]['MYM'][1] * w for d, w in zip(ds, stu_w)) / sum(stu_w)) if sum(stu_w) > 0 else _mean('MYM', 1)
    inputs['TOTAL JEMBER BONDOWOSO'] = {'STU': (_g(rt, 'T. STU'), _g(rt, 'R. STU')), 'MAX': (_g(rt, 'T. MAXI'), _g(rt, 'R. MAXI')), 'CLS': (_g(rt, 'T. CLASSY'), _g(rt, 'R. CLASSY')),
                     'OPR': (sum(inputs[d]['OPR'][0] for d in ds), sum(inputs[d]['OPR'][1] for d in ds)), 'SAL': (_g(rt, 'T. MP'), _g(rt, 'R. MP')),
                     'APV': (_g(rt, 'T. ACC%') / 100.0, _g(rt, 'R. ACC%') / 100.0), 'INC': (_g(rt, 'T. INC'), _g(rt, 'R. INC')),
                     'MYM': (_mean('MYM', 0), mym_r), 'MYD': (_g(rt, 'T. MYDR'), _g(rt, 'R. MYDR')),
                     'GRT': (_mean('GRT', 0), _mean('GRT', 1)), 'GRV': (_g(rt, 'T. G REV'), _g(rt, 'R. G REV')), 'OPS': ops_pair}
    return inputs

# ===================== DASHBOARD KPI: tampilan HTML =====================
DASH_NAVY, DASH_TEAL = '#0B2A5B', '#1B5E75'
DASH_ABBR = {1:'JAN',2:'FEB',3:'MAR',4:'APR',5:'MEI',6:'JUN',7:'JUL',8:'AGU',9:'SEP',10:'OKT',11:'NOV',12:'DES'}
DASH_GRADE_COL = {'A':'#00FF66','B':'#00E5FF','C':'#FFEA00','D':'#FF4D8D'}

_BATCH_PREFIX = ('Review_1_Penjualan_', 'Review_2_Proses_', 'Review_3_Leasing_', 'Review_4_Type_', 'Review_5_Salesman_', 'Review_Aktifitas_PT_', 'Review_Aktifitas_Sub_')

def bsg_kpi_capture_html(sel_dash):
    """HTML Dashboard KPI (tanpa tabel ranking) — sama dengan isi tombol Download JPG di menu DASHBOARD KPI."""
    scheme = 'MDS' if sel_dash == 'TOTAL JEMBER BONDOWOSO' else 'DEALER'
    inputs = get_dashboard_kpi_inputs(); rows_s, s_s, p_s, t_s = kpi_score_entity(inputs[sel_dash], scheme); d_tbl = []
    for d_k in daftar_dealer:
        r_k, s_k, p_k, t_k = kpi_score_entity(inputs[d_k], scheme)
        d_tbl.append({'name': d_k, 'code': KPI_DEALER_CODES.get(d_k, '-'), 'point': t_k, 'grade': kpi_grade(t_k), 'sales': s_k, 'process': p_k, 'achv': {c: r_k[c]['achv'] for c in KPI_CODES}, 'note': kpi_note(p_k)})
    d_tbl.sort(key=lambda x_d: (-x_d['point'], -x_d['process']))
    date_lbl = f'{tanggal_akhir} {DASH_ABBR.get(current_month, "")} {current_year}'; date_short = f'{tanggal_akhir:02d}/{current_month:02d}/{current_year}'
    note = ('Bobot skema MDS (sesuai gambar): OPS level supervisor, nilainya sama untuk semua dealer.' if scheme == 'MDS' else 'Bobot skema Dealer (Excel R-MDS): OPR 15 & OPS 0; STU/MAX/CLS memakai 75% poin dasar.') + ' Kolom NOTE: PROCESS ≥45 tanpa potongan, 30–44 = CUT 30%, <30 = CUT 50%. CONCLUTION tampil mulai tgl 27.'
    h_ = build_dashboard_kpi_html(sel_dash, date_lbl, date_short, rows_s, s_s, p_s, t_s, d_tbl, sel_dash, tanggal_akhir >= 27, note, with_rank=False)
    return f'<div style="width:1320px;background-color:#0E1117;padding:18px 20px;box-sizing:border-box;">{h_}</div>'

_BATCH_HTML = """<html><head><script src="https://cdnjs.cloudflare.com/ajax/libs/html2canvas/1.4.1/html2canvas.min.js"></script><script src="https://cdnjs.cloudflare.com/ajax/libs/jszip/3.10.1/jszip.min.js"></script>
<style>body{margin:0;padding:2px 0;background:transparent;font-family:'Source Sans Pro',sans-serif;display:flex;align-items:center;gap:10px;}
.b{background:#4318FF;color:#fff;border:1px solid #4318FF;border-radius:6px;padding:8px 16px;font-size:13px;font-weight:700;cursor:pointer;}.b:hover{background:#3311CC;}.b:disabled{opacity:.6;cursor:wait;}
#st{font-size:12.5px;color:#FFEA00;font-weight:600;}</style></head><body>
<button class="b" id="go" onclick="runAll()">📥 Download SEMUA gambar untuk PPT (__N__ gambar · 1 file ZIP)</button><span id="st"></span><div id="hold" style="position:fixed;left:-20000px;top:0;"></div>
<script>
const ITEMS=__ITEMS__, ZIPNAME=__ZIPNAME__;
const sleep=ms=>new Promise(r=>setTimeout(r,ms));
async function waitFor(fn,ms){const t0=Date.now();while(Date.now()-t0<ms){try{if(fn())return true;}catch(e){}await sleep(150);}return false;}
async function runAll(){
 const st=document.getElementById('st'),btn=document.getElementById('go');btn.disabled=true;
 try{
  if(typeof html2canvas==='undefined'||typeof JSZip==='undefined')throw new Error('library belum termuat (cek koneksi internet)');
  const zip=new JSZip();const fail=[];let ok=0;
  for(let i=0;i<ITEMS.length;i++){const it=ITEMS[i];st.textContent='Memproses '+(i+1)+'/'+ITEMS.length+': '+it.fn+' ...';
   try{let url;
    if(it.kind==='panel'){const f=document.createElement('iframe');f.style.cssText='position:fixed;left:-20000px;top:0;width:1500px;height:1400px;border:0;';f.srcdoc=it.html;document.body.appendChild(f);
     const good=await waitFor(()=>f.contentWindow&&f.contentWindow.pnCapData&&typeof f.contentWindow.html2canvas!=='undefined',25000);if(!good)throw new Error('panel tidak termuat');
     await sleep(600);url=await f.contentWindow.pnCapData();f.remove();}
    else{const h=document.getElementById('hold');h.innerHTML=it.html;await sleep(300);const c=await html2canvas(h.firstElementChild,{backgroundColor:'#0E1117',scale:2,useCORS:true,logging:false});h.innerHTML='';url=c.toDataURL('image/jpeg',0.95);}
    zip.file(it.fn+'.jpg',url.split(',')[1],{base64:true});ok++;
   }catch(e){fail.push(it.fn+' ('+e.message+')');}}
  if(!ok)throw new Error('tidak ada gambar yang berhasil dibuat');
  const blob=await zip.generateAsync({type:'blob'});const a=document.createElement('a');a.href=URL.createObjectURL(blob);a.download=ZIPNAME;document.body.appendChild(a);a.click();a.remove();
  st.style.color=fail.length?'#FFEA00':'#00FF66';st.textContent='✅ '+ok+' gambar tersimpan di '+ZIPNAME+(fail.length?' · gagal: '+fail.join('; '):'')+' — unggah file ZIP ini di bawah.';
 }catch(e){st.style.color='#FF4D8D';st.textContent='Gagal: '+e.message;}
 btn.disabled=false;}
</script></body></html>"""

def bsg_entity_kpi_capture_html(ent):
    """HTML Dashboard KPI satu koordinator / counter (tanpa tabel ranking) — sama dengan isi tombol Download JPG di menu DASHBOARD KPI. Return (kind, html)."""
    kind = 'counter' if str(ent).upper().startswith('COUNTER') else 'koor'
    entities = get_focus_entities()[1 if kind == 'counter' else 0]
    ctype = detect_counter_type(ent) if kind == 'counter' else 'ADMIN'
    rows_e, s_e, p_e, t_e = kpi_score_entity_sim(get_entity_kpi_inputs(kind, ent), kind, ctype)
    if kind == 'koor':
        note_e = (f'Dealer: {kpi_entity_dealer(ent)}. Bobot & rumus poin sesuai file MASTER_KPI_KOOR (Result 40 + Process 60 = 100). Target & Result dari tabel MASTER KPI Koordinator; MYDR & Google Review mengikuti dealer induk. '
                  f'Promosi Achiev dikunci 100%. Productivity = STU ÷ QTY Salesman, target {KPI_ENT_PROD_TARGET}.')
    else:
        note_e = (f'Dealer: {kpi_entity_dealer(ent)} | Tipe: COUNTER {ctype}. Bobot & rumus poin sesuai file MASTER_KPI_COUNTER (Result 40 + Process 60 = 100); poin STU / Classy / PRM dari tabel lookup Sheet5. '
                  'Target & Result dari tabel MASTER KPI Counter; MYDR & Google Review mengikuti dealer induk; Walk In dari kolom Prospect Type (Walk In). Promosi & Member Point Achiev dikunci 100%.')
    rank_tbl = []
    for e_k in entities:
        ct_k = detect_counter_type(e_k) if kind == 'counter' else 'ADMIN'
        r_k, s_k, p_k, t_k = kpi_score_entity_sim(get_entity_kpi_inputs(kind, e_k), kind, ct_k)
        rank_tbl.append({'name': e_k, 'code': (kpi_entity_dealer(e_k) or '-').replace('SIP ', ''), 'point': t_k, 'grade': kpi_grade(t_k), 'sales': s_k, 'process': p_k,
                         'achv': {c: r_k[c]['achv'] for c in KPI_ENT_CODES[kind]}, 'note': kpi_note(p_k)})
    rank_tbl.sort(key=lambda x_d: (-x_d['point'], -x_d['process'], x_d['name']))
    date_lbl = f'{tanggal_akhir} {DASH_ABBR.get(current_month, "")} {current_year}'; date_short = f'{tanggal_akhir:02d}/{current_month:02d}/{current_year}'
    h_ = build_dashboard_kpi_html(ent, date_lbl, date_short, rows_e, s_e, p_e, t_e, rank_tbl, ent, False, note_e, codes=KPI_ENT_CODES[kind], rank_kind=kind, with_rank=False)
    return kind, f'<div style="width:1320px;background-color:#0E1117;padding:18px 20px;box-sizing:border-box;">{h_}</div>'

def bsg_render_ppt_batch(zip_label, kpi_dealer=None, kpi_entities=None):
    """Satu tombol: unduh semua gambar panel yang sedang tampil (+ Dashboard KPI) sebagai ZIP, siap diunggah ke kotak gambar PPT."""
    items = [{'fn': k_, 'kind': 'panel', 'html': v_} for k_, v_ in _PANEL_REG.items() if k_.startswith(_BATCH_PREFIX)]
    if kpi_dealer:
        try: items.append({'fn': f'Dashboard_KPI_{_tp_san(kpi_dealer)}_Tgl_{tanggal_akhir}', 'kind': 'frag', 'html': bsg_kpi_capture_html(kpi_dealer)})
        except Exception: pass
    for ent_ in (kpi_entities or []):   # Dashboard KPI koordinator / counter terlemah (slide 14 & 18 template)
        try:
            kind_, h_ = bsg_entity_kpi_capture_html(ent_)
            items.append({'fn': f"Dashboard_KPI_{'Counter' if kind_ == 'counter' else 'Koordinator'}_{_tp_san(ent_)}_Tgl_{tanggal_akhir}", 'kind': 'frag', 'html': h_})
        except Exception as e_k: st.caption(f'⚠️ Dashboard KPI {ent_} tidak dapat dibuat ({type(e_k).__name__}: {str(e_k)[:90]}).')
    if not items: st.caption('Belum ada panel yang dirender untuk diunduh.'); return
    j = lambda o: json.dumps(o, ensure_ascii=False).replace('</', '<\\/')
    components.html(_BATCH_HTML.replace('__N__', str(len(items))).replace('__ITEMS__', j(items)).replace('__ZIPNAME__', j(f"JEMBER_BONDOWOSO_Gambar_PPT_{_tp_san(zip_label)}.zip")), height=64)

def _ppt_collect_images(files):
    """Gabungkan unggahan: file JPG/PNG langsung atau isi file ZIP hasil tombol Download SEMUA."""
    import io, zipfile
    out = {}
    for f_ in (files or []):
        n_, b_ = f_.name, f_.getvalue()
        if n_.lower().endswith('.zip'):
            try:
                with zipfile.ZipFile(io.BytesIO(b_)) as z_:
                    for zi in z_.namelist():
                        if zi.lower().endswith(('.jpg', '.jpeg', '.png')): out[os.path.basename(zi)] = z_.read(zi)
            except Exception: pass
        else: out[n_] = b_
    return out

def bsg_ppt_kpi_payload(sel_dash):
    """Data KPI (Result/Process) dealer terpilih — dipakai slide 17 (tabel Achievement KPI)."""
    scheme = 'MDS' if sel_dash == 'TOTAL JEMBER BONDOWOSO' else 'DEALER'
    inputs = get_dashboard_kpi_inputs()
    rows_s, s_s, p_s, t_s = kpi_score_entity(inputs[sel_dash], scheme)
    d_tbl = []
    for d_k in daftar_dealer:
        r_k, s_k, p_k, t_k = kpi_score_entity(inputs[d_k], scheme)
        d_tbl.append({'name': d_k, 'code': KPI_DEALER_CODES.get(d_k, '-'), 'point': t_k, 'grade': kpi_grade(t_k), 'sales': s_k, 'process': p_k, 'achv': {c: r_k[c]['achv'] for c in KPI_CODES}, 'note': kpi_note(p_k)})
    d_tbl.sort(key=lambda x_d: (-x_d['point'], -x_d['process']))
    date_lbl = f'{tanggal_akhir} {DASH_ABBR.get(current_month, "")} {current_year}'
    note = ('Bobot skema MDS: OPS level supervisor, nilainya sama untuk semua dealer.' if scheme == 'MDS' else 'Bobot skema Dealer (Excel R-MDS): OPR 15 & OPS 0; STU/MAX/CLS memakai 75% poin dasar.') + ' PROCESS ≥45 tanpa potongan, 30–44 = CUT 30%, <30 = CUT 50%.'
    return bsg_collect_ppt_payload(sel_dash, scheme, rows_s, s_s, p_s, t_s, d_tbl, inputs, date_lbl, note, with_teams=False)

def _ppt_ai_cache(dealer):
    """Narasi Condition/Plan hasil tab 'Analisa & Action Plan' (Gemini / Smart Analyzer) bila sudah pernah dibuat untuk dealer & tanggal ini."""
    pre = f'ai_tbl6_v10_35w_{dealer}_{int(tanggal_akhir)}_'
    try:
        for k_ in reversed(list(st.session_state.keys())):
            if isinstance(k_, str) and k_.startswith(pre) and isinstance(st.session_state.get(k_), dict) and st.session_state[k_]: return st.session_state[k_]
    except Exception: pass
    return {}

def bsg_ppt_for_dealer(sel_dash, tpl_bytes=None, images=None):
    """Bangun PPT dari TEMPLATE_PPT_JEMBER_BONDOWOSO.pptx untuk dealer terpilih (atau TOTAL JEMBER BONDOWOSO). Return (bytes, nama_file, daftar_peringatan)."""
    src = tpl_bytes or _tp_find_template()
    if src is None: raise FileNotFoundError(f'File template {_TPL_NAME} tidak ditemukan. Letakkan di folder yang sama dengan app.py (root repo GitHub) atau unggah lewat kotak "Template PPT" di bawah tombol.')
    is_total = (sel_dash == 'TOTAL JEMBER BONDOWOSO'); dlist = list(daftar_dealer) if is_total else [sel_dash]
    tgl = int(tanggal_akhir); tgl_s = 30 if tgl > 21 else min(30, tgl + 7)
    sisa = max(1, sum(1 for d_n in range(tgl + 1, min(tgl_s, max_day_in_month) + 1) if is_workday_day(d_n))); pct_s = tabel_target_harian.get(tgl_s, 100.0) / 100.0
    H = dict(tgl_sasaran=tgl_s, timeline=f'S.d Tgl {tgl_s} {periode_bulan_tahun} ({sisa} hari kerja)')
    addf = lambda curr, targ, mp: calc_rational_7d_addition(curr, targ, tgl, tgl_s, persen_desimal, pct_s, active_mp=max(1, mp))
    S_d = _ppt_scope(dlist, None, sel_dash, is_total)
    def _weakest(counter):
        c_ = [r_ for r_ in S_d['ent_p'] if str(r_[_TP_KN]).upper().startswith('COUNTER') == counter]
        c_ = [r_ for r_ in c_ if r_['TARGET'] > 0] or c_
        return sorted(c_, key=lambda r_: (r_['_pct'], r_['STU']))[0][_TP_KN] if c_ else None
    def _ent_scope(ent):
        if not ent: return None
        dl = kpi_entity_dealer(ent) or (None if is_total else sel_dash); S_ = _ppt_scope([dl] if dl else dlist, ent, ent, False); S_['dealer_lbl'] = dl or sel_dash; return S_
    S_k, S_c = _ent_scope(_weakest(False)), _ent_scope(_weakest(True))
    try: P = bsg_ppt_kpi_payload(sel_dash)
    except Exception: P = None
    D = dict(scope=sel_dash, tgl=tgl, periode=periode_bulan_tahun, S_dealer=S_d, S_koor=S_k, S_counter=S_c, H=H, addf=addf, ai=_ppt_ai_cache(sel_dash), kpi=P, images=images or {})
    data, warn = bsg_ppt_tpl_build(D, src)
    nm = f"JEMBER_BONDOWOSO_Review_{re.sub(r'[^A-Za-z0-9]+', '_', sel_dash)}_Tgl{tgl}_{re.sub(r'[^A-Za-z0-9]+', '_', periode_bulan_tahun)}.pptx"
    return data, nm, warn

def _dash_fmt(code, v):
    if code in ('APV', 'MYM', 'YCN'): return f'{v * 100:.1f}%'
    if code == 'GRT': return f'{v:.1f}'
    if code == 'PRD': return f'{v:.2f}'
    return f'{v:,.0f}'

DASH_ST_COL = {'ok': '#22C55E', 'warn': '#FACC15', 'bad': '#FB7185', 'na': '#64748B'}
DASH_ST_BG = {'ok': 'rgba(34,197,94,0.15)', 'warn': 'rgba(250,204,21,0.14)', 'bad': 'rgba(251,113,133,0.17)', 'na': 'rgba(148,163,184,0.10)'}
DASH_GRADE_BG = {'A': 'rgba(0,255,102,0.14)', 'B': 'rgba(0,229,255,0.14)', 'C': 'rgba(255,234,0,0.14)', 'D': 'rgba(255,77,141,0.16)'}
DASH_CONCL = {'VERY GOOD': ('#22C55E', 'rgba(34,197,94,0.15)', 'rgba(34,197,94,0.45)'), 'GOOD': ('#FACC15', 'rgba(250,204,21,0.14)', 'rgba(250,204,21,0.45)'),
              'NOT GOOD': ('#FB7185', 'rgba(251,113,133,0.16)', 'rgba(251,113,133,0.45)')}

def _dash_status(r):
    if r['weight'] == 0: return 'na'
    if r.get('fixed') or r['point'] >= r['weight']: return 'ok'
    return 'warn' if r['point'] > 0 else 'bad'

def _dash_chip(txt, fg, bg, bd='transparent'):
    return f'<span style="display:inline-block;padding:2px 10px;border-radius:999px;font-size:11px;font-weight:800;letter-spacing:.3px;color:{fg};background:{bg};border:1px solid {bd};white-space:nowrap;">{txt}</span>'

def _dash_tile(label, value, suffix, color, prog=None, concl=None):
    bar = ''
    if prog is not None:
        p = max(0.0, min(100.0, prog))
        bar = f'<div style="height:7px;border-radius:4px;background:#1E293B;margin-top:9px;overflow:hidden;"><div style="width:{p:.1f}%;height:100%;border-radius:4px;background:{color};"></div></div>'
    ch = ''
    if concl:
        fg, bg, bd = DASH_CONCL.get(concl, ('#94A3B8', 'rgba(148,163,184,0.10)', 'rgba(148,163,184,0.3)'))
        ch = f'<div style="margin-top:8px;">{_dash_chip(concl, fg, bg, bd)}</div>'
    return (f'<div style="flex:1 1 128px;min-width:120px;background:linear-gradient(160deg,#151B2B 0%,#0E1117 100%);border:1px solid #232B3E;border-top:3px solid {color};border-radius:12px;padding:10px 14px 12px 14px;box-sizing:border-box;">'
            f'<div style="font-size:10.5px;font-weight:800;letter-spacing:.9px;color:#94A3B8;">{label}</div>'
            f'<div style="font-size:36px;font-weight:900;color:{color};line-height:1.1;margin-top:2px;">{value}<span style="font-size:14px;font-weight:700;color:#64748B;"> {suffix}</span></div>{bar}{ch}</div>')

def build_dashboard_kpi_html(sel_name, date_lbl, date_short, rows, s_pt, p_pt, tot_pt, dealer_tbl, hl_dealer, show_concl, scheme_note, codes=None, rank_kind=None, with_rank=True):
    codes_l = list(codes) if codes else list(KPI_CODES)
    grade = kpi_grade(tot_pt); gcol = DASH_GRADE_COL.get(grade, '#FFFFFF')
    res_set = set(KPI_ENT_RESULT_CODES) if rank_kind else set(KPI_SALES_CODES)
    max_res = sum(rows[c]['weight'] for c in codes_l if c in res_set); max_prc = sum(rows[c]['weight'] for c in codes_l if c not in res_set)
    res_c = ('VERY GOOD' if s_pt >= 38 else ('GOOD' if s_pt >= 20 else 'NOT GOOD')) if show_concl else None
    prc_c = ('VERY GOOD' if p_pt >= 57 else ('GOOD' if p_pt >= 45 else 'NOT GOOD')) if show_concl else None
    # ---------- baris hero: identitas + 4 tile ----------
    identity = (f'<div style="flex:1.35 1 340px;background:linear-gradient(120deg,{DASH_NAVY} 0%,#1E3A8A 100%);border:1px solid #1E3A8A;border-radius:12px;padding:14px 20px;display:flex;flex-direction:column;justify-content:center;box-sizing:border-box;">'
                f'<div style="font-size:11px;font-weight:800;letter-spacing:1.4px;color:#93C5FD;">KEY PERFORMANCE INDICATOR REPORT</div>'
                f'<div style="font-size:27px;font-weight:900;color:#FFFFFF;line-height:1.15;margin:5px 0 9px 0;">{sel_name}</div>'
                f'<div>{_dash_chip("📅 " + date_lbl, "#E0F2FE", "rgba(255,255,255,0.12)", "rgba(255,255,255,0.25)")}</div></div>')
    grade_tile = (f'<div style="flex:0.75 1 118px;min-width:110px;background:linear-gradient(160deg,#151B2B 0%,#0E1117 100%);border:1px solid #232B3E;border-top:3px solid {gcol};border-radius:12px;padding:10px 12px;display:flex;flex-direction:column;align-items:center;justify-content:center;box-sizing:border-box;">'
                  f'<div style="font-size:10.5px;font-weight:800;letter-spacing:.9px;color:#94A3B8;">GRADE</div>'
                  f'<div style="width:62px;height:62px;border-radius:16px;background:{DASH_GRADE_BG.get(grade, "rgba(255,255,255,0.08)")};border:2px solid {gcol};color:{gcol};font-size:40px;font-weight:900;display:flex;align-items:center;justify-content:center;margin-top:6px;">{grade}</div></div>')
    tiles = (grade_tile + _dash_tile('TOTAL POINT', f'{tot_pt}', '/ 100', '#E2E8F0', tot_pt)
             + _dash_tile('RESULT', f'{s_pt}', f'/ {max_res}', '#60A5FA', (s_pt / max_res * 100) if max_res else 0, res_c)
             + _dash_tile('PROCESS', f'{p_pt}', f'/ {max_prc}', '#A78BFA', (p_pt / max_prc * 100) if max_prc else 0, prc_c))
    hero = f'<div style="display:flex;gap:12px;flex-wrap:wrap;align-items:stretch;margin-bottom:12px;">{identity}<div style="flex:2.4 1 560px;display:flex;gap:12px;flex-wrap:wrap;">{tiles}</div></div>'
    # ---------- sorotan ----------
    lost = sorted([(c, rows[c]['weight'] - rows[c]['point']) for c in codes_l if rows[c]['weight'] > 0 and not rows[c].get('fixed') and rows[c]['point'] < rows[c]['weight']], key=lambda t: -t[1])
    zero = [c for c, _ in lost if rows[c]['point'] == 0]
    full = [c for c in codes_l if rows[c]['weight'] > 0 and (rows[c].get('fixed') or rows[c]['point'] >= rows[c]['weight'])]
    chips = ''
    if zero: chips += _dash_chip(f'🚨 {len(zero)} KPI belum menghasilkan poin: ' + ', '.join(zero), '#FB7185', 'rgba(251,113,133,0.14)', 'rgba(251,113,133,0.4)') + ' '
    if lost: chips += _dash_chip('📉 Poin hilang terbesar: ' + ' · '.join(f'{rows[c]["name"]} −{l}' for c, l in lost[:3]), '#FACC15', 'rgba(250,204,21,0.12)', 'rgba(250,204,21,0.4)') + ' '
    chips += _dash_chip(f'✅ {len(full)} KPI poin penuh', '#22C55E', 'rgba(34,197,94,0.14)', 'rgba(34,197,94,0.4)')
    spot = f'<div style="display:flex;align-items:center;gap:8px;flex-wrap:wrap;margin-bottom:14px;"><span style="font-size:10.5px;font-weight:800;letter-spacing:1px;color:#94A3B8;">SOROTAN</span>{chips}</div>'
    # ---------- tabel pencapaian + diagram (SATU tabel: tiap baris KPI punya sel diagram di kanan, otomatis sejajar) ----------
    th = f'background:{DASH_NAVY};color:#fff;font-size:10.5px;font-weight:800;letter-spacing:.3px;padding:8px 6px;border-bottom:2px solid #1E3A8A;'
    heads = ['CODE', 'PERFORMANCE CATEGORY', 'WEIGHT', 'MIN ACV', 'TARGET', 'RESULT', 'ACHIEVEMENT', 'POINT']
    td = 'padding:6px 6px;border-bottom:1px solid #1E293B;font-size:12.5px;color:#E2E8F0;'
    max_a = max([rows[c]['achv'] for c in codes_l] + [1.0]); axis = max(1.2, max_a * 1.25); mark = 100.0 / axis
    NCOL = len(heads) + 1
    body = ''
    for g_lbl, g_codes, g_pt, g_max in (('RESULT', [c for c in codes_l if c in res_set], s_pt, max_res), ('PROCESS', [c for c in codes_l if c not in res_set], p_pt, max_prc)):
        if not g_codes: continue
        body += (f'<tr><td colspan="{NCOL}" style="background:#12304F;padding:6px 10px;border-bottom:1px solid #1E3A8A;"><span style="font-size:11px;font-weight:900;letter-spacing:1px;color:#93C5FD;">{g_lbl}</span>'
                 f'<span style="float:right;font-size:11.5px;font-weight:800;color:#E2E8F0;">{g_pt} <span style="color:#64748B;">/ {g_max} poin</span></span></td></tr>')
        for c in g_codes:
            r = rows[c]; stt = _dash_status(r); col, bgc = DASH_ST_COL[stt], DASH_ST_BG[stt]
            pb = max(0.0, min(r['achv'], 1.0)) * 100
            w = max(0.0, min(r['achv'] / axis, 1.0)) * 100
            tgt_txt = '—' if r.get('fixed') else _dash_fmt(c, r['target']); res_txt = '—' if r.get('fixed') else _dash_fmt(c, r['result'])
            min_txt = '—' if (r['weight'] == 0 or r.get('no_min')) else f'{r["min_acv"] * 100:.0f}%'
            chart_cell = (f'<td style="{td}width:36%;min-width:230px;padding:6px 16px 6px 12px;border-left:2px solid #1E3A8A;vertical-align:middle;">'
                          f'<div style="position:relative;height:18px;background:#141C2B;border-radius:9px;"><div style="position:absolute;left:0;top:0;height:18px;width:{w:.1f}%;background:{col};border-radius:9px;"></div>'
                          f'<div style="position:absolute;left:{mark:.1f}%;top:-4px;height:26px;border-left:1px dashed rgba(255,255,255,0.55);"></div>'
                          f'<div style="position:absolute;left:calc({w:.1f}% + 7px);top:1px;font-size:11px;font-weight:800;color:#E2E8F0;white-space:nowrap;">{r["achv"] * 100:.1f}%</div></div></td>')
            body += (f'<tr style="background:{bgc};"><td style="{td}width:44px;padding:6px 6px;"><span style="display:inline-block;min-width:34px;text-align:center;padding:2px 4px;border-radius:6px;background:rgba(255,255,255,0.08);color:{col};font-weight:900;font-size:11px;">{c}</span></td>'
                     f'<td style="{td}line-height:1.25;">{r["name"]}</td><td style="{td}text-align:right;color:#94A3B8;">{r["weight"]}</td><td style="{td}text-align:right;color:#94A3B8;">{min_txt}</td>'
                     f'<td style="{td}text-align:right;">{tgt_txt}</td><td style="{td}text-align:right;">{res_txt}</td>'
                     f'<td style="{td}text-align:right;white-space:nowrap;"><span style="display:inline-block;width:56px;height:7px;border-radius:4px;background:#1E293B;vertical-align:middle;margin-right:7px;overflow:hidden;"><span style="display:block;width:{pb:.1f}%;height:100%;background:{col};"></span></span>'
                     f'<span style="display:inline-block;min-width:46px;font-weight:800;color:{col};">{r["achv"] * 100:.1f}%</span></td>'
                     f'<td style="{td}text-align:right;white-space:nowrap;"><span style="font-weight:900;color:#FFFFFF;font-size:14px;">{r["point"]}</span><span style="color:#64748B;font-size:11px;">/{r["weight"]}</span></td>{chart_cell}</tr>')
    head1 = (f'<tr><th colspan="{len(heads)}" style="background:{DASH_TEAL};color:#fff;font-size:12px;font-weight:800;letter-spacing:.7px;padding:8px 14px;text-align:left;border-radius:10px 0 0 0;">ACHIEVEMENT BY PERFORMANCE CATEGORY</th>'
             f'<th style="background:{DASH_TEAL};color:#fff;font-size:12px;font-weight:800;letter-spacing:.7px;padding:8px 14px;text-align:left;white-space:nowrap;border-radius:0 10px 0 0;">ACHIEVEMENT VS TARGET PER KPI'
             f'<span style="float:right;margin-left:14px;">{date_short}</span></th></tr>')
    head2 = '<tr>' + ''.join(f'<th style="{th}text-align:{"left" if i < 2 else "right"};">{h}</th>' for i, h in enumerate(heads)) + f'<th style="{th}text-align:left;border-left:2px solid #1E3A8A;padding-left:12px;">CAPAIAN (garis putus-putus = 100% target)</th></tr>'
    legend = ('<div style="display:flex;gap:14px;flex-wrap:wrap;justify-content:flex-end;font-size:10.5px;color:#94A3B8;padding:6px 14px;">'
              '<span><span style="color:#22C55E;">●</span> poin penuh</span><span><span style="color:#FACC15;">●</span> sebagian poin</span><span><span style="color:#FB7185;">●</span> 0 poin</span></div>')
    table_html = (f'<div style="overflow-x:auto;border:1px solid #232B3E;border-radius:10px;background:#0E1117;"><table style="width:100%;border-collapse:collapse;"><thead>{head1}{head2}</thead><tbody>{body}</tbody></table>{legend}</div>')
    mid = f'<div style="margin-bottom:14px;">{table_html}</div>'
    if not dealer_tbl or not with_rank:   # dashboard koordinator / counter, atau versi capture: tanpa tabel ranking
        return f'<div style="font-family:Source Sans Pro,sans-serif;color:#E2E8F0;">{hero}{spot}{mid}<div style="font-size:11px;color:#94A3B8;margin-top:2px;">{scheme_note}</div></div>'
    # ---------- tabel ranking (dealer / koordinator / counter) ----------
    rk_codes = list(KPI_ENT_CODES[rank_kind]) if rank_kind else ['STU', 'MAX', 'CLS', 'OPR', 'SAL', 'APV', 'INC', 'MYM', 'MYD', 'GRT', 'GRV']
    rk_min = (lambda c: KPI_ENT_MIN_ACV[rank_kind].get(c, 0.0)) if rank_kind else (lambda c: KPI_MIN_ACV[c])
    rk_title = {'koor': 'RANKING KOORDINATOR', 'counter': 'RANKING COUNTER'}.get(rank_kind, 'RANKING DEALER')
    rk_name_hdr = {'koor': 'KOORDINATOR NAME', 'counter': 'COUNTER NAME'}.get(rank_kind, 'DEALER NAME')
    dh = ['#', 'CODE', rk_name_hdr, 'POINT', 'GRADE', 'SALES', 'PROCESS'] + rk_codes + ['NOTE']
    dth = f'background:{DASH_NAVY};color:#fff;font-size:11px;font-weight:800;padding:8px 6px;border-bottom:2px solid #1E3A8A;text-align:center;white-space:nowrap;'
    medal = {1: '#FACC15', 2: '#CBD5E1', 3: '#FB923C'}
    drows = ''
    for i, dr in enumerate(dealer_tbl, 1):
        gc = DASH_GRADE_COL.get(dr['grade'], '#FFFFFF'); hl = (dr['name'] == hl_dealer)
        outl = 'outline:2px solid #00E5FF;outline-offset:-2px;' if hl else ''
        dtd = 'padding:6px 6px;border-bottom:1px solid #1E293B;font-size:12.5px;text-align:center;white-space:nowrap;color:#E2E8F0;'
        rk_col = medal.get(i, '#64748B')
        pbar = f'<span style="display:inline-block;width:46px;height:6px;border-radius:3px;background:#1E293B;vertical-align:middle;margin-left:6px;overflow:hidden;"><span style="display:block;width:{max(0, min(100, dr["point"])):.0f}%;height:100%;background:{gc};"></span></span>'
        cells = (f'<td style="{dtd}"><span style="display:inline-block;min-width:22px;padding:1px 0;border-radius:50%;font-weight:900;color:{rk_col};">{i}</span></td><td style="{dtd}color:#94A3B8;">{dr["code"]}</td>'
                 f'<td style="{dtd}text-align:left;font-weight:800;color:#FFFFFF;">{dr["name"]}</td><td style="{dtd}font-weight:900;color:#FFFFFF;">{dr["point"]}{pbar}</td>'
                 f'<td style="{dtd}"><span style="display:inline-block;min-width:26px;padding:1px 0;border-radius:8px;background:{DASH_GRADE_BG.get(dr["grade"], "transparent")};border:1px solid {gc};font-weight:900;color:{gc};">{dr["grade"]}</span></td>'
                 f'<td style="{dtd}">{dr["sales"]}</td><td style="{dtd}">{dr["process"]}</td>')
        for c in rk_codes:
            av = dr['achv'][c]; stt = 'ok' if av >= 1 else ('warn' if av >= rk_min(c) else 'bad')
            cells += f'<td style="{dtd}padding:4px 3px;"><span style="display:inline-block;min-width:42px;padding:2px 4px;border-radius:6px;font-weight:800;font-size:12px;color:{DASH_ST_COL[stt]};background:{DASH_ST_BG[stt]};">{av * 100:.0f}%</span></td>'
        cells += f'<td style="{dtd}font-weight:800;color:#FF80AB;">{dr["note"]}</td>'
        drows += f'<tr style="{outl}">{cells}</tr>'
    dealer_html = (f'<div style="background:{DASH_TEAL};color:#fff;font-size:12px;font-weight:800;letter-spacing:.7px;padding:8px 14px;border-radius:10px 10px 0 0;">{rk_title} — POINT, GRADE & ACHIEVEMENT PER KPI</div>'
                   f'<div style="overflow-x:auto;border:1px solid #232B3E;border-radius:0 0 10px 10px;background:#0E1117;"><table style="width:100%;border-collapse:collapse;"><thead><tr>'
                   + ''.join(f'<th style="{dth}">{h}</th>' for h in dh) + f'</tr></thead><tbody>{drows}</tbody></table></div>'
                   f'<div style="font-size:11px;color:#94A3B8;margin-top:6px;">{scheme_note}</div>')
    return f'<div style="font-family:Source Sans Pro,sans-serif;color:#E2E8F0;">{hero}{spot}{mid}{dealer_html}</div>'


# ===================== DASHBOARD KPI KOORDINATOR & COUNTER (rumus & bobot = sheet Dashboard MASTER_KPI_KOOR / MASTER_KPI_COUNTER) =====================
KPI_ENT_CODES = {'koor':    ['STU', 'CLS', 'MAX', 'PRM', 'SAL', 'PRD', 'DBS', 'INC', 'MYD', 'GRV'],
                 'counter': ['STU', 'CLS', 'MAX', 'PRM', 'WLK', 'MYM', 'DBS', 'INC', 'MYD', 'GRV']}
KPI_ENT_RESULT_CODES = ['STU', 'CLS', 'MAX']
KPI_ENT_NAMES = {'STU': 'STU Total', 'CLS': 'STU AT Classy', 'MAX': 'STU AT PRM (Maxi)', 'PRM': 'Promosi', 'SAL': 'QTY Salesman', 'PRD': 'Productivity Salesman',
                 'WLK': 'Walk In', 'MYM': 'Member Point', 'DBS': 'Database', 'INC': 'Incoming', 'MYD': 'My Yamaha Dealer Rating (MYDR)', 'GRV': 'Google Review'}
KPI_ENT_NAMES_KIND = {'koor': {'PRM': 'Promosi - Aktifitas Online 2X & Offline 2X'}, 'counter': {'PRM': 'Promosi - Live TikTok/IG'}}
KPI_ENT_WEIGHT = {'koor':    {'STU': 20, 'CLS': 10, 'MAX': 10, 'PRM': 15, 'SAL': 10, 'PRD': 5, 'DBS': 10, 'INC': 10, 'MYD': 5, 'GRV': 5},
                  'counter': {'STU': 20, 'CLS': 10, 'MAX': 10, 'PRM': 15, 'WLK': 10, 'MYM': 5, 'DBS': 10, 'INC': 10, 'MYD': 5, 'GRV': 5}}
KPI_ENT_MIN_ACV = {'koor':    {'STU': .70, 'CLS': .70, 'MAX': .70, 'PRM': .85, 'SAL': .85, 'PRD': .85, 'DBS': .85, 'INC': .85, 'MYD': .85, 'GRV': .85},
                   'counter': {'PRM': .85, 'WLK': .60, 'MYM': .85, 'DBS': .85, 'INC': .75, 'MYD': .85, 'GRV': .85}}   # counter STU/CLS/MAX: poin dari tabel lookup (tanpa min ACV)
KPI_ENT_FIXED_ACHV = {'koor': {'PRM'}, 'counter': {'PRM', 'MYM'}}   # Achiev dikunci 100% (permintaan: Promosi semua & Member Point counter = poin penuh)
KPI_ENT_PROD_TARGET = 2.5                                           # target Productivity Salesman koordinator (sheet Dashboard koor, sel F16)
# Tabel poin counter (sheet Sheet5 MASTER_KPI_COUNTER): index = jumlah result unit (0..30) -> (poin COUNTER ADMIN, poin COUNTER MURNI/SALES)
KPI_COUNTER_TABLE = {
    'STU': [(0,0),(0,0),(0,0),(0,0),(0,0),(0,0),(0,0),(0,0),(0,0),(10,0),(11,0),(13,0),(15,10),(17,11),(18,12),(20,14),(20,15),(20,16),(20,18),(20,19),(20,20)] + [(20,20)] * 10,
    'MAX': [(0,0),(0,0),(0,0),(0,0),(6,0),(8,5),(10,7),(10,8),(10,10)] + [(10,10)] * 22,
    'CLS': [(0,0),(5,0),(10,5),(10,10)] + [(10,10)] * 27,
}
KPI_COUNTER_MURNI = ['NUR AFIFAH APRILIANA', 'EMILIA DIANA', 'SITI NURJANAH', 'KOESMIATI', 'ALBERTINE CHRISANTINA PUTRI', 'IIN BARIROH', 'TRIA DAMAYANTI']
KPI_COUNTER_ADMIN = ['NURUL FADILAH', 'PUTRI ZAIN MAZDA', 'RIKA DWI ASTUTI', 'RISKA MELIYANA', 'MUHAMMAD HAMDAN WAHDI', 'NICKY LINTANG PERTIWI', 'DESVIN NURUL PRISTANTINI',
                     'DIAJENG AYUDHIA UTAMI PUTRI', 'NOVIN MARIA PURWANITA', 'INDAH RETUWATI R.', 'DINI KOYIMATUL HIKMAH', 'HERDIK RIZAL UTAMANDA', 'ANNISA ROISNAINI', 'LAILI NISFIATURROHMA']

def _kpi_num(x): return int(round(x)) if abs(x - round(x)) < 1e-9 else x

def detect_counter_type(ent):
    """'MURNI' / 'ADMIN' dari daftar nama di file MASTER_KPI_COUNTER; tidak dikenali -> ADMIN (bisa diganti manual di dashboard)."""
    e = alnum_str(ent).replace('COUNTER', '')
    def hit(names): return any(alnum_str(n) and (alnum_str(n) in e or (len(e) >= 5 and e in alnum_str(n))) for n in names)
    if hit(KPI_COUNTER_MURNI): return 'MURNI'
    return 'ADMIN'

def kpi_point_ent(kind, code, target, result, weight, ctype='ADMIN'):
    """Poin per item persis rumus kolom Point sheet Dashboard (koor: MASTER_KPI_KOOR, counter: MASTER_KPI_COUNTER). Dibatasi maksimal bobot."""
    t, r = float(target or 0), float(result or 0)
    achv = (r / t) if t > 0 else 0.0
    w = weight
    half_up = int(math.ceil(0.5 * w - 1e-9))
    def mid(mn, mult): return min(w, half_up + int(math.floor((achv - mn) / 0.15 * mult + 1e-9)))
    def near_miss(): return _kpi_num(0.5 * w) if (t <= 8 and r == t - 1 and r > 0) else 0
    if kind == 'counter' and code in ('STU', 'CLS', 'MAX'):
        return KPI_COUNTER_TABLE[code][max(0, min(int(round(r)), 30))][0 if ctype == 'ADMIN' else 1]
    if achv >= 1 - 1e-9: return w
    if code in ('STU', 'CLS', 'MAX'): return mid(0.70, 4) if achv >= 0.70 else near_miss()
    if code in ('PRM', 'DBS'):        return mid(0.85, 4) if achv >= 0.85 else near_miss()
    if code == 'SAL':                 return mid(0.85, 4) if achv >= 0.85 else (_kpi_num(0.5 * w) if (t == 6 and r == 5) else 0)
    if code == 'PRD':                 return mid(0.85, 4) if achv >= 0.85 else 0
    if code == 'WLK':                 return mid(0.60, 2) if achv >= 0.60 else (_kpi_num(0.6 * w) if (t == 6 and r == 5) else 0)
    if code == 'INC':                 return mid(0.75 if kind == 'counter' else 0.85, 2) if achv >= (0.75 if kind == 'counter' else 0.85) else 0
    return mid(0.85, 2) if achv >= 0.85 else 0   # MYD, GRV, MYM

def kpi_score_entity_sim(tr, kind, ctype='ADMIN'):
    rows = {}
    for c in KPI_ENT_CODES[kind]:
        t, r = tr.get(c, (0, 0))
        fixed = c in KPI_ENT_FIXED_ACHV[kind]
        if fixed: t, r = 1.0, 1.0
        t, r = float(t or 0), float(r or 0)
        achv = (r / t) if t > 0 else 0.0
        w = KPI_ENT_WEIGHT[kind][c]
        pt = kpi_point_ent(kind, c, t, r, w, ctype)
        no_min = kind == 'counter' and c in ('STU', 'CLS', 'MAX')
        rows[c] = dict(code=c, name=KPI_ENT_NAMES_KIND[kind].get(c, KPI_ENT_NAMES[c]), weight=w, min_acv=KPI_ENT_MIN_ACV[kind].get(c, KPI_ENT_MIN_ACV[kind].get('STU', 0.0)) if not no_min else 0.0,
                       target=t, result=r, achv=achv, point=pt, fixed=fixed, no_min=no_min)
    sales = sum(rows[c]['point'] for c in KPI_ENT_RESULT_CODES)
    process = sum(rows[c]['point'] for c in KPI_ENT_CODES[kind] if c not in KPI_ENT_RESULT_CODES)
    return rows, _kpi_num(sales), _kpi_num(process), _kpi_num(sales + process)

def get_entity_kpi_inputs(kind, ent):
    """Target & Result per item. Semua angka dari tabel MASTER KPI (koordinator / counter); MYDR & Google Review = dealer induk."""
    df_it, _ = build_kpi_section_df([ent], 'NAMA KOORDINATOR' if kind == 'koor' else 'NAMA COUNTER', kind, kind)
    r = df_it.iloc[0] if (df_it is not None and not df_it.empty) else None
    def g(k): return safe_float(r[k]) if (r is not None and k in r.index) else 0.0
    t_stu, r_stu, t_mp, r_mp = g('T. STU'), g('R. STU'), g('T. MP'), g('R. MP')
    inp = {'STU': (t_stu, r_stu), 'CLS': (g('T. CLASSY'), g('R. CLASSY')), 'MAX': (g('T. MAXI'), g('R. MAXI')), 'PRM': (1.0, 1.0),
           'INC': (g('T. INC'), g('R. INC')), 'DBS': (g('T. DB'), g('R. DB')), 'MYD': (g('T. MYDR'), g('R. MYDR')), 'GRV': (g('T. G REV'), g('R. G REV'))}
    if kind == 'koor':
        inp['SAL'] = (t_mp, r_mp)
        inp['PRD'] = (KPI_ENT_PROD_TARGET, (r_stu / r_mp) if r_mp > 0 else 0.0)
    else:
        inp['WLK'] = (g('T. Walk In'), g('R. Walk In'))
        inp['MYM'] = (1.0, 1.0)
    return inp

def get_focus_entities():
    """Daftar koordinator & counter yang bertugas di 10 dealer fokus (urut sesuai daftar_dealer)."""
    ent_all = set(clean_str(v) for v in mapping_salesman_koor.values() if clean_str(v)) | set(k for k in saved_tk.keys() if k != 'TOTAL')
    for d_s, ck in [(df_faktur, 'Koor'), (df_nota, 'Koor'), (df_database, 'Koor'), (df_spk, 'Koor_SPK'), (df_rating, 'Koor')]:
        if d_s is not None and ck in d_s.columns:
            ent_all.update(clean_str(k) for k in d_s[ck].dropna().unique() if clean_str(k) and clean_str(k) not in ['LAIN-LAIN', 'NAN', 'NONE'])
    focus = [e for e in ent_all if kpi_entity_dealer(e)]
    key_fn = lambda e: (daftar_dealer.index(kpi_entity_dealer(e)), e)
    return sorted([e for e in focus if not e.upper().startswith('COUNTER')], key=key_fn), sorted([e for e in focus if e.upper().startswith('COUNTER')], key=key_fn)

def render_entity_kpi_dashboard(kind, entities, title_txt, sel_key, date_lbl, date_short):
    st.markdown(f'<div class="section-title">{title_txt}</div>', unsafe_allow_html=True)
    lbl = 'Koordinator' if kind == 'koor' else 'Counter'
    if not entities:
        st.info(f'ℹ️ Belum ada {lbl} dari 10 dealer fokus. Pastikan file Data Lookup & Database sudah diupload.')
        return
    sel = st.selectbox(f'👤 Pilih {lbl}:', entities, index=0, key=sel_key, format_func=lambda e: f'{e}  —  {kpi_entity_dealer(e)}')
    ctype = 'ADMIN'
    if kind == 'counter':
        auto_t = detect_counter_type(sel)
        ctype = st.radio('Tipe Counter (menentukan tabel poin STU / Classy / PRM):', ['ADMIN', 'MURNI'], index=['ADMIN', 'MURNI'].index(auto_t), horizontal=True, key=f'ctype_{sel}')
    rows_e, s_e, p_e, t_e = kpi_score_entity_sim(get_entity_kpi_inputs(kind, sel), kind, ctype)
    if kind == 'koor':
        note_e = (f'Dealer: {kpi_entity_dealer(sel)}. Bobot & rumus poin sesuai file MASTER_KPI_KOOR (Result 40 + Process 60 = 100). Target & Result dari tabel MASTER KPI Koordinator; MYDR & Google Review mengikuti dealer induk. '
                  f'Promosi Achiev dikunci 100%. Productivity = STU ÷ QTY Salesman, target {KPI_ENT_PROD_TARGET}.')
    else:
        note_e = (f'Dealer: {kpi_entity_dealer(sel)} | Tipe: COUNTER {ctype}. Bobot & rumus poin sesuai file MASTER_KPI_COUNTER (Result 40 + Process 60 = 100); poin STU / Classy / PRM dari tabel lookup Sheet5. '
                  'Target & Result dari tabel MASTER KPI Counter; MYDR & Google Review mengikuti dealer induk; Walk In dari kolom Prospect Type (Walk In). Promosi & Member Point Achiev dikunci 100%.')
    rank_tbl = []
    for e_k in entities:
        ct_k = detect_counter_type(e_k) if kind == 'counter' else 'ADMIN'
        if e_k == sel: ct_k = ctype
        r_k, s_k, p_k, t_k = kpi_score_entity_sim(get_entity_kpi_inputs(kind, e_k), kind, ct_k)
        rank_tbl.append({'name': e_k, 'code': (kpi_entity_dealer(e_k) or '-').replace('SIP ', ''), 'point': t_k, 'grade': kpi_grade(t_k), 'sales': s_k, 'process': p_k,
                         'achv': {c: r_k[c]['achv'] for c in KPI_ENT_CODES[kind]}, 'note': kpi_note(p_k)})
    rank_tbl.sort(key=lambda x_d: (-x_d['point'], -x_d['process'], x_d['name']))
    html_e = build_dashboard_kpi_html(sel, date_lbl, date_short, rows_e, s_e, p_e, t_e, rank_tbl, sel, False, note_e, codes=KPI_ENT_CODES[kind], rank_kind=kind)
    st.markdown(html_e, unsafe_allow_html=True)
    html_e_cap = build_dashboard_kpi_html(sel, date_lbl, date_short, rows_e, s_e, p_e, t_e, rank_tbl, sel, False, note_e, codes=KPI_ENT_CODES[kind], rank_kind=kind, with_rank=False)   # copy/download HD: dashboard saja, tanpa tabel ranking
    cap_e = f'<div style="width:1320px;background-color:#0E1117;padding:18px 20px;box-sizing:border-box;">{html_e_cap}</div>'
    render_table_image_buttons(f'cap_dashboard_kpi_{kind}', f'Dashboard_KPI_{lbl}_{sel}_Tgl_{tanggal_akhir}', standalone_html=cap_e)


def bsg_render_stu_per_leasing_faktur():
    """STU PER LEASING — FAKTUR (semua dealer) — kini di tab STU › Analisa & Action Plan."""
    jdl_stu_lsg_fak = f'STU PER LEASING — FAKTUR (1 s.d {tanggal_akhir} {periode_bulan_tahun} vs Bulan Lalu)'
    st.markdown(f'<div class="section-title">{jdl_stu_lsg_fak}</div>', unsafe_allow_html=True)

    AREA_DEALERS_STU_LEASING = [
        ('JEMBER', ['SENTRAL YAMAHA JEMBER', 'SIP AMBULU', 'SIP BALUNG', 'SIP GAJAH MADA', 'SIP KEBONSARI', 'SIP RAMBIPUJI', 'SIP SUMBERSARI', 'SIP TRUNOJOYO']),
        ('BONDOWOSO', ['SIP BONDOWOSO', 'SIP NANGKAAN'])
    ]
    ALL_28_DEALERS = [d for _, d_list in AREA_DEALERS_STU_LEASING for d in d_list]

    def match_dealer_stu_leasing(val):
        s = clean_str(val)
        if not s or s in ['NAN', 'NONE']: return ''
        if s in ALL_28_DEALERS: return s
        s_ns = alnum_str(s)
        for d in ALL_28_DEALERS:
            if alnum_str(d) == s_ns: return d
        for k1, k2, res in [('KARUNIA PUTRA', '', 'KARUNIA PUTRA MOTOR / SIT'), ('BUDI JAYA', 'PHJ', 'BUDI JAYA MOTOR/ PHJ'), ('BUDI JAYA', '', 'BUDI JAYA MOTOR (BWI)'), ('KARUNIA MOTOR', 'JJG', 'KARUNIA MOTOR (JJG)'), ('KARUNIA MOTOR', 'RGJ', 'KARUNIA MOTOR (RGJ)'), ('INDO PERKASA', 'KENCONG', 'INDO PERKASA KENCONG'), ('INDO PERKASA', 'CLURING', 'INDO PERKASA CLURING'), ('INDO PERKASA', 'JEMBER', 'INDO PERKASA JEMBER'), ('AMIN MOTOR', 'GENTENG', 'AMIN MOTOR GENTENG'), ('AMIN MOTOR', '', 'AMIN MOTOR BANYUWANGI'), ('SENTRAL YAMAHA', '', 'SENTRAL YAMAHA JEMBER'), ('SUMBER MAS', '', 'SUMBER MAS MOTOR')]:
            if k1 in s and (not k2 or k2 in s): return res
        return normalize_dealer_name(s)

    df_source_faktur_lsg = df_faktur_all if (df_faktur_all is not None and not df_faktur_all.empty) else df_faktur
    if df_source_faktur_lsg is not None and not df_source_faktur_lsg.empty:
        df_fa = df_source_faktur_lsg.copy()
        dlr_col_fa = 'Nama Dealer' if 'Nama Dealer' in df_fa.columns else next((c for c in df_fa.columns if 'DEALER' in str(c).upper()), df_fa.columns[1] if df_fa.shape[1] > 1 else df_fa.columns[0])
        df_fa['Dealer_Map_Leasing'] = df_fa[dlr_col_fa].apply(match_dealer_stu_leasing)
        fa_curr = df_fa[(df_fa['bulan'] == current_month) & (df_fa['tahun'] == current_year) & (df_fa['hari'] >= 1) & (df_fa['hari'] <= tanggal_akhir)] if 'bulan' in df_fa.columns else pd.DataFrame()
        fa_lm = df_fa[(df_fa['bulan'] == last_month) & (df_fa['tahun'] == last_month_year) & (df_fa['hari'] >= 1) & (df_fa['hari'] <= tanggal_akhir)] if 'bulan' in df_fa.columns else pd.DataFrame()
        KREDIT_CATS = ['BAF', 'ADIRA', 'SOF', 'IMFI', 'OTHERS']

        def calc_stu_per_leasing_row(label_row, dealer_subset):
            sub_c = fa_curr[fa_curr['Dealer_Map_Leasing'].isin(dealer_subset)] if not fa_curr.empty else pd.DataFrame()
            sub_l = fa_lm[fa_lm['Dealer_Map_Leasing'].isin(dealer_subset)] if not fa_lm.empty else pd.DataFrame()
            row_d, kredit_c, kredit_l = {'NAMA DEALER': label_row}, 0, 0
            for lsg in KREDIT_CATS:
                c_val = int((sub_c['Leasing_Faktur_Clean'] == lsg).sum()) if not sub_c.empty else 0
                l_val = int((sub_l['Leasing_Faktur_Clean'] == lsg).sum()) if not sub_l.empty else 0
                row_d[lsg], row_d[f'{lsg} LM'] = c_val, l_val
                kredit_c += c_val; kredit_l += l_val
                if lsg == 'BAF': row_d['_baf_grw'], row_d['_baf_val'] = ((c_val / l_val * 100.0) if l_val > 0 else 0.0), c_val
            row_d['BCS'] = f'{(row_d["BAF"] / kredit_c * 100.0 if kredit_c > 0 else 0.0):.1f}%'
            row_d['KREDIT'], row_d['KREDIT LM'] = kredit_c, kredit_l
            cash_c = int((sub_c['Leasing_Faktur_Clean'] == 'CASH').sum()) if not sub_c.empty else 0
            cash_l = int((sub_l['Leasing_Faktur_Clean'] == 'CASH').sum()) if not sub_l.empty else 0
            row_d['CASH'], row_d['CASH LM'], row_d['TOTAL'], row_d['TOTAL LM'] = cash_c, cash_l, kredit_c + cash_c, kredit_l + cash_l
            return row_d

        full_rows_stu_lsg, item_rows_stu_lsg = [], []
        for area_name, d_list in AREA_DEALERS_STU_LEASING:
            area_dealer_rows = sorted([calc_stu_per_leasing_row(d_nm, [d_nm]) for d_nm in d_list], key=lambda r: (r['_baf_grw'], r['_baf_val']))
            full_rows_stu_lsg.extend(area_dealer_rows); item_rows_stu_lsg.extend(area_dealer_rows)
            full_rows_stu_lsg.append(calc_stu_per_leasing_row(area_name, d_list))

        full_rows_stu_lsg.append(calc_stu_per_leasing_row('TOTAL JEMBER', ALL_28_DEALERS))
        row_tot_jember_bondowoso = calc_stu_per_leasing_row('TOTAL JEMBER BONDOWOSO', daftar_dealer)
        full_rows_stu_lsg.append(row_tot_jember_bondowoso)

        kredit_jember_bondowoso_c, kredit_jember_bondowoso_l = safe_int(row_tot_jember_bondowoso.get('KREDIT', 0)), safe_int(row_tot_jember_bondowoso.get('KREDIT LM', 0))
        tot_jember_bondowoso_c, tot_jember_bondowoso_l = safe_int(row_tot_jember_bondowoso.get('TOTAL', 0)), safe_int(row_tot_jember_bondowoso.get('TOTAL LM', 0))
        cash_jember_bondowoso_c, cash_jember_bondowoso_l = safe_int(row_tot_jember_bondowoso.get('CASH', 0)), safe_int(row_tot_jember_bondowoso.get('CASH LM', 0))

        row_cs_jember_bondowoso = {'NAMA DEALER': 'CREDIT SHARE JEMBER BONDOWOSO', 'BCS': row_tot_jember_bondowoso.get('BCS', '0.0%'), '_baf_grw': 0.0, '_baf_val': 0}
        for lsg in KREDIT_CATS:
            val_c, val_l = safe_int(row_tot_jember_bondowoso.get(lsg, 0)), safe_int(row_tot_jember_bondowoso.get(f'{lsg} LM', 0))
            row_cs_jember_bondowoso[lsg] = f'{(val_c / kredit_jember_bondowoso_c * 100.0):.1f}%' if kredit_jember_bondowoso_c > 0 else '0.0%'
            row_cs_jember_bondowoso[f'{lsg} LM'] = f'{(val_l / kredit_jember_bondowoso_l * 100.0):.1f}%' if kredit_jember_bondowoso_l > 0 else '0.0%'
        row_cs_jember_bondowoso['KREDIT'] = f'{(kredit_jember_bondowoso_c / tot_jember_bondowoso_c * 100.0):.1f}%' if tot_jember_bondowoso_c > 0 else '0.0%'
        row_cs_jember_bondowoso['KREDIT LM'] = f'{(kredit_jember_bondowoso_l / tot_jember_bondowoso_l * 100.0):.1f}%' if tot_jember_bondowoso_l > 0 else '0.0%'
        row_cs_jember_bondowoso['CASH'] = f'{(cash_jember_bondowoso_c / tot_jember_bondowoso_c * 100.0):.1f}%' if tot_jember_bondowoso_c > 0 else '0.0%'
        row_cs_jember_bondowoso['CASH LM'] = f'{(cash_jember_bondowoso_l / tot_jember_bondowoso_l * 100.0):.1f}%' if tot_jember_bondowoso_l > 0 else '0.0%'
        row_cs_jember_bondowoso['TOTAL'], row_cs_jember_bondowoso['TOTAL LM'] = ('100.0%' if tot_jember_bondowoso_c > 0 else '0.0%'), ('100.0%' if tot_jember_bondowoso_l > 0 else '0.0%')
        full_rows_stu_lsg.append(row_cs_jember_bondowoso)

        kol_stu_per_lsg = ['NAMA DEALER', 'BCS', 'BAF', 'BAF LM', 'ADIRA', 'ADIRA LM', 'SOF', 'SOF LM', 'IMFI', 'IMFI LM', 'OTHERS', 'OTHERS LM', 'KREDIT', 'KREDIT LM', 'CASH', 'CASH LM', 'TOTAL', 'TOTAL LM']
        render_compact_html_table(pd.DataFrame(item_rows_stu_lsg), pd.DataFrame(full_rows_stu_lsg), kol_stu_per_lsg, 'NAMA DEALER', '_baf_grw', key_suffix=f'STU_Per_Leasing_Tgl_{tanggal_akhir}', header_labels=kol_stu_per_lsg, ultra_compact=True, title_text=jdl_stu_lsg_fak)
    else: st.warning('⚠️ File Faktur belum tersedia. Silakan unggah file Faktur di menu **📁 UPLOAD FILE**.')


# ============ RUNNING TEXT (pesan dari admin/SPV, tampil di atas dropdown dealer) ============
RUNNING_TEXT_FILE = _ABS('running_text.json')

def bsg_render_running_text():
    """Tampilkan pesan berjalan (marquee) bila admin sudah mengisi & mengaktifkannya di tab Analisa & Action Plan."""
    d = load_json(RUNNING_TEXT_FILE)
    txt = str(d.get('text', '') or '').strip()
    if not txt or not d.get('active', True): return
    msg = ' &nbsp;&nbsp;•&nbsp;&nbsp; '.join(_esc(x_.strip()) for x_ in txt.splitlines() if x_.strip())
    dur = max(18, int(len(txt) * 0.22))   # makin panjang pesan, makin lama durasi agar tetap terbaca
    st.markdown(
        '<style>@keyframes bsgmq{0%{transform:translateX(0)}100%{transform:translateX(-100%)}}'
        '.bsg-mq{background:linear-gradient(90deg,#4318FF 0%,#1E293B 100%);border:1px solid #2D3748;border-radius:10px;overflow:hidden;white-space:nowrap;margin:0 0 12px 0;padding:8px 0;}'
        f'.bsg-mq span{{display:inline-block;padding-left:100%;font-size:15px;font-weight:700;color:#FFFFFF;animation:bsgmq {dur}s linear infinite;}}'
        '.bsg-mq:hover span{animation-play-state:paused;}</style>'
        f'<div class="bsg-mq"><span>📢 {msg}</span></div>', unsafe_allow_html=True)

def bsg_admin_running_text_editor():
    """Hanya dipanggil dari tab Analisa & Action Plan (tab ini hanya tampil untuk SPV/admin)."""
    d = load_json(RUNNING_TEXT_FILE)
    st.markdown('<div class="section-title">📢 RUNNING TEXT — Pesan untuk Semua Tim</div>', unsafe_allow_html=True)
    st.caption('Pesan ini tampil berjalan di atas dropdown "Pilih Dealer untuk Review" untuk semua pengguna. Hanya akun SPV/admin yang melihat kotak ini. Baris baru dipisahkan dengan tanda •.')
    txt_new = st.text_area('Isi pesan:', value=str(d.get('text', '') or ''), height=100, max_chars=600, key='rt_text_in', placeholder='Contoh: Briefing pagi 08.30 · Laporan rekrutmen 11.00 · Visit market 13.00 · Daily report 16.00')
    act_new = st.checkbox('Tampilkan running text', value=bool(d.get('active', True)), key='rt_active_in')
    c_s, c_h, _ = st.columns([1, 1, 4])
    if c_s.button('💾 Simpan pesan', key='rt_save', type='primary', use_container_width=True):
        save_json({'text': txt_new.strip(), 'active': bool(act_new)}, RUNNING_TEXT_FILE); st.success('✅ Running text disimpan.'); st.rerun()
    if c_h.button('🗑️ Hapus pesan', key='rt_clear', use_container_width=True):
        save_json({'text': '', 'active': False}, RUNNING_TEXT_FILE); st.success('✅ Running text dihapus.'); st.rerun()
    st.divider()

if menu_pilihan == '🎯 MASTER KPI':
    st.markdown(f'<div class="main-title">DASHBOARD KPI (1 s.d {tanggal_akhir} {periode_bulan_tahun})</div><div class="sub-title">Monitoring Target (T.) vs Result (R.) — KPI Dealer, KPI Koordinator & KPI Counter</div>', unsafe_allow_html=True)

    saved_kpi = load_json(KPI_SAVED_FILE)
    saved_td = load_json(TARGET_DEALER_FILE)
    saved_tk = load_json(TARGET_KOOR_FILE)

    KPI_METRIC_COLS = ['T. STU', 'R. STU', 'T. MAXI', 'R. MAXI', 'T. CLASSY', 'R. CLASSY', 'T. MP', 'R. MP', 'T. INC', 'R. INC', 'T. ACC%', 'R. ACC%', 'T. MYDR', 'R. MYDR', 'T. G REV', 'R. G REV']
    KPI_EDITABLE_COLS = ['T. STU', 'T. MAXI', 'T. CLASSY', 'T. MP', 'T. INC', 'T. ACC%', 'T. G REV', 'R. G REV']
    KPI_PAIRS_ALL = [('T. STU', 'R. STU'), ('T. MAXI', 'R. MAXI'), ('T. CLASSY', 'R. CLASSY'), ('T. MP', 'R. MP'), ('T. INC', 'R. INC'), ('T. ACC%', 'R. ACC%'), ('T. MYDR', 'R. MYDR'), ('T. G REV', 'R. G REV'), ('T. DB', 'R. DB'), ('T. Walk In', 'R. Walk In')]

    def kpi_metric_cols(section_key):
        # Koordinator: + DB | Counter: + DB & Walk In | Dealer: kolom dasar
        return KPI_METRIC_COLS + {'koor': ['T. DB', 'R. DB'], 'counter': ['T. DB', 'R. DB', 'T. Walk In', 'R. Walk In']}.get(section_key, [])

    def kpi_editable_cols(section_key):
        if section_key == 'dealer': return list(KPI_EDITABLE_COLS)
        # G REV koordinator/counter mengikuti dealer -> tidak diedit di sini
        return [c for c in KPI_EDITABLE_COLS if c not in ('T. G REV', 'R. G REV')] + ['T. DB'] + (['T. Walk In'] if section_key == 'counter' else [])

    def build_kpi_capture_html(df_full, name_col, title_txt, dom_id, metric_cols=None):
        cols_all = [name_col] + (metric_cols or KPI_METRIC_COLS)
        pairs = [p for p in KPI_PAIRS_ALL if p[0] in cols_all]
        r_to_t = {r_c: t_c for t_c, r_c in pairs}
        html = [f'<div id="{dom_id}" class="compact-table-wrap"><div class="cap-title-banner">🐺 {title_txt}</div><table class="compact-tbl"><thead><tr>']
        for c in cols_all:
            bg_th = 'background-color:#1E3A8A;color:#FFFFFF;' if c == name_col else ('background-color:#0F172A;color:#F8FAFC;' if c.startswith('T.') else 'background-color:#1E40AF;color:#FFFFFF;')
            html.append(f'<th style="{bg_th}">{c}</th>')
        html.append('</tr></thead><tbody>')
        for _, row in df_full.iterrows():
            nm = str(row[name_col])
            is_tot = nm == 'TOTAL'
            tr_style = 'background-color:#4318FF;color:#FFFFFF;font-weight:800;' if is_tot else ''
            html.append(f'<tr style="{tr_style}">')
            for c in cols_all:
                val = row[c]
                disp_v = f"{val:.1f}%" if c in ['T. ACC%', 'R. ACC%'] else str(val)
                cell_st = ''
                if not is_tot:
                    if c == name_col:
                        cell_st = 'color:#FFFFFF;font-weight:700;'
                    elif c.startswith('T.'):
                        cell_st = 'color:#CBD5E1;background-color:#111827;'
                    elif c in r_to_t:
                        t_col = r_to_t[c]
                        tv, rv = safe_float(row[t_col]), safe_float(row[c])
                        ideal_cmp = tv if c in ['R. MP', 'R. ACC%', 'R. MYDR'] else (tv * persen_desimal)
                        if tv > 0 and rv >= ideal_cmp:
                            cell_st = 'background-color:rgba(0,255,102,0.16);color:#00FF66;font-weight:800;'
                        elif tv > 0:
                            cell_st = 'background-color:rgba(255,0,127,0.22);color:#FF80AB;font-weight:800;'
                        else:
                            cell_st = 'color:#E2E8F0;'
                html.append(f'<td style="{cell_st}">{disp_v}</td>')
            html.append('</tr>')
        html.append('</tbody></table></div>')
        return ''.join(html)

    def render_kpi_section(df_items, df_full, name_col, section_key, title_txt, btn_label):
        if df_full is None or df_full.empty:
            st.info(f"ℹ️ Belum ada data untuk {title_txt}.")
            return
        metric_cols_sec, edit_cols_sec = kpi_metric_cols(section_key), kpi_editable_cols(section_key)
        cols_show = [name_col] + metric_cols_sec
        pairs = [p for p in KPI_PAIRS_ALL if p[0] in cols_show]
        r_to_t = {r_c: t_c for t_c, r_c in pairs}

        def sty_kpi(row):
            if str(row[name_col]) == 'TOTAL':
                return ['background-color: #4318FF; color: #FFFFFF; font-weight: 800'] * len(row)
            styles = []
            for c in cols_show:
                if c == name_col:
                    styles.append('font-weight: 700; color: #FFFFFF')
                elif c.startswith('T.'):
                    styles.append('color: #CBD5E1; background-color: #141A26')
                elif c in r_to_t:
                    t_c = r_to_t[c]
                    tv, rv = safe_float(row[t_c]), safe_float(row[c])
                    ideal_cmp = tv if c in ['R. MP', 'R. ACC%', 'R. MYDR'] else (tv * persen_desimal)
                    if tv > 0 and rv >= ideal_cmp:
                        styles.append('background-color: rgba(0, 255, 102, 0.16); color: #00FF66; font-weight: 800')
                    elif tv > 0:
                        styles.append('background-color: rgba(255, 0, 127, 0.22); color: #FF4081; font-weight: 800')
                    else:
                        styles.append('color: #E2E8F0')
                else:
                    styles.append('')
            return styles

        col_cfg = {name_col: st.column_config.TextColumn(name_col, width="medium")}
        for c in metric_cols_sec:
            if c in ['T. ACC%', 'R. ACC%']:
                col_cfg[c] = st.column_config.NumberColumn(c, width="small", format="%.1f%%")
            else:
                col_cfg[c] = st.column_config.NumberColumn(c, width="small", format="%d")

        disabled_cols = [c for c in cols_show if c not in edit_cols_sec]
        ed_df = st.data_editor(
            df_full[cols_show].style.apply(sty_kpi, axis=1),
            hide_index=True,
            use_container_width=True,
            height=calc_table_height(df_full),
            column_config=col_cfg,
            disabled=disabled_cols,
            key=f'editor_kpi_{section_key}'
        )
        c_b1, c_b2 = st.columns([1, 3])
        with c_b1:
            if st.button(btn_label, type='primary', key=f'btn_save_kpi_{section_key}'):
                cur_all = load_json(KPI_SAVED_FILE)
                sec_dict = cur_all.get(section_key, {})
                td_sync = load_json(TARGET_DEALER_FILE) if section_key == 'dealer' else load_json(TARGET_KOOR_FILE)
                for _, r in ed_df.iterrows():
                    enm = str(r[name_col])
                    if enm == 'TOTAL': continue
                    sec_dict[enm] = {k_ed: (safe_float(r[k_ed]) if 'ACC%' in k_ed else safe_int(r[k_ed])) for k_ed in edit_cols_sec}
                    td_sync[enm] = safe_int(r['T. STU'])
                cur_all[section_key] = sec_dict
                save_json(cur_all, KPI_SAVED_FILE)
                save_json(td_sync, TARGET_DEALER_FILE if section_key == 'dealer' else TARGET_KOOR_FILE)
                st.success(f'✅ Target & Data {title_txt} berhasil disimpan!')
                st.rerun()
        with c_b2:
            cap_id = f'cap_kpi_{section_key}'
            cap_html = build_kpi_capture_html(df_full, name_col, title_txt, cap_id, metric_cols_sec)
            render_table_image_buttons(cap_id, f'{title_txt}_Tgl_{tanggal_akhir}', standalone_html=cap_html)

    df_kpi_dlr_items, df_kpi_dlr_full = build_kpi_section_df(daftar_dealer, 'NAMA DEALER', 'dealer', 'dealer')
    if not df_kpi_dlr_full.empty:
        tot_d = df_kpi_dlr_full.iloc[-1]
        tot_d = df_kpi_dlr_full.iloc[-1]
        bsg_safe(bsg_render_master_kpi_cards, tot_d, target_persen_harian, label="Kartu KPI")

    st.divider()
    jdl_kpi_1 = f'1. TABEL KPI DEALER (1 s.d {tanggal_akhir} {periode_bulan_tahun})'
    st.markdown(f'<div class="section-title">{jdl_kpi_1}</div>', unsafe_allow_html=True)
    render_kpi_section(df_kpi_dlr_items, df_kpi_dlr_full, 'NAMA DEALER', 'dealer', jdl_kpi_1, '💾 Simpan KPI Dealer')

    ent_all_kpi = set(clean_str(v) for v in mapping_salesman_koor.values() if clean_str(v)) | set(k for k in saved_tk.keys() if k != 'TOTAL')
    for d_s, ck in [(df_faktur, 'Koor'), (df_nota, 'Koor'), (df_database, 'Koor'), (df_spk, 'Koor_SPK'), (df_rating, 'Koor')]:
        if d_s is not None and ck in d_s.columns:
            ent_all_kpi.update(clean_str(k) for k in d_s[ck].dropna().unique() if clean_str(k) and clean_str(k) not in ['LAIN-LAIN', 'NAN', 'NONE'])

    koor_entities = sorted([e for e in ent_all_kpi if not e.upper().startswith('COUNTER')])
    cnt_entities = sorted([e for e in ent_all_kpi if e.upper().startswith('COUNTER')])

    st.divider()
    jdl_kpi_2 = f'2. TABEL KPI KOORDINATOR (1 s.d {tanggal_akhir} {periode_bulan_tahun})'
    st.markdown(f'<div class="section-title">{jdl_kpi_2}</div>', unsafe_allow_html=True)
    render_unmapped_salesman_alert('kpi_koor')
    df_kpi_kr_items, df_kpi_kr_full = build_kpi_section_df(koor_entities, 'NAMA KOORDINATOR', 'koor', 'koor')
    render_kpi_section(df_kpi_kr_items, df_kpi_kr_full, 'NAMA KOORDINATOR', 'koor', jdl_kpi_2, '💾 Simpan KPI Koordinator')

    st.divider()
    jdl_kpi_3 = f'3. TABEL KPI COUNTER (1 s.d {tanggal_akhir} {periode_bulan_tahun})'
    st.markdown(f'<div class="section-title">{jdl_kpi_3}</div>', unsafe_allow_html=True)
    df_kpi_cnt_items, df_kpi_cnt_full = build_kpi_section_df(cnt_entities, 'NAMA COUNTER', 'counter', 'counter')
    render_kpi_section(df_kpi_cnt_items, df_kpi_cnt_full, 'NAMA COUNTER', 'counter', jdl_kpi_3, '💾 Simpan KPI Counter')

elif menu_pilihan == '📊 DASHBOARD KPI':
    st.markdown(f'<div class="main-title">DASHBOARD KPI (1 s.d {tanggal_akhir} {periode_bulan_tahun})</div><div class="sub-title">Laporan Key Performance Indicator — poin, grade & capaian tiap KPI (Target & Result dari MASTER KPI)</div>', unsafe_allow_html=True)
    col_d1, col_d2 = st.columns([1, 2])
    with col_d1: sel_dash = st.selectbox('🏢 Pilih Dealer:', ['TOTAL JEMBER BONDOWOSO'] + daftar_dealer, index=0, key='sel_dash_kpi')
    # Bobot otomatis: TOTAL JEMBER BONDOWOSO -> bobot MDS (sesuai gambar) | pilih dealer -> bobot Dealer (sesuai Excel R-MDS)
    scheme_dash = 'MDS' if sel_dash == 'TOTAL JEMBER BONDOWOSO' else 'DEALER'
    with col_d2:
        st.write('')
        st.caption('⚖️ Bobot poin otomatis: **' + ('MDS (sesuai gambar)' if scheme_dash == 'MDS' else 'Dealer (sesuai Excel R-MDS)') + '**')

    man_cur = load_json(KPI_MANUAL_FILE)
    man_d_cur = man_cur.get('dealers', {})
    man_ops_cur = {**DEFAULT_KPI_OPS, **man_cur.get('ops', {})}
    with st.expander('✏️ Input KPI Manual — OPR, MYM, GRT & OPS (belum ada di MASTER KPI)'):
        st.caption('Nilai awal diambil dari file KPI Report (15 Okt 2026). Update sesuai data terbaru lalu klik Simpan. MYM dalam persen (contoh 97.5). OPS bersifat level supervisor dan dipakai untuk semua pilihan.')
        df_man = pd.DataFrame([{'DEALER': d, **{k: safe_float(({**DEFAULT_KPI_MANUAL.get(d, {}), **man_d_cur.get(d, {})}).get(k, 0)) for k in KPI_MANUAL_COLS}} for d in daftar_dealer])
        cfg_man = {'DEALER': st.column_config.TextColumn('DEALER', width='medium')}
        for k in KPI_MANUAL_COLS: cfg_man[k] = st.column_config.NumberColumn(k, format='%.2f' if ('MYM' in k or 'GRT' in k) else '%d')
        ed_man = st.data_editor(df_man, hide_index=True, use_container_width=True, disabled=['DEALER'], column_config=cfg_man, key='editor_kpi_manual')
        c_o1, c_o2, c_o3 = st.columns([1, 1, 2])
        with c_o1: ops_t_in = st.number_input('T. OPS (Operational Supervisor)', value=float(safe_float(man_ops_cur.get('T. OPS'))), step=1.0, key='inp_t_ops')
        with c_o2: ops_r_in = st.number_input('R. OPS (Operational Supervisor)', value=float(safe_float(man_ops_cur.get('R. OPS'))), step=1.0, key='inp_r_ops')
        with c_o3:
            st.write('')
            if st.button('💾 Simpan Input Manual', type='primary', key='btn_save_kpi_manual'):
                cur_m = load_json(KPI_MANUAL_FILE)
                cur_m['dealers'] = {str(r_m['DEALER']): {k: safe_float(r_m[k]) for k in KPI_MANUAL_COLS} for _, r_m in ed_man.iterrows()}
                cur_m['ops'] = {'T. OPS': float(ops_t_in), 'R. OPS': float(ops_r_in)}
                save_json(cur_m, KPI_MANUAL_FILE)
                st.success('✅ Input manual KPI berhasil disimpan!')
                st.rerun()

    inputs_dash = get_dashboard_kpi_inputs()
    rows_sel, s_sel, p_sel, t_sel = kpi_score_entity(inputs_dash[sel_dash], scheme_dash)
    dealer_tbl = []
    for d_k in daftar_dealer:
        r_k, s_k, p_k, t_k = kpi_score_entity(inputs_dash[d_k], scheme_dash)
        dealer_tbl.append({'name': d_k, 'code': KPI_DEALER_CODES.get(d_k, '-'), 'point': t_k, 'grade': kpi_grade(t_k), 'sales': s_k, 'process': p_k, 'achv': {c: r_k[c]['achv'] for c in KPI_CODES}, 'note': kpi_note(p_k)})
    dealer_tbl.sort(key=lambda x_d: (-x_d['point'], -x_d['process']))
    date_lbl_dash = f'{tanggal_akhir} {DASH_ABBR.get(current_month, "")} {current_year}'
    date_short_dash = f'{tanggal_akhir:02d}/{current_month:02d}/{current_year}'
    note_dash = ('Bobot skema MDS (sesuai gambar): OPS level supervisor, nilainya sama untuk semua dealer.' if scheme_dash == 'MDS' else 'Bobot skema Dealer (Excel R-MDS): OPR 15 & OPS 0; STU/MAX/CLS memakai 75% poin dasar.') + ' Kolom NOTE: PROCESS ≥45 tanpa potongan, 30–44 = CUT 30%, <30 = CUT 50%. CONCLUTION tampil mulai tgl 27.'
    dash_html = build_dashboard_kpi_html(sel_dash, date_lbl_dash, date_short_dash, rows_sel, s_sel, p_sel, t_sel, dealer_tbl, sel_dash, tanggal_akhir >= 27, note_dash)
    st.markdown(dash_html, unsafe_allow_html=True)
    dash_html_cap = build_dashboard_kpi_html(sel_dash, date_lbl_dash, date_short_dash, rows_sel, s_sel, p_sel, t_sel, dealer_tbl, sel_dash, tanggal_akhir >= 27, note_dash, with_rank=False)   # copy/download HD: dashboard saja, tanpa tabel ranking
    cap_dash = f'<div style="width:1320px;background-color:#0E1117;padding:18px 20px;box-sizing:border-box;">{dash_html_cap}</div>'
    render_table_image_buttons('cap_dashboard_kpi', f'Dashboard_KPI_{sel_dash}_Tgl_{tanggal_akhir}', standalone_html=cap_dash)

    koor_focus, cnt_focus = get_focus_entities()
    st.divider()
    render_entity_kpi_dashboard('koor', koor_focus, f'DASHBOARD KPI KOORDINATOR (1 s.d {tanggal_akhir} {periode_bulan_tahun})', 'sel_dash_kpi_koor', date_lbl_dash, date_short_dash)
    st.divider()
    render_entity_kpi_dashboard('counter', cnt_focus, f'DASHBOARD KPI COUNTER (1 s.d {tanggal_akhir} {periode_bulan_tahun})', 'sel_dash_kpi_counter', date_lbl_dash, date_short_dash)

elif menu_pilihan == '🏢 STU':
    bsg_render_running_text()
    col_sel_d, col_sel_k, col_sel_s = st.columns(3)
    with col_sel_d: dlr_pil = st.selectbox('🏢 Pilih Dealer untuk Review:', ['TOTAL JEMBER BONDOWOSO'] + daftar_dealer, index=0)
    _ppt_dealer = dlr_pil   # dealer murni dari dropdown (sebelum label ditambah koordinator) → dasar isi PPT
    is_total = (dlr_pil == 'TOTAL JEMBER BONDOWOSO')
    dlr_list = daftar_dealer if is_total else [dlr_pil]

    # --- Dropdown ke-2: filter Koordinator / Counter (default ALL = seluruh tim dealer) ---
    _opt_ent = set()
    if not is_total:
        for _ds, _dc, _kc in [(df_faktur, 'Nama Dealer', 'Koor'), (df_nota, ('Nama Dealer' if (df_nota is not None and 'Nama Dealer' in df_nota.columns) else 'Dealer'), 'Koor'),
                              (df_database, 'Nama_Dealer_Clean', 'Koor'), (df_spk, 'Nama_Dealer_SPK', 'Koor_SPK')]:
            if _ds is not None and not _ds.empty and _dc in _ds.columns and _kc in _ds.columns:
                _opt_ent.update(clean_str(v) for v in _ds.loc[_ds[_dc].isin(dlr_list), _kc].dropna().unique() if clean_str(v) and clean_str(v) not in ['NAN', 'NONE', 'LAIN-LAIN'])
    _opt_ent = sorted([e for e in _opt_ent if not e.startswith('COUNTER')]) + sorted([e for e in _opt_ent if e.startswith('COUNTER')])
    _lbl_all = 'ALL (SELURUH DEALER)' if is_total else f'ALL (TIM {dlr_pil})'
    with col_sel_k:
        _sel_ent = st.selectbox('👤 Pilih Koordinator / Counter:', [_lbl_all] + _opt_ent, index=0, key=f'sel_koor_stu_{dlr_pil}', disabled=is_total)
    koor_pil = None if (is_total or _sel_ent == _lbl_all) else _sel_ent

    # --- Dropdown ke-3: filter Salesman (hanya salesman di bawah koordinator terpilih; bila ALL = seluruh salesman dealer) ---
    _opt_sm = set()
    if not is_total:
        for _ds, _dc, _kc in [(df_faktur, 'Nama Dealer', 'Koor'), (df_nota, ('Nama Dealer' if (df_nota is not None and 'Nama Dealer' in df_nota.columns) else 'Dealer'), 'Koor'),
                              (df_database, 'Nama_Dealer_Clean', 'Koor'), (df_spk, 'Nama_Dealer_SPK', 'Koor_SPK')]:
            if _ds is not None and not _ds.empty and _dc in _ds.columns and 'Salesman_Clean' in _ds.columns:
                _msk = _ds[_dc].isin(dlr_list)
                if koor_pil and _kc in _ds.columns: _msk = _msk & (_ds[_kc] == koor_pil)
                _opt_sm.update(clean_str(v) for v in _ds.loc[_msk, 'Salesman_Clean'].dropna().unique() if clean_str(v) and clean_str(v) not in ['NAN', 'NONE', 'LAIN-LAIN', 'TOTAL', 'JUMLAH', '-'])
    _opt_sm = sorted(_opt_sm)
    _lbl_sm_all = 'ALL SALESMAN' if is_total else (f'ALL (SALESMAN {koor_pil})' if koor_pil else f'ALL (SALESMAN {dlr_pil})')
    with col_sel_s:
        _sel_sm = st.selectbox('🧑‍💼 Pilih Salesman:', [_lbl_sm_all] + _opt_sm, index=0, key=f'sel_sales_stu_{dlr_pil}_{koor_pil}', disabled=(is_total or not _opt_sm))
    sales_pil = None if (is_total or _sel_sm == _lbl_sm_all) else _sel_sm

    scope_txt = 'Seluruh Dealer JEMBER BONDOWOSO' if is_total else (f'Khusus Dealer {dlr_pil}' if not koor_pil else f'Khusus {koor_pil} (Dealer {dlr_pil})')
    if sales_pil: scope_txt = f'Khusus Salesman {sales_pil}' + (f' (Koordinator {koor_pil}, Dealer {dlr_pil})' if koor_pil else f' (Dealer {dlr_pil})')
    if koor_pil: dlr_pil = f'{dlr_pil} — {koor_pil}'   # label judul tabel, key widget & cache AI ikut berubah per koordinator/counter
    if sales_pil: dlr_pil = f'{dlr_pil} — {sales_pil}'   # idem untuk salesman terpilih
    main_hdr_text = f'DETAIL DEALER — {dlr_pil} (1 s.d {tanggal_akhir} {periode_bulan_tahun})'
    sub_hdr_text = f'Rekapitulasi Terpadu Penjualan, Proses Kerja, SPK Leasing, Penjualan Per Type & Detail Salesman {scope_txt} (Acuan Ideal: DB Quality 30% | STU% 10% | ACC% 50%)'
    _hd_l, _hd_r = st.columns([4, 1.25])
    with _hd_l: st.markdown(f'<div class="main-title">{main_hdr_text}</div><div class="sub-title">{sub_hdr_text}</div>', unsafe_allow_html=True)
    with _hd_r:
        st.markdown('<div style="height:14px"></div>', unsafe_allow_html=True)
        _ppt_click = st.button('📑 CREATE PPT', key='btn_create_ppt', type='primary', use_container_width=True, help=f'Buat PowerPoint dari {_TPL_NAME} untuk: {_ppt_dealer} (sesuai dropdown dealer). Isi mengikuti hasil analisa dashboard.')
        _ppt_slot = st.container()
        _ppt_tpl_up = st.file_uploader(f'Template PPT tidak ditemukan di folder app — unggah {_TPL_NAME}', type=['pptx'], key='ppt_tpl_up') if _tp_find_template() is None else None
    with st.expander('🖼️ Gambar dashboard untuk PPT (opsional) — agar tampilan PPT sama persis dengan dashboard'):
        st.caption('1) Klik **Download SEMUA gambar** → terunduh 1 file ZIP berisi gambar panel dealer yang sedang tampil + Dashboard KPI. '
                   '2) Untuk slide 9–18 (koordinator & counter terlemah): Dashboard KPI keduanya sudah ikut di ZIP pertama. Agar panel penjualan/salesman/aktifitas keduanya juga ada, pilih koordinator terlemah lalu counter terlemah di dropdown **Koordinator / Counter**, klik tombol yang sama lagi (nama keduanya disebut di bawah). '
                   '3) Unggah semua file ZIP di kotak bawah, lalu klik CREATE PPT. Slide yang gambarnya lengkap memakai gambar dashboard; sisanya memakai tabel otomatis.')
        _ppt_dl_ph = st.empty()
        _ppt_dl_ph.caption('⏳ Tombol Download SEMUA gambar disiapkan setelah seluruh tabel selesai dimuat. Jika tombol tidak muncul setelah halaman selesai, ada error di bagian bawah halaman — scroll ke bawah dan kirim pesan errornya.')
        _ppt_imgs_up = st.file_uploader('Unggah file ZIP (atau JPG/PNG)', type=['zip', 'jpg', 'jpeg', 'png'], accept_multiple_files=True, key='ppt_imgs_up')
    if _ppt_click:
        with _ppt_slot:
            with st.spinner(f'Menyusun PowerPoint {_ppt_dealer}…'):
                try:
                    _ppt_bytes, _ppt_name, _ppt_warn = bsg_ppt_for_dealer(_ppt_dealer, _ppt_tpl_up.getvalue() if _ppt_tpl_up is not None else None, _ppt_collect_images(_ppt_imgs_up))
                    st.session_state['ppt_file'] = (_ppt_name, _ppt_bytes, _ppt_dealer, _ppt_warn)
                except ImportError as _e_imp:
                    st.session_state['ppt_file'] = None
                    st.error(f'Library untuk membuat PowerPoint belum bisa dimuat: **{_e_imp}**. Pasang dengan: `python -m pip install python-pptx` (di Python yang sama dengan yang menjalankan Streamlit), lalu restart aplikasi. Di Streamlit Cloud: tambahkan `python-pptx` ke requirements.txt lalu Reboot app.')
                except FileNotFoundError as _e_nf:
                    st.session_state['ppt_file'] = None
                    st.error(str(_e_nf))
                except Exception as _e_ppt:
                    st.session_state['ppt_file'] = None
                    st.error(f'Gagal membuat PPT: {_e_ppt}')
    _ppt_saved = st.session_state.get('ppt_file')
    if _ppt_saved and len(_ppt_saved) > 2 and _ppt_saved[2] == _ppt_dealer:   # tampil hanya bila PPT dibuat untuk dealer yang sedang dipilih
        with _ppt_slot:
            st.download_button('⬇️ Unduh PPT', data=_ppt_saved[1], file_name=_ppt_saved[0], mime='application/vnd.openxmlformats-officedocument.presentationml.presentation', use_container_width=True, key='dl_ppt', on_click=lambda: st.session_state.update({'ppt_file': None}))
            st.caption(_ppt_saved[0])
            for _w in (_ppt_saved[3] if len(_ppt_saved) > 3 else []): (st.info if _w.startswith('ℹ️') else st.warning)(_w)

    saved_td, saved_tk = load_json(TARGET_DEALER_FILE), load_json(TARGET_KOOR_FILE)
    cd_nota = 'Nama Dealer' if (df_nota is not None and 'Nama Dealer' in df_nota.columns) else 'Dealer'

    f_all = df_faktur[df_faktur['Nama Dealer'].isin(dlr_list)] if (df_faktur is not None and 'Nama Dealer' in df_faktur.columns) else pd.DataFrame()
    f_c = df_fak_curr_g[df_fak_curr_g['Nama Dealer'].isin(dlr_list)] if not df_fak_curr_g.empty else pd.DataFrame()
    f_l = df_fak_lm_g[df_fak_lm_g['Nama Dealer'].isin(dlr_list)] if not df_fak_lm_g.empty else pd.DataFrame()
    n_all = df_nota[df_nota[cd_nota].isin(dlr_list)] if (df_nota is not None and cd_nota in df_nota.columns) else pd.DataFrame()
    n_c = df_nota_curr_g[df_nota_curr_g[cd_nota].isin(dlr_list)] if not df_nota_curr_g.empty else pd.DataFrame()
    n_t = df_nota_tod_g[df_nota_tod_g[cd_nota].isin(dlr_list)] if not df_nota_tod_g.empty else pd.DataFrame()
    db_all = df_database[df_database['Nama_Dealer_Clean'].isin(dlr_list)] if (df_database is not None and 'Nama_Dealer_Clean' in df_database.columns) else pd.DataFrame()
    db_c = df_db_curr_g[df_db_curr_g['Nama_Dealer_Clean'].isin(dlr_list)] if not df_db_curr_g.empty else pd.DataFrame()
    s_all = df_spk[df_spk['Nama_Dealer_SPK'].isin(dlr_list)] if (df_spk is not None and 'Nama_Dealer_SPK' in df_spk.columns) else pd.DataFrame()
    s_c = df_spk_curr_g[df_spk_curr_g['Nama_Dealer_SPK'].isin(dlr_list)] if not df_spk_curr_g.empty else pd.DataFrame()
    s_l = df_spk_lm_g[df_spk_lm_g['Nama_Dealer_SPK'].isin(dlr_list)] if not df_spk_lm_g.empty else pd.DataFrame()

    if koor_pil:   # filter seluruh data tab STU hanya untuk Koordinator / Counter terpilih
        def _kfilt(df_x, col_x): return df_x[df_x[col_x] == koor_pil] if (df_x is not None and not df_x.empty and col_x in df_x.columns) else (df_x if df_x is None else df_x.iloc[0:0])
        f_all, f_c, f_l = _kfilt(f_all, 'Koor'), _kfilt(f_c, 'Koor'), _kfilt(f_l, 'Koor')
        n_all, n_c, n_t = _kfilt(n_all, 'Koor'), _kfilt(n_c, 'Koor'), _kfilt(n_t, 'Koor')
        db_all, db_c = _kfilt(db_all, 'Koor'), _kfilt(db_c, 'Koor')
        s_all, s_c, s_l = _kfilt(s_all, 'Koor_SPK'), _kfilt(s_c, 'Koor_SPK'), _kfilt(s_l, 'Koor_SPK')
    if sales_pil:   # filter lanjutan: seluruh data tab STU hanya untuk salesman terpilih (target = target salesman)
        def _sfilt(df_x): return df_x[df_x['Salesman_Clean'] == sales_pil] if (df_x is not None and not df_x.empty and 'Salesman_Clean' in df_x.columns) else (df_x if df_x is None else df_x.iloc[0:0])
        f_all, f_c, f_l = _sfilt(f_all), _sfilt(f_c), _sfilt(f_l)
        n_all, n_c, n_t = _sfilt(n_all), _sfilt(n_c), _sfilt(n_t)
        db_all, db_c = _sfilt(db_all), _sfilt(db_c)
        s_all, s_c, s_l = _sfilt(s_all), _sfilt(s_c), _sfilt(s_l)

    ent_dlr = set()
    for ds, ck in [(f_all, 'Koor'), (n_all, 'Koor'), (db_all, 'Koor'), (s_all, 'Koor_SPK')]:
        if not ds.empty and ck in ds.columns: ent_dlr.update(clean_str(v) for v in ds[ck].dropna().unique() if clean_str(v) and clean_str(v) not in ['NAN', 'NONE', 'LAIN-LAIN'])
    gab_ent, lbl_tot, kol_nm = sorted(list(ent_dlr)), ('TOTAL JEMBER BONDOWOSO' if is_total else f'TOTAL ({dlr_pil})'), 'KOOR / COUNTER / DEALER'

    r_tot_p = calc_sales_row(lbl_tot, kol_nm, (get_auto_sales_target(sales_pil) if sales_pil else get_koor_or_counter_target(koor_pil, saved_tk) if koor_pil else sum(safe_int(saved_td.get(d_t, 0)) for d_t in dlr_list)), f_c, f_l, n_c, n_t)
    r_tot_pr = calc_proses_row(lbl_tot, kol_nm, db_c, len(f_c), n_c)
    r_tot_s = calc_leasing_row(lbl_tot, s_c, s_l, kol_nm)

    rp_all = [calc_sales_row(e, kol_nm, (get_auto_sales_target(sales_pil) if sales_pil else get_koor_or_counter_target(e, saved_tk)), f_c[f_c['Koor'] == e] if not f_c.empty else pd.DataFrame(), f_l[f_l['Koor'] == e] if not f_l.empty else pd.DataFrame(), n_c[n_c['Koor'] == e] if not n_c.empty else pd.DataFrame(), n_t[n_t['Koor'] == e] if not n_t.empty else pd.DataFrame()) for e in gab_ent]
    df_p_it = pd.DataFrame(rp_all).sort_values(by=['_pct', 'STU'], ascending=[True, True]).reset_index(drop=True) if rp_all else pd.DataFrame()

    rpr_all = [calc_proses_row(e, kol_nm, db_c[db_c['Koor'] == e] if not db_c.empty else pd.DataFrame(), len(f_c[f_c['Koor'] == e]) if not f_c.empty else 0, (n_c[n_c['Koor'] == e] if not n_c.empty else pd.DataFrame())) for e in gab_ent]
    df_pr_it = pd.DataFrame(rpr_all).sort_values(by=['_stu_pct', 'STU'], ascending=[True, True]).reset_index(drop=True) if rpr_all else pd.DataFrame()

    rs_all = [calc_leasing_row(e, s_c[s_c['Koor_SPK'] == e] if not s_c.empty else pd.DataFrame(), s_l[s_l['Koor_SPK'] == e] if not s_l.empty else pd.DataFrame(), kol_nm) for e in gab_ent]
    df_s_it = pd.DataFrame(rs_all).sort_values(by=['_acc_val', '_spk_val'], ascending=[True, True]).reset_index(drop=True) if rs_all else pd.DataFrame()
    # Frame tampilan Tabel 1-3: jika TOTAL JEMBER BONDOWOSO, baris = per Dealer (kolom sama dengan tabel Koordinator)
    if is_total:
        def _dsub(df_x, col_x, d_x): return df_x[df_x[col_x] == d_x] if (df_x is not None and not df_x.empty and col_x in df_x.columns) else pd.DataFrame()
        rp_d = [calc_sales_row(d, kol_nm, safe_int(saved_td.get(d, 0)), _dsub(f_c, 'Nama Dealer', d), _dsub(f_l, 'Nama Dealer', d), _dsub(n_c, cd_nota, d), _dsub(n_t, cd_nota, d)) for d in daftar_dealer]
        df_p_disp = pd.DataFrame(rp_d).sort_values(by=['_pct', 'STU'], ascending=[True, True]).reset_index(drop=True)
        rpr_d = [calc_proses_row(d, kol_nm, _dsub(db_c, 'Nama_Dealer_Clean', d), len(_dsub(f_c, 'Nama Dealer', d)), _dsub(n_c, cd_nota, d)) for d in daftar_dealer]
        df_pr_disp = pd.DataFrame(rpr_d).sort_values(by=['_stu_pct', 'STU'], ascending=[True, True]).reset_index(drop=True)
        rs_d = [calc_leasing_row(d, _dsub(s_c, 'Nama_Dealer_SPK', d), _dsub(s_l, 'Nama_Dealer_SPK', d), kol_nm) for d in daftar_dealer]
        df_s_disp = pd.DataFrame(rs_d).sort_values(by=['_acc_val', '_spk_val'], ascending=[True, True]).reset_index(drop=True)
    else:
        df_p_disp, df_pr_disp, df_s_disp = df_p_it, df_pr_it, df_s_it

    # Tabel 2 (TOTAL JEMBER BONDOWOSO): tambah kolom SPK & ACC% setelah DB Q
    def _spk_acc_vals(df_x):
        n_spk = len(df_x) if df_x is not None else 0
        ok_x = (df_x is not None and not df_x.empty and 'Status_SPK_Clean' in df_x.columns)
        a_x = int((df_x['Status_SPK_Clean'] == 'DISETUJUI').sum()) if ok_x else 0
        r_x = int((df_x['Status_SPK_Clean'] == 'DITOLAK').sum()) if ok_x else 0
        acc_x = (a_x / (a_x + r_x) * 100.0) if (a_x + r_x) > 0 else 0.0
        return n_spk, f'{acc_x:.2f}%'
    r_tot_pr_disp = dict(r_tot_pr)
    if is_total:
        _pr_spk = {d_x: _spk_acc_vals(_dsub(s_c, 'Nama_Dealer_SPK', d_x)) for d_x in daftar_dealer}
        df_pr_disp = df_pr_disp.copy()
        df_pr_disp['SPK'] = df_pr_disp[kol_nm].map(lambda x_n: _pr_spk.get(x_n, (0, '0.00%'))[0])
        df_pr_disp['ACC%'] = df_pr_disp[kol_nm].map(lambda x_n: _pr_spk.get(x_n, (0, '0.00%'))[1])
        r_tot_pr_disp['SPK'], r_tot_pr_disp['ACC%'] = _spk_acc_vals(s_c)
        _lm_d = {d_x: len(_dsub(s_l, 'Nama_Dealer_SPK', d_x)) for d_x in daftar_dealer}
        _grw = lambda a_c, a_l: f'{(a_c / a_l * 100.0):.2f}%' if a_l > 0 else '0.00%'
        df_pr_disp['SPK LM'] = df_pr_disp[kol_nm].map(lambda x_n: _lm_d.get(x_n, 0))
        df_pr_disp['GROWTH'] = df_pr_disp[kol_nm].map(lambda x_n: _grw(_pr_spk.get(x_n, (0, ''))[0], _lm_d.get(x_n, 0)))
        r_tot_pr_disp['SPK LM'], r_tot_pr_disp['GROWTH'] = len(s_l), _grw(len(s_c), len(s_l))

    sm_set = set()
    for ds in [f_c, n_c, db_c, s_c]:
        if not ds.empty and 'Salesman_Clean' in ds.columns: sm_set.update(clean_str(v) for v in ds['Salesman_Clean'].dropna().unique() if clean_str(v) and clean_str(v) not in ['NAN', 'NONE', 'TOTAL'])
    if not sm_set:
        for ds in [f_all, n_all, db_all, s_all]:
            if not ds.empty and 'Salesman_Clean' in ds.columns: sm_set.update(clean_str(v) for v in ds['Salesman_Clean'].dropna().unique() if clean_str(v) and clean_str(v) not in ['NAN', 'NONE', 'TOTAL'])

    _act_sm = set()   # acuan file Database: salesman yang sudah tidak ada di Database (resign) tidak ditampilkan walau masih ada di Master Employee / Faktur / SPK
    if df_database is not None and not df_database.empty and 'Salesman_Clean' in df_database.columns:
        _act_sm = set(clean_str(v) for v in df_database['Salesman_Clean'].dropna().unique() if clean_str(v))
    if _act_sm: sm_set = {m for m in sm_set if m in _act_sm}

    r_sm, t_sa, t_sr = [], 0, 0
    for sm in sorted(list(sm_set)):
        targ, stu = get_auto_sales_target(sm), (len(f_c[f_c['Salesman_Clean'] == sm]) if (not f_c.empty and 'Salesman_Clean' in f_c.columns) else 0)
        pct_sm = (stu / targ * 100.0) if targ > 0 else 0.0
        sdb = db_c[db_c['Salesman_Clean'] == sm] if (not db_c.empty and 'Salesman_Clean' in db_c.columns) else pd.DataFrame()
        db_v, hot_v = len(sdb), (len(sdb[sdb['Data_Type_Clean'] == 'DEAL PROSPECT']) if (not sdb.empty and 'Data_Type_Clean' in sdb.columns) else 0)
        sspk = s_c[s_c['Salesman_Clean'] == sm] if (not s_c.empty and 'Salesman_Clean' in s_c.columns) else pd.DataFrame()
        spk_v, acc_v, rej_v = len(sspk), (len(sspk[sspk['Status_SPK_Clean'] == 'DISETUJUI']) if (not sspk.empty and 'Status_SPK_Clean' in sspk.columns) else 0), (len(sspk[sspk['Status_SPK_Clean'] == 'DITOLAK']) if (not sspk.empty and 'Status_SPK_Clean' in sspk.columns) else 0)
        t_sa += acc_v; t_sr += rej_v
        dbq_v, stup_v, acc_p_v = (hot_v / db_v * 100.0 if db_v > 0 else 0.0), (stu / db_v * 100.0 if db_v > 0 else 0.0), (acc_v / (acc_v + rej_v) * 100.0 if (acc_v + rej_v) > 0 else 0.0)
        lm_v = len(f_l[f_l['Salesman_Clean'] == sm]) if (not f_l.empty and 'Salesman_Clean' in f_l.columns) else 0
        r_sm.append({'NAMA SALESMAN': sm, 'TARGET': targ, 'STU': stu, col_name_pct: f'{pct_sm:.2f}%', 'AVG DB': f'{(db_v/hari_aktif):.2f}', 'DB': db_v, 'HOT': hot_v, 'DB Q': f'{dbq_v:.2f}%', 'SPK': spk_v, 'ACC%': f'{acc_p_v:.2f}%', 'STU%': f'{stup_v:.2f}%', 'LM': lm_v, '_lm_val': lm_v, '_gap_lm': stu - lm_v, '_stu_val': stu, '_db_val': db_v, '_pct_val': pct_sm, '_dbq_val': dbq_v, '_stup_val': stup_v, '_acc_val': acc_p_v})

    forc_pct_dlr = (r_tot_p['FORC'] / r_tot_p['TARGET'] * 100.0) if r_tot_p['TARGET'] > 0 else 0.0
    nota_tot_now = safe_int(r_tot_p['NOTA'])
    _nm_tot = nota_metrics(n_c)
    div_7, div_3, n7_cnt, n3_cnt = _nm_tot['d7'], _nm_tot['d3'], _nm_tot['n7'], _nm_tot['n3']
    avg_nota_v, avg7_v, avg3_v = _nm_tot['avg'], _nm_tot['avg7'], _nm_tot['avg3']
    sisa_hari_bln = SISA_HARI_KERJA
    forc_nota_v, forc7_v, forc3_v = _nm_tot['f_nota'], _nm_tot['f7'], _nm_tot['f3']

    C_GRN, C_YEL, C_RED, C_CYN = '#00FF66', '#FFEA00', '#FF4D8D', '#00E5FF'
    def _stat_col(v, good, warn): return C_GRN if v >= good else (C_YEL if v >= warn else C_RED)
    def _kpi_card(icon, label, value, sub, color):
        _pg = _KPI_PROG.get(label, (None, None))
        return bsg_card(icon, label, value, sub, color, _pg[0], _pg[1])
    def _kpi_grid(cards): return '<div style="display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px;margin-bottom:12px;">' + ''.join(cards) + '</div>'
    def _unit(txt): return f'<span style="font-size:13px;font-weight:600;color:#94A3B8;"> {txt}</span>'
    def _trend(v, ref, ref_lbl='AVG NOTA'):
        if v > ref: return f'<span style="color:{C_GRN};font-weight:800;">▲ naik</span> vs {ref_lbl}'
        if v < ref: return f'<span style="color:{C_RED};font-weight:800;">▼ turun</span> vs {ref_lbl}'
        return f'▬ sama dengan {ref_lbl}'
    def _avg_fc(avg_v, forc_v):   # tampilan ringkas: angka AVG ➜ angka forecast closing (warna forecast mengikuti % dari target)
        tg = safe_int(r_tot_p['TARGET'])
        col_f = _stat_col(forc_v / tg * 100.0, 100.0, 85.0) if tg > 0 else C_CYN
        return f'{avg_v:.2f} <span style="font-size:20px;font-weight:700;color:#94A3B8;">➜</span> <span style="color:{col_f};">{forc_v:,}</span>'

    _dbq_n, _stup_n, _acc_n = safe_float(r_tot_pr['DB Q']), safe_float(r_tot_pr['STU%']), safe_float(r_tot_s['ACC%'])
    _grw_n = safe_float(r_tot_s['GROWTH'])
    _KPI_PROG = {'Target vs Result': (min(r_tot_p['_pct'], 100.0), target_persen_harian), 'DB Quality': bsg_ach_bar(_dbq_n, 30.0), 'STU Ratio': bsg_ach_bar(_stup_n, 10.0),
                 'SPK Growth': bsg_ach_bar(_grw_n, 100.0), 'Approval': bsg_ach_bar(_acc_n, 50.0), 'Forecast Closing': (min(forc_pct_dlr, 100.0), None)}
    _db_tod = db_c[db_c['hari'] == tanggal_akhir] if (db_c is not None and not db_c.empty and 'hari' in db_c.columns) else pd.DataFrame()   # database prospek tanggal terakhir (hari ini)
    _db_tod_n = len(_db_tod)
    _hot_tod_n = int((_db_tod['Data_Type_Clean'] == 'DEAL PROSPECT').sum()) if (not _db_tod.empty and 'Data_Type_Clean' in _db_tod.columns) else 0
    _dbq_tod = (_hot_tod_n / _db_tod_n * 100.0) if _db_tod_n > 0 else 0.0
    _nota_tod_n = safe_int(r_tot_p['TODAY'])
    _db_tot_n, _hot_tot_n = safe_int(r_tot_pr['DB']), safe_int(r_tot_pr['HOT'])
    kpi_cards_row1 = [
        _kpi_card('🎯', 'Target vs Result', f"{r_tot_p['TARGET']:,} / {r_tot_p['STU']:,} <span style='color:#FFEA00;'>!</span> <span style='color:#FFEA00;'>{r_tot_p['IDEAL']:,}</span>", f"Result <b>{r_tot_p['_pct']:.1f}%</b> · Ideal <b>{target_persen_harian:.0f}%</b>", C_GRN if r_tot_p['STU'] >= r_tot_p['IDEAL'] else (C_YEL if r_tot_p['STU'] >= 0.85 * r_tot_p['IDEAL'] else C_RED)),
        _kpi_card('✨', 'DB Quality', f"{r_tot_pr['DB Q']} <span style='color:#FFEA00;'>!</span> <span style='color:#FFEA00;'>30%</span>", f"{r_tot_pr['DB']:,} DB ➜ {r_tot_pr['HOT']:,} HOT", _stat_col(_dbq_n, 30.0, 20.0)),
        _kpi_card('📈', 'STU Ratio', f"{r_tot_pr['STU%']} <span style='color:#FFEA00;'>!</span> <span style='color:#FFEA00;'>10%</span>", f"{r_tot_pr['DB']:,} DB ➜ {r_tot_pr['STU']:,} STU", _stat_col(_stup_n, 10.0, 7.0)),
        _kpi_card('📊', 'AVG NOTA', _avg_fc(avg_nota_v, forc_nota_v), f'{nota_tot_now:,} nota ÷ {hari_aktif} hari kerja', C_CYN),
        _kpi_card('⚡', 'AVG 7', _avg_fc(avg7_v, forc7_v), f'{n7_cnt:,} nota ÷ {div_7} hari kerja · {_trend(avg7_v, avg_nota_v)}', C_GRN if avg7_v >= avg_nota_v else C_RED),
        _kpi_card('🔥', 'AVG 3', _avg_fc(avg3_v, forc3_v), f'{n3_cnt:,} nota ÷ {div_3} hari kerja · {_trend(avg3_v, avg7_v, "AVG 7")}', C_GRN if avg3_v >= avg7_v else C_RED),
    ]
    kpi_cards_row2 = [
        _kpi_card('🚀', 'SPK Growth', r_tot_s['GROWTH'], f"SPK {r_tot_s['SPK']:,} vs LM {r_tot_s['SPK LM']:,}", _stat_col(_grw_n, 100.0, 80.0)),
        _kpi_card('✅', 'Approval', f"{r_tot_s['ACC%']} <span style='color:#FFEA00;'>!</span> <span style='color:#FFEA00;'>50%</span>", f"SPK terbuka <b>{safe_int(r_tot_s['TERBUKA']):,}</b> · SPK {safe_int(r_tot_s['SPK']):,} vs LM {safe_int(r_tot_s['SPK LM']):,}", _stat_col(_acc_n, 50.0, 40.0)),
        _kpi_card('🏁', 'Forecast Closing', f"{r_tot_p['FORC']:,}" + _unit('unit'), f"<b>{forc_pct_dlr:.1f}%</b> dari target", _stat_col(forc_pct_dlr, 100.0, 85.0)),
        _kpi_card('🗂️', 'DB Hari Ini', f'{_db_tod_n:,}' + _unit('DB'), f"tgl {tanggal_akhir} · {(_db_tod_n / _db_tot_n * 100.0 if _db_tot_n > 0 else 0.0):.1f}% dari total DB<br>Total DB s.d hari ini {_db_tot_n:,}", C_CYN),
        _kpi_card('🔥', 'HOT Hari Ini', f'{_hot_tod_n:,}' + _unit('HOT'), f"tgl {tanggal_akhir} · DB Q hari ini <b>{_dbq_tod:.1f}%</b><br>Total HOT s.d hari ini {_hot_tot_n:,} · Ideal 30%", _stat_col(_dbq_tod, 30.0, 20.0) if _db_tod_n > 0 else C_CYN),
        _kpi_card('🧾', 'Nota Hari Ini', f'{_nota_tod_n:,}' + _unit('nota'), f"tgl {tanggal_akhir}<br>{_trend(_nota_tod_n, avg_nota_v)}", C_GRN if _nota_tod_n >= avg_nota_v else C_RED),
    ]
    kpi_cards_html = _kpi_grid(kpi_cards_row1) + _kpi_grid(kpi_cards_row2)
    st.markdown(kpi_cards_html, unsafe_allow_html=True)

    df_koor_only, df_cnt_only = pd.DataFrame(), pd.DataFrame()
    if not df_p_it.empty:
        df_koor_only = df_p_it[~df_p_it[kol_nm].str.upper().str.startswith('COUNTER')].copy()
        df_cnt_only = df_p_it[df_p_it[kol_nm].str.upper().str.startswith('COUNTER')].sort_values(by=['_pct', 'STU'], ascending=[True, True])

    def _focus_ent_html(df_ent, empty_msg):
        if df_ent is None or df_ent.empty: return empty_msg
        ev_list = []
        for _, rk in df_ent.iterrows():
            knm = str(rk[kol_nm])
            k_pr = df_pr_it[df_pr_it[kol_nm] == knm].iloc[0] if (not df_pr_it.empty and knm in df_pr_it[kol_nm].values) else {}
            k_sp = df_s_it[df_s_it[kol_nm] == knm].iloc[0] if (not df_s_it.empty and knm in df_s_it[kol_nm].values) else {}
            ev_list.append({'nama': knm, 'target': safe_int(rk['TARGET']), 'ideal': safe_int(rk['IDEAL']), 'stu': safe_int(rk['STU']), 'pct': safe_float(rk['_pct']), 'db': safe_int(k_pr.get('DB', 0)), 'dbq_str': k_pr.get('DB Q', '0.00%'), 'stup_str': k_pr.get('STU%', '0.00%'), 'spk_grw_str': k_sp.get('GROWTH', '0.00%'), 'acc_str': k_sp.get('ACC%', '0.00%'), 'dbq_val': safe_float(k_pr.get('DB Q', 0.0)), 'stup_val': safe_float(k_pr.get('STU%', 0.0)), 'acc_val': safe_float(k_sp.get('_acc_val', 0.0))})
        ev_list.sort(key=lambda x: (0 if x['target'] > 0 else 1, x['pct'], x['stu'], x['stup_val'], x['dbq_val'], x['acc_val']))
        fk = ev_list[0]
        return f"<span style='color:#FFEA00;font-weight:800;'>{fk['nama']}</span> — Target vs Result: <b>{fk['target']}/{fk['stu']} ({fk['pct']:.2f}% | Ideal: {fk['ideal']})</b> | DB Quality: <b>{fk['dbq_str']}</b> | STU Ratio: <b>{fk['stup_str']}</b> | SPK Growth: <b>{fk['spk_grw_str']}</b> | Approval: <b>{fk['acc_str']}</b>"
    focus_koor_html = _focus_ent_html(df_koor_only, "Tidak ada data Koordinator")
    focus_cnt_html = _focus_ent_html(df_cnt_only, "Tidak ada data Counter")

    # ---- Kebutuhan salesman baru & grade salesman (dipakai di notif utama + badge bawah Tabel Detail Salesman) ----
    _mp_tgt_dlr = safe_int(r_tot_p['TARGET'])   # target sesuai dropbox (dealer / TOTAL JEMBER BONDOWOSO / koordinator)
    _mp_tgt_sm = int(sum(safe_int(r_['TARGET']) for r_ in r_sm)) if r_sm else 0   # target total man power yang ada
    _mp_raw = (_mp_tgt_dlr * 1.25 - _mp_tgt_sm) / 2.0   # ((target x 125%) - target total MP) / 2
    _mp_need = int(-(-_mp_raw // 1)) if _mp_raw > 0 else 0   # dibulatkan ke atas
    _mp_ideal = float(target_persen_harian)
    _mp_grn, _mp_red = [], []
    for _r_ in (r_sm or []):
        if bsg_num(_r_.get('TARGET')) <= 0: continue   # grade hanya utk salesman yang punya target (sama dgn grafik)
        _pc_ = bsg_num(_r_.get('_pct_val'))
        if _pc_ >= _mp_ideal: _mp_grn.append(str(_r_['NAMA SALESMAN']))
        elif _pc_ < 0.85 * _mp_ideal: _mp_red.append(str(_r_['NAMA SALESMAN']))
    _mp_calc_txt = f"(Target {_mp_tgt_dlr:,} × 125% − Target total MP {_mp_tgt_sm:,}) ÷ 2 = {max(_mp_raw, 0):.2f}"
    if _mp_need > 0: focus_need_html = f"<span style='color:#FFEA00;font-weight:800;'>{_mp_need} orang</span> salesman baru — {_mp_calc_txt} → dibulatkan ke atas"
    else: focus_need_html = f"<span style='color:#00FF66;font-weight:800;'>Tidak diperlukan</span> penambahan salesman baru — {_mp_calc_txt}"
    _mp_nm_grn = ', '.join(f"<span style='color:{BSG_GRN};font-weight:800;'>{_esc(n_)}</span>" for n_ in _mp_grn)
    _mp_nm_red = ', '.join(f"<span style='color:{BSG_RED};font-weight:800;'>{_esc(n_)}</span>" for n_ in _mp_red)
    _mp_n1 = (f"Dengan kondisi man power saat ini, dibutuhkan <b style='color:{BSG_YEL};'>{_mp_need} orang</b> salesman baru untuk bisa mencapai target <b>{_esc(dlr_pil)}</b> di bulan ini." if _mp_need > 0
              else f"Dengan kondisi man power saat ini, <b style='color:{BSG_GRN};'>tidak diperlukan</b> penambahan salesman baru untuk mencapai target <b>{_esc(dlr_pil)}</b> di bulan ini.")
    _mp_n2 = f"Maksimalkan {_mp_nm_grn} sambil menunggu penambahan salesman baru." if _mp_grn else "Belum ada salesman dengan grade hijau — maksimalkan seluruh salesman yang ada sambil menunggu penambahan salesman baru."
    _mp_n3 = f"Segera lakukan pendampingan dan coaching harian terhadap {_mp_nm_red}." if _mp_red else "Tidak ada salesman dengan grade merah saat ini."
    sm_notif_html = (f'<div style="margin-top:12px;padding:11px 16px;background:linear-gradient(90deg,rgba(255,0,127,0.14) 0%,rgba(14,17,23,0.95) 100%);border-left:4px solid #FF007F;border:1px solid #2D3748;border-left:4px solid #FF007F;border-radius:6px;font-size:13.5px;line-height:1.65;color:#E2E8F0;font-family:\'Source Sans Pro\',\'Segoe UI\',Arial,sans-serif;">'
                     f'<div>📢 <b>1.</b> {_mp_n1}</div><div style="margin-top:4px;">📢 <b>2.</b> {_mp_n2}</div><div style="margin-top:4px;">📢 <b>3.</b> {_mp_n3}</div>'
                     f'<div style="margin-top:6px;font-size:11px;color:#94A3B8;">Perhitungan: {_mp_calc_txt} → {_mp_need} orang · Grade: hijau ≥ ideal {_mp_ideal:.0f}% · merah &lt; 85% dari ideal.</div></div>')

    focus_sm_html = "Tidak ada Salesman aktif (Non-Counter dengan DB >= 2)"
    if r_sm:
        df_sm_all = pd.DataFrame(r_sm)
        df_sm_valid = df_sm_all[(~df_sm_all['NAMA SALESMAN'].apply(is_counter_personnel)) & (df_sm_all['_db_val'] >= 2)].copy()
        if not df_sm_valid.empty:
            df_sm_valid['has_t'] = df_sm_valid['TARGET'].apply(lambda x: 0 if x > 0 else 1)
            df_sm_valid = df_sm_valid.sort_values(by=['has_t', '_pct_val', '_stu_val', '_stup_val', '_dbq_val', '_acc_val', '_db_val'], ascending=[True]*7).reset_index(drop=True)
            fs = df_sm_valid.iloc[0]
            fs_idl = round(safe_int(fs['TARGET']) * persen_desimal)
            focus_sm_html = f"<span style='color:#FFEA00;font-weight:800;'>{fs['NAMA SALESMAN']}</span> — Target vs Result: <b>{fs['TARGET']}/{fs['STU']} ({fs[col_name_pct]} | Ideal: {fs_idl})</b> | DB: <b>{fs['DB']}</b> (DB Quality: <b>{fs['DB Q']}</b>) | STU Ratio: <b>{fs['STU%']}</b> | SPK: <b>{fs['SPK']}</b> (Approval: <b>{fs['ACC%']}</b>)"

    focus_walk_html = 'Tidak ada data Walk In pada periode ini'
    if db_c is not None and not db_c.empty and db_c.shape[1] > 6:
        _wi_all = db_c[db_c.iloc[:, 6].apply(alnum_str) == 'WALKIN']
        if not _wi_all.empty:   # seluruh Walk In (tidak per koordinator / counter)
            _t_n, _t_s = len(_wi_all), (int((_wi_all['STU'] == 'STU').sum()) if 'STU' in _wi_all.columns else 0)
            focus_walk_html = f"DB Walk In: <b>{_t_n}</b> | STU Walk In: <b>{_t_s}</b> | Need FU: <span style='color:#FFEA00;font-weight:800;'>{_t_n - _t_s}</span> <span style='color:#94A3B8;'>(DB Walk In − STU Walk In)</span>"

    st.markdown(f"""<div class="focus-alert-box"><div style="margin-top:4px;">🚶 <b>Focus Walk In :</b> {focus_walk_html}</div><div style="margin-top:4px;">🎯 <b>Focus Koordinator :</b> {focus_koor_html}</div><div>🏪 <b>Focus Counter :</b> {focus_cnt_html}</div><div style="margin-top:4px;">👤 <b>Focus Salesman :</b> {focus_sm_html}</div><div style="margin-top:4px;">🧑‍💼 <b>MP Health : Kebutuhan Salesman Baru</b> → {focus_need_html}</div></div>""", unsafe_allow_html=True)

    _tbl_counter = [0]
    def nxt_no():
        _tbl_counter[0] += 1
        return _tbl_counter[0]

    _tab_names = ['💼 Penjualan & Proses', '👤 Salesman & Konsumen', '🏦 SPK & Leasing', '🏍️ Penjualan Per Type', '📌 Ringkasan Visual'] + (['🧠 Analisa & Action Plan'] if IS_SPV else [])
    _tabs_all = st.tabs(_tab_names)
    tab_pp, tab_sm, tab_sl, tab_ty, tab_rk = _tabs_all[:5]
    tab_an = _tabs_all[5] if IS_SPV else None
    with tab_rk:
        bsg_safe(bsg_render_ringkasan, dict(kol_nm=kol_nm, ideal_pct=target_persen_harian, tanggal_akhir=tanggal_akhir, max_day=max_day_in_month, df_p_disp=df_p_disp, df_pr_disp=df_pr_disp,
                 df_s_disp=df_s_disp, r_tot_p=r_tot_p, r_tot_pr=r_tot_pr, r_tot_s=r_tot_s, f_c=f_c, s_c=s_c, is_total=is_total, scope_key=dlr_pil,
                 period_key=f'{current_year}-{current_month}', dlr_label=dlr_pil, sisa_hari=SISA_HARI_KERJA, hari_aktif=hari_aktif, forc_nota=forc_nota_v, forc7=forc7_v, forc3=forc3_v), label='Ringkasan visual')
    with tab_pp:
        jdl_rev_1 = f'{nxt_no()}. TABEL PENJUALAN — {dlr_pil} (1 s.d {tanggal_akhir} {periode_bulan_tahun})'
        st.markdown(f'<div class="section-title">{jdl_rev_1}</div>', unsafe_allow_html=True)
        kol_pj = [kol_nm, 'TARGET', 'IDEAL', 'STU', 'NOTA', 'TODAY', col_name_pct, 'GAP TG', 'STU LM', '% LM', 'GAP LM']   # kolom FORC sudah tampil di kartu atas & garis kurva
        df_p_full = pd.concat([df_p_disp, pd.DataFrame([r_tot_p])], ignore_index=True)

        tbl1_only_html = build_compact_html_string(df_p_disp, df_p_full, kol_pj, kol_nm, '_pct', f'cap_tbl1_inner_{dlr_pil}', title_text=None)
        custom_cap_tbl1 = f"""<div style="width:1320px;background-color:#0E1117;padding:20px 24px;font-family:'Source Sans Pro',sans-serif;color:#FFFFFF;box-sizing:border-box;">
        <div style="font-size:38px;font-weight:800;color:#4318FF;margin-bottom:4px;line-height:1.1;">{main_hdr_text}</div>
        <div style="font-size:14px;font-weight:600;color:#64748B;margin-bottom:18px;">{sub_hdr_text}</div>
        {kpi_cards_html}
        <div style="margin-top:12px;margin-bottom:18px;padding:11px 16px;background:linear-gradient(90deg,rgba(255,0,127,0.14) 0%,rgba(14,17,23,0.95) 100%);border-left:4px solid #FF007F;border-top:1px solid #2D3748;border-right:1px solid #2D3748;border-bottom:1px solid #2D3748;border-radius:6px;font-size:13.5px;line-height:1.65;color:#E2E8F0;">
            <div style="margin-top:4px;">🚶 <b>Focus Walk In :</b> {focus_walk_html}</div>
            <div style="margin-top:4px;">🎯 <b>Focus Koordinator :</b> {focus_koor_html}</div>
            <div>🏪 <b>Focus Counter :</b> {focus_cnt_html}</div>
            <div style="margin-top:4px;">👤 <b>Focus Salesman :</b> {focus_sm_html}</div>
            <div style="margin-top:4px;">🧑‍💼 <b>MP Health : Kebutuhan Salesman Baru</b> → {focus_need_html}</div>
        </div>
        <hr style="border:none;border-top:1px solid #2D3748;margin:16px 0;">
        <div style="font-size:21px;font-weight:700;color:#4318FF;margin-bottom:10px;">{jdl_rev_1}</div>
        {tbl1_only_html}
    </div>"""
        _cap_hdr_tbl1 = (f'<div style="font-size:30px;font-weight:800;color:#4318FF;margin-bottom:4px;line-height:1.1;">{main_hdr_text}</div><div style="font-size:13px;font-weight:600;color:#64748B;margin-bottom:14px;">{sub_hdr_text}</div>{kpi_cards_html}'
                         f'<div style="margin-top:12px;padding:11px 16px;background:linear-gradient(90deg,rgba(255,0,127,0.14) 0%,rgba(14,17,23,0.95) 100%);border-left:4px solid #FF007F;border-top:1px solid #2D3748;border-right:1px solid #2D3748;border-bottom:1px solid #2D3748;border-radius:6px;font-size:13.5px;line-height:1.65;color:#E2E8F0;">'
                         f'<div style="margin-top:4px;">🚶 <b>Focus Walk In :</b> {focus_walk_html}</div><div style="margin-top:4px;">🎯 <b>Focus Koordinator :</b> {focus_koor_html}</div><div>🏪 <b>Focus Counter :</b> {focus_cnt_html}</div><div style="margin-top:4px;">👤 <b>Focus Salesman :</b> {focus_sm_html}</div><div style="margin-top:4px;">🧑‍💼 <b>MP Health : Kebutuhan Salesman Baru</b> → {focus_need_html}</div></div>')
        bsg_panel_table(df_p_disp, df_p_full, kol_pj, kol_nm, '_pct', f'Review_1_Penjualan_{dlr_pil}', jdl_rev_1, bsg_pn_curve_charts, (f_c, r_tot_p['TARGET'], tanggal_akhir, max_day_in_month, r_tot_p['FORC'], dlr_pil, (forc_nota_v, forc7_v, forc3_v)), layout='side2', cap_header=_cap_hdr_tbl1)

        if is_total:
            st.divider()
            jdl_rev_1b = f'{nxt_no()}. TABEL PENJUALAN KOORDINATOR & COUNTER PER DEALER — {dlr_pil} (1 s.d {tanggal_akhir} {periode_bulan_tahun})'
            st.markdown(f'<div class="section-title">{jdl_rev_1b}</div>', unsafe_allow_html=True)
            rows_1b_items, rows_1b_full = [], []
            for d_g in list(df_p_disp[kol_nm]):
                f_c_d, f_l_d = _dsub(f_c, 'Nama Dealer', d_g), _dsub(f_l, 'Nama Dealer', d_g)
                n_c_d, n_t_d = _dsub(n_c, cd_nota, d_g), _dsub(n_t, cd_nota, d_g)
                ents_d = set()
                for ds_g, ck_g, dc_g in [(f_all, 'Koor', 'Nama Dealer'), (n_all, 'Koor', cd_nota), (db_all, 'Koor', 'Nama_Dealer_Clean'), (s_all, 'Koor_SPK', 'Nama_Dealer_SPK')]:
                    sub_g = _dsub(ds_g, dc_g, d_g)
                    if not sub_g.empty and ck_g in sub_g.columns: ents_d.update(clean_str(v) for v in sub_g[ck_g].dropna().unique() if clean_str(v) and clean_str(v) not in ['NAN', 'NONE'])
                rows_e = []
                for e_g in sorted(ents_d):
                    r_e = calc_sales_row(e_g, kol_nm, get_koor_or_counter_target(e_g, saved_tk), _dsub(f_c_d, 'Koor', e_g), _dsub(f_l_d, 'Koor', e_g), _dsub(n_c_d, 'Koor', e_g), _dsub(n_t_d, 'Koor', e_g))
                    if r_e['TARGET'] == 0 and r_e['STU'] == 0 and r_e['NOTA'] == 0 and r_e['STU LM'] == 0: continue
                    rows_e.append(r_e)
                rows_e.sort(key=lambda r_z: (0 if r_z['TARGET'] > 0 else 1, r_z['_pct'], r_z['STU']))
                rows_1b_items.extend(rows_e)
                rows_1b_full.extend(rows_e)
                rows_1b_full.append(calc_sales_row(f'SUB TOTAL {d_g}', kol_nm, safe_int(saved_td.get(d_g, 0)), f_c_d, f_l_d, n_c_d, n_t_d))
            if rows_1b_items:
                df_1b_items = pd.DataFrame(rows_1b_items)
                df_1b_full = pd.concat([pd.DataFrame(rows_1b_full), pd.DataFrame([r_tot_p])], ignore_index=True)
                render_compact_html_table(df_1b_items, df_1b_full, kol_pj, kol_nm, '_pct', key_suffix=f'Review_1B_Koor_Counter_Per_Dealer_{dlr_pil}', title_text=jdl_rev_1b)
            else:
                st.info('ℹ️ Belum ada data Koordinator / Counter pada periode ini.')

        st.divider()
        jdl_rev_2 = f'{nxt_no()}. TABEL PROSES KERJA — {dlr_pil} (1 s.d {tanggal_akhir} {periode_bulan_tahun})'
        st.markdown(f'<div class="section-title">{jdl_rev_2}</div>', unsafe_allow_html=True)
        kol_pr = [kol_nm, 'JML MP', 'AVG DB', 'DB', 'HOT', 'DB Q', 'SPK', 'ACC%', 'SPK LM', 'GROWTH', 'STU', 'STU%', 'AVG NOTA', 'AVG 7', 'AVG 3'] if is_total else [kol_nm, 'JML MP', 'AVG DB', 'DB', 'HOT', 'DB Q', 'STU', 'STU%', 'AVG NOTA', 'AVG 7', 'AVG 3']
        bsg_panel_table(df_pr_disp, pd.concat([df_pr_disp, pd.DataFrame([r_tot_pr_disp])], ignore_index=True), kol_pr, kol_nm, '_stu_pct', f'Review_2_Proses_{dlr_pil}', jdl_rev_2, bsg_pn_process_charts, (df_pr_disp, kol_nm), layout='stack_eq', bottom_fn=bsg_pn_process_bottom, bottom_args=(df_pr_disp, df_s_disp, kol_nm, hari_aktif))

        with st.expander('🔍 Rekonsiliasi Jumlah Database (cek selisih KPI vs Tabel Proses vs file asli)', expanded=False):
            _rc_rows = []
            _raw_by_dlr = {}
            for p_db in (paths_database or []):
                d_n = normalize_dealer_name(os.path.splitext(os.path.basename(p_db))[0]) or 'TIDAK DIKETAHUI'
                _raw_by_dlr[d_n] = _raw_by_dlr.get(d_n, 0) + db_file_rawcount(p_db, get_file_mtime(p_db))
            _da = df_database_alldlr if (df_database_alldlr is not None and not df_database_alldlr.empty) else pd.DataFrame()
            for d_n in sorted(set(list(_raw_by_dlr.keys()) + (list(_da['Nama_Dealer_Clean'].unique()) if not _da.empty else []))):
                g_d = _da[_da['Nama_Dealer_Clean'] == d_n] if not _da.empty else pd.DataFrame()
                in_m = g_d[(g_d['bulan'] == current_month) & (g_d['tahun'] == current_year)] if (not g_d.empty and 'bulan' in g_d.columns) else g_d
                in_p = in_m[in_m['hari'] <= tanggal_akhir] if (not in_m.empty and 'hari' in in_m.columns) else in_m
                _rc_rows.append({'DEALER': d_n + ('' if d_n in daftar_dealer else '  (di luar daftar dealer → tidak dihitung)'), 'BARIS FILE (MENTAH)': _raw_by_dlr.get(d_n, 0), 'TERBACA': len(g_d),
                                 'TANGGAL TAK TERBACA': int(g_d['Tanggal_Parsed'].isna().sum()) if (not g_d.empty and 'Tanggal_Parsed' in g_d.columns) else 0,
                                 'BULAN INI (SEMUA TGL)': len(in_m), f'S.D TGL {tanggal_akhir}': len(in_p),
                                 'KOOR TAK TERPETAKAN (LAIN-LAIN)': int((in_p['Koor'] == 'LAIN-LAIN').sum()) if (not in_p.empty and 'Koor' in in_p.columns) else 0})
            if _rc_rows:
                _rc_df = pd.DataFrame(_rc_rows)
                _rc_tot = {c: (_rc_df[c].sum() if c != 'DEALER' else 'TOTAL') for c in _rc_df.columns}
                st.dataframe(pd.concat([_rc_df, pd.DataFrame([_rc_tot])], ignore_index=True), hide_index=True, use_container_width=True)
            _kpi_sum_koor = sum(len(df_db_curr_g[df_db_curr_g['Koor'] == e_k]) for e_k in gab_ent) if not df_db_curr_g.empty else 0
            st.caption(f'Tabel Proses (TOTAL) = {len(db_c):,} baris | jumlah per koordinator/counter (dipakai KPI) = {_kpi_sum_koor:,} baris. Selisih keduanya = baris yang koordinatornya tidak terpetakan (LAIN-LAIN) — salesman perlu dilengkapi di Lookup. '
                       'Baris tanpa Customer Name tetap dihitung bila tanggal prospeknya valid; baris dengan tanggal tak terbaca tidak masuk periode.')

        st.divider()
        jdl_rev_pt = f'{nxt_no()}. TABEL REVIEW AKTIFITAS (PROSPECT TYPE) — {dlr_pil} (1 s.d {tanggal_akhir} {periode_bulan_tahun})'
        st.markdown(f'<div class="section-title">{jdl_rev_pt}</div>', unsafe_allow_html=True)
        def _act_vals(d_x):
            n_db = len(d_x)
            n_hot = int(d_x['Data_Type_Clean'].isin(['DEAL PROSPECT', 'SEAL PROSPECT']).sum()) if (n_db and 'Data_Type_Clean' in d_x.columns) else 0
            n_stu = int((d_x['STU'] == 'STU').sum()) if (n_db and 'STU' in d_x.columns) else 0
            return n_db, n_hot, (n_hot / n_db * 100.0 if n_db else 0.0), n_stu, (n_stu / n_db * 100.0 if n_db else 0.0)
        _is_cnt_view = bool(koor_pil) and str(koor_pil).upper().startswith('COUNTER')   # khusus COUNTER: Walk In 60% | Media Sosial 30% | Event 7% | Service 3%
        _ACT_TGT_PCT = {'WALKIN': 60, 'MEDIASOSIAL': 30, 'EVENT': 7, 'SERVICE': 3} if _is_cnt_view else {'EVENT': 30, 'WALKIN': 20, 'MEDIASOSIAL': 25, 'EVENTNONNACS': 10, 'KUNJUNGANRUMAH': 10, 'SERVICE': 5}   # % dari target total (counter / koordinator / dealer)
        def _act_pct(nm):
            a_n = alnum_str(nm)
            if a_n in _ACT_TGT_PCT: return _ACT_TGT_PCT[a_n]
            if 'SOSMED' in a_n or 'MEDIASOSIAL' in a_n: return 30 if _is_cnt_view else 25
            if _is_cnt_view: return None
            if 'NONNACS' in a_n: return 10
            return None
        _act_tgt_total = safe_int(r_tot_p['TARGET'])
        def _act_table(db_df, grp, name_col, tot_label, with_target=False):
            rows = []
            for g_v in grp.unique():
                v_a = _act_vals(db_df[grp == g_v])
                r_a = {name_col: g_v, 'DB': v_a[0], 'HOT': v_a[1], 'DB Q': f'{v_a[2]:.2f}%', 'STU': v_a[3], 'STU%': f'{v_a[4]:.2f}%', '_stu_pct': v_a[4]}
                if with_target:
                    pc = _act_pct(g_v); tg = int(round(_act_tgt_total * pc / 100.0)) if pc is not None else 0
                    r_a['TARGET'] = tg if pc is not None else '-'
                    r_a['ACHIEV'] = (f'{(v_a[3] / tg * 100.0):.2f}%' if tg > 0 else '0.00%') if pc is not None else '-'
                rows.append(r_a)
            df_i = pd.DataFrame(rows).sort_values(by=['DB'], ascending=[False]).reset_index(drop=True)
            v_t = _act_vals(db_df)
            tot = {name_col: tot_label, 'DB': v_t[0], 'HOT': v_t[1], 'DB Q': f'{v_t[2]:.2f}%', 'STU': v_t[3], 'STU%': f'{v_t[4]:.2f}%', '_stu_pct': v_t[4]}
            if with_target: tot['TARGET'], tot['ACHIEV'] = _act_tgt_total, (f'{(v_t[3] / _act_tgt_total * 100.0):.2f}%' if _act_tgt_total > 0 else '0.00%')
            return df_i, pd.concat([df_i, pd.DataFrame([tot])], ignore_index=True)
        _kol_act = ['DB', 'HOT', 'DB Q', 'STU', 'STU%']
        _kol_act_pt = ['TARGET', 'STU', 'ACHIEV', 'DB', 'HOT', 'DB Q', 'STU%']
        _jdl_sub_base = f'{nxt_no()}. TABEL SUB PROSPECT TYPE — {dlr_pil} (1 s.d {tanggal_akhir} {periode_bulan_tahun})'
        if db_c is None or db_c.empty or db_c.shape[1] < 8:
            st.info('ℹ️ Belum ada data Database pada periode ini.')
            st.divider()
            st.markdown(f'<div class="section-title">{_jdl_sub_base}</div>', unsafe_allow_html=True)
            st.info('ℹ️ Belum ada data Database pada periode ini.')
        else:
            _pt_g = db_c.iloc[:, 6].apply(lambda v: clean_str(v) or '(KOSONG)')   # kolom G = Prospect Type
            _st_g = db_c.iloc[:, 7].apply(lambda v: strip_sub_code(v) or '(KOSONG)')   # kolom H = Sub Prospect Type (tanpa kode di depan)
            _df_pt_i, _df_pt_f = _act_table(db_c, _pt_g, 'PROSPECT TYPE', lbl_tot, with_target=True)
            bsg_panel_table(_df_pt_i, _df_pt_f, ['PROSPECT TYPE'] + _kol_act_pt, 'PROSPECT TYPE', '_stu_pct', f'Review_Aktifitas_PT_{dlr_pil}', jdl_rev_pt, bsg_pn_activity_charts, (_df_pt_i, target_persen_harian), layout='stack3')

            st.divider()
            st.markdown(f'<div class="section-title">{_jdl_sub_base}</div>', unsafe_allow_html=True)
            _opt_pt = list(_df_pt_i['PROSPECT TYPE'])
            _sel_pt = st.selectbox('📂 Pilih Prospect Type:', _opt_pt, index=next((i_ for i_, v_ in enumerate(_opt_pt) if alnum_str(v_) == 'EVENT'), 0), key=f'sel_sub_pt_{dlr_pil}')
            _db_sel = db_c[_pt_g == _sel_pt]
            _df_sub_i, _df_sub_f = _act_table(_db_sel, _st_g[_pt_g == _sel_pt], 'SUB PROSPECT TYPE', f'TOTAL {_sel_pt}')
            bsg_panel_table(_df_sub_i, _df_sub_f, ['SUB PROSPECT TYPE'] + _kol_act, 'SUB PROSPECT TYPE', '_stu_pct', f'Review_Aktifitas_Sub_{dlr_pil}_{_sel_pt}', f'{_jdl_sub_base} — {_sel_pt}', bsg_pn_sub_charts, (_df_sub_i, _sel_pt), layout='side', wrap_first=True)

    with tab_sm:
        jdl_rev_5 = f'{nxt_no()}. TABEL DETAIL SALESMAN — {dlr_pil} (1 s.d {tanggal_akhir} {periode_bulan_tahun})'
        st.markdown(f'<div class="section-title">{jdl_rev_5}</div>', unsafe_allow_html=True)
        if len(mapping_emp_target) <= len(FIXED_COUNTER_TARGETS): st.warning('⚠️ File Master Employee belum diunggah. Silakan unggah di menu **📁 UPLOAD FILE** agar kolom TARGET otomatis terisi.')
        kol_sm = ['NAMA SALESMAN', 'TARGET', 'STU'] + (['LM'] if is_total else []) + [col_name_pct, 'AVG DB', 'DB', 'HOT', 'DB Q', 'SPK', 'ACC%', 'STU%']
        if r_sm:
            df_smi = pd.DataFrame(r_sm)
            if is_total:
                df_smi['_no_target'] = (df_smi['TARGET'] <= 0).astype(int)
                df_smi = df_smi.sort_values(by=['_no_target', '_pct_val', '_gap_lm', '_stu_val', '_db_val'], ascending=[True, True, True, True, True]).reset_index(drop=True)
            else:
                df_smi = df_smi.sort_values(by=['TARGET', '_pct_val', '_stu_val', '_db_val'], ascending=[False, True, True, True]).reset_index(drop=True)
            tt, ts, tdb, thot, tspk, n_sm = int(df_smi['TARGET'].sum()), int(df_smi['STU'].sum()), int(df_smi['DB'].sum()), int(df_smi['HOT'].sum()), int(df_smi['SPK'].sum()), len(df_smi)
            tpct_sm = (ts / tt * 100.0) if tt > 0 else 0.0
            tot_sm = {'NAMA SALESMAN': lbl_tot, 'TARGET': tt, 'STU': ts, col_name_pct: f'{tpct_sm:.2f}%', 'AVG DB': f'{(tdb/n_sm/hari_aktif if n_sm>0 else 0):.2f}', 'DB': tdb, 'HOT': thot, 'DB Q': f'{(thot/tdb*100 if tdb>0 else 0):.2f}%', 'SPK': tspk, 'ACC%': f'{(t_sa/(t_sa+t_sr)*100 if (t_sa+t_sr)>0 else 0):.2f}%', 'STU%': f'{(ts/tdb*100 if tdb>0 else 0):.2f}%', '_stu_val': ts, '_pct_val': tpct_sm}
            tot_sm['LM'] = int(df_smi['LM'].sum())
            bsg_panel_table(df_smi, pd.concat([df_smi, pd.DataFrame([tot_sm])], ignore_index=True), kol_sm, 'NAMA SALESMAN', '_pct_val', f'Review_5_Salesman_{dlr_pil}', jdl_rev_5, bsg_pn_salesman_charts, (df_smi, target_persen_harian), layout='stack_eq', footer_html=sm_notif_html)   # tabel + 2 grafik, 1 tombol Copy/Download HD
        else: st.info(f'ℹ️ Belum ada data Salesman untuk dealer {dlr_pil}.')

        st.divider()
        jdl_rev_unreg = f'{nxt_no()}. TABEL KONSUMEN BELUM REGISTRASI RATING (MYDR) — {dlr_pil} (1 s.d {tanggal_akhir} {periode_bulan_tahun})'
        st.markdown(f'<div class="section-title">{jdl_rev_unreg}</div>', unsafe_allow_html=True)
        if not f_c.empty and 'Status_Registrasi_Rating' in f_c.columns:
            tot_stu_dlr = len(f_c)
            df_sudah_reg = f_c[f_c['Status_Registrasi_Rating'] == 'SUDAH REGISTRASI'].copy()
            df_belum_reg = f_c[f_c['Status_Registrasi_Rating'] == 'BELUM REGISTRASI'].copy()
            jml_sudah, jml_belum = len(df_sudah_reg), len(df_belum_reg)
            pct_reg_dlr = (jml_sudah / tot_stu_dlr * 100.0) if tot_stu_dlr > 0 else 0.0
            st.markdown(f'''<div class="unmapped-sales-box" style="border-left-color:#00E5FF; background:linear-gradient(90deg, rgba(0,229,255,0.10) 0%, rgba(14,17,23,0.96) 100%); margin-bottom:10px;">
            📊 <b>Rekapitulasi Status Registrasi Rating (MYDR) — {dlr_pil}:</b> Total Faktur (STU): <b>{tot_stu_dlr} Konsumen</b> &nbsp;|&nbsp;
            <span style="color:#00FF66;font-weight:800;">✅ Sudah Registrasi: {jml_sudah} Konsumen ({pct_reg_dlr:.1f}%)</span> &nbsp;|&nbsp;
            <span style="color:#FF007F;font-weight:800;">⚠️ Belum Registrasi: {jml_belum} Konsumen ({100.0 - pct_reg_dlr:.1f}%)</span>
        </div>''', unsafe_allow_html=True)
            bsg_safe(bsg_render_reg_panel, jml_sudah, jml_belum, f_c, dlr_pil, label="Grafik registrasi")

            if not df_belum_reg.empty:
                df_belum_reg = df_belum_reg.sort_values(by=['Tanggal_Parsed', 'Salesman_Clean', 'Nama_Konsumen_Clean'], ascending=[True, True, True]).reset_index(drop=True)
                rows_unreg = []
                for idx_u, r_u in df_belum_reg.iterrows():
                    tgl_str = r_u['Tanggal_Parsed'].strftime('%d/%m/%Y') if (pd.notna(r_u.get('Tanggal_Parsed')) and hasattr(r_u.get('Tanggal_Parsed'), 'strftime')) else '-'
                    nm_kons = clean_str(r_u.get('Nama_Konsumen_Clean', '')) or '-'
                    no_hp = clean_str(r_u.get('No_HP_Clean', '')) or '-'
                    tp_unit = clean_str(r_u.get('Tipe_Unit_Display', '')) or clean_str(r_u.get('Type_Clean', '')) or '-'
                    lsg_u = clean_str(r_u.get('Leasing_Faktur_Clean', '')) or '-'
                    sm_u = clean_str(r_u.get('Salesman_Clean', '')) or '-'
                    kr_u = clean_str(r_u.get('Koor', '')) or '-'
                    st_u = clean_str(r_u.get('Status_Registrasi_Rating', 'BELUM REGISTRASI'))
                    rows_unreg.append({
                        'NO': idx_u + 1,
                        'TGL FAKTUR': tgl_str,
                        'NAMA KONSUMEN': nm_kons,
                        'NO. HP': no_hp,
                        'TIPE UNIT': tp_unit,
                        'PEMBAYARAN': lsg_u,
                        'NAMA SALESMAN': sm_u,
                        'KOORDINATOR': kr_u,
                        'STATUS REGISTRASI': st_u,
                        '_sort_val': 0.0
                    })
                df_unreg_items = pd.DataFrame(rows_unreg)
                tot_unreg_row = {
                    'NO': 'TOTAL',
                    '_summary': f'TOTAL STU: {tot_stu_dlr} KONSUMEN &nbsp;|&nbsp; SUDAH REGISTRASI: {jml_sudah} &nbsp;|&nbsp; BELUM REGISTRASI: {jml_belum} &nbsp;|&nbsp; CAPAIAN: {pct_reg_dlr:.1f}%',
                    '_sort_val': 0.0
                }
                kol_unreg = ['NO', 'TGL FAKTUR', 'NAMA KONSUMEN', 'NO. HP', 'TIPE UNIT', 'PEMBAYARAN', 'NAMA SALESMAN', 'KOORDINATOR', 'STATUS REGISTRASI']
                df_unreg_full = pd.concat([df_unreg_items, pd.DataFrame([tot_unreg_row])], ignore_index=True)
                render_compact_html_table(df_unreg_items, df_unreg_full, kol_unreg, 'NO', '_sort_val', key_suffix=f'Review_6_Belum_Registrasi_{dlr_pil}', title_text=jdl_rev_unreg, wrap_first='all')

                with st.expander(f"📋 Lihat Seluruh Data Faktur ({tot_stu_dlr} Konsumen: Sudah & Belum Registrasi) — {dlr_pil}"):
                    all_f_rows = []
                    for idx_a, r_a in f_c.sort_values(by=['Status_Registrasi_Rating', 'Tanggal_Parsed', 'Nama_Konsumen_Clean']).reset_index(drop=True).iterrows():
                        tgl_a = r_a['Tanggal_Parsed'].strftime('%d/%m/%Y') if (pd.notna(r_a.get('Tanggal_Parsed')) and hasattr(r_a.get('Tanggal_Parsed'), 'strftime')) else '-'
                        all_f_rows.append({
                            'NO': idx_a + 1,
                            'TGL FAKTUR': tgl_a,
                            'NAMA KONSUMEN': clean_str(r_a.get('Nama_Konsumen_Clean', '')) or '-',
                            'NO. HP': clean_str(r_a.get('No_HP_Clean', '')) or '-',
                            'TIPE UNIT': clean_str(r_a.get('Tipe_Unit_Display', '')) or '-',
                            'PEMBAYARAN': clean_str(r_a.get('Leasing_Faktur_Clean', '')) or '-',
                            'NAMA SALESMAN': clean_str(r_a.get('Salesman_Clean', '')) or '-',
                            'KOORDINATOR': clean_str(r_a.get('Koor', '')) or '-',
                            'STATUS REGISTRASI': clean_str(r_a.get('Status_Registrasi_Rating', ''))
                        })
                    st.dataframe(pd.DataFrame(all_f_rows), hide_index=True, use_container_width=True)
            else:
                st.success(f'🎉 Luar Biasa! Seluruh konsumen faktur ({tot_stu_dlr} konsumen) di {dlr_pil} SUDAH melakukan registrasi rating (MYDR).')
        else:
            st.info(f'ℹ️ Belum ada data Faktur untuk dealer {dlr_pil} pada periode ini.')

    with tab_sl:
        jdl_rev_3 = f'{nxt_no()}. TABEL SPK & LEASING — {dlr_pil} (1 s.d {tanggal_akhir} {periode_bulan_tahun})'
        st.markdown(f'<div class="section-title">{jdl_rev_3}</div>', unsafe_allow_html=True)
        kol_sp = [kol_nm, 'SPK', 'SPK LM', 'GROWTH', 'ACC%', 'TERBUKA', 'BAF', '%BAF', 'ADIRA', '%ADIRA', 'SOF', '%SOF', 'IMFI', '%IMFI']
        _lsg_grp = 'Nama Dealer' if is_total else 'Koor'
        def _stu_lsg(df_x, l_x): return int((df_x['Leasing_Faktur_Clean'] == l_x).sum()) if (df_x is not None and not df_x.empty and 'Leasing_Faktur_Clean' in df_x.columns) else 0
        df_s_tbl, r_tot_s_tbl = df_s_disp.copy(), dict(r_tot_s)
        for _l in ['BAF', 'ADIRA', 'SOF']:
            df_s_tbl[f'STU {_l}'] = [_stu_lsg(f_c[f_c[_lsg_grp] == str(n_x)] if (not f_c.empty and _lsg_grp in f_c.columns) else None, _l) for n_x in df_s_tbl[kol_nm]]
            r_tot_s_tbl[f'STU {_l}'] = _stu_lsg(f_c, _l)
        bsg_panel_table(df_s_tbl, pd.concat([df_s_tbl, pd.DataFrame([r_tot_s_tbl])], ignore_index=True), kol_sp + ['STU BAF', 'STU ADIRA', 'STU SOF'], kol_nm, '_acc_val', f'Review_3_Leasing_{dlr_pil}', jdl_rev_3, bsg_pn_leasing_charts, (df_s_tbl, r_tot_s_tbl, kol_nm, f_c), layout='stack3s',
                        chart_cols='minmax(0,1.15fr) minmax(0,1fr) minmax(0,1fr)', bottom_fn=bsg_pn_leasing_bottom, bottom_args=(r_tot_s_tbl, s_c), bottom_cols='minmax(0,1fr) minmax(0,2fr)')


    with tab_ty:
        sfx4 = ' & PER DEALER' if is_total else ''
        jdl_rev_4 = f'{nxt_no()}. TABEL PENJUALAN PER TYPE{sfx4} — {dlr_pil} (1 s.d {tanggal_akhir} {periode_bulan_tahun})'
        st.markdown(f'<div class="section-title">{jdl_rev_4}</div>', unsafe_allow_html=True)
        TYPE_TGT_PCT = {'AT PREMIUM': 60, 'AT CLASSY': 25, 'AT LPM': 10, 'MOPED': 3, 'OTHERS': 2}   # % dari target total (dealer / total sesuai dropbox)
        TYPE_ORDER = list(TYPE_TGT_PCT.keys())
        kol_t_d = ['TYPE', 'TARGET', 'STU', 'ACHIEV', 'LM', 'GROWTH', 'DB', 'HOT', 'DB Q', 'STU CASH', 'STU BAF', 'STU ADIRA', 'STU SOF']
        _tgt_total = safe_int(r_tot_p['TARGET'])
        def _grp_type(d_x):
            if d_x is None or d_x.empty or 'Type_Clean' not in d_x.columns: return d_x
            d_x = d_x.copy(); d_x['Type_Grp'] = d_x['Type_Clean'].where(d_x['Type_Clean'].isin(['AT PREMIUM', 'AT CLASSY', 'AT LPM', 'MOPED']), 'OTHERS')
            return d_x
        def _add_tgt(df_x, tgt_fn):
            if df_x is None or df_x.empty: return df_x
            df_x = df_x.copy()
            df_x['TARGET'] = [tgt_fn(str(v)) for v in df_x['TYPE']]
            df_x['ACHIEV'] = [f'{(safe_int(a) / t * 100.0):.2f}%' if t > 0 else '0.00%' for a, t in zip(df_x['STU'], df_x['TARGET'])]
            df_x['LM'] = df_x['STU LM']   # 'STU LM' tetap dipertahankan (dipakai tabel Analisa)
            return df_x
        f_cg, f_lg, db_cg, s_cg = _grp_type(f_c), _grp_type(f_l), _grp_type(db_c), _grp_type(s_c)
        df_td_i, df_td_f, _ = build_type_table(f_cg, f_lg, db_cg, s_cg, [(k, 'Type_Grp', 'Type_Grp', 'Type_Grp') for k in TYPE_ORDER], 'TYPE', tot_label=lbl_tot, sort_by_growth=True)
        _tfn = lambda nm: _tgt_total if nm == lbl_tot else int(round(_tgt_total * TYPE_TGT_PCT.get(nm, 0) / 100.0))
        df_td_i, df_td_f = _add_tgt(df_td_i, _tfn), _add_tgt(df_td_f, _tfn)
        if not df_td_f.empty and is_total:
            rows_t_full = []
            for _, tr_row in df_td_i.iterrows():
                kt_n = str(tr_row['TYPE'])
                r_type = tr_row.to_dict(); r_type['_row_ratio'] = float('nan')
                rows_t_full.append(r_type)
                di_t, _, _ = build_type_table(f_cg[f_cg['Type_Grp'] == kt_n] if not f_cg.empty else pd.DataFrame(), f_lg[f_lg['Type_Grp'] == kt_n] if not f_lg.empty else pd.DataFrame(), db_cg[db_cg['Type_Grp'] == kt_n] if not db_cg.empty else pd.DataFrame(), s_cg[s_cg['Type_Grp'] == kt_n] if not s_cg.empty else pd.DataFrame(), [(d_x, 'Nama Dealer', 'Nama_Dealer_Clean', 'Nama_Dealer_SPK') for d_x in daftar_dealer], 'TYPE', sort_by_growth=True)
                if di_t.empty: continue
                di_t = _add_tgt(di_t, lambda nm, _k=kt_n: int(round(safe_int(saved_td.get(nm, 0)) * TYPE_TGT_PCT.get(_k, 0) / 100.0)))
                g_min, g_max = float(di_t['_growth_val'].min()), float(di_t['_growth_val'].max())
                for j_t, r_dt in di_t.iterrows():
                    rd_t = r_dt.to_dict()
                    rd_t['_row_ratio'] = ((float(r_dt['_growth_val']) - g_min) / (g_max - g_min)) if g_max > g_min else (j_t / max(1, len(di_t) - 1))
                    rows_t_full.append(rd_t)
            r_tot_type = df_td_f.iloc[-1].to_dict(); r_tot_type['_row_ratio'] = float('nan')
            rows_t_full.append(r_tot_type)
            bsg_panel_table(df_td_i, pd.DataFrame(rows_t_full), kol_t_d, 'TYPE', '_growth_val', f'Review_4_Type_{dlr_pil}', jdl_rev_4, bsg_pn_type_charts, (df_td_i, target_persen_harian), layout='side', highlight_names=set(TYPE_ORDER), grid_cols='minmax(0,2.2fr) minmax(0,1fr)', bottom_fn=bsg_pn_type_bottom, bottom_args=(df_td_i, f_cg), bottom_left=True)
        elif not df_td_f.empty: bsg_panel_table(df_td_i, df_td_f, kol_t_d, 'TYPE', '_growth_val', f'Review_4_Type_{dlr_pil}', jdl_rev_4, bsg_pn_type_charts, (df_td_i, target_persen_harian), layout='side', grid_cols='minmax(0,2.2fr) minmax(0,1fr)', bottom_fn=bsg_pn_type_bottom, bottom_args=(df_td_i, f_cg), bottom_left=True)

    # ---- tombol Download SEMUA gambar untuk PPT (diisi setelah semua panel tab Penjualan/Salesman/Leasing dirender, sebelum tab Analisa) ----
    with _ppt_dl_ph.container():
        _wk_k = _wk_c = None
        try:
            if not koor_pil and not df_p_it.empty:
                _c_k = df_p_it[~df_p_it[kol_nm].str.upper().str.startswith('COUNTER')]; _c_k = _c_k[_c_k['TARGET'] > 0] if (_c_k['TARGET'] > 0).any() else _c_k
                _c_c = df_p_it[df_p_it[kol_nm].str.upper().str.startswith('COUNTER')]; _c_c = _c_c[_c_c['TARGET'] > 0] if (_c_c['TARGET'] > 0).any() else _c_c
                _wk_k, _wk_c = (str(_c_k.iloc[0][kol_nm]) if not _c_k.empty else None), (str(_c_c.iloc[0][kol_nm]) if not _c_c.empty else None)
                st.caption('Terlemah untuk slide 9–18 → Koordinator: **' + (_wk_k or '-') + '** · Counter: **' + (_wk_c or '-') + '** (Dashboard KPI keduanya ikut masuk ke ZIP)')
        except Exception: pass
        try: bsg_render_ppt_batch(dlr_pil, kpi_dealer=(None if koor_pil else _ppt_dealer), kpi_entities=([koor_pil] if koor_pil else [x_ for x_ in (_wk_k, _wk_c) if x_]))
        except Exception as _e_bt: st.error(f'Tombol Download SEMUA gagal dibuat: {type(_e_bt).__name__}: {_e_bt}')

    if not IS_SPV: st.stop()   # tanpa kode akses SPV: tab Analisa & Action Plan tidak dirender sama sekali
    with tab_an:
        bsg_admin_running_text_editor()
        tgl_sasaran = 30 if tanggal_akhir > 21 else min(30, tanggal_akhir + 7)
        sisa_hari = max(1, sum(1 for d_n in range(tanggal_akhir + 1, min(tgl_sasaran, max_day_in_month) + 1) if is_workday_day(d_n)))
        horizon_lbl = f"Optimalisasi Closing s.d Tgl 30 {periode_bulan_tahun} ({sisa_hari} Hari Kerja)" if tanggal_akhir > 21 else f"Akselerasi 7 Hari ke Depan s.d Tgl {tgl_sasaran} {periode_bulan_tahun} ({sisa_hari} Hari Kerja)"

        jdl_rev_6 = f'{nxt_no()}. TABEL ANALISA & ACTION PLAN — {dlr_pil} ({horizon_lbl})'
        st.markdown(f'<div class="section-title">{jdl_rev_6}</div>', unsafe_allow_html=True)

        dlr_targ, dlr_ideal, dlr_stu, dlr_nota, dlr_stulm = safe_int(r_tot_p.get('TARGET', 0)), safe_int(r_tot_p.get('IDEAL', 0)), safe_int(r_tot_p.get('STU', 0)), safe_int(r_tot_p.get('NOTA', 0)), safe_int(r_tot_p.get('STU LM', 0))
        dlr_pct, dlr_gap_ideal, dlr_gap_lm = safe_float(r_tot_p.get('_pct', 0.0)), dlr_stu - dlr_ideal, dlr_stu - dlr_stulm
        cur_mp, cur_db, cur_hot = (safe_int(r_tot_pr.get('JML MP', 0)) or max(1, len(r_sm))), safe_int(r_tot_pr.get('DB', 0)), safe_int(r_tot_pr.get('HOT', 0))
        cur_avg_db, cur_dbq, cur_stup = ((cur_db / cur_mp / hari_aktif) if cur_mp > 0 else 0.0), ((cur_hot / cur_db * 100.0) if cur_db > 0 else 0.0), ((dlr_stu / cur_db * 100.0) if cur_db > 0 else 0.0)
        cur_spk, cur_spk_lm, cur_open_spk, cur_acc_pct = safe_int(r_tot_s.get('SPK', 0)), safe_int(r_tot_s.get('SPK LM', 0)), safe_int(r_tot_s.get('TERBUKA', 0)), safe_float(r_tot_s.get('_acc_val', 0.0))

        pct_sasaran = tabel_target_harian.get(tgl_sasaran, 100.0) / 100.0
        tambahan_stu_wajib, stu_sasaran_rasional = calc_rational_7d_addition(dlr_stu, dlr_targ, tanggal_akhir, tgl_sasaran, persen_desimal, pct_sasaran, active_mp=cur_mp)
        target_stu_per_hari = tambahan_stu_wajib / sisa_hari

        fak_kredit_curr = {l_k: (len(f_c[f_c['Leasing_Faktur_Clean'] == l_k]) if (not f_c.empty and 'Leasing_Faktur_Clean' in f_c.columns) else 0) for l_k in ['BAF', 'ADIRA', 'SOF']}
        tot_fak_kredit_curr = sum(fak_kredit_curr.values())
        lsg_stats = [(l_nm, safe_int(r_tot_s.get(l_nm, 0)), safe_float(r_tot_s.get(f'%{l_nm}', 0.0))) for l_nm in ['BAF', 'ADIRA', 'SOF', 'IMFI']]
        lsg_active = [x for x in lsg_stats if x[1] > 0]
        best_lsg, best_fak_lsg = (max(lsg_active, key=lambda x: (x[2], x[1])) if lsg_active else ('BAF', 0, 0.0)), max(['BAF', 'ADIRA', 'SOF'], key=lambda k: fak_kredit_curr[k])

        df_rej_dlr = s_c[s_c['Status_SPK_Clean'] == 'DITOLAK'] if (not s_c.empty and 'Status_SPK_Clean' in s_c.columns) else pd.DataFrame()
        top_rej_reasons_str, top_1_reason = "Belum ada SPK ditolak", "Kapasitas/SLIK"
        if not df_rej_dlr.empty and 'Alasan_Reject_Clean' in df_rej_dlr.columns:
            vc_rej = df_rej_dlr['Alasan_Reject_Clean'].value_counts()
            if not vc_rej.empty: top_1_reason, top_rej_reasons_str = str(vc_rej.index[0]), ", ".join([f"{k} ({v} SPK)" for k, v in vc_rej.head(4).items()])

        eff_acc_rate = 0.50 if cur_acc_pct >= 99.9 else (max(IDEAL_ACC_PCT, cur_acc_pct) / 100.0)
        potensi_stu_open = min(tambahan_stu_wajib, round(cur_open_spk * eff_acc_rate))
        kebutuhan_spk_baru = max(2, round(max(1, tambahan_stu_wajib - round(potensi_stu_open * 0.5)) / eff_acc_rate))
        target_db_baru_total = max(round(tambahan_stu_wajib / (IDEAL_STU_PCT / 100.0)), int(cur_mp * 3 * sisa_hari))
        target_db_per_mp_hari = min(5.0, max(3.0, round((target_db_baru_total / sisa_hari) / cur_mp, 1))) if cur_mp > 0 else 3.5
        tot_db_plan_dlr = round(target_db_per_mp_hari * cur_mp * sisa_hari)

        default_cond_dealer = limit_words(f"STU {dlr_stu} unit ({dlr_pct:.1f}% dari target {dlr_targ}), gap ideal {dlr_gap_ideal:+d} unit, vs LM {dlr_gap_lm:+d} unit. Input {cur_avg_db:.1f} DB/MP/hari ({cur_db} DB), DB Q {cur_dbq:.1f}% (ideal 30%), STU% {cur_stup:.1f}% (ideal 10%).")
        default_plan_dealer = limit_words(f"Target rasional s.d tgl {tgl_sasaran}: kejar +{tambahan_stu_wajib} STU (menjadi {stu_sasaran_rasional} unit / {target_stu_per_hari:.1f} STU/hari). Input {target_db_per_mp_hari:.1f} DB/MP/hari (+{tot_db_plan_dlr} DB) TANPA MENURUNKAN KUALITAS DATABASE.")

        default_cond_type, default_plan_type, wt_rational_str = "Data type belum tersedia.", "Belum ada rencana aksi per type.", ""
        wt_name, wt_stu, wt_lm, wt_db, wt_hot, wt_spk, wt_stu_add, wt_spk_add, wt_db_add = "AT PREMIUM", 0, 0, 0, 0, 0, 2, 4, 20
        _df_td_focus = df_td_i[df_td_i['TYPE'].isin(FOKUS_KATEGORI_TYPE)] if not df_td_i.empty else df_td_i   # analisa tetap fokus AT Premium/Classy/LPM
        if not _df_td_focus.empty:
            wt_row = _df_td_focus.iloc[0]
            wt_name, wt_stu, wt_lm, wt_db, wt_hot, wt_spk = str(wt_row['TYPE']), safe_int(wt_row['STU']), safe_int(wt_row['STU LM']), safe_int(wt_row['DB']), safe_int(wt_row['HOT']), safe_int(wt_row['SPK'])
            default_cond_type = limit_words(f"Segmen terendah {wt_name}: STU {wt_stu} vs LM {wt_lm} ({wt_stu - wt_lm:+d} unit, Growth {wt_row['GROWTH']}). Input {wt_db} DB, {wt_hot} HOT (DB Q {wt_row['DB Q']} vs 30%), {wt_spk} SPK (ACC% {wt_row['ACC%']}), STU% {wt_row['STU%']}.")
            wt_stu_add = max(2, min(round(tambahan_stu_wajib * 0.45), max(3, round((wt_stu / hari_aktif) * sisa_hari * 1.8) + 2)))
            wt_spk_add, wt_db_add = max(wt_stu_add + 1, round(wt_stu_add / 0.5)), max(round(wt_stu_add / 0.10), sisa_hari * cur_mp)
            wt_rational_str = f"Segmen terburuk {wt_name}: target rasional +{wt_stu_add} STU (menjadi {wt_stu + wt_stu_add} unit), +{wt_spk_add} SPK, +{wt_db_add} DB."
            default_plan_type = limit_words(f"Fokus {wt_name}: kejar +{wt_stu_add} STU (menjadi {wt_stu + wt_stu_add} unit) dan +{wt_spk_add} SPK (ACC% >=50%). Tambah +{wt_db_add} DB {wt_name} TANPA MENURUNKAN KUALITAS DATABASE (DB Q >=30%, STU% >=10%).")

        default_cond_leasing = limit_words(f"Total {cur_spk} SPK (vs LM {cur_spk_lm}), ACC% 100% belum valid karena {cur_open_spk} SPK masih Terbuka. Faktur kredit cair {tot_fak_kredit_curr} unit (BAF {fak_kredit_curr['BAF']}, ADIRA {fak_kredit_curr['ADIRA']}, SOF {fak_kredit_curr['SOF']}).") if cur_acc_pct >= 99.9 else limit_words(f"Total {cur_spk} SPK (vs LM {cur_spk_lm}), ACC% {cur_acc_pct:.1f}% (ideal 50%), {cur_open_spk} SPK Terbuka, reject tertinggi: {top_1_reason}. Faktur kredit {tot_fak_kredit_curr} unit (terkuat {best_lsg[0]}).")
        default_plan_leasing = limit_words(f"Eksekusi {cur_open_spk} SPK Terbuka dalam 1x24 jam (target cair +{potensi_stu_open} STU) dan tambah +{kebutuhan_spk_baru} SPK baru. Prioritaskan 60% aplikasi ke {best_fak_lsg} serta perketat screening {top_1_reason} agar ACC% >=50%.")

        default_cond_koor, default_plan_koor, wk_rational_str = "Belum ada data Koordinator.", "Belum ada rencana aksi Koordinator.", ""
        wk_nm, wk_t, wk_s, wk_add, wk_goal_rasional, wk_spk_need, wk_db_per_sm, wk_mp = "-", 0, 0, 2, 2, 4, 4, 1
        if not df_koor_only.empty:
            wk = df_koor_only.sort_values(by=['_pct', 'STU'], ascending=[True, True]).iloc[0]
            wk_nm, wk_t, wk_i, wk_s, wk_pct = str(wk[kol_nm]), safe_int(wk['TARGET']), safe_int(wk['IDEAL']), safe_int(wk['STU']), safe_float(wk['_pct'])
            wk_pr = df_pr_it[df_pr_it[kol_nm] == wk_nm].iloc[0] if (not df_pr_it.empty and wk_nm in df_pr_it[kol_nm].values) else {}
            wk_sp = df_s_it[df_s_it[kol_nm] == wk_nm].iloc[0] if (not df_s_it.empty and wk_nm in df_s_it[kol_nm].values) else {}
            wk_mp = max(1, safe_int(wk_pr.get('JML MP', 1)))
            default_cond_koor = limit_words(f"Koordinator terendah {wk_nm}: STU {wk_s}/{wk_t} unit ({wk_pct:.1f}%, gap ideal {wk_s - wk_i:+d}). Input {wk_pr.get('AVG DB', '0')} DB/MP/hari, DB Q {wk_pr.get('DB Q', '0%')} (ideal 30%), STU% {wk_pr.get('STU%', '0%')}, ACC% {wk_sp.get('ACC%', '0%')} .")
            wk_add, wk_goal_rasional = calc_rational_7d_addition(wk_s, wk_t, tanggal_akhir, tgl_sasaran, persen_desimal, pct_sasaran, active_mp=wk_mp)
            wk_spk_need, wk_db_per_sm = max(wk_add + 1, round(wk_add / 0.5)), min(5, max(3, round((wk_add / 0.10) / max(1, wk_mp * sisa_hari))))
            wk_rational_str = f"Koordinator terburuk {wk_nm} (STU saat ini {wk_s} dari target {wk_t}): target rasional 7 hari adalah +{wk_add} STU (menjadi {wk_goal_rasional} unit), +{wk_spk_need} SPK, {wk_db_per_sm} DB/MP/hari."
            default_plan_koor = limit_words(f"Tim {wk_nm} kejar target rasional +{wk_add} STU (menjadi {wk_goal_rasional} unit) dan +{wk_spk_need} SPK baru. Disiplinkan {wk_mp} MP input {wk_db_per_sm} DB/hari TANPA MENURUNKAN KUALITAS DATABASE (DB Q >=30%, STU% >=10%).")

        default_cond_counter, default_plan_counter, wc_rational_str = "Belum ada data Counter.", "Belum ada rencana aksi Counter.", ""
        wc_nm, wc_t, wc_s, wc_add, wc_goal_rasional, wc_spk_need, wc_db_hari = "-", 0, 0, 2, 2, 4, 4
        if not df_cnt_only.empty:
            wc = df_cnt_only.iloc[0]
            wc_nm, wc_t, wc_i, wc_s, wc_pct = str(wc[kol_nm]), safe_int(wc['TARGET']), safe_int(wc['IDEAL']), safe_int(wc['STU']), safe_float(wc['_pct'])
            wc_pr = df_pr_it[df_pr_it[kol_nm] == wc_nm].iloc[0] if (not df_pr_it.empty and wc_nm in df_pr_it[kol_nm].values) else {}
            wc_sp = df_s_it[df_s_it[kol_nm] == wc_nm].iloc[0] if (not df_s_it.empty and wc_nm in df_s_it[kol_nm].values) else {}
            default_cond_counter = limit_words(f"Counter terendah {wc_nm}: STU {wc_s}/{wc_t} unit ({wc_pct:.1f}%, gap ideal {wc_s - wc_i:+d}). Prospek {wc_pr.get('DB', 0)} DB, DB Q {wc_pr.get('DB Q', '0%')} (ideal 30%), STU% {wc_pr.get('STU%', '0%')}, ACC% {wc_sp.get('ACC%', '0%')} .")
            wc_add, wc_goal_rasional = calc_rational_7d_addition(wc_s, wc_t, tanggal_akhir, tgl_sasaran, persen_desimal, pct_sasaran, active_mp=1)
            wc_spk_need, wc_db_hari = max(wc_add + 1, round(wc_add / 0.5)), min(6, max(4, round((wc_add / 0.10) / sisa_hari)))
            wc_rational_str = f"Counter terburuk {wc_nm} (STU saat ini {wc_s} dari target {wc_t}): target rasional 7 hari adalah +{wc_add} STU (menjadi {wc_goal_rasional} unit), +{wc_spk_need} SPK, {wc_db_hari} DB/hari."
            default_plan_counter = limit_words(f"{wc_nm} kejar target rasional +{wc_add} STU (menjadi {wc_goal_rasional} unit), cairkan {safe_int(wc_sp.get('TERBUKA', 0))} SPK Terbuka, cetak +{wc_spk_need} SPK. Input {wc_db_hari} DB/hari TANPA MENURUNKAN KUALITAS DATABASE (DB Q >=30%).")

        default_cond_mp = "Seluruh Salesman lapangan aktif sudah di zona aman."
        default_plan_mp = "Pertahankan input database harian TANPA MENURUNKAN KUALITAS DATABASE sesuai standar DB Q 30%, STU% 10%, ACC% 50%."
        top_bad_raw_str, names_kritis, cnt_kritis = "Tidak ada salesman aktif di zona kritis.", "-", 0
        if r_sm:
            df_sm_eval = pd.DataFrame(r_sm)
            min_p_tbl, max_p_tbl = float(df_sm_eval['_pct_val'].min()), float(df_sm_eval['_pct_val'].max())
            def is_zona_kritis(row_s):
                v_p, t_g, s_u = safe_float(row_s['_pct_val']), safe_int(row_s['TARGET']), safe_int(row_s['STU'])
                idl = round(t_g * persen_desimal)
                if (s_u >= idl and idl > 0) or (v_p >= target_persen_harian and target_persen_harian > 0): return False
                ratio_c = ((v_p - min_p_tbl) / (max_p_tbl - min_p_tbl)) if max_p_tbl > min_p_tbl else 0.0
                return (ratio_c <= 0.15) or (v_p < target_persen_harian)
            df_sm_non_cnt = df_sm_eval[(~df_sm_eval['NAMA SALESMAN'].apply(is_counter_personnel)) & (df_sm_eval['_db_val'] >= 2)].copy()
            df_kritis_with_target = df_sm_non_cnt[(df_sm_non_cnt['TARGET'] > 0) & (df_sm_non_cnt.apply(is_zona_kritis, axis=1))].sort_values(by=['TARGET', '_pct_val', '_stu_val', '_db_val'], ascending=[False, True, True, True])
            if not df_kritis_with_target.empty:
                max_kritis_target = int(df_kritis_with_target['TARGET'].max())
                top_bad = df_kritis_with_target[df_kritis_with_target['TARGET'] == max_kritis_target].head(4)
                if len(top_bad) < 3: top_bad = df_kritis_with_target.head(4)
            else: top_bad = df_sm_non_cnt[df_sm_non_cnt.apply(is_zona_kritis, axis=1)].sort_values(by=['TARGET', '_pct_val', '_stu_val', '_db_val'], ascending=[False, True, True, True]).head(4)
            if not top_bad.empty:
                cnt_kritis = len(top_bad)
                names_kritis = ", ".join([f"{rb['NAMA SALESMAN']} (STU {safe_int(rb['STU'])}/{safe_int(rb['TARGET'])}, DB {safe_int(rb['DB'])}, SPK {safe_int(rb['SPK'])})" for _, rb in top_bad.iterrows()])
                top_bad_raw_str = "; ".join([f"{rb['NAMA SALESMAN']}: Target {rb['TARGET']}, STU {rb['STU']} ({rb[col_name_pct]}), DB {rb['DB']} (AVG {rb['AVG DB']}), HOT {rb['HOT']} (DB Q {rb['DB Q']}), SPK {rb['SPK']} (ACC% {rb['ACC%']}), STU% {rb['STU%']} -> Target Rasional 7 Hari: +1 s.d +2 STU" for _, rb in top_bad.iterrows()])
                default_cond_mp = limit_words(f"{cnt_kritis} Salesman aktif zona kritis: {names_kritis}. Rata-rata DB Quality dan konversi STU% masih di bawah standar ideal 30% dan 10%.")
                default_plan_mp = limit_words(f"Kawal joint canvassing {cnt_kritis} MP kritis untuk kejar target rasional +1 s.d +2 STU dan +3 SPK/orang. Wajib input 4 DB/hari/MP TANPA MENURUNKAN KUALITAS DATABASE (DB Q >=30%).")

        def _prospect_stats(db_df, kws, exact=False):
            if db_df is None or db_df.empty or db_df.shape[1] < 8: return None
            m_p = db_df.iloc[:, 6].apply(alnum_str).apply(lambda v: (v in kws) if exact else any(k_w in v for k_w in kws))
            d_p = db_df[m_p]
            if d_p.empty: return None
            n_p, h_p, q_p, s_p, sp_p = _act_vals(d_p)
            sub_nm = d_p.iloc[:, 7].apply(lambda v: strip_sub_code(v) or '(KOSONG)')
            subs = []
            for g_v in sub_nm.unique():
                v_s = _act_vals(d_p[sub_nm == g_v]); subs.append((g_v, v_s[0], v_s[1], v_s[2], v_s[3], v_s[4]))
            subs.sort(key=lambda x: -x[1])
            weak = sorted([x for x in subs if x[1] >= 3] or subs, key=lambda x: (x[5], x[3], -x[1]))
            return {'n': n_p, 'hot': h_p, 'dbq': q_p, 'stu': s_p, 'stup': sp_p, 'share': (n_p / len(db_df) * 100.0 if len(db_df) else 0.0), 'subs': subs, 'weak': weak[0] if weak else None}
        def _md_texts(label, st_p, with_sub):
            if not st_p:
                return (f'Belum ada database {label} pada periode ini (1 s.d {tanggal_akhir}).', f'Aktifkan kegiatan {label} dan wajib input database tiap kegiatan TANPA MENURUNKAN KUALITAS DATABASE.')
            top = st_p['subs'][0]
            if with_sub and st_p['weak']:   # Aktifitas MD: khusus prospect EVENT -> sub terlemah
                w = st_p['weak']
                cond = f"Event total {st_p['n']} DB (STU {st_p['stu']}, STU% {st_p['stup']:.1f}%). Sub terlemah {w[0]}: {w[1]} DB, HOT {w[2]} (DB Q {w[3]:.1f}% vs 30%), STU {w[4]} (STU% {w[5]:.1f}% vs 10%)."
                need_w = max(0, math.ceil(w[1] * IDEAL_STU_PCT / 100.0) - w[4])
                plan = (f"Evaluasi event {w[0]}: follow-up {w[2]} HOT 1x24 jam, kejar +{need_w} STU (STU% 10%), perketat penjaringan agar DB Q >=30% TANPA MENURUNKAN KUALITAS DATABASE." if need_w > 0
                        else f"Pertahankan konversi {w[0]} di atas 10% dan replikasi pola terbaiknya ke sub Event lain TANPA MENURUNKAN KUALITAS DATABASE.")
                return cond, plan
            cond = f"{label} {st_p['n']} DB ({st_p['share']:.1f}% dari total DB), HOT {st_p['hot']} (DB Q {st_p['dbq']:.1f}% vs 30%), STU {st_p['stu']} (STU% {st_p['stup']:.1f}% vs 10%). Sub terbesar {top[0]} ({top[1]} DB, STU% {top[5]:.1f}%)."
            need = max(0, math.ceil(st_p['n'] * IDEAL_STU_PCT / 100.0) - st_p['stu'])
            plan = (f"Fokus {top[0]}: follow-up {st_p['hot']} HOT 1x24 jam, kejar +{need} STU agar STU% mencapai 10%" if need > 0 else f"Pertahankan STU% {label} di atas 10%, perluas {top[0]}") + ", saring data agar DB Q >=30% TANPA MENURUNKAN KUALITAS DATABASE."
            return cond, plan
        _st_event, _st_sosmed = _prospect_stats(db_c, ['EVENT'], exact=True), _prospect_stats(db_c, ['MEDIASOSIAL', 'SOSMED'])
        default_cond_md, default_plan_md = _md_texts('Event', _st_event, True)
        default_cond_sosmed, default_plan_sosmed = _md_texts('Media Sosial', _st_sosmed, False)

        def build_smart_dynamic_narrative():
            v_idx = random.randint(0, 2)
            b_dlr = "kualitas prospek DB Q rendah" if cur_dbq < 25 else "konversi closing HOT ke STU macet"
            c_dlr = [
                f"Bottleneck {dlr_pil} pada {b_dlr} (DB Q {cur_dbq:.1f}% vs 30%, STU% {cur_stup:.1f}% vs 10%), memicu defisit ideal {dlr_gap_ideal:+d} unit (STU {dlr_stu}/{dlr_targ}, vs LM {dlr_gap_lm:+d}).",
                f"Capaian {dlr_pil} baru {dlr_stu} unit ({dlr_pct:.1f}%), tertahan {b_dlr}: dari {cur_db} DB hanya jadi {cur_hot} HOT ({cur_dbq:.1f}%) dan {dlr_stu} STU ({cur_stup:.1f}%).",
                f"Realisasi {dlr_stu}/{dlr_targ} unit (gap ideal {dlr_gap_ideal:+d}) akibat lemahnya filter prospek ({cur_avg_db:.1f} DB/MP/hr, DB Q {cur_dbq:.1f}%, STU% {cur_stup:.1f}%)."
            ][v_idx]
            p_dlr = [
                f"Akselerasi s.d tgl {tgl_sasaran}: kunci +{tambahan_stu_wajib} STU (total {stu_sasaran_rasional} unit / {target_stu_per_hari:.1f} STU/hr) dan {target_db_per_mp_hari:.1f} DB/MP/hr (+{tot_db_plan_dlr} DB) TANPA MENURUNKAN KUALITAS DATABASE.",
                f"Kejar kuota rasional +{tambahan_stu_wajib} STU ({target_stu_per_hari:.1f} unit/hr) lewat bedah {cur_hot} HOT prospek serta disiplin {target_db_per_mp_hari:.1f} DB/MP/hr TANPA MENURUNKAN KUALITAS DATABASE.",
                f"Wajib tambah +{tambahan_stu_wajib} STU (menjadi {stu_sasaran_rasional} unit) dan +{tot_db_plan_dlr} DB baru (DB Q >=30%, STU% >=10%) TANPA MENURUNKAN KUALITAS DATABASE."
            ][v_idx]
            c_tp = f"Kontraksi terdalam di {wt_name} (STU {wt_stu} vs LM {wt_lm}); dari {wt_db} DB hanya menghasilkan {wt_hot} HOT dan {wt_spk} SPK sehingga konversi terhambat." if not df_td_i.empty else default_cond_type
            p_tp = f"Genjot segmen {wt_name}: sasar +{wt_stu_add} STU (menjadi {wt_stu + wt_stu_add} unit), +{wt_spk_add} SPK, dan +{wt_db_add} DB tertarget TANPA MENURUNKAN KUALITAS DATABASE." if not df_td_i.empty else default_plan_type
            c_lsg = f"Dari {cur_spk} SPK (LM {cur_spk_lm}), faktur kredit cair {tot_fak_kredit_curr} unit; tertahan {cur_open_spk} SPK Terbuka dan reject dominan karena {top_1_reason} (ACC% {cur_acc_pct:.1f}%)."
            p_lsg = f"Selesaikan {cur_open_spk} SPK Terbuka dalam 1x24 jam (cair +{potensi_stu_open} STU), cetak +{kebutuhan_spk_baru} SPK baru, alihkan aplikasi ke {best_fak_lsg} & perketat pra-screening {top_1_reason}."
            c_kr = f"Tim {wk_nm} di zona kritis (STU {wk_s}/{wk_t}); produktivitas DB lapangan dan rasio konversi HOT ke STU masih di bawah standar 30% & 10%." if not df_koor_only.empty else default_cond_koor
            p_kr = f"Tim {wk_nm} wajib amankan +{wk_add} STU (menjadi {wk_goal_rasional} unit), +{wk_spk_need} SPK, dan {wk_db_per_sm} DB/MP/hr TANPA MENURUNKAN KUALITAS DATABASE." if not df_koor_only.empty else default_plan_koor
            c_cn = f"Konversi walk-in/leads digital {wc_nm} belum optimal (STU {wc_s}/{wc_t}), dipicu rendahnya follow-up HOT prospek ke pengajuan SPK." if not df_cnt_only.empty else default_cond_counter
            p_cn = f"Fokus {wc_nm} kejar +{wc_add} STU (total {wc_goal_rasional} unit), cetak +{wc_spk_need} SPK, serta input {wc_db_hari} DB/hr TANPA MENURUNKAN KUALITAS DATABASE." if not df_cnt_only.empty else default_plan_counter
            c_mp = f"Terdeteksi {cnt_kritis} salesman zona kritis ({names_kritis}) dengan target besar namun minim konversi SPK dan rendah DB Quality." if cnt_kritis > 0 else default_cond_mp
            p_mp = f"Supervisi harian {cnt_kritis} MP kritis untuk mengunci +1 s.d +2 STU, +3 SPK, dan 4-5 DB/hr/orang TANPA MENURUNKAN KUALITAS DATABASE." if cnt_kritis > 0 else default_plan_mp
            return {k: limit_words(v, MAX_NARRATIVE_WORDS) for k, v in {'cond_dealer': c_dlr, 'plan_dealer': p_dlr, 'cond_type': c_tp, 'plan_type': p_tp, 'cond_leasing': c_lsg, 'plan_leasing': p_lsg, 'cond_koor': c_kr, 'plan_koor': p_kr, 'cond_counter': c_cn, 'plan_counter': p_cn, 'cond_mp': c_mp, 'plan_mp': p_mp, 'cond_md': default_cond_md, 'plan_md': default_plan_md, 'cond_sosmed': default_cond_sosmed, 'plan_sosmed': default_plan_sosmed}.items()}

        active_api_key = resolve_gemini_api_key()
        cache_key = f"ai_tbl6_v10_35w_{dlr_pil}_{tanggal_akhir}_{dlr_stu}_{cur_db}_{cur_spk}"

        col_btn1, col_btn2 = st.columns([1, 2])
        with col_btn1: gen_ai_clicked = st.button('✨ Generate Narasi Dinamis dengan AI (Gemini)', type='primary', use_container_width=True)
        with col_btn2:
            with st.expander('🤖 Pengaturan / Generate API Key Share GitHub (Klik untuk buka)'):
                st.caption('💡 Untuk memanggil server Gemini secara langsung, gunakan API Key gratis dari Google AI Studio (diawali **`AIzaSy...`**). Jika menggunakan token lain, sistem otomatis memakai *Smart Dynamic Analyzer*:')
                inp_api_key = st.text_input('Masukkan / Ganti Gemini API Key (Format AIzaSy...):', value=active_api_key, type='password')
                c_k1, c_k2 = st.columns(2)
                with c_k1:
                    if st.button('💾 Terapkan & Simpan API Key', use_container_width=True):
                        clean_inp = inp_api_key.strip() or DEFAULT_SHARED_API_KEY
                        st.session_state['shared_gemini_api_key'] = clean_inp
                        save_json({'gemini_api_key': clean_inp}, AI_CONFIG_FILE)
                        active_api_key = clean_inp
                        st.success('✅ API Key disimpan & aktif!')
                with c_k2:
                    st.markdown('<a href="https://aistudio.google.com/app/apikey" target="_blank" style="display:inline-block;text-align:center;width:100%;padding:6px 10px;background:#1E293B;color:#00E5FF;border:1px solid #334155;border-radius:6px;text-decoration:none;font-size:12.5px;font-weight:700;">🔑 Ambil Key AIzaSy Gratis (Google AI Studio)</a>', unsafe_allow_html=True)
                st.code('GEMINI_API_KEY = "AIzaSy...isi-key-Anda..."', language='toml'); st.caption('🔒 Simpan key di .streamlit/secrets.toml (atau env GEMINI_API_KEY), jangan ditulis di kode.')

        if gen_ai_clicked:
            with st.spinner(f'🤖 Sedang membedah data mentah & menyusun Target Rasional 7 Hari ke depan untuk {dlr_pil} (Maks 35 kata/sel)...'):
                raw_t1 = df_p_it[[kol_nm, 'TARGET', 'IDEAL', 'STU', 'NOTA', col_name_pct, 'STU LM', '% LM']].to_string(index=False) if not df_p_it.empty else "Kosong"
                raw_t2 = df_pr_it[[kol_nm, 'JML MP', 'AVG DB', 'DB', 'HOT', 'DB Q', 'STU', 'STU%']].to_string(index=False) if not df_pr_it.empty else "Kosong"
                raw_t3 = df_s_it[[kol_nm, 'SPK', 'SPK LM', 'GROWTH', 'ACC%', 'TERBUKA', 'BAF', '%BAF', 'ADIRA', '%ADIRA', 'SOF', '%SOF', 'IMFI', '%IMFI']].to_string(index=False) if not df_s_it.empty else "Kosong"
                raw_t4 = df_td_i[['TYPE', 'STU', 'STU LM', 'GROWTH', 'DB', 'HOT', 'DB Q', 'SPK', 'ACC%', 'STU%']].to_string(index=False) if not df_td_i.empty else "Kosong"
                raw_context_for_ai = (
                    f"Dealer: {dlr_pil} | Periode: 1 s.d {tanggal_akhir} {periode_bulan_tahun} | Horizon Action Plan: {horizon_lbl}.\n"
                    f"ACUAN TARGET RASIONAL 7 HARI KE DEPAN (WAJIB DIPAKAI AI AGAR TARGET MASUK AKAL & MENANTANG):\n"
                    f"- Dealer {dlr_pil} (STU saat ini {dlr_stu} dari target bulanan {dlr_targ}, Ideal tgl {tanggal_akhir}={dlr_ideal}): Target Rasional s.d tgl {tgl_sasaran} adalah tambahan +{tambahan_stu_wajib} STU (sehingga total mencapai {stu_sasaran_rasional} unit, atau {target_stu_per_hari:.1f} STU/hari), input {target_db_per_mp_hari:.1f} DB/MP/hari.\n"
                    f"- {wt_rational_str}\n- {wk_rational_str}\n- {wc_rational_str}\n"
                    f"TOTAL LEASING & FAKTUR KREDIT: Total SPK={cur_spk} (LM={cur_spk_lm}), ACC%={cur_acc_pct:.1f}%, SPK Terbuka={cur_open_spk} (target cair rasional +{potensi_stu_open} STU, tambah +{kebutuhan_spk_baru} SPK baru). Faktur Kredit Cair={tot_fak_kredit_curr} unit (BAF={fak_kredit_curr['BAF']}, ADIRA={fak_kredit_curr['ADIRA']}, SOF={fak_kredit_curr['SOF']}). Alasan Reject SPK Tertinggi: {top_rej_reasons_str}.\n"
                    f"--- DATA TABEL 1 (PENJUALAN KOOR/COUNTER) ---\n{raw_t1}\n--- DATA TABEL 2 (PROSES KERJA KOOR/COUNTER) ---\n{raw_t2}\n--- DATA TABEL 3 (LEASING KOOR/COUNTER) ---\n{raw_t3}\n--- DATA TABEL 4 (PER TYPE) ---\n{raw_t4}\n--- DATA SALESMAN AKTIF ZONA KRITIS (NON-COUNTER, DB>=2) ---\n{top_bad_raw_str}"
                )
                raw_context_for_ai += f"\n--- AKTIFITAS MD: SUB PROSPECT TYPE EVENT ---\n{default_cond_md}\n--- PROSPECT MEDIA SOSIAL ---\n{default_cond_sosmed}"
                try:
                    st.session_state[cache_key] = call_gemini_ai_table(active_api_key, raw_context_for_ai)
                    st.success('✅ Narasi Dinamis AI Gemini dengan Target Rasional (Maks 35 kata per kotak) berhasil dibuat!')
                except Exception:
                    st.session_state[cache_key] = build_smart_dynamic_narrative()
                    st.success('✅ Narasi Dinamis berbasis Analisa Silang Tabel 1–5 berhasil diperbarui! (Tips: Gunakan key berawalan `AIzaSy...` dari Google AI Studio pada menu pengaturan di kanan atas jika ingin terhubung ke Cloud LLM).')

        ai_res = st.session_state.get(cache_key, {})
        analisa_rows_data = [
            ('🏢 STU Dealer', limit_words(ai_res.get('cond_dealer', default_cond_dealer)), limit_words(ai_res.get('plan_dealer', default_plan_dealer))),
            ('🏍️ STU Per Type', limit_words(ai_res.get('cond_type', default_cond_type)), limit_words(ai_res.get('plan_type', default_plan_type))),
            ('💼 Leasing', limit_words(ai_res.get('cond_leasing', default_cond_leasing)), limit_words(ai_res.get('plan_leasing', default_plan_leasing))),
            ('👔 STU Koordinator', limit_words(ai_res.get('cond_koor', default_cond_koor)), limit_words(ai_res.get('plan_koor', default_plan_koor))),
            ('🏪 STU Counter', limit_words(ai_res.get('cond_counter', default_cond_counter)), limit_words(ai_res.get('plan_counter', default_plan_counter))),
            ('👥 STU Man Power', limit_words(ai_res.get('cond_mp', default_cond_mp)), limit_words(ai_res.get('plan_mp', default_plan_mp))),
            ('📣 Aktifitas MD (Event)', limit_words(ai_res.get('cond_md', default_cond_md)), limit_words(ai_res.get('plan_md', default_plan_md))),
            ('📱 Media Sosial', limit_words(ai_res.get('cond_sosmed', default_cond_sosmed)), limit_words(ai_res.get('plan_sosmed', default_plan_sosmed))),
        ]

        rows_html_list, copy_texts_dict = [], {}
        for idx_r, (item_lbl, c_txt, p_txt) in enumerate(analisa_rows_data):
            cid, pid = f"cond_txt_{idx_r}", f"plan_txt_{idx_r}"
            copy_texts_dict[cid], copy_texts_dict[pid] = c_txt, p_txt
            rows_html_list.append(f'<tr><td class="item-col">{item_lbl}</td><td class="cond-col">{format_paragraphs_html(c_txt)}<div class="cell-btn-wrap no-capture"><button class="cell-copy-btn" onclick="copyCellText(\'{cid}\', this)">📋 Copy Condition</button></div></td><td class="plan-col">{format_paragraphs_html(p_txt)}<div class="cell-btn-wrap no-capture"><button class="cell-copy-btn" onclick="copyCellText(\'{pid}\', this)">📋 Copy Action Plan</button></div></td></tr>')

        tabel_analisa_dom_id = f"Capture_Analisa_{re.sub(r'[^a-zA-Z0-9]', '_', dlr_pil)}"
        copy_payload_json, safe_dlr_fname = json.dumps(copy_texts_dict), re.sub(r'[^a-zA-Z0-9_-]', '_', f"Analisa_ActionPlan_{dlr_pil}_Tgl_{tanggal_akhir}")

        components.html(f"""<html><head><script src="https://cdnjs.cloudflare.com/ajax/libs/html2canvas/1.4.1/html2canvas.min.js"></script><style>
        body {{ margin: 0; padding: 0; background: transparent; font-family: 'Source Sans Pro', sans-serif; color: #E2E8F0; }}
        .compact-table-wrap {{ width: 100%; overflow-x: auto; border-radius: 8px; border: 1px solid #2D3748; background-color: #0E1117; }}
        .cap-title-banner {{ background: linear-gradient(90deg, #4318FF 0%, #1E293B 100%); color: #FFFFFF !important; font-family: 'Source Sans Pro', sans-serif; font-size: 15px; font-weight: 800; padding: 9px 14px; border-bottom: 2px solid #3B18FF; letter-spacing: 0.4px; text-align: left; }}
        table.narrative-tbl {{ width: 100%; border-collapse: collapse; font-family: 'Source Sans Pro', sans-serif; font-size: 13.5px; background-color: #0E1117; }}
        table.narrative-tbl th {{ background-color: #4318FF; color: #FFFFFF; font-weight: 700; padding: 10px 14px; border: 1px solid #2D3748; text-align: left; font-size: 13.5px; }}
        table.narrative-tbl td {{ padding: 10px 14px; border: 1px solid #2D3748; text-align: left; vertical-align: top; color: #E2E8F0; line-height: 1.55; font-weight: 400; }}
        table.narrative-tbl td.item-col {{ width: 15%; color: #E2E8F0; text-align: left; background-color: #161B26; white-space: nowrap; font-weight: 600; }}
        table.narrative-tbl td.cond-col {{ width: 42.5%; background-color: #0E1117; }}
        table.narrative-tbl td.plan-col {{ width: 42.5%; background-color: #111622; }}
        p.nar-p {{ margin: 0 0 6px 0; padding: 0; color: #E2E8F0; font-weight: 400; line-height: 1.55; text-align: left; }}
        p.nar-p:last-of-type {{ margin-bottom: 4px; }}
        .cell-btn-wrap {{ display: flex; justify-content: flex-end; margin-top: 4px; }}
        .cell-copy-btn {{ background-color: #1E293B; color: #CBD5E1; border: 1px solid #334155; border-radius: 5px; padding: 3px 9px; font-size: 11.5px; font-weight: 600; cursor: pointer; transition: all 0.15s; }}
        .cell-copy-btn:hover {{ background-color: #4318FF; color: #FFFFFF; border-color: #4318FF; }}
        .bottom-toolbar {{ display: flex; align-items: center; gap: 10px; margin-top: 10px; padding-bottom: 6px; flex-wrap: wrap; }}
        .act-btn {{ background-color: #1E293B; color: #E2E8F0; border: 1px solid #334155; border-radius: 6px; padding: 6px 13px; font-size: 12.5px; font-weight: 600; cursor: pointer; display: inline-flex; align-items: center; gap: 6px; transition: all 0.15s; }}
        .act-btn:hover {{ background-color: #4318FF; color: #FFFFFF; border-color: #4318FF; }}
        .print-btn {{ background-color: #4318FF; color: #FFFFFF; border: 1px solid #4318FF; }}
        .print-btn:hover {{ background-color: #3311CC; }}
        #status_analisa {{ font-size: 12.5px; color: #00FF66; font-weight: 600; }}
    </style></head><body>
    <div id="{tabel_analisa_dom_id}" class="compact-table-wrap">
        <div class="cap-title-banner capture-only-banner" style="display:none;">🐺 {jdl_rev_6}</div>
        <table class="narrative-tbl">
            <thead><tr><th style="width:15%;">ITEM</th><th style="width:42.5%;">CONDITION (ANALISA SINGKAT MAKS 35 KATA — IDEAL: DB Q 30% | STU% 10% | ACC% 50%)</th><th style="width:42.5%;">ACTION PLAN (TARGET RASIONAL MAKS 35 KATA — {horizon_lbl.upper()})</th></tr></thead>
            <tbody>{''.join(rows_html_list)}</tbody>
        </table>
    </div>
    <div class="bottom-toolbar">
        <button class="act-btn" onclick="captureAnalisaTbl('copy')">📸 Copy Tabel (Gambar)</button>
        <button class="act-btn" onclick="captureAnalisaTbl('jpg')">📥 Download JPG</button>
        <button class="act-btn print-btn" onclick="printAnalisaTbl()">🖨️ Print / Print Preview Tabel Analisa</button>
        <span id="status_analisa"></span>
    </div>
    <script>
    const cellTexts = {copy_payload_json};
    function fallbackCopyText(text) {{
        const ta = document.createElement('textarea'); ta.value = text; ta.style.position = 'fixed'; ta.style.left = '-9999px';
        document.body.appendChild(ta); ta.select(); document.execCommand('copy'); document.body.removeChild(ta);
    }}
    async function copyCellText(keyId, btnEl) {{
        const txt = cellTexts[keyId] || ''; const origLabel = btnEl.innerText;
        try {{ await navigator.clipboard.writeText(txt); }} catch (e) {{ fallbackCopyText(txt); }}
        btnEl.innerText = '✅ Tersalin!'; btnEl.style.backgroundColor = '#00C853'; btnEl.style.color = '#FFFFFF';
        setTimeout(() => {{ btnEl.innerText = origLabel; btnEl.style.backgroundColor = ''; btnEl.style.color = ''; }}, 2000);
    }}
    async function captureAnalisaTbl(mode) {{
        const msg = document.getElementById('status_analisa'); msg.style.color = '#FFEA00'; msg.innerText = 'Memproses gambar tabel analisa...';
        const tblEl = document.getElementById('{tabel_analisa_dom_id}');
        const btns = tblEl.querySelectorAll('.no-capture');
        const bannerEl = tblEl.querySelector('.capture-only-banner');
        btns.forEach(b => b.style.display = 'none');
        if (bannerEl) bannerEl.style.display = 'block';
        try {{
            const canvas = await html2canvas(tblEl, {{ backgroundColor: '#0E1117', scale: 2, useCORS: true, logging: false }});
            btns.forEach(b => b.style.display = 'flex');
            if (bannerEl) bannerEl.style.display = 'none';
            if (mode === 'copy') {{
                canvas.toBlob(async (blob) => {{
                    try {{
                        await navigator.clipboard.write([new ClipboardItem({{ 'image/png': blob }})]);
                        msg.style.color = '#00FF66'; msg.innerText = '✅ Gambar Tabel Analisa + Judul disalin ke Clipboard! (Siap Ctrl+V di WA/Word/Excel)';
                    }} catch (err) {{
                        const a = document.createElement('a'); a.href = canvas.toDataURL('image/jpeg', 0.95); a.download = '{safe_dlr_fname}.jpg'; a.click();
                        msg.style.color = '#00E5FF'; msg.innerText = '📥 Gambar otomatis diunduh sebagai JPG.';
                    }}
                    setTimeout(() => msg.innerText = '', 4500);
                }}, 'image/png');
            }} else {{
                const a = document.createElement('a'); a.href = canvas.toDataURL('image/jpeg', 0.95); a.download = '{safe_dlr_fname}.jpg'; a.click();
                msg.style.color = '#00FF66'; msg.innerText = '✅ File {safe_dlr_fname}.jpg berhasil diunduh!';
                setTimeout(() => msg.innerText = '', 4000);
            }}
        }} catch (e) {{
            btns.forEach(b => b.style.display = 'flex');
            if (bannerEl) bannerEl.style.display = 'none';
            msg.style.color = '#FF007F'; msg.innerText = 'Gagal: ' + e.message;
        }}
    }}
    function printAnalisaTbl() {{
        const tblClone = document.getElementById('{tabel_analisa_dom_id}').cloneNode(true);
        tblClone.querySelectorAll('.no-capture').forEach(el => el.remove());
        tblClone.querySelectorAll('.capture-only-banner').forEach(el => el.remove());
        const w = window.open('', '_blank', 'width=1100,height=800');
        w.document.write(`
            <html><head><title>Print Review Analisa — DEALER {dlr_pil} (1 s.d {tanggal_akhir} {periode_bulan_tahun})</title>
            <style>
                @page {{ size: A4 landscape; margin: 12mm; }}
                body {{ font-family: Arial, sans-serif; color: #111827; background: #FFFFFF; margin: 0; padding: 10px; }}
                h2 {{ color: #1E3A8A; margin: 0 0 4px 0; font-size: 18px; }}
                .sub {{ color: #4B5563; font-size: 12.5px; margin-bottom: 12px; font-weight: 600; }}
                table.narrative-tbl {{ width: 100%; border-collapse: collapse; font-size: 12px; }}
                table.narrative-tbl th {{ background-color: #1E3A8A !important; color: #FFFFFF !important; -webkit-print-color-adjust: exact; print-color-adjust: exact; padding: 8px 10px; border: 1px solid #1F2937; text-align: left; }}
                table.narrative-tbl td {{ padding: 8px 10px; border: 1px solid #374151; vertical-align: top; text-align: left; line-height: 1.5; color: #111827; }}
                table.narrative-tbl td.item-col {{ font-weight: bold; width: 14%; background-color: #F3F4F6 !important; -webkit-print-color-adjust: exact; print-color-adjust: exact; text-align: left; white-space: nowrap; }}
                p.nar-p {{ margin: 0; }}
                .no-print-bar {{ margin-bottom: 12px; padding: 8px 12px; background: #EEF2FF; border: 1px solid #4318FF; border-radius: 6px; display: flex; justify-content: space-between; align-items: center; }}
                .no-print-bar button {{ background: #4318FF; color: #fff; border: none; padding: 6px 14px; border-radius: 5px; font-weight: bold; cursor: pointer; }}
                @media print {{ .no-print-bar {{ display: none !important; }} }}
            </style></head><body>
            <div class="no-print-bar">
                <span><strong>🖨️ Print Preview:</strong> Gunakan pengaturan <em>Landscape</em> pada dialog printer untuk hasil cetak terbaik.</span>
                <button onclick="window.print()">🖨️ Cetak Sekarang / Simpan PDF</button>
            </div>
            <h2>TABEL ANALISA & ACTION PLAN — DEALER {dlr_pil}</h2>
            <div class="sub">Periode Analisa: 1 s.d {tanggal_akhir} {periode_bulan_tahun} | Fokus Horizon: {horizon_lbl} | Acuan Ideal: DB Quality 30%, STU% 10%, ACC% 50%</div>
            ` + tblClone.innerHTML + `
            <script>setTimeout(() => window.print(), 400);<\\/script></body></html>
        `);
        w.document.close();
    }}
    </script></body></html>""", height=620, scrolling=True)

        st.divider()
        bsg_render_stu_per_leasing_faktur()
